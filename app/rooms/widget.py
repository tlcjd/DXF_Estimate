from __future__ import annotations

import traceback
import hashlib
import math
import json
import re
import uuid
from copy import deepcopy
from datetime import datetime
from collections import defaultdict
from pathlib import Path
from PySide6.QtCore import Qt, Signal, QTimer, QItemSelectionModel, QEvent, QMimeData
from PySide6.QtGui import QColor, QBrush, QKeySequence, QCursor, QFont, QShortcut, QDrag
from PySide6.QtWidgets import (QWidget,QVBoxLayout,QHBoxLayout,QLabel,QComboBox,QPushButton,QSpinBox,QDoubleSpinBox,QCheckBox,QSplitter,QTableWidget,QTableWidgetItem,QHeaderView,QAbstractItemView,QGroupBox,QFormLayout,QGridLayout,QFileDialog,QMessageBox,QDialog,QSizePolicy,QDialogButtonBox,QApplication,QLineEdit,QInputDialog,QToolTip,QToolButton,QRadioButton,QButtonGroup,QListWidget,QListWidgetItem,QColorDialog,QStyledItemDelegate,QTabWidget,QFrame)
from app.common.models import RoomRecord, normalize_room_name, room_name_is_missing
from app.common.room_categories import ROOM_CATEGORIES, classify_room_name, normalize_room_category_rules
from app.core.paths import MATERIAL_FINISH_PATH, WALL_FINISH_ASSIGNMENTS_DIR, WORKING_DXF_PATH, TEMP_PROJECT_DIR
from app.rooms.core import (
    validate_room_numbers, save_rooms_to_dxf, create_room_blocks, refresh_room_blocks,
    create_wall_finish_lines, ROOM_WALL_LINK_LAYERS, room_wall_link_layers, write_room_wall_link_dxf,
    read_room_wall_link_dxf, validate_room_wall_link_items, match_room_wall_link_items,
    room_program_id, wall_program_id, link_object_id,
)
from app.rooms.finish_panel import (
    RoomFinishPanel, calculate_room_finish_value, normalize_room_finish_assignment,
    normalize_room_finish_section, normalize_finish_group_type, normalize_finish_group_division, FINISH_GROUP_TYPES,
)
from app.common.finish_calculations import normalize_formula_text
from app.common.material_units import normalize_material_unit
from app.cad.colors import aci_to_hex, color_to_aci
from app.common.aci_color_dialog import get_aci_color
from app.rooms.preview import RoomPreview
from app.common.geometry import point_in_polygon

GROUP_COLOR_PALETTE = (
    "#4E79A7", "#F28E2B", "#59A14F", "#E15759", "#B07AA1",
    "#76B7B2", "#EDC948", "#FF9DA7", "#9C755F", "#6B8E23",
)


def finish_group_color(group_id: str, value: str = "") -> str:
    color = QColor(str(value or ""))
    if color.isValid():
        return color.name().upper()
    digest = hashlib.sha1(str(group_id or "group").encode("utf-8")).digest()
    return GROUP_COLOR_PALETTE[int.from_bytes(digest[:2], "big") % len(GROUP_COLOR_PALETTE)]

def finish_group_aci(group_id: str, color: str = "", aci=None) -> int:
    """그룹의 기준색은 ACI 번호로 관리한다. 구형 RGB 그룹은 최근접 ACI로 이관한다."""
    if aci is not None and str(aci).strip() != "":
        return color_to_aci(color, aci)
    return color_to_aci(finish_group_color(group_id, color))


def wall_key_room_handle(key: str) -> str:
    return str(key or "").split(":", 1)[0]


def wall_key_length_m(room: RoomRecord, key: str) -> float:
    parts = str(key or "").split(":")
    if not parts or parts[0] != str(room.boundary_handle or ""):
        return 0.0
    try:
        if len(parts) == 2:
            points = room.points
            index = int(parts[1])
        elif len(parts) == 3 and parts[1].lower().startswith("d"):
            deduction_index = int(parts[1][1:])
            points = list(getattr(room, "deduction_points", []) or [])[deduction_index]
            index = int(parts[2])
        else:
            return 0.0
        x1, y1 = points[index]
        x2, y2 = points[(index + 1) % len(points)]
        return ((x2 - x1) ** 2 + (y2 - y1) ** 2) ** 0.5 / 1000.0
    except (ValueError, IndexError, TypeError):
        return 0.0



def floor_sort_key(floor: str):
    """건축 층 표기를 실제 층 순서로 정렬한다.

    기본 순서: 깊은 지하층 -> B1F -> 각 층의 중간층(M1...) -> 지상층 -> RF/R -> PH -> PHR.
    알 수 없는 표기는 맨 뒤에서 문자순으로 정렬한다.
    """
    raw = (floor or "").strip().upper().replace(" ", "")
    if not raw:
        return (9, 0, 0, "")

    m = re.fullmatch(r"B(\d+)F?(?:-?M(\d+))?", raw)
    if m:
        base = int(m.group(1))
        mid = int(m.group(2) or 0)
        return (0, -base, mid, raw)

    m = re.fullmatch(r"(\d+)F?(?:-?M(\d+))?", raw)
    if m:
        base = int(m.group(1))
        mid = int(m.group(2) or 0)
        return (1, base, mid, raw)

    if raw in {"R", "RF"}:
        return (2, 0, 0, raw)

    m = re.fullmatch(r"PH(\d+)", raw)
    if m:
        return (3, int(m.group(1)), 0, raw)

    if raw in {"PHR", "PHRF", "PH-R"}:
        return (4, 0, 0, raw)

    return (8, 0, 0, raw)


class ClipboardTableWidget(QTableWidget):
    """임시 엑셀 편집창용: Excel과 유사한 복사/붙여넣기/Delete 지원."""
    def keyPressEvent(self, event):
        if event.matches(QKeySequence.Copy):
            ranges = self.selectedRanges()
            if ranges:
                r = ranges[0]
                lines = []
                for row in range(r.topRow(), r.bottomRow() + 1):
                    vals = []
                    for col in range(r.leftColumn(), r.rightColumn() + 1):
                        item = self.item(row, col)
                        vals.append(item.text() if item else "")
                    lines.append("\t".join(vals))
                QApplication.clipboard().setText("\n".join(lines))
            return
        if event.matches(QKeySequence.Paste):
            text = QApplication.clipboard().text()
            if not text:
                return
            start_row = self.currentRow() if self.currentRow() >= 0 else 0
            start_col = self.currentColumn() if self.currentColumn() >= 0 else 0
            for r_off, line in enumerate(text.replace("\r\n", "\n").replace("\r", "\n").split("\n")):
                if line == "" and r_off == len(text.splitlines()):
                    continue
                for c_off, value in enumerate(line.split("\t")):
                    row, col = start_row + r_off, start_col + c_off
                    if row >= self.rowCount() or col >= self.columnCount():
                        continue
                    item = self.item(row, col)
                    if item is not None and (item.flags() & Qt.ItemIsEditable):
                        item.setText(value)
            return
        if event.key() in (Qt.Key_Delete, Qt.Key_Backspace):
            for item in self.selectedItems():
                if item.flags() & Qt.ItemIsEditable:
                    item.setText("")
            return
        super().keyPressEvent(event)


class OpeningReadTable(QTableWidget):
    delete_requested = Signal()

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.setAcceptDrops(True)

    def keyPressEvent(self, event):
        if event.key() in (Qt.Key_Delete, Qt.Key_Backspace, Qt.Key_F4) and not self.state() == QAbstractItemView.State.EditingState:
            if event.key() == Qt.Key_F4:
                rows = {index.row() for index in self.selectionModel().selectedRows()}
                if not rows and self.currentRow() >= 0:
                    rows = {self.currentRow()}
                count = max(1, len(rows))
                text = "현재 행을 삭제하시겠습니까?" if count == 1 else f"선택한 {count}개 행을 삭제하시겠습니까?"
                if QMessageBox.question(
                    self, "목록 삭제", text,
                    QMessageBox.Yes | QMessageBox.No, QMessageBox.Yes
                ) != QMessageBox.Yes:
                    return
            self.delete_requested.emit()
            return
        super().keyPressEvent(event)

    def dragEnterEvent(self, event):
        if event.mimeData().hasFormat("application/x-dxf-window-code"):
            event.acceptProposedAction(); return
        super().dragEnterEvent(event)

    def dragMoveEvent(self, event):
        if event.mimeData().hasFormat("application/x-dxf-window-code"):
            event.acceptProposedAction(); return
        super().dragMoveEvent(event)

    def dropEvent(self, event):
        if event.mimeData().hasFormat("application/x-dxf-window-code"):
            try:
                code = bytes(event.mimeData().data("application/x-dxf-window-code")).decode("utf-8").strip().upper()
            except Exception:
                code = ""
            dialog = self.parent()
            if code and hasattr(dialog, "apply_dragged_code"):
                dialog.apply_dragged_code(code)
                event.acceptProposedAction(); return
        super().dropEvent(event)


class OpeningReadDialog(QDialog):
    """연결 CAD에서 읽은 문/창 후보를 층별로 검토하고 최종 부호/길이를 보정한다."""

    HEADERS = ["층", "읽은 부호", "검출 창길이(mm)", "적용 부호", "적용 창길이(mm)", "판정"]

    def __init__(self, owner, filter_ids=None, assignment_mode: bool = False):
        super().__init__(owner)
        self.owner = owner
        self.filter_ids = {str(v or "") for v in (filter_ids or []) if str(v or "")}
        self.assignment_mode = bool(assignment_mode)
        self._updating = False
        # 적용부호/적용길이는 편집창 하단 '적용'을 누르기 전까지 원본에 쓰지 않는다.
        self._pending_edits: dict[tuple[str, str], dict[str, object]] = {}
        self.collapsed_floors: set[str] = set()
        self.sort_column = 0
        self.sort_ascending = True
        self.setWindowTitle("창호 지정목록" if self.assignment_mode else "창호 리스트")
        # 전체 창호 리스트도 도면을 덜 가리도록 폭을 크게 줄이고, 열은 최초 1회 내용 기준 Fit 후 수동 조절한다.
        self.resize(620 if self.assignment_mode else 620, 760 if self.assignment_mode else 620)

        root = QVBoxLayout(self)
        top = QHBoxLayout()
        self.read_btn = QPushButton("창호 읽기")
        self.read_btn.setToolTip("현재층/전층 범위와 연동 CAD/베이스 DXF 소스를 선택하여 도곽 내부 문/창 후보를 읽습니다.")
        self.assign_btn = QPushButton("창호지정")
        self.assign_btn.setToolTip("선택 행에 창호 탭의 창호 리스트 부호를 지정합니다. F5와 동일합니다.")
        self.delete_btn = QPushButton("선택 삭제")
        self.delete_btn.setToolTip("선택한 검출 항목만 목록/도곽 표시에서 제거합니다. CAD 원본 객체는 삭제하지 않습니다.")
        self.collapse_all_btn = QPushButton("전체 접기")
        self.expand_all_btn = QPushButton("전체 펼치기")
        top.addWidget(self.read_btn)
        top.addWidget(self.assign_btn)
        top.addWidget(self.delete_btn)
        top.addSpacing(10)
        top.addWidget(self.collapse_all_btn)
        top.addWidget(self.expand_all_btn)
        top.addStretch(1)
        self.status_label = QLabel("")
        self.status_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        top.addWidget(self.status_label)
        root.addLayout(top)

        hint = QLabel(
            "층 행을 더블클릭하면 해당 층을 접거나 펼칩니다. 열 머리글을 누르면 평면 실목록처럼 정렬됩니다. "
            "여러 행 선택 후 '적용 부호' 또는 '적용 창길이' 셀 하나를 수정하면 선택행에 같은 값이 입력됩니다. "
            "실제 평면 반영은 하단 '적용'을 눌러 확정합니다. Delete: 선택 삭제"
        )
        hint.setWordWrap(True)
        root.addWidget(hint)

        self.table = OpeningReadTable(0, len(self.HEADERS), self)
        self.table.setHorizontalHeaderLabels(self.HEADERS)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.ExtendedSelection)
        self.table.setAlternatingRowColors(True)
        self.table.setEditTriggers(QAbstractItemView.DoubleClicked | QAbstractItemView.EditKeyPressed | QAbstractItemView.SelectedClicked)
        header = self.table.horizontalHeader()
        header.setSectionsClickable(True)
        header.setSortIndicatorShown(True)
        header.setSortIndicator(self.sort_column, Qt.AscendingOrder)
        # 전체 목록/지정목록 모두 최초 1회 내용 기준 Fit 후 사용자가 직접 열너비를 조절할 수 있게 한다.
        header.setSectionResizeMode(QHeaderView.Interactive)
        header.setStretchLastSection(False)
        self._assignment_columns_fitted = False
        # 목록 행높이는 기존 대비 약 85%로 줄인다.
        current_row_h = max(1, self.table.verticalHeader().defaultSectionSize())
        self.table.verticalHeader().setDefaultSectionSize(max(18, int(round(current_row_h * 0.85))))
        self.table.verticalHeader().setMinimumSectionSize(16)
        root.addWidget(self.table, 1)

        buttons = QDialogButtonBox(QDialogButtonBox.Close)
        self.apply_btn = buttons.addButton("적용", QDialogButtonBox.ActionRole)
        if buttons.button(QDialogButtonBox.Close) is not None:
            buttons.button(QDialogButtonBox.Close).setText("닫기")
        self.apply_btn.setToolTip("편집창에서 수정한 적용부호/적용 창길이를 평면 창호에 반영합니다.")
        self.apply_btn.clicked.connect(self.apply_pending_changes)
        buttons.rejected.connect(self.close)
        root.addWidget(buttons)

        self.read_btn.clicked.connect(self._request_opening_read)
        self.assign_btn.clicked.connect(self.open_catalog_for_selected)
        self.delete_btn.clicked.connect(self.delete_selected)
        self.assign_shortcut = QShortcut(QKeySequence(Qt.Key_F5), self)
        self.assign_shortcut.activated.connect(self.open_catalog_for_selected)
        self.collapse_all_btn.clicked.connect(self.collapse_all_floors)
        self.expand_all_btn.clicked.connect(self.expand_all_floors)
        self.table.delete_requested.connect(self.delete_selected)
        self.table.itemChanged.connect(self._on_item_changed)
        self.table.itemSelectionChanged.connect(self._sync_preview_selection_from_table)
        self.table.cellDoubleClicked.connect(self._on_cell_double_clicked)
        header.sectionClicked.connect(self._on_header_clicked)
        if self.assignment_mode:
            # F5로 선택 창호만 호출한 창은 재검출 관리가 아니라 지정 전용이다.
            self.read_btn.hide()
            self.collapse_all_btn.hide()
            self.expand_all_btn.hide()
            hint.setText(
                "현재 도면에서 선택한 창호만 표시합니다. '적용 부호'를 직접 입력하거나 F5/창호지정으로 창호 선택창을 열 수 있습니다. "
                "수정 후 하단 '적용'을 눌러 실제 평면에 반영합니다."
            )
        self.refresh()

    def _sync_preview_selection_from_table(self) -> None:
        """목록 선택을 도면 창호 선택 강조와 즉시 동기화한다."""
        if self._updating:
            return
        ids = {
            self._row_key(index.row())[1]
            for index in self.table.selectionModel().selectedRows()
            if self._row_key(index.row())[1]
        }
        self.owner._opening_edit_selected_ids = set(ids)
        self.owner.preview.set_selected_opening_ids(ids)
        self.owner._update_opening_edit_status(ids)
        if self.owner._opening_catalog_dialog is not None:
            self.owner._opening_catalog_dialog.refresh()
        if self.assignment_mode:
            self.status_label.setText(f"목록 선택 {len(ids):,}개 · 도면 강조 연동")

    def _fit_assignment_columns_once(self) -> None:
        if getattr(self, "_assignment_columns_fitted", False):
            return
        self.table.resizeColumnsToContents()
        # 지나치게 좁거나 긴 값 하나 때문에 표 전체가 벌어지지 않도록 실사용 폭으로 제한한다.
        mins = [52, 72, 96, 72, 96, 72]
        maxs = [82, 118, 132, 118, 132, 165]
        for col in range(self.table.columnCount()):
            width = self.table.columnWidth(col)
            self.table.setColumnWidth(col, max(mins[col], min(width + 10, maxs[col])))
        self._assignment_columns_fitted = True

    def _request_opening_read(self) -> None:
        # 새로 읽으면 현재 미적용 편집값은 버리고 새 후보를 기준으로 시작한다.
        self._pending_edits.clear()
        self.owner.request_opening_read_from_dialog()

    def _pending_value(self, floor: str, read_id: str, field: str, default):
        return self._pending_edits.get((str(floor or ""), str(read_id or "")), {}).get(field, default)

    def _set_pending(self, floor: str, read_id: str, field: str, value) -> None:
        key = (str(floor or ""), str(read_id or ""))
        if not key[1]:
            return
        self._pending_edits.setdefault(key, {})[str(field)] = value

    def stage_catalog_code(self, code: str, ids=None) -> None:
        code = str(code or "").strip().upper()
        master = self.owner._window_catalog_by_code().get(code)
        if master is None:
            QMessageBox.warning(self, "창호 지정", "창호 리스트에 없는 부호는 입력할 수 없습니다.")
            return
        targets = {str(v or "") for v in (ids if ids is not None else self.owner._opening_edit_selected_ids) if str(v or "")}
        if not targets:
            QMessageBox.information(self, "창호 지정", "적용할 창호를 먼저 선택해 주세요.")
            return
        try:
            master_length = float(master.get("applied_length_mm", 0.0) or 0.0)
        except Exception:
            master_length = 0.0
        changed = 0
        for row in range(self.table.rowCount()):
            floor, read_id = self._row_key(row)
            if read_id not in targets:
                continue
            self._set_pending(floor, read_id, "applied_code", code)
            self._set_pending(floor, read_id, "applied_length_mm", master_length if master_length > 0 else "")
            changed += 1
        self.refresh(f"{code} · {changed:,}개 입력 대기 · 하단 적용 필요")

    def apply_pending_changes(self) -> None:
        if not self._pending_edits:
            self.status_label.setText("적용할 변경사항 없음")
            return
        changed = self.owner.apply_opening_read_pending(self._pending_edits)
        self._pending_edits.clear()
        self.refresh(f"{changed:,}개 변경 반영" if changed else "변경사항 확인 완료")

    @staticmethod
    def _friendly_source(marker: dict) -> str:
        source = str(marker.get("source", "") or "")
        kind = str(marker.get("kind", "") or "")
        labels = {
            "block": "블록",
            "block+color": "블록 + 색상선",
            "arc90": "90° ARC",
            "color-cluster": "색상선 묶음",
        }
        text = labels.get(source, source or "기타")
        if kind == "door":
            return f"{text} / 문"
        if kind == "window":
            return f"{text} / 창"
        return text

    @staticmethod
    def _length_text(value) -> str:
        try:
            number = float(value)
        except (TypeError, ValueError):
            return ""
        if number <= 0:
            return ""
        if abs(number - round(number)) < 0.01:
            return f"{int(round(number)):,}"
        return f"{number:,.1f}"

    @staticmethod
    def _number_key(value):
        try:
            return float(value)
        except (TypeError, ValueError):
            return -1.0

    def _marker_sort_key(self, marker: dict):
        col = int(self.sort_column)
        if col == 1:
            return str(marker.get("read_label", marker.get("label", "")) or "").casefold()
        if col == 2:
            return self._number_key(marker.get("detected_length_mm"))
        if col == 3:
            return str(marker.get("applied_code", "") or "").casefold()
        if col == 4:
            return self._number_key(marker.get("applied_length_mm"))
        if col == 5:
            return self._friendly_source(marker).casefold()
        # 층 정렬일 때 같은 층 안에서는 도면상 위→아래, 왼쪽→오른쪽 순서를 유지한다.
        return (-self._number_key(marker.get("y")), self._number_key(marker.get("x")))

    def refresh(self, status_text: str | None = None) -> None:
        previous_selected_ids = set()
        if self.assignment_mode and hasattr(self, "table"):
            try:
                previous_selected_ids = {
                    self._row_key(index.row())[1]
                    for index in self.table.selectionModel().selectedRows()
                    if self._row_key(index.row())[1]
                }
            except Exception:
                previous_selected_ids = set()
        self._updating = True
        try:
            self.table.setRowCount(0)
            floor_map: dict[str, list[dict]] = {}
            for floor, markers in dict(self.owner.opening_read_results or {}).items():
                values = [dict(marker or {}) for marker in (markers or [])]
                if self.filter_ids:
                    values = [marker for marker in values if str(marker.get("read_id") or "") in self.filter_ids]
                if values or not self.filter_ids:
                    floor_map[str(floor or "")] = values

            floor_names = sorted(floor_map, key=floor_sort_key, reverse=(self.sort_column == 0 and not self.sort_ascending))
            total_items = sum(len(values) for values in floor_map.values())
            for floor in floor_names:
                markers = list(floor_map.get(floor, []) or [])
                if self.sort_column != 0:
                    markers.sort(key=self._marker_sort_key, reverse=not self.sort_ascending)
                else:
                    markers.sort(key=self._marker_sort_key)

                collapsed = floor in self.collapsed_floors
                header_row = self.table.rowCount()
                self.table.insertRow(header_row)
                header_values = [("▶ " if collapsed else "▼ ") + floor, f"{len(markers):,}개", "", "", "", ""]
                for col, value in enumerate(header_values):
                    item = QTableWidgetItem(str(value))
                    item.setFlags(item.flags() & ~Qt.ItemIsEditable)
                    item.setBackground(QBrush(QColor("#DCEEFF")))
                    font = item.font(); font.setBold(True); item.setFont(font)
                    if col == 0:
                        item.setData(Qt.UserRole, [floor, ""])
                        item.setData(Qt.UserRole + 1, "floor_header")
                        item.setToolTip("더블클릭하면 이 층 창호 목록을 접거나 펼칩니다.")
                    self.table.setItem(header_row, col, item)

                for marker in markers:
                    row = self.table.rowCount()
                    self.table.insertRow(row)
                    read_id = str(marker.get("read_id", "") or "")
                    values = [
                        floor,
                        str(marker.get("read_label", marker.get("label", "")) or ""),
                        self._length_text(marker.get("detected_length_mm")),
                        str(self._pending_value(floor, read_id, "applied_code", marker.get("applied_code", "")) or ""),
                        self._length_text(self._pending_value(floor, read_id, "applied_length_mm", marker.get("applied_length_mm"))),
                        self._friendly_source(marker),
                    ]
                    applied_code = str(values[3] or "").strip()
                    # 적용부호가 있는 창호와 미지정 창호를 목록 글씨색으로 즉시 구분한다.
                    row_foreground = QBrush(QColor("#0B6E2E" if applied_code else "#B42318"))
                    for col, value in enumerate(values):
                        item = QTableWidgetItem(value)
                        if col == 0:
                            item.setData(Qt.UserRole, [floor, read_id])
                            item.setData(Qt.UserRole + 1, "data")
                        if col not in (3, 4):
                            item.setFlags(item.flags() & ~Qt.ItemIsEditable)
                        item.setForeground(row_foreground)
                        self.table.setItem(row, col, item)
                    self.table.setRowHidden(row, collapsed)

            header = self.table.horizontalHeader()
            header.setSortIndicator(
                self.sort_column,
                Qt.AscendingOrder if self.sort_ascending else Qt.DescendingOrder,
            )
            self.collapse_all_btn.setEnabled(bool(floor_names))
            self.expand_all_btn.setEnabled(bool(floor_names))
            if status_text is not None:
                self.status_label.setText(str(status_text))
            elif total_items:
                self.status_label.setText(
                    f"선택 창호 {total_items:,}개" if self.assignment_mode else f"총 {total_items:,}개 / {len(floor_names):,}개 층"
                )
            else:
                self.status_label.setText("선택 창호 없음" if self.assignment_mode else "읽은 창호 없음")
            if self.assignment_mode and total_items:
                # 최초 호출은 도면에서 넘겨받은 모든 항목을 선택하고, 이후 refresh에서는 사용자가 고른 부분선택을 유지한다.
                target_ids = previous_selected_ids or set(self.filter_ids)
                self.table.clearSelection()
                selection = self.table.selectionModel()
                for row in range(self.table.rowCount()):
                    read_id = self._row_key(row)[1]
                    if read_id and read_id in target_ids:
                        selection.select(
                            self.table.model().index(row, 0),
                            QItemSelectionModel.Select | QItemSelectionModel.Rows,
                        )
            if not getattr(self, "_assignment_columns_fitted", False):
                QTimer.singleShot(0, self._fit_assignment_columns_once)
        finally:
            self._updating = False

    def _row_key(self, row: int):
        item = self.table.item(row, 0)
        data = item.data(Qt.UserRole) if item is not None else None
        if isinstance(data, (list, tuple)) and len(data) >= 2 and str(data[1] or ""):
            return str(data[0] or ""), str(data[1] or "")
        return "", ""

    def _update_assignment_row_color(self, row: int) -> None:
        if row < 0 or not self._row_key(row)[1]:
            return
        code_item = self.table.item(row, 3)
        assigned = bool(str(code_item.text() if code_item is not None else "").strip())
        brush = QBrush(QColor("#0B6E2E" if assigned else "#B42318"))
        for col in range(self.table.columnCount()):
            item = self.table.item(row, col)
            if item is not None:
                item.setForeground(brush)

    def _floor_header_at_row(self, row: int) -> str:
        item = self.table.item(row, 0)
        if item is None or item.data(Qt.UserRole + 1) != "floor_header":
            return ""
        data = item.data(Qt.UserRole)
        if isinstance(data, (list, tuple)) and data:
            return str(data[0] or "")
        return ""

    def _selected_rows(self, include_row: int | None = None) -> list[int]:
        rows = {index.row() for index in self.table.selectionModel().selectedRows() if self._row_key(index.row())[1]}
        if include_row is not None and include_row >= 0 and self._row_key(int(include_row))[1]:
            rows.add(int(include_row))
        return sorted(rows)

    def _on_header_clicked(self, column: int) -> None:
        column = int(column)
        if column == self.sort_column:
            self.sort_ascending = not self.sort_ascending
        else:
            self.sort_column = column
            self.sort_ascending = True
        self.refresh()

    def _on_cell_double_clicked(self, row: int, _column: int) -> None:
        floor = self._floor_header_at_row(int(row))
        if not floor:
            return
        self._set_floor_collapsed(floor, floor not in self.collapsed_floors)

    def _set_floor_collapsed(self, floor: str, collapsed: bool) -> None:
        floor = str(floor or "")
        if not floor:
            return
        if collapsed:
            self.collapsed_floors.add(floor)
        else:
            self.collapsed_floors.discard(floor)
        # 전체 재구성 대신 현재 표에서 바로 숨겨 목록 선택상태의 불필요한 흔들림을 줄인다.
        active = False
        for row in range(self.table.rowCount()):
            header_floor = self._floor_header_at_row(row)
            if header_floor:
                active = header_floor == floor
                if active:
                    item = self.table.item(row, 0)
                    if item is not None:
                        item.setText(("▶ " if collapsed else "▼ ") + floor)
                continue
            row_floor, read_id = self._row_key(row)
            if active and row_floor == floor and read_id:
                self.table.setRowHidden(row, collapsed)
            elif active and row_floor != floor:
                active = False

    def collapse_all_floors(self) -> None:
        floors = {str(floor or "") for floor in dict(self.owner.opening_read_results or {}).keys() if str(floor or "")}
        self.collapsed_floors.update(floors)
        self.refresh()

    def expand_all_floors(self) -> None:
        self.collapsed_floors.clear()
        self.refresh()

    def _on_item_changed(self, item: QTableWidgetItem) -> None:
        if self._updating or item.column() not in (3, 4):
            return
        floor, read_id = self._row_key(item.row())
        if not read_id:
            return
        rows = self._selected_rows(item.row())
        raw = str(item.text() or "").strip()
        if item.column() == 3:
            value = raw.upper()
            if value and value not in self.owner._window_catalog_by_code():
                QMessageBox.warning(self, "창호 지정", "창호 리스트에 없는 부호는 입력할 수 없습니다.")
                self.refresh()
                return
            if item.text() != value:
                self._updating = True; item.setText(value); self._updating = False
            master = self.owner._window_catalog_by_code().get(value) if value else None
            try:
                master_length = float((master or {}).get("applied_length_mm", 0.0) or 0.0)
            except Exception:
                master_length = 0.0
            pending_length = master_length if master_length > 0 else ""
            for row in rows:
                r_floor, r_id = self._row_key(row)
                if r_id:
                    self._set_pending(r_floor, r_id, "applied_code", value)
                    self._set_pending(r_floor, r_id, "applied_length_mm", pending_length)
                    target = self.table.item(row, 3)
                    length_target = self.table.item(row, 4)
                    if target is not None and target.text() != value:
                        self._updating = True; target.setText(value); self._updating = False
                    length_text = self._length_text(pending_length)
                    if length_target is not None and length_target.text() != length_text:
                        self._updating = True; length_target.setText(length_text); self._updating = False
                    self._update_assignment_row_color(row)
            self.status_label.setText(f"적용부호 수정 대기 · {len(rows):,}개 · 하단 적용 필요")
            return

        normalized = raw.replace(",", "")
        if normalized:
            try:
                number = float(normalized)
                if number < 0:
                    raise ValueError
            except ValueError:
                QMessageBox.warning(self, "창호 편집", "적용 창길이는 0 이상의 숫자로 입력해 주세요. 단위는 mm입니다.")
                self.refresh()
                return
            value = number
        else:
            value = ""
        for row in rows:
            r_floor, r_id = self._row_key(row)
            if r_id:
                self._set_pending(r_floor, r_id, "applied_length_mm", value)
                target = self.table.item(row, 4)
                text = self._length_text(value)
                if target is not None and target.text() != text:
                    self._updating = True; target.setText(text); self._updating = False
        self.status_label.setText(f"적용 창길이 수정 대기 · {len(rows):,}개 · 하단 적용 필요")

    def open_catalog_for_selected(self) -> None:
        rows = self._selected_rows()
        ids = {self._row_key(row)[1] for row in rows if self._row_key(row)[1]}
        if not ids:
            QMessageBox.information(self, "창호 지정", "창호 편집 목록에서 적용할 창호를 먼저 선택해 주세요.")
            return
        self.owner._opening_edit_selected_ids = set(ids)
        self.owner.preview.set_selected_opening_ids(ids)
        self.owner.open_opening_catalog_dialog(self)

    def apply_dragged_code(self, code: str) -> None:
        rows = self._selected_rows()
        ids = {self._row_key(row)[1] for row in rows if self._row_key(row)[1]}
        if not ids:
            QMessageBox.information(self, "창호 지정", "창호 편집 목록에서 적용할 창호를 먼저 선택해 주세요.")
            return
        self.owner._opening_edit_selected_ids = set(ids)
        self.owner.preview.set_selected_opening_ids(ids)
        self.owner.open_opening_catalog_dialog(self)
        if self.owner._opening_catalog_dialog is not None:
            self.owner._opening_catalog_dialog.stage_code(str(code or ""))

    def delete_selected(self) -> None:
        rows = self._selected_rows()
        if not rows:
            return
        self.owner.delete_opening_read_rows(rows, self)

    def moveEvent(self, event):
        super().moveEvent(event)
        try:
            dialog = self.owner._opening_catalog_dialog
            if dialog is not None and dialog.isVisible():
                self.owner._position_opening_catalog_dialog(dialog, self)
        except Exception:
            pass

    def resizeEvent(self, event):
        super().resizeEvent(event)
        try:
            dialog = self.owner._opening_catalog_dialog
            if dialog is not None and dialog.isVisible():
                self.owner._position_opening_catalog_dialog(dialog, self)
        except Exception:
            pass

    def closeEvent(self, event):
        try:
            dialog = self.owner._opening_catalog_dialog
            if dialog is not None:
                dialog.close()
        except Exception:
            pass
        super().closeEvent(event)


class OpeningCatalogTable(QTableWidget):
    """평면 창호 편집에서 선택 부호에 드롭할 수 있는 창호 마스터 목록."""
    def startDrag(self, supportedActions):
        row = self.currentRow()
        if row < 0:
            return
        item = self.item(row, 1)
        code = str(item.text() if item else "").strip().upper()
        if not code:
            return
        mime = QMimeData()
        mime.setData("application/x-dxf-window-code", code.encode("utf-8"))
        drag = QDrag(self)
        drag.setMimeData(mime)
        drag.exec(Qt.CopyAction)


class OpeningCatalogDialog(QDialog):
    """창호 편집창에서 적용할 마스터 부호를 고르는 종속 보조창."""
    HEADERS = ["구분", "창호명", "가로", "세로", "D/W", "내/외"]

    def __init__(self, owner, parent_dialog=None):
        super().__init__(parent_dialog or owner)
        self.owner = owner
        self.parent_dialog = parent_dialog
        self._staged_code = ""
        self.setWindowTitle("창호 선택창")
        self.resize(430, 560)
        self.setModal(False)
        self.setWindowFlags(self.windowFlags() | Qt.Tool)
        root = QVBoxLayout(self)
        self.info = QLabel("현재 선택 창호: 0개")
        root.addWidget(self.info)
        self.table = OpeningCatalogTable(0, len(self.HEADERS), self)
        self.table.setHorizontalHeaderLabels(self.HEADERS)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.SingleSelection)
        self.table.setDragEnabled(True)
        self.table.setDragDropMode(QAbstractItemView.DragOnly)
        self.table.setDefaultDropAction(Qt.CopyAction)
        self.table.verticalHeader().setVisible(False)
        header = self.table.horizontalHeader(); header.setSectionResizeMode(QHeaderView.Interactive)
        header.setStretchLastSection(False)
        current_row_h = max(1, self.table.verticalHeader().defaultSectionSize())
        self.table.verticalHeader().setDefaultSectionSize(max(18, int(round(current_row_h * 0.85))))
        self.table.verticalHeader().setMinimumSectionSize(16)
        self._columns_fitted = False
        root.addWidget(self.table, 1)
        row = QHBoxLayout()
        row.addWidget(QLabel("적용부호"))
        self.code_edit = QLineEdit(); self.code_edit.setPlaceholderText("목록에 있는 창호명만 입력")
        row.addWidget(self.code_edit, 1)
        root.addLayout(row)
        hint = QLabel("목록에서 부호를 선택하거나 직접 입력한 뒤 '적용'을 누르면 창호 편집창의 적용부호에 입력됩니다. 실제 평면 반영은 창호 편집창 하단 '적용'에서 확정합니다.")
        hint.setWordWrap(True); hint.setStyleSheet("color:#555;"); root.addWidget(hint)
        bottom = QHBoxLayout(); bottom.addStretch(1)
        self.apply_btn = QPushButton("적용"); self.close_btn = QPushButton("닫기")
        bottom.addWidget(self.apply_btn); bottom.addWidget(self.close_btn); root.addLayout(bottom)
        self.apply_btn.clicked.connect(self.apply_staged)
        self.close_btn.clicked.connect(self.close)
        self.code_edit.returnPressed.connect(lambda: self.apply_btn.setFocus(Qt.TabFocusReason))
        self.table.cellDoubleClicked.connect(lambda row, _col: self.stage_row(row))
        self.table.itemSelectionChanged.connect(self._sync_input_from_row)
        self.refresh()

    def _fit_columns_once(self):
        if getattr(self, "_columns_fitted", False):
            return
        self.table.resizeColumnsToContents()
        mins = [68, 72, 52, 52, 42, 42]
        maxs = [92, 105, 70, 70, 55, 55]
        for col in range(self.table.columnCount()):
            width = self.table.columnWidth(col)
            self.table.setColumnWidth(col, max(mins[col], min(width + 8, maxs[col])))
        self._columns_fitted = True

    def refresh(self):
        catalog = list(self.owner._window_catalog() or [])
        current = str(self.code_edit.text() or self._staged_code or "").strip().upper()
        self.table.setRowCount(0)
        for data in catalog:
            row = self.table.rowCount(); self.table.insertRow(row)
            vals = [data.get("category", ""), data.get("code", ""), data.get("width", ""), data.get("height", ""), data.get("dw", ""), data.get("inout", "")]
            for col, value in enumerate(vals):
                item = QTableWidgetItem(str(value or "")); item.setFlags(item.flags() & ~Qt.ItemIsEditable); self.table.setItem(row, col, item)
            if str(data.get("code") or "").upper() == current:
                self.table.selectRow(row)
        self.info.setText(f"현재 선택 창호: {len(self.owner._opening_edit_selected_ids):,}개 · 마스터 {len(catalog):,}개")
        if current: self.code_edit.setText(current)
        if not getattr(self, "_columns_fitted", False):
            QTimer.singleShot(0, self._fit_columns_once)

    def _sync_input_from_row(self):
        row = self.table.currentRow()
        if row >= 0 and self.table.item(row, 1):
            self.stage_code(self.table.item(row, 1).text())

    def stage_row(self, row: int):
        item = self.table.item(int(row), 1)
        if item is not None: self.stage_code(item.text())

    def stage_code(self, code: str):
        code = str(code or "").strip().upper()
        self._staged_code = code
        self.code_edit.setText(code)

    def apply_staged(self):
        code = str(self.code_edit.text() or "").strip().upper()
        if not code: return
        if code not in self.owner._window_catalog_by_code():
            QMessageBox.warning(self, "창호 지정", "창호 리스트에 없는 부호는 입력할 수 없습니다.")
            return
        if not self.owner._opening_edit_selected_ids:
            QMessageBox.information(self, "창호 지정", "적용할 창호를 먼저 선택해 주세요.")
            return
        self._staged_code = code
        if isinstance(self.parent_dialog, OpeningReadDialog):
            self.parent_dialog.stage_catalog_code(code)
        else:
            self.owner.apply_opening_catalog_code(code)
        self.refresh()


class RoomTableWidget(QTableWidget):
    """내부 목록에서 Enter로 현재 셀을 바로 편집한다.

    V23.9: 셀 위젯(구분 콤보) 위에서 발생하는 hover / 창 재활성화 계열
    이벤트는 행 선택을 절대 바꾸지 않는다. 실제 마우스 클릭/키보드 선택만 허용한다.
    """

    def selectionCommand(self, index, event=None):
        if event is not None:
            try:
                event_type = event.type()
                if event_type in (
                    QEvent.MouseMove,
                    QEvent.HoverMove,
                    QEvent.HoverEnter,
                    QEvent.Enter,
                    QEvent.FocusIn,
                    QEvent.WindowActivate,
                    QEvent.ActivationChange,
                ):
                    return QItemSelectionModel.NoUpdate
            except Exception:
                pass
        return super().selectionCommand(index, event)

    def keyPressEvent(self, event):
        if event.key() in (Qt.Key_Return, Qt.Key_Enter):
            row = self.currentRow()
            column = self.currentColumn()
            if row >= 0 and column >= 0:
                cell_widget = self.cellWidget(row, column)
                if isinstance(cell_widget, QComboBox) and cell_widget.isEnabled():
                    # 구분 콤보는 NoFocus로 운용한다. Enter에서는 포커스를 넘기지 않고
                    # 팝업만 직접 열어 창 활성화/hover가 행 선택으로 번지는 경로를 막는다.
                    cell_widget.showPopup()
                    return
                item = self.item(row, column)
                if item is not None and item.flags() & Qt.ItemIsEditable:
                    self.editItem(item)
                    return
        super().keyPressEvent(event)


class RoomCategoryDelegate(QStyledItemDelegate):
    """실 구분은 일반 셀로 표시하고 실제 클릭 편집 때만 드롭다운을 생성한다."""

    def __init__(self, rooms_tab: "RoomsTab", parent=None):
        super().__init__(parent)
        self.rooms_tab = rooms_tab

    def createEditor(self, parent, option, index):
        combo = QComboBox(parent)
        combo.addItem("")
        combo.addItems(list(ROOM_CATEGORIES))
        combo.setMaxVisibleItems(16)
        font = combo.font(); font.setBold(False); font.setWeight(QFont.Normal); combo.setFont(font)
        combo.activated.connect(lambda _idx, editor=combo: (self.commitData.emit(editor), self.closeEditor.emit(editor)))
        return combo

    def setEditorData(self, editor, index):
        if not isinstance(editor, QComboBox):
            return
        value = str(index.data(Qt.EditRole) or "").strip()
        pos = editor.findText(value, Qt.MatchFixedString)
        editor.setCurrentIndex(max(0, pos))
        editor.showPopup()

    def setModelData(self, editor, model, index):
        if not isinstance(editor, QComboBox):
            return
        category = editor.currentText().strip()
        room_index = index.data(Qt.UserRole)
        try:
            room_index = int(room_index)
        except Exception:
            return
        model.setData(index, category, Qt.EditRole)
        self.rooms_tab._on_room_category_changed(room_index, category)


class RoomBatchEditDialog(QDialog):
    def __init__(self, parent, floor: str, room_entries: list[tuple[int, RoomRecord]]):
        super().__init__(parent)
        self.setWindowTitle(f"{floor} 실 목록 임시 엑셀")
        self.resize(720, 560)
        self.room_entries = room_entries
        layout = QVBoxLayout(self)
        guide = QLabel("No와 실명만 수정됩니다. Excel에서 복사한 여러 셀도 Ctrl+V로 붙여넣을 수 있습니다.")
        layout.addWidget(guide)
        self.table = ClipboardTableWidget(len(room_entries), 5)
        self.table.setHorizontalHeaderLabels(["층", "No", "실명", "면적(M2)", "둘레(m)"])
        self.table.horizontalHeader().setSectionResizeMode(2, QHeaderView.Stretch)
        self.table.setColumnWidth(0, 80); self.table.setColumnWidth(1, 100); self.table.setColumnWidth(3, 100); self.table.setColumnWidth(4, 100)
        self.table.setSelectionMode(QAbstractItemView.ExtendedSelection)
        for row, (room_index, room) in enumerate(room_entries):
            values = [room.floor, room.room_number, normalize_room_name(room.room_name), f"{room.area_m2:.3f}", f"{room.perimeter_m:.3f}"]
            for col, value in enumerate(values):
                item = QTableWidgetItem(str(value))
                item.setData(Qt.UserRole, room_index)
                if col not in (1, 2):
                    item.setFlags(item.flags() & ~Qt.ItemIsEditable)
                    item.setBackground(QBrush(QColor("#F2F2F2")))
                self.table.setItem(row, col, item)
        layout.addWidget(self.table, 1)
        buttons = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        buttons.button(QDialogButtonBox.Ok).setText("목록에 적용")
        buttons.button(QDialogButtonBox.Cancel).setText("취소")
        buttons.accepted.connect(self.accept); buttons.rejected.connect(self.reject)
        layout.addWidget(buttons)

    def updates(self):
        result = []
        for row in range(self.table.rowCount()):
            idx_item = self.table.item(row, 0)
            room_index = idx_item.data(Qt.UserRole) if idx_item else None
            if room_index is None:
                continue
            no_item = self.table.item(row, 1)
            name_item = self.table.item(row, 2)
            result.append((int(room_index), (no_item.text() if no_item else "").strip(), (name_item.text() if name_item else "").strip()))
        return result


class SelectedRoomsBatchChangeDialog(QDialog):
    """선택한 실 목록의 층/구분/천정고를 체크 항목별로 일괄 변경한다."""

    def __init__(self, parent, source_room: RoomRecord, selected_rooms: list[RoomRecord]):
        super().__init__(parent)
        self.setWindowTitle("선택한 목록 일괄변경")
        self.resize(650, 250)
        self.source_room = source_room
        self.selected_rooms = list(selected_rooms or [])
        self._checks: dict[str, QCheckBox] = {}
        self._editors: dict[str, QWidget] = {}

        layout = QVBoxLayout(self)
        guide = QLabel(
            f"선택한 {len(self.selected_rooms):,}개 실에 체크한 항목만 적용합니다. "
            "변경값은 현재 행의 값을 미리 표시합니다."
        )
        layout.addWidget(guide)

        self.table = QTableWidget(3, 4)
        self.table.setHorizontalHeaderLabels(["적용", "대상", "기존값", "변경값"])
        self.table.verticalHeader().setVisible(False)
        self.table.setSelectionMode(QAbstractItemView.NoSelection)
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.horizontalHeader().setSectionResizeMode(0, QHeaderView.Fixed)
        self.table.horizontalHeader().setSectionResizeMode(1, QHeaderView.Fixed)
        self.table.horizontalHeader().setSectionResizeMode(2, QHeaderView.Stretch)
        self.table.horizontalHeader().setSectionResizeMode(3, QHeaderView.Stretch)
        self.table.setColumnWidth(0, 64)
        self.table.setColumnWidth(1, 92)
        self.table.verticalHeader().setDefaultSectionSize(34)
        layout.addWidget(self.table)

        def display_text(value):
            text = str(value or "").strip()
            return text if text else "(공란)"

        def existing_summary(values):
            normalized = [display_text(value) for value in values]
            unique = list(dict.fromkeys(normalized))
            if len(unique) <= 1:
                return (unique[0] if unique else "(공란)", "")
            tooltip = " / ".join(unique[:12])
            if len(unique) > 12:
                tooltip += f" / ... 외 {len(unique) - 12}종"
            return (f"여러 값 ({len(unique)}종)", tooltip)

        ceiling_formatter = getattr(parent, "_format_ceiling_height", lambda value: "" if value is None else str(value))
        specs = [
            ("floor", "층", [room.floor for room in self.selected_rooms]),
            ("category", "구분", [room.room_category for room in self.selected_rooms]),
            ("ceiling_height_m", "천정고", [ceiling_formatter(room.ceiling_height_m) for room in self.selected_rooms]),
        ]
        for row, (key, label, existing_values) in enumerate(specs):
            check_host = QWidget()
            check_layout = QHBoxLayout(check_host)
            check_layout.setContentsMargins(0, 0, 0, 0)
            check_layout.setAlignment(Qt.AlignCenter)
            check = QCheckBox()
            check.setToolTip("체크한 항목만 선택한 실에 적용합니다.")
            check_layout.addWidget(check)
            self.table.setCellWidget(row, 0, check_host)
            self._checks[key] = check

            target_item = QTableWidgetItem(label)
            target_item.setTextAlignment(Qt.AlignCenter)
            target_item.setFlags(target_item.flags() & ~Qt.ItemIsEditable)
            self.table.setItem(row, 1, target_item)

            existing_text, tooltip = existing_summary(existing_values)
            existing_item = QTableWidgetItem(existing_text)
            existing_item.setFlags(existing_item.flags() & ~Qt.ItemIsEditable)
            if tooltip:
                existing_item.setToolTip(tooltip)
            self.table.setItem(row, 2, existing_item)

            if key == "category":
                editor = QComboBox()
                editor.addItem("")
                editor.addItems(list(ROOM_CATEGORIES))
                source_value = str(source_room.room_category or "").strip()
                if source_value and editor.findText(source_value, Qt.MatchFixedString) < 0:
                    editor.addItem(source_value)
                editor.setCurrentIndex(max(0, editor.findText(source_value, Qt.MatchFixedString)))
            else:
                editor = QLineEdit()
                if key == "floor":
                    editor.setText(str(source_room.floor or ""))
                else:
                    editor.setText(ceiling_formatter(source_room.ceiling_height_m))
                    editor.setPlaceholderText("공란 또는 0 이상의 숫자(m)")
            self.table.setCellWidget(row, 3, editor)
            self._editors[key] = editor

        buttons = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        buttons.button(QDialogButtonBox.Ok).setText("선택 목록에 적용")
        buttons.button(QDialogButtonBox.Cancel).setText("취소")
        buttons.accepted.connect(self._validate_and_accept)
        buttons.rejected.connect(self.reject)
        layout.addWidget(buttons)

    def _validate_and_accept(self):
        if not any(check.isChecked() for check in self._checks.values()):
            QMessageBox.information(self, "선택한 목록 일괄변경", "적용할 항목을 하나 이상 체크해 주세요.")
            return
        if self._checks["ceiling_height_m"].isChecked():
            text = str(self._editors["ceiling_height_m"].text() or "").strip()
            if text:
                try:
                    value = float(text.replace(",", ""))
                    if value < 0:
                        raise ValueError
                except Exception:
                    QMessageBox.warning(self, "천정고 입력", "천정고는 0 이상의 숫자(m) 또는 공란으로 입력해 주세요.")
                    self._editors["ceiling_height_m"].setFocus()
                    return
        self.accept()

    def changes(self) -> dict:
        result: dict = {}
        if self._checks["floor"].isChecked():
            result["floor"] = str(self._editors["floor"].text() or "").strip() or "(층 미입력)"
        if self._checks["category"].isChecked():
            result["category"] = str(self._editors["category"].currentText() or "").strip()
        if self._checks["ceiling_height_m"].isChecked():
            text = str(self._editors["ceiling_height_m"].text() or "").strip()
            result["ceiling_height_m"] = None if not text else float(text.replace(",", ""))
        return result


class RoomCategoryComboBox(QComboBox):
    def __init__(self, parent=None):
        super().__init__(parent)
        # 표의 기본 격자선만 남기고 셀 안쪽 콤보박스 사각 테두리는 표시하지 않는다.
        self.setFrame(False)
        # 표 셀 위에 상시 배치된 콤보가 단순 마우스 진입만으로 부모 QTableWidget의
        # 현재/선택 행을 바꾸지 않도록 한다. 실제 클릭은 콤보에서 처리한다.
        # V23.9: ClickFocus도 제거한다. 창이 다시 활성화될 때 Qt가 셀 위젯의
        # 포커스를 복원하면서 해당 index를 현재행으로 만드는 현상을 원천 차단한다.
        # 실제 클릭 시 행 선택은 mousePressEvent에서 명시적으로 처리한다.
        self.setFocusPolicy(Qt.NoFocus)
        self.setAttribute(Qt.WA_NoMousePropagation, True)
        font = self.font()
        font.setBold(False)
        font.setWeight(QFont.Normal)
        self.setFont(font)
        self.setStyleSheet(
            "QComboBox { border:none; background:transparent; color:#111111; padding:0 16px 0 3px; font-weight:normal; }"
            "QComboBox::drop-down { border:none; background:transparent; width:16px; }"
        )
        popup = self.view()
        popup.setMouseTracking(True)
        popup.viewport().setMouseTracking(True)
        popup.setStyleSheet(
            "QAbstractItemView { border:1px solid #A0A0A0; background:#FFFFFF; color:#111111; "
            "selection-background-color:#17365D; selection-color:#FFFFFF; font-weight:normal; outline:0; }"
            "QAbstractItemView::item { background:#FFFFFF; color:#111111; padding:2px 4px; min-height:20px; }"
            "QAbstractItemView::item:hover { background:#FFF2CC; color:#111111; }"
            "QAbstractItemView::item:selected { background:#17365D; color:#FFFFFF; }"
        )

    def mousePressEvent(self, event):
        # hover는 행 선택에 관여하지 않고, 실제 클릭했을 때만 해당 행을 현재 행으로 맞춘다.
        # Ctrl/Shift 다중선택 중에는 기존 선택 묶음을 깨지 않는다.
        table = self.parentWidget()
        while table is not None and not isinstance(table, QTableWidget):
            table = table.parentWidget()
        if table is not None and not (event.modifiers() & (Qt.ControlModifier | Qt.ShiftModifier)):
            try:
                point = table.viewport().mapFromGlobal(event.globalPosition().toPoint())
                index = table.indexAt(point)
                if index.isValid():
                    table.setCurrentCell(index.row(), index.column(), QItemSelectionModel.NoUpdate)
                    if not table.selectionModel().isRowSelected(index.row(), table.rootIndex()):
                        table.clearSelection()
                        table.selectRow(index.row())
            except Exception:
                pass
        super().mousePressEvent(event)


class RoomsTab(QWidget):
    data_changed = Signal()
    background_regenerate_requested = Signal()
    background_filter_requested = Signal()
    cad_create_blocks_requested = Signal(object)
    cad_refresh_blocks_requested = Signal(object)
    background_settings_requested = Signal()
    dimension_settings_changed = Signal(object)
    label_size_changed = Signal(int)
    original_text_toggled = Signal(bool)
    cad_background_mode_changed = Signal(str)
    group_by_floor_toggled = Signal(bool)
    number_entry_settings_changed = Signal(object)
    drawing_library_requested = Signal()
    room_column_widths_changed = Signal(object)
    finish_summary_column_order_changed = Signal(object)
    finish_group_type_labels_changed = Signal(object)
    finish_group_tab_settings_changed = Signal(object)
    room_wall_link_cad_requested = Signal(object)
    opening_read_requested = Signal(object)

    def __init__(self, background_max_pixels: int = 1800, edge_length_font_size: int = 45, edge_length_offset: int = 150, edge_length_width_ratio: int = 80, label_font_size: int = 100, number_entry_settings: dict | None = None, room_column_widths=None, room_category_rules=None, finish_formula_codes=None, finish_categories=None, material_trades=None, material_units=None, finish_summary_column_order=None):
        super().__init__()
        self.rooms: list[RoomRecord] = []
        self.errors: list[str] = []
        self.current_source_path = ""
        self.original_source_path = ""
        self.working_dxf_path = ""
        self.base_dxf_path = ""
        self.cad_background = []
        self.loading_table = False
        self.cad_connected_mode = False
        self.room_rules = {}
        self.wall_finish_assignments: dict[str, dict] = {}
        self.room_finish_assignments: dict[str, dict] = {}
        self.finish_groups: dict[str, dict] = {}
        self.summary_markups: dict[str, float] = {}
        self.opening_read_results: dict[str, list[dict]] = {}
        self._opening_read_dialog = None
        self._opening_assign_dialog = None
        self._opening_catalog_dialog = None
        self._window_catalog_provider = lambda: []
        self._opening_edit_selected_ids: set[str] = set()
        self._opening_edit_mode = False
        self._opening_last_move = None
        self._active_finish_group_id = ""
        self.selected_wall_keys: set[str] = set()
        self._wall_group_dialog = None
        self._room_group_dialog = None
        # 바닥/천정/벽 지정모드도 창호 편집처럼 도곽 상단 상태창으로 안내한다.
        self._finish_group_overlay_section = ""
        self._finish_group_overlay_label = ""
        self._finish_group_overlay_kind = ""
        self._finish_group_overlay_confirm_callback = None
        self._finish_group_overlay_cancel_callback = None
        self._finish_group_overlay_detail_dialog = None
        self._room_group_list = None
        self._selected_group_room_handles: set[str] = set()
        self._group_mode_blocked_room_handles: set[str] = set()
        self._group_mode_blocked_wall_keys: set[str] = set()
        self.current_combined_errors: list[str] = []
        self.cad_validation_errors: list[str] = []
        self.room_wall_validation_errors: list[str] = []
        self.table_row_room_indices: list[int | None] = []
        self.table_row_floors: list[str | None] = []
        self.room_category_combos: dict[int, RoomCategoryComboBox] = {}
        self.collapsed_floors: set[str] = set()
        self.group_by_floor = True
        # 기본 목록은 층별 No 오름차순이며 No 공백은 항상 번호가 있는 실 아래에 둔다.
        self.room_sort_column: int | None = 1
        self.room_sort_ascending = True
        self._direct_number_mode = False
        self._direct_number_settings = None
        self._direct_number_next = 1
        self._direct_number_floor = None
        self._direct_number_write_name = False
        self._name_input_mode = False
        self.number_entry_settings = self._normalize_number_entry_settings(number_entry_settings)
        self.room_category_rules = normalize_room_category_rules(room_category_rules)

        root = QVBoxLayout(self)

        toolbar = QHBoxLayout()
        self.view_mode = QComboBox()
        self.view_mode.addItems(["층 전체 보기", "선택 실만 보기"])
        self.auto_number_btn = QPushButton("번호 기입")
        self.create_block_btn = QPushButton("실명블록 생성")
        self.cad_link_btn = QPushButton("CAD 연동")
        self.opening_test_btn = QPushButton("창호 리스트")
        self.opening_test_btn.setToolTip("창호 리스트를 엽니다. CAD에서 새로 읽을 때는 창 안의 창호 읽기 버튼을 사용합니다.")
        self.opening_test_btn.setEnabled(False)
        self.opening_edit_btn = QPushButton("창호 편집")
        self.opening_edit_btn.setCheckable(True)
        self.opening_edit_btn.setToolTip("도곽의 창호 부호를 클릭/드래그로 선택하고 Delete 또는 상태창 삭제 버튼으로 수동 제거합니다.")
        self.opening_edit_btn.setEnabled(False)
        self.regenerate_background_btn = QPushButton("CAD 배경 재생성")
        self.vector_filter_btn = QPushButton("도면 정리")
        self.vector_filter_btn.setToolTip("원본 정리 캐시와 프로그램 최적화 도면을 나눠 관리합니다.")
        self.vector_stage_label = QLabel("표시: 없음")
        self.vector_stage_label.setToolTip("현재 도곽에 표시 중인 벡터 단계입니다.")
        self.vector_stage_label.setStyleSheet("font-weight: 600; color: #555555;")
        self.background_settings_btn = QPushButton("표시 설정")
        self.error_check_btn = QPushButton("오류/확인사항")
        self.save_dxf_btn = QPushButton("수정내용 DXF 저장")
        self.save_dxf_btn.setEnabled(False)
        toolbar.addWidget(QLabel("도면 보기:"))
        toolbar.addWidget(self.view_mode)
        toolbar.addSpacing(12)
        toolbar.addWidget(QLabel("문자 크기:"))
        self.label_size_spin = QSpinBox()
        self.label_size_spin.setRange(6, 500)
        self.label_size_spin.setValue(int(label_font_size))
        self.label_size_spin.setSingleStep(10)
        self.label_size_spin.setSuffix(" pt")
        self.label_size_spin.setFixedWidth(90)
        toolbar.addWidget(self.label_size_spin)
        toolbar.addSpacing(12)
        toolbar.addWidget(QLabel("CAD 배경:"))
        self.cad_background_combo = QComboBox()
        self.cad_background_combo.addItem("없음", "none")
        self.cad_background_combo.addItem("검은색", "black")
        self.cad_background_combo.addItem("하얀색", "white")
        self.cad_background_combo.setCurrentIndex(self.cad_background_combo.findData("black"))
        self.cad_background_combo.setFixedWidth(88)
        self.cad_background_combo.setEnabled(False)
        toolbar.addWidget(self.cad_background_combo)
        self.group_overlay_check = QCheckBox("그룹 표시")
        self.group_overlay_check.setChecked(True)
        toolbar.addWidget(self.group_overlay_check)
        self.room_line_check = QCheckBox("실 라인")
        self.room_line_check.setChecked(True)
        self.room_line_check.setEnabled(False)
        toolbar.addWidget(self.room_line_check)
        # 준비됨/미생성 상태문구는 상단 도구줄을 복잡하게 만들어 표시하지 않는다.
        self.cad_background_status_label = QLabel("")
        self.cad_background_status_label.hide()
        self.original_text_check = QCheckBox("원본 텍스트")
        self.original_text_check.setChecked(True)
        self.original_text_check.setEnabled(False)
        toolbar.addWidget(self.original_text_check)
        toolbar.addStretch(1)
        toolbar.addWidget(self.cad_link_btn)
        toolbar.addWidget(self.opening_test_btn)
        toolbar.addWidget(self.regenerate_background_btn)
        toolbar.addWidget(self.vector_filter_btn)
        toolbar.addWidget(self.error_check_btn)
        toolbar.addWidget(self.save_dxf_btn)
        root.addLayout(toolbar)

        self.main_splitter = QSplitter(Qt.Horizontal)
        self.main_splitter.setChildrenCollapsible(False)
        self.main_splitter.setHandleWidth(4)

        left = QWidget()
        left.setMinimumWidth(250)
        left.setSizePolicy(QSizePolicy.Preferred, QSizePolicy.Expanding)
        left_layout = QVBoxLayout(left)
        self.room_table = RoomTableWidget(0, 8)
        self.room_table.setHorizontalHeaderLabels(["층", "No", "실명", "구분", "적용층", "면적(M2)", "둘레(m)", "천정고(m)"])
        self.room_table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.room_table.setSelectionMode(QAbstractItemView.ExtendedSelection)
        self.room_table.setAlternatingRowColors(False)
        self.room_table.verticalHeader().setVisible(False)
        self.room_table.verticalHeader().setDefaultSectionSize(22)
        self.room_table.verticalHeader().setMinimumSectionSize(18)
        self.room_table.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
        self.room_table.horizontalHeader().setSortIndicatorShown(True)
        self.room_table.horizontalHeader().setSortIndicator(1, Qt.AscendingOrder)
        widths = self._normalize_room_column_widths(room_column_widths)
        for i, width in enumerate(widths):
            self.room_table.setColumnWidth(i, width)
        # V23.15: 구분은 상시 QComboBox 셀 위젯을 제거하고 실제 클릭 때만
        # delegate 편집기를 띄운다. 마우스 hover/드래그/창 활성화로 값이
        # 바뀌지 않고 선택행 배경도 다른 열과 동일하게 표시된다.
        self.room_table.setItemDelegateForColumn(3, RoomCategoryDelegate(self, self.room_table))
        self.room_table.cellClicked.connect(self._edit_room_category_on_click)
        self._room_column_save_timer = QTimer(self)
        self._room_column_save_timer.setSingleShot(True)
        self._room_column_save_timer.setInterval(300)
        self._room_column_save_timer.timeout.connect(self._emit_room_column_widths)
        self.room_table.horizontalHeader().sectionResized.connect(self._queue_room_column_width_save)
        floor_group_row = QHBoxLayout()
        floor_group_row.addWidget(QLabel("층 목록"))
        self.group_by_floor_check = QCheckBox("층별 묶기")
        self.group_by_floor_check.setChecked(True)
        floor_group_row.addWidget(self.group_by_floor_check)
        floor_group_row.addStretch(1)
        self.expand_all_floors_btn = QPushButton("전체 펼치기")
        self.collapse_all_floors_btn = QPushButton("전체 접기")
        self.expand_all_floors_btn.setMaximumHeight(25)
        self.collapse_all_floors_btn.setMaximumHeight(25)
        floor_group_row.addWidget(self.expand_all_floors_btn)
        floor_group_row.addWidget(self.collapse_all_floors_btn)
        left_layout.addLayout(floor_group_row)
        left_layout.addWidget(self.room_table, 1)
        room_action_row = QHBoxLayout()
        self.create_block_btn.setFixedSize(92, 24)
        self.auto_number_btn.setFixedSize(62, 24)
        room_action_row.addWidget(self.create_block_btn)
        room_action_row.addWidget(self.auto_number_btn)
        self.direct_name_btn = QPushButton("실명 직접입력")
        self.direct_name_btn.setFixedSize(88, 24)
        room_action_row.addWidget(self.direct_name_btn)
        room_action_row.addStretch(1)
        self.apply_selected_btn = QPushButton("일괄변경")
        self.apply_selected_btn.setFixedSize(68, 24)
        self.apply_selected_btn.setToolTip("선택한 실의 층·구분·천정고를 체크 항목별로 일괄 변경합니다.")
        room_action_row.addWidget(self.apply_selected_btn)
        left_layout.addLayout(room_action_row)
        self.batch_edit_btn = QPushButton("현재 층 목록 일괄편집")
        self.batch_edit_btn.setMinimumHeight(26)
        self.batch_edit_btn.setMaximumHeight(30)
        left_layout.addWidget(self.batch_edit_btn)
        self.main_splitter.addWidget(left)

        center = QWidget()
        center.setMinimumWidth(320)
        center.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        center_layout = QVBoxLayout(center)
        center_layout.setContentsMargins(0, 0, 0, 0)

        self.center_splitter = QSplitter(Qt.Vertical)
        self.center_splitter.setChildrenCollapsible(False)
        self.center_splitter.setHandleWidth(4)

        preview_section = QWidget()
        preview_section_layout = QVBoxLayout(preview_section)
        preview_section_layout.setContentsMargins(0, 0, 0, 0)

        # 도면과 조절 상자를 같은 레이아웃 셀에 겹쳐 배치한다.
        # 조절 상자는 그래픽 장면이 아닌 일반 위젯 레이아웃 소속이므로
        # 확대·이동과 무관하게 항상 화면 우측 하단에 고정된다.
        preview_host = QWidget()
        preview_stack = QGridLayout(preview_host)
        preview_stack.setContentsMargins(0, 0, 0, 0)
        preview_stack.setSpacing(0)
        self.preview = RoomPreview()
        self.preview.set_label_font_size(int(label_font_size))
        self.preview.set_background_max_pixels(background_max_pixels)
        self.preview.set_edge_length_font_size(edge_length_font_size)
        self.preview.set_edge_length_offset(edge_length_offset)
        self.preview.set_edge_length_width_ratio(edge_length_width_ratio)
        preview_stack.addWidget(self.preview, 0, 0)

        # V23.58: 층 전체보기 전용 도면 편집 패널. V23.54의 100% 크기를 유지하고
        # 접기/펼치기 + 구획벽 나누기/합치기 + 창호 편집을 한 곳에 모은다.
        self.floor_tools_overlay = QFrame()
        self.floor_tools_overlay.setFrameShape(QFrame.StyledPanel)
        self.floor_tools_overlay.setAutoFillBackground(True)
        self.floor_tools_overlay.setStyleSheet(
            "QFrame { background:#F4F6F8; border:1px solid #7A8793; border-radius:4px; }"
            "QLabel { color:#222; border:none; font-weight:bold; }"
            "QPushButton { min-height:23px; padding:1px 8px; }"
            "QPushButton:checked { background:#FFE3C2; border:1px solid #E27600; }"
            "QToolButton { min-width:23px; min-height:23px; padding:0px 3px; }"
        )
        floor_tools_row = QHBoxLayout(self.floor_tools_overlay)
        floor_tools_row.setContentsMargins(7, 4, 7, 4)
        floor_tools_row.setSpacing(6)
        floor_tools_row.addWidget(QLabel("도면 편집"))

        self.floor_tools_content = QWidget()
        floor_tools_content_row = QHBoxLayout(self.floor_tools_content)
        floor_tools_content_row.setContentsMargins(0, 0, 0, 0)
        floor_tools_content_row.setSpacing(6)

        self.boundary_split_btn = QPushButton("구획벽 나누기")
        self.boundary_split_btn.setCheckable(True)
        self.boundary_split_btn.setEnabled(False)
        self.boundary_split_btn.setToolTip("실 구분선에 마우스를 올리면 10mm 단위 위치가 표시되고 클릭한 위치에 바로 절점을 추가합니다.")
        floor_tools_content_row.addWidget(self.boundary_split_btn)

        self.boundary_merge_btn = QPushButton("구획벽 합치기")
        self.boundary_merge_btn.setCheckable(True)
        self.boundary_merge_btn.setEnabled(False)
        self.boundary_merge_btn.setToolTip("일직선으로 이어진 두 구획벽의 공유 절점에 마우스를 올리고 클릭하면 절점을 제거해 한 선으로 합칩니다.")
        floor_tools_content_row.addWidget(self.boundary_merge_btn)

        # 기존 상단 툴바의 창호 편집 버튼을 도면 편집 패널로 이동한다.
        floor_tools_content_row.addWidget(self.opening_edit_btn)

        self.boundary_split_status_label = QLabel("10mm 스냅")
        self.boundary_split_status_label.setStyleSheet("font-weight:normal; color:#A84400;")
        self.boundary_split_status_label.hide()
        floor_tools_content_row.addWidget(self.boundary_split_status_label)
        floor_tools_row.addWidget(self.floor_tools_content)

        self.floor_tools_collapse_btn = QToolButton()
        self.floor_tools_collapse_btn.setText("◀")
        self.floor_tools_collapse_btn.setToolTip("도면 편집 도구를 접습니다.")
        floor_tools_row.addWidget(self.floor_tools_collapse_btn)

        self.floor_tools_overlay.setSizePolicy(QSizePolicy.Maximum, QSizePolicy.Fixed)
        preview_stack.addWidget(self.floor_tools_overlay, 0, 0, Qt.AlignLeft | Qt.AlignTop)

        # V15.4: 선택실 보기에서 자주 조정하는 치수 표시값은 도면 우하단에 바로 노출한다.
        self._syncing_dimension_controls = False
        self.dimension_overlay = QGroupBox("치수 표시")
        dim_grid = QGridLayout(self.dimension_overlay)
        dim_grid.setContentsMargins(8, 6, 8, 6)
        dim_grid.setHorizontalSpacing(6)
        dim_grid.setVerticalSpacing(3)
        self.dimension_font_spin = QSpinBox(); self.dimension_font_spin.setRange(6, 300); self.dimension_font_spin.setSingleStep(10); self.dimension_font_spin.setSuffix(" pt"); self.dimension_font_spin.setValue(edge_length_font_size)
        self.dimension_width_spin = QSpinBox(); self.dimension_width_spin.setRange(10, 200); self.dimension_width_spin.setSingleStep(10); self.dimension_width_spin.setSuffix(" %"); self.dimension_width_spin.setValue(edge_length_width_ratio)
        self.dimension_offset_spin = QSpinBox(); self.dimension_offset_spin.setRange(-5000, 5000); self.dimension_offset_spin.setSingleStep(10); self.dimension_offset_spin.setSuffix(" mm"); self.dimension_offset_spin.setValue(edge_length_offset)
        self.dimension_show_check = QCheckBox("치수표시"); self.dimension_show_check.setChecked(True)
        self.dimension_unit_m_check = QCheckBox("m표시"); self.dimension_unit_m_check.setChecked(True)
        dim_grid.addWidget(QLabel("문자 크기"), 0, 0); dim_grid.addWidget(self.dimension_font_spin, 0, 1)
        dim_grid.addWidget(QLabel("폭"), 1, 0); dim_grid.addWidget(self.dimension_width_spin, 1, 1)
        dim_grid.addWidget(QLabel("옵셋"), 2, 0); dim_grid.addWidget(self.dimension_offset_spin, 2, 1)
        dim_check_row = QHBoxLayout()
        dim_check_row.setContentsMargins(0, 0, 0, 0)
        dim_check_row.setSpacing(10)
        dim_check_row.addWidget(self.dimension_unit_m_check)
        dim_check_row.addWidget(self.dimension_show_check)
        dim_check_row.addStretch(1)
        dim_grid.addLayout(dim_check_row, 3, 0, 1, 2)
        self.dimension_overlay.setFixedWidth(190)
        self.dimension_overlay.setSizePolicy(QSizePolicy.Fixed, QSizePolicy.Fixed)
        self.dimension_overlay.setAutoFillBackground(True)
        self.dimension_overlay.setStyleSheet(
            "QGroupBox { background-color:#FFFFFF; color:#111111; border:1px solid #8A8A8A; "
            "border-radius:3px; margin-top:8px; font-weight:bold; }"
            "QGroupBox::title { subcontrol-origin:margin; left:8px; padding:0 3px; background:#FFFFFF; }"
            "QLabel, QCheckBox { background-color:#FFFFFF; color:#111111; font-weight:normal; }"
            "QSpinBox { background-color:#FFFFFF; color:#111111; }"
        )
        self.dimension_overlay.adjustSize()
        self.dimension_overlay.hide()
        preview_stack.addWidget(self.dimension_overlay, 0, 0, Qt.AlignRight | Qt.AlignBottom)

        # 창호 편집모드 상태창: 도곽 상단에 고정하고 삭제 후에도 모드를 유지한다.
        self.opening_edit_overlay = QFrame()
        self.opening_edit_overlay.setFrameShape(QFrame.StyledPanel)
        self.opening_edit_overlay.setAutoFillBackground(True)
        self.opening_edit_overlay.setStyleSheet(
            "QFrame { background:#FFF7D6; border:1px solid #C79A00; border-radius:4px; }"
            "QLabel { color:#222; border:none; font-weight:bold; }"
            "QPushButton { min-height:22px; padding:1px 9px; }"
        )
        opening_edit_row = QHBoxLayout(self.opening_edit_overlay)
        opening_edit_row.setContentsMargins(8, 4, 8, 4)
        opening_edit_row.setSpacing(6)
        self.opening_edit_status_label = QLabel("창호 편집 | 선택 0개 | 클릭 선택 · 드래그 이동(10mm) · Delete 삭제 · ESC 종료")
        self.opening_move_distance_spin = QDoubleSpinBox()
        self.opening_move_distance_spin.setDecimals(0)
        self.opening_move_distance_spin.setRange(0.0, 1000000.0)
        self.opening_move_distance_spin.setSingleStep(10.0)
        self.opening_move_distance_spin.setSuffix(" mm")
        self.opening_move_distance_spin.setFixedWidth(100)
        self.opening_move_distance_spin.setEnabled(False)
        self.opening_move_exact_btn = QPushButton("거리 적용")
        self.opening_move_exact_btn.setEnabled(False)
        self.opening_catalog_assign_btn = QPushButton("창호지정")
        self.opening_catalog_assign_btn.setToolTip("F5와 동일: 현재 선택 창호만 들어있는 창호 지정목록을 엽니다.")
        self.opening_edit_delete_btn = QPushButton("삭제")
        self.opening_edit_close_btn = QPushButton("종료")
        opening_edit_row.addWidget(self.opening_edit_status_label)
        opening_edit_row.addWidget(self.opening_catalog_assign_btn)
        opening_edit_row.addWidget(QLabel("이동"))
        opening_edit_row.addWidget(self.opening_move_distance_spin)
        opening_edit_row.addWidget(self.opening_move_exact_btn)
        opening_edit_row.addWidget(self.opening_edit_delete_btn)
        opening_edit_row.addWidget(self.opening_edit_close_btn)
        self.opening_edit_overlay.setSizePolicy(QSizePolicy.Maximum, QSizePolicy.Fixed)
        self.opening_edit_overlay.hide()
        preview_stack.addWidget(self.opening_edit_overlay, 0, 0, Qt.AlignHCenter | Qt.AlignTop)

        # 바닥/천정/벽 그룹 지정 상태창. 상세 실목록은 필요할 때만 목록 버튼으로 연다.
        self.finish_group_edit_overlay = QFrame()
        self.finish_group_edit_overlay.setFrameShape(QFrame.StyledPanel)
        self.finish_group_edit_overlay.setAutoFillBackground(True)
        self.finish_group_edit_overlay.setStyleSheet(
            "QFrame { background:#EAF4FF; border:1px solid #4B8AC9; border-radius:4px; }"
            "QLabel { color:#1E2D3D; border:none; font-weight:bold; }"
            "QPushButton { min-height:22px; padding:1px 9px; }"
        )
        finish_group_edit_row = QHBoxLayout(self.finish_group_edit_overlay)
        finish_group_edit_row.setContentsMargins(8, 4, 8, 4)
        finish_group_edit_row.setSpacing(6)
        self.finish_group_edit_status_label = QLabel("그룹 지정 | 선택 0개")
        self.finish_group_edit_list_btn = QPushButton("목록")
        self.finish_group_edit_confirm_btn = QPushButton("확인")
        self.finish_group_edit_cancel_btn = QPushButton("취소")
        finish_group_edit_row.addWidget(self.finish_group_edit_status_label)
        finish_group_edit_row.addWidget(self.finish_group_edit_list_btn)
        finish_group_edit_row.addWidget(self.finish_group_edit_confirm_btn)
        finish_group_edit_row.addWidget(self.finish_group_edit_cancel_btn)
        self.finish_group_edit_overlay.setSizePolicy(QSizePolicy.Maximum, QSizePolicy.Fixed)
        self.finish_group_edit_overlay.hide()
        preview_stack.addWidget(self.finish_group_edit_overlay, 0, 0, Qt.AlignHCenter | Qt.AlignTop)

        preview_stack.setRowStretch(0, 1)
        preview_stack.setColumnStretch(0, 1)
        preview_section_layout.addWidget(preview_host, 1)
        preview_help_row = QHBoxLayout()
        preview_help_row.setContentsMargins(0, 0, 0, 0)
        self.preview_help = QLabel(
            "마우스 휠: 확대·축소 | 일반모드 드래그: 화면 이동 | 벽 선택모드 드래그: 박스 선택 | 가운데버튼 더블클릭: 전체보기"
        )
        self.preview_help.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        self.preview_help.setStyleSheet("color:#0066CC;")
        preview_help_row.addWidget(self.preview_help, 1)
        self.lower_panel_toggle_btn = QToolButton()
        self.lower_panel_toggle_btn.setText("▼")
        self.lower_panel_toggle_btn.setToolTip("선택 실/층 정보와 재료마감 집계를 아래로 접습니다. · Ctrl+Num2 / Ctrl+Num5")
        self.finish_panel_toggle_btn = QToolButton()
        self.finish_panel_toggle_btn.setText("▶")
        self.finish_panel_toggle_btn.setToolTip("오른쪽 재료마감 패널을 접습니다. · Ctrl+Num6")
        preview_help_row.addWidget(self.lower_panel_toggle_btn, 0)
        preview_help_row.addWidget(self.finish_panel_toggle_btn, 0)
        # 도면표시는 현재 재료마감 탭의 사용자 구분(예: 바탕/마감) 기준으로 바뀐다.
        # 콤보 itemData는 JSON [내부탭키, 구분]이며 구분 '*'는 현재 탭 전체다.
        self.finish_group_filter_combo = QComboBox()
        # V23.54: 없음/바탕/마감 표시 구분칸은 화면 점유를 줄이기 위해 기존 폭의 약 절반으로 고정.
        self.finish_group_filter_combo.setFixedWidth(68)
        self.finish_group_filter_combo.setToolTip(
            "없음 / 현재 탭 구분별 / 전체를 선택합니다. 탭별 설정에서 다른 탭 구분도 표시할 수 있습니다."
        )
        preview_help_row.addWidget(self.finish_group_filter_combo, 0)
        self.dxf_library_btn = QPushButton("DXF 불러오기")
        preview_help_row.addWidget(self.dxf_library_btn, 0)
        preview_help_row.addWidget(self.background_settings_btn, 0)
        # V24.21: 현재 벡터 표시 단계는 메인 창 최하단 상태/알림 라인의 우측 끝에서 표시한다.
        preview_section_layout.addLayout(preview_help_row)
        self.center_splitter.addWidget(preview_section)

        info_group = QGroupBox("선택 실 / 층 정보")
        info_group.setMinimumWidth(0)
        # 하단 패널을 키워도 선택 실/층 정보는 상단의 내용 높이만 유지한다.
        # 남는 세로 공간은 오른쪽 재료마감 집계표가 사용한다.
        info_group.setSizePolicy(QSizePolicy.Ignored, QSizePolicy.Fixed)
        form = QGridLayout(info_group)
        self.floor_label = QLabel("-")
        self.number_label = QLabel("-")
        self.name_label = QLabel("-")
        self.perimeter_label = QLabel("-")
        self.area_label = QLabel("-")
        self.boundary_handle_label = QLabel("-")
        self.tag_handle_label = QLabel("-")
        self.closed_label = QLabel("-")
        info_rows = [
            ("층", self.floor_label, "No", self.number_label),
            ("실명", self.name_label, "면적", self.area_label),
            ("둘레", self.perimeter_label, "닫힘", self.closed_label),
        ]
        for row, (label1, value1, label2, value2) in enumerate(info_rows):
            form.addWidget(QLabel(label1), row, 0)
            form.addWidget(value1, row, 1)
            form.addWidget(QLabel(label2), row, 2)
            form.addWidget(value2, row, 3)
        form.setColumnStretch(1, 1)
        form.setColumnStretch(3, 1)
        # 내용 3줄이 보이는 높이로 고정하고 하단 패널 상단에 붙여 둔다.
        info_group.adjustSize()
        info_group.setFixedHeight(info_group.sizeHint().height())

        self.finish_summary_group = QGroupBox("층별 적용 재료마감 집계")
        self.finish_summary_group.setMinimumWidth(0)
        self.finish_summary_group.setSizePolicy(QSizePolicy.Ignored, QSizePolicy.Expanding)
        summary_layout = QVBoxLayout(self.finish_summary_group)
        self.finish_summary_table = QTableWidget(0, 4)
        self.finish_summary_table.setHorizontalHeaderLabels(["품명", "규격", "단위", "수량"])
        self.finish_summary_table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.finish_summary_table.setSelectionMode(QAbstractItemView.NoSelection)
        self.finish_summary_table.setToolTip("열 머리글을 좌우로 드래그하면 열 순서를 바꿀 수 있습니다.")
        self.finish_summary_table.setAlternatingRowColors(False)
        self.finish_summary_table.verticalHeader().setVisible(False)
        summary_header = self.finish_summary_table.horizontalHeader()
        summary_header.setSectionResizeMode(QHeaderView.Interactive)
        summary_header.setSectionsMovable(True)
        summary_header.setStretchLastSection(True)
        for column, width in enumerate((190, 180, 70, 105)):
            self.finish_summary_table.setColumnWidth(column, width)
        # 중앙 하단 재료마감 집계는 헤더를 직접 드래그해 열 순서를 바꿀 수 있다.
        # 논리 열(품명/규격/단위/수량)은 그대로 두고 visual index만 이동하므로
        # 집계 데이터 작성 로직에는 영향을 주지 않는다. 사용자 순서는 settings2.ini에 저장한다.
        self._restore_finish_summary_column_order(finish_summary_column_order)
        summary_header.sectionMoved.connect(self._emit_finish_summary_column_order)
        self.finish_summary_table.setMinimumWidth(0)
        self.finish_summary_table.setSizePolicy(QSizePolicy.Ignored, QSizePolicy.Expanding)
        self.finish_summary_table.setMinimumHeight(110)
        # 최대 높이를 두지 않는다. 가운데 하단 splitter를 늘리면 표가 남는 높이를 전부 사용한다.
        summary_layout.addWidget(self.finish_summary_table, 1)

        # 하단은 선택 실/층 정보 : 재료마감 집계를 약 25 : 75 비율로 사용한다.
        # 최대폭을 고정하면 창 폭이 바뀔 때 한쪽이 찌그러지거나 겹쳐 보일 수 있으므로
        # 두 패널 모두 자유롭게 늘어나게 두고 layout stretch로 비율만 유지한다.
        info_group.setMinimumWidth(0)
        info_group.setMaximumWidth(16777215)
        self.lower_info_panel = QWidget()
        self.lower_info_panel.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        lower_info = QHBoxLayout(self.lower_info_panel)
        lower_info.setContentsMargins(0, 0, 0, 0)
        lower_info.setSpacing(6)
        lower_info.addWidget(info_group, 1, Qt.AlignTop)
        lower_info.addWidget(self.finish_summary_group, 3)
        lower_info.setStretch(0, 1)
        lower_info.setStretch(1, 3)
        self.center_splitter.addWidget(self.lower_info_panel)
        self.center_splitter.setStretchFactor(0, 1)
        self.center_splitter.setStretchFactor(1, 0)
        self.center_splitter.setSizes([650, 160])
        self._lower_panel_last_size = 160
        center_layout.addWidget(self.center_splitter)
        self.main_splitter.addWidget(center)

        materials_path = MATERIAL_FINISH_PATH
        self.finish_panel = RoomFinishPanel(
            materials_path,
            formula_codes=finish_formula_codes,
            finish_categories=finish_categories,
            material_trades=material_trades,
            material_units=material_units,
        )
        self.finish_panel.setMinimumWidth(330)
        self.finish_panel.setSizePolicy(QSizePolicy.Preferred, QSizePolicy.Expanding)
        self.main_splitter.addWidget(self.finish_panel)

        self.main_splitter.setStretchFactor(0, 0)
        self.main_splitter.setStretchFactor(1, 1)
        self.main_splitter.setStretchFactor(2, 0)
        # 구분 드롭다운까지 첫 화면에서 확인하기 쉽도록 목록 기본 폭을 넓힌다.
        self.main_splitter.setSizes([390, 720, 420])
        self._finish_panel_last_size = 420
        root.addWidget(self.main_splitter)

        self.room_table.itemSelectionChanged.connect(self.on_table_select)
        self.room_table.itemChanged.connect(self.on_table_changed)
        self.room_table.cellDoubleClicked.connect(self._on_room_table_double_clicked)
        self.room_table.horizontalHeader().sectionClicked.connect(self._on_room_header_clicked)
        self.group_by_floor_check.toggled.connect(self._on_group_by_floor_toggled)
        self.expand_all_floors_btn.clicked.connect(self.expand_all_floors)
        self.collapse_all_floors_btn.clicked.connect(self.collapse_all_floors)
        self.batch_edit_btn.clicked.connect(self.open_batch_edit_choice)
        self.view_mode.currentIndexChanged.connect(self._on_view_mode_changed)
        self.label_size_spin.valueChanged.connect(self._on_label_size_changed)
        self.preview.room_clicked.connect(self.on_preview_room_clicked)
        self.preview.empty_clicked.connect(self.on_preview_empty_clicked)
        self.preview.room_edit_requested.connect(self.on_preview_room_edit_requested)
        self.preview.escape_pressed.connect(self._on_preview_escape_pressed)
        self.preview.opening_selection_changed.connect(self._on_preview_opening_selection_changed)
        self.preview.opening_delete_requested.connect(self._on_preview_opening_delete_requested)
        self.preview.opening_move_preview_changed.connect(self._on_preview_opening_move_preview)
        self.preview.opening_move_requested.connect(self._on_preview_opening_move_requested)
        self.preview.opening_nudge_requested.connect(self._on_preview_opening_move_requested)
        self.preview.opening_code_edit_requested.connect(self._edit_selected_opening_codes)
        self.preview.opening_catalog_requested.connect(self.open_selected_opening_assign_dialog)
        self.preview.opening_catalog_drop_requested.connect(self.apply_opening_catalog_code)
        self.preview.boundary_split_hover_changed.connect(self._on_boundary_split_hover_changed)
        self.preview.boundary_segment_split_requested.connect(self._on_boundary_segment_split_requested)
        self.preview.boundary_merge_hover_changed.connect(self._on_boundary_merge_hover_changed)
        self.preview.boundary_vertex_merge_requested.connect(self._on_boundary_vertex_merge_requested)
        self.preview.wall_segment_clicked.connect(self.on_wall_segment_clicked)
        self.preview.wall_segments_box_selected.connect(self.on_wall_segments_box_selected)
        self.preview.group_room_clicked.connect(self.on_group_room_clicked)
        self.preview.group_rooms_box_selected.connect(self.on_group_rooms_box_selected)
        self.finish_panel.assignments_changed.connect(self._on_room_finish_panel_changed)
        self.finish_panel.material_delete_requested.connect(self.handle_material_deleted)
        self.finish_panel.group_create_requested.connect(self.create_finish_group)
        self.finish_panel.group_delete_requested.connect(self.delete_finish_group)
        self.finish_panel.group_break_requested.connect(self.break_finish_group)
        self.finish_panel.group_color_requested.connect(self.change_finish_group_color)
        self.finish_panel.group_rename_requested.connect(self.rename_finish_group)
        self.finish_panel.group_division_requested.connect(self.change_finish_group_division)
        self.finish_panel.group_formula_requested.connect(self.change_finish_group_formula)
        self.finish_panel.material_formula_requested.connect(self.change_finish_material_formula)
        self.finish_panel.group_selected.connect(self.on_finish_group_selected)
        self.finish_panel.group_materials_add_requested.connect(self.add_materials_to_finish_group)
        self.finish_panel.wall_group_mode_requested.connect(self.start_wall_group_mode)
        self.finish_panel.group_type_changed.connect(self._on_finish_group_type_changed)
        self.finish_panel.group_type_labels_changed.connect(self._on_finish_group_type_labels_changed)
        self.finish_panel.group_tab_settings_changed.connect(self._on_finish_group_tab_settings_changed)
        self.finish_panel.group_section_move_requested.connect(self.move_finish_group_section)
        self.finish_panel.group_order_changed.connect(self.reorder_finish_groups)
        self.finish_picker_shortcut = QShortcut(QKeySequence(Qt.Key_F5), self)
        self.finish_picker_shortcut.setContext(Qt.WidgetWithChildrenShortcut)
        self.finish_picker_shortcut.activated.connect(self._handle_f5_shortcut)
        self.cad_background_combo.currentIndexChanged.connect(self._on_cad_background_mode_changed)
        self.group_overlay_check.toggled.connect(self.preview.set_show_group_overlay)
        self.finish_group_filter_combo.currentIndexChanged.connect(self._on_finish_group_filter_changed)
        self._rebuild_finish_group_filter_combo(preferred=(self.finish_panel.current_group_type(), "*"))
        self.room_line_check.toggled.connect(self.preview.set_show_room_overlay)
        self.original_text_check.toggled.connect(self.preview.set_show_original_text)
        self.original_text_check.toggled.connect(self.original_text_toggled.emit)
        self.background_settings_btn.clicked.connect(self.background_settings_requested.emit)
        self.direct_name_btn.clicked.connect(self.start_direct_name_input)
        self.apply_selected_btn.clicked.connect(self.apply_selected_room_fields)
        self.dxf_library_btn.clicked.connect(self.drawing_library_requested.emit)
        self.dimension_font_spin.valueChanged.connect(self._on_dimension_quick_changed)
        self.dimension_width_spin.valueChanged.connect(self._on_dimension_quick_changed)
        self.dimension_offset_spin.valueChanged.connect(self._on_dimension_quick_changed)
        self.dimension_unit_m_check.toggled.connect(self._on_dimension_quick_changed)
        self.dimension_show_check.toggled.connect(self._on_dimension_visibility_changed)
        self.boundary_split_btn.toggled.connect(self._toggle_boundary_split_mode)
        self.boundary_merge_btn.toggled.connect(self._toggle_boundary_merge_mode)
        self.floor_tools_collapse_btn.clicked.connect(self._toggle_floor_tools_overlay)
        self.lower_panel_toggle_btn.clicked.connect(self._toggle_lower_panel)
        self.finish_panel_toggle_btn.clicked.connect(self._toggle_finish_panel)
        self.auto_number_btn.clicked.connect(self.open_number_entry_dialog)
        self.create_block_btn.clicked.connect(self.request_create_blocks)
        self.cad_link_btn.clicked.connect(self.open_cad_link_dialog)
        self.opening_test_btn.clicked.connect(self.open_opening_read_dialog)
        self.opening_edit_btn.toggled.connect(self._toggle_opening_edit_mode)
        self.finish_group_edit_list_btn.clicked.connect(self._toggle_finish_group_detail_dialog)
        self.finish_group_edit_confirm_btn.clicked.connect(self._confirm_finish_group_overlay)
        self.finish_group_edit_cancel_btn.clicked.connect(self._cancel_finish_group_overlay)
        self.opening_catalog_assign_btn.clicked.connect(self.open_selected_opening_assign_dialog)
        self.opening_edit_delete_btn.clicked.connect(self._delete_selected_opening_markers)
        self.opening_move_exact_btn.clicked.connect(self._apply_opening_exact_distance)
        self.opening_edit_close_btn.clicked.connect(lambda: self.opening_edit_btn.setChecked(False))
        app = QApplication.instance()
        if app is not None:
            app.installEventFilter(self)
        self.regenerate_background_btn.clicked.connect(self.background_regenerate_requested.emit)
        self.vector_filter_btn.clicked.connect(self.background_filter_requested.emit)
        self.error_check_btn.clicked.connect(self.show_error_dialog)
        self.save_dxf_btn.clicked.connect(self.request_save_dxf)
        self._update_opening_edit_controls()


    @staticmethod
    def _finish_filter_data(group_type: str, division: str) -> str:
        return json.dumps([str(group_type or ""), str(division or "")], ensure_ascii=False)

    @staticmethod
    def _parse_finish_filter_data(value) -> tuple[str, str]:
        try:
            data = json.loads(str(value or ""))
            if isinstance(data, list) and len(data) >= 2:
                group_type = str(data[0] or "").strip()
                division = str(data[1] or "").strip()
                if group_type in FINISH_GROUP_TYPES:
                    return group_type, division or "*"
        except Exception:
            pass
        return "", ""

    def _on_finish_group_filter_changed(self, _index: int = -1):
        group_type, division = self._current_finish_group_display_filter()
        self.preview.set_finish_group_filter(group_type, division)

    def _current_finish_group_display_filter(self) -> tuple[str, str]:
        return self._parse_finish_filter_data(self.finish_group_filter_combo.currentData())

    def _current_finish_group_display_type(self, allow_none: bool = False) -> str:
        # 구형 호출부 호환: 현재 도면표시가 없으면 우측 활성 탭 내부키를 돌려준다.
        group_type, _division = self._current_finish_group_display_filter()
        if group_type in FINISH_GROUP_TYPES:
            return group_type
        return "" if allow_none else self.finish_panel.current_group_type()

    def _current_finish_group_display_division(self, group_type: str | None = None) -> str:
        wanted_type = normalize_finish_group_type(group_type or self.finish_panel.current_group_type())
        current_type, division = self._current_finish_group_display_filter()
        if current_type == wanted_type and division not in {"", "*"}:
            return division
        return ""

    def _set_finish_group_display_filter(self, group_type: str, division: str = "*") -> None:
        group_type = normalize_finish_group_type(group_type)
        division = str(division or "*").strip() or "*"
        target = self._finish_filter_data(group_type, division)
        index = self.finish_group_filter_combo.findData(target)
        if index < 0:
            self._rebuild_finish_group_filter_combo(preferred=(group_type, division))
            index = self.finish_group_filter_combo.findData(target)
        if index >= 0 and self.finish_group_filter_combo.currentIndex() != index:
            self.finish_group_filter_combo.setCurrentIndex(index)
        else:
            self.preview.set_finish_group_filter(group_type, division)

    def _set_finish_group_display_type(self, group_type: str):
        # 기존 호출부는 해당 탭의 '전체' 표시로 해석한다.
        self._set_finish_group_display_filter(group_type, "*")

    def _rebuild_finish_group_filter_combo(self, preferred=None) -> None:
        current_type = self.finish_panel.current_group_type()
        if preferred is None:
            previous = self._current_finish_group_display_filter()
            preferred = previous if previous[0] else (current_type, "*")
        preferred_type, preferred_division = preferred
        labels = self.finish_panel.group_type_labels()
        show_other = self.finish_panel.show_other_group_filters(current_type)

        entries: list[tuple[str, str]] = [("없음", self._finish_filter_data("", ""))]
        for division in self.finish_panel.group_divisions(current_type):
            entries.append((division, self._finish_filter_data(current_type, division)))
        entries.append(("전체", self._finish_filter_data(current_type, "*")))

        if show_other:
            for other_type in FINISH_GROUP_TYPES:
                if other_type == current_type:
                    continue
                tab_label = str(labels.get(other_type, other_type) or other_type)
                for division in self.finish_panel.group_divisions(other_type):
                    entries.append((f"{tab_label} - {division}", self._finish_filter_data(other_type, division)))
                entries.append((f"{tab_label} - 전체", self._finish_filter_data(other_type, "*")))

        target = self._finish_filter_data(preferred_type, preferred_division) if preferred_type else self._finish_filter_data("", "")
        self.finish_group_filter_combo.blockSignals(True)
        try:
            self.finish_group_filter_combo.clear()
            for text, data in entries:
                self.finish_group_filter_combo.addItem(text, data)
            index = self.finish_group_filter_combo.findData(target)
            if index < 0:
                index = self.finish_group_filter_combo.findData(self._finish_filter_data(current_type, "*"))
            self.finish_group_filter_combo.setCurrentIndex(max(0, index))
        finally:
            self.finish_group_filter_combo.blockSignals(False)
        group_type, division = self._current_finish_group_display_filter()
        self.preview.set_finish_group_filter(group_type, division)

    def set_finish_group_type_labels(self, labels=None) -> None:
        self.finish_panel.set_group_type_labels(labels or {})
        self._rebuild_finish_group_filter_combo()

    def set_finish_group_tab_settings(self, settings=None) -> None:
        self.finish_panel.set_group_tab_settings(settings or {})
        self._rebuild_finish_group_filter_combo(preferred=(self.finish_panel.current_group_type(), "*"))

    def _on_finish_group_type_changed(self, group_type: str) -> None:
        # 탭을 옮기면 도면표시도 그 탭의 구분목록으로 바꾸고 기본은 '전체'로 둔다.
        self._rebuild_finish_group_filter_combo(preferred=(normalize_finish_group_type(group_type), "*"))

    def _on_finish_group_type_labels_changed(self, labels) -> None:
        self._rebuild_finish_group_filter_combo()
        self.finish_group_type_labels_changed.emit(dict(labels or {}))

    def _on_finish_group_tab_settings_changed(self, settings) -> None:
        payload = dict(settings or {})

        # V23.35: 한 설정창에서 5개 탭을 동시에 수정하므로 탭별 rename map을 한 번에 받는다.
        rename_maps = {
            str(key): dict(value or {})
            for key, value in dict(payload.get("rename_maps", {}) or {}).items()
            if str(key) in FINISH_GROUP_TYPES
        }
        # V23.29 단일탭 payload도 계속 읽는다.
        changed_type = str(payload.get("changed_group_type", "") or "")
        rename_map = dict(payload.get("rename_map", {}) or {})
        if changed_type in FINISH_GROUP_TYPES and rename_map and changed_type not in rename_maps:
            rename_maps[changed_type] = rename_map

        renamed_any = False
        for group in self.finish_groups.values():
            group_type = normalize_finish_group_type(group.get("group_type", ""))
            type_map = rename_maps.get(group_type, {})
            if not type_map:
                continue
            division = normalize_finish_group_division(group.get("group_division", ""), group_type)
            renamed = str(type_map.get(division.casefold(), "") or "").strip()
            if renamed:
                group["group_division"] = renamed
                renamed_any = True

        if renamed_any:
            # 그룹 기준값을 실별 assignment까지 다시 복사해 모든 화면/집계가 같은 구분명을 쓴다.
            self._ensure_finish_group_colors()
            self.save_finish_assignments()
            self.preview.set_room_finish_assignments(self.room_finish_assignments)
            self.preview.set_finish_groups(self.finish_groups)
            self._refresh_current_finish_panel()

        # 현재 프로젝트에서 실제 사용 중인 구분은 설정에서 빠져도 숨기지 않고 자동 복원한다.
        divisions = dict(payload.get("divisions", {}) or {})
        changed = False
        for group in self.finish_groups.values():
            group_type = normalize_finish_group_type(group.get("group_type", ""))
            division = normalize_finish_group_division(group.get("group_division", ""), group_type)
            values = list(divisions.get(group_type, self.finish_panel.group_divisions(group_type)) or [])
            if division.casefold() not in {str(item).casefold() for item in values}:
                values.append(division)
                divisions[group_type] = values
                changed = True
        if changed:
            payload["divisions"] = divisions
            self.finish_panel.set_group_tab_settings(payload)
        self._rebuild_finish_group_filter_combo(preferred=(self.finish_panel.current_group_type(), "*"))
        self.finish_group_tab_settings_changed.emit(self.finish_panel.group_tab_settings())

    def _on_cad_background_mode_changed(self, _index: int = -1):
        mode = str(self.cad_background_combo.currentData() or "none")
        self.preview.set_cad_background_mode(mode)
        has_ready = bool(self.preview.floor_background_files or self.preview.floor_vector_cache or self.preview.cad_background)
        self.room_line_check.setEnabled(mode != "none" and has_ready)
        self.cad_background_mode_changed.emit(mode)
        # 배경 없음에서는 실 라인 체크상태 자체는 보존하고 조작만 비활성화한다.

    def set_cad_background_mode(self, mode: str):
        mode = str(mode or "black").strip().lower()
        if mode not in {"none", "black", "white"}:
            mode = "black"
        index = self.cad_background_combo.findData(mode)
        self.cad_background_combo.blockSignals(True)
        if index >= 0:
            self.cad_background_combo.setCurrentIndex(index)
        self.cad_background_combo.blockSignals(False)
        self.preview.set_cad_background_mode(mode)
        has_ready = bool(self.preview.floor_background_files or self.preview.floor_vector_cache or self.preview.cad_background)
        self.room_line_check.setEnabled(mode != "none" and has_ready)

    @staticmethod
    def _normalize_room_column_widths(values) -> list[int]:
        defaults = [82, 85, 150, 95, 130, 82, 82, 85]
        if isinstance(values, str):
            values = [part.strip() for part in values.split(",")]
        try:
            widths = [max(30, min(2000, int(value))) for value in list(values or [])]
        except Exception:
            return defaults
        if len(widths) == 6:
            # V19.4: 층/No/실명/면적/둘레/적용층 → V20.0 새 순서
            return [*widths[:3], defaults[3], widths[5], widths[3], widths[4], defaults[7]]
        if len(widths) == 7:
            # V20.0~V20.3 순서 끝에 천정고 열을 추가한다.
            return [*widths, defaults[7]]
        return widths if len(widths) == len(defaults) else defaults

    @staticmethod
    def _migrate_legacy_room_column_widths(values) -> list[int]:
        """과거 열 순서를 V20.4(천정고 포함) 순서로 한 번만 변환한다."""
        if isinstance(values, str):
            values = [part.strip() for part in values.split(",")]
        try:
            widths = [max(30, min(2000, int(value))) for value in list(values or [])]
        except Exception:
            widths = []
        if len(widths) == 6:
            return [*widths[:3], 95, widths[5], widths[3], widths[4], 85]
        if len(widths) == 7:
            # V19.5~V19.9: 층/No/실명/구분/면적/둘레/적용층 → V20.4 새 순서
            return [*widths[:4], widths[6], widths[4], widths[5], 85]
        if len(widths) == 8:
            return widths
        return [82, 85, 150, 95, 130, 82, 82, 85]

    def _queue_room_column_width_save(self, *_args):
        self._room_column_save_timer.start()

    def _emit_room_column_widths(self):
        self.room_column_widths_changed.emit([
            self.room_table.columnWidth(index)
            for index in range(self.room_table.columnCount())
        ])

    @staticmethod
    def _normalize_finish_summary_column_order(values) -> list[int]:
        """중앙 하단 집계표의 visual 열 순서를 0~3 논리열 순열로 정규화한다."""
        defaults = [0, 1, 2, 3]
        if isinstance(values, str):
            values = [part.strip() for part in values.split(",") if part.strip() != ""]
        try:
            order = [int(value) for value in list(values or [])]
        except Exception:
            return defaults
        return order if len(order) == 4 and sorted(order) == defaults else defaults

    def _restore_finish_summary_column_order(self, values):
        header = self.finish_summary_table.horizontalHeader()
        order = self._normalize_finish_summary_column_order(values)
        for target_visual, logical_index in enumerate(order):
            current_visual = header.visualIndex(logical_index)
            if current_visual >= 0 and current_visual != target_visual:
                header.moveSection(current_visual, target_visual)

    def _emit_finish_summary_column_order(self, *_args):
        header = self.finish_summary_table.horizontalHeader()
        order = [
            header.logicalIndex(visual_index)
            for visual_index in range(header.count())
        ]
        self.finish_summary_column_order_changed.emit(order)


    def apply_preview_settings(self, background_max_pixels: int, edge_length_font_size: int, edge_length_offset: int = 150, edge_length_width_ratio: int = 80, background_style="gray", background_opacity=100, boundary_color="#0066FF", boundary_width=9, render_quality="high", selection_color="#FF0000", selection_width=18, background_line_width: float = 1.5, background_aci_widths: str = "", background_keep_red: bool = False, edge_length_show_unit_m: bool = True, label_full_view_weight: float = 0.70, label_selected_view_weight: float = 1.00, room_label_scale_mode: str = "room_area", deduction_color="#FF00FF", deduction_width=1, wall_selection_color="#FF00FF", wall_selection_width=2, group_overlay_opacity=35, wall_selection_zoom_scale: float = 0.35, edge_length_visible: bool = True, opening_line_color="#8ED4E3", opening_line_width=4, background_zoomout_width_weight: float = 0.25):
        self.preview.set_background_max_pixels(background_max_pixels)
        self.preview.set_edge_length_font_size(edge_length_font_size)
        self.preview.set_edge_length_offset(edge_length_offset)
        self.preview.set_edge_length_width_ratio(edge_length_width_ratio)
        self.preview.set_edge_length_show_unit_m(edge_length_show_unit_m)
        self.preview.set_label_scale_mode(room_label_scale_mode)
        self.preview.set_label_view_weights(label_full_view_weight, label_selected_view_weight)
        self.set_dimension_controls(edge_length_font_size, edge_length_width_ratio, edge_length_offset, edge_length_show_unit_m, edge_length_visible)
        self.preview.set_deduction_style(deduction_color, deduction_width, redraw=False)
        self.preview.set_wall_selection_style(wall_selection_color, wall_selection_width, wall_selection_zoom_scale, redraw=False)
        self.preview.set_opening_style(opening_line_color, opening_line_width, redraw=False)
        self.preview.set_group_overlay_opacity(group_overlay_opacity, redraw=False)
        self.preview.apply_cad_style(background_style, background_opacity, boundary_color, boundary_width, render_quality, selection_color, selection_width, background_line_width, background_aci_widths, background_keep_red, background_zoomout_width_weight)

    def set_dimension_controls(self, font_size: int, width_ratio: int, offset: int, show_unit_m: bool, visible: bool = True):
        self._syncing_dimension_controls = True
        try:
            self.dimension_font_spin.setValue(int(font_size))
            self.dimension_width_spin.setValue(int(width_ratio))
            self.dimension_offset_spin.setValue(int(offset))
            self.dimension_unit_m_check.setChecked(bool(show_unit_m))
            self.dimension_show_check.setChecked(bool(visible))
        finally:
            self._syncing_dimension_controls = False

    def _on_label_size_changed(self, value: int):
        value = int(value)
        self.preview.set_label_font_size(value)
        self.label_size_changed.emit(value)

    def _on_dimension_quick_changed(self, *_args):
        if self._syncing_dimension_controls:
            return
        values = {
            "edge_length_font_size": self.dimension_font_spin.value(),
            "edge_length_width_ratio": self.dimension_width_spin.value(),
            "edge_length_offset": self.dimension_offset_spin.value(),
            "edge_length_show_unit_m": self.dimension_unit_m_check.isChecked(),
        }
        self.preview.set_edge_length_font_size(values["edge_length_font_size"])
        self.preview.set_edge_length_width_ratio(values["edge_length_width_ratio"])
        self.preview.set_edge_length_offset(values["edge_length_offset"])
        self.preview.set_edge_length_show_unit_m(values["edge_length_show_unit_m"])
        self.dimension_settings_changed.emit(values)

    def _on_dimension_visibility_changed(self, checked: bool):
        if self._syncing_dimension_controls:
            return
        self.dimension_settings_changed.emit({"edge_length_visible": bool(checked)})
        # 치수 표시 ON/OFF는 같은 도면의 표시 변경일 뿐이므로 현재 zoom/PAN을 유지한다.
        # 전체 refresh를 거치면 치수 유무에 따라 view scope 키가 달라져 화면이 재맞춤될 수 있어
        # 현재 장면만 동일 위치에서 다시 그리는 전용 경로를 사용한다.
        self.preview.set_edge_lengths_visible(bool(checked))

    @staticmethod
    def _numpad_digit_from_event(event):
        """Windows 숫자키패드는 NumLock 상태에 따라 Qt.Key_2 대신 Down 등으로 들어올 수 있다.

        nativeVirtualKey(VK_NUMPAD0~9)를 우선 사용하고, 그 값이 없을 때만
        Qt.KeypadModifier + Qt.Key_0~9 조합을 사용한다.
        """
        try:
            native = int(event.nativeVirtualKey())
        except Exception:
            native = 0
        if 0x60 <= native <= 0x69:  # VK_NUMPAD0 ~ VK_NUMPAD9
            return native - 0x60
        try:
            if event.modifiers() & Qt.KeypadModifier:
                key = int(event.key())
                for digit in range(10):
                    if key == int(getattr(Qt, f"Key_{digit}")):
                        return digit
        except Exception:
            pass
        return None

    def eventFilter(self, watched, event):
        try:
            if event.type() == QEvent.KeyPress and self.isVisible():
                # 창호 편집은 평면 도곽에서 계속 유지되는 모드다. Delete로 지워도 종료하지 않는다.
                if self._opening_edit_mode and event.key() in (Qt.Key_Delete, Qt.Key_Backspace):
                    focus = QApplication.focusWidget()
                    # 문자/표 셀 편집 중 Delete는 원래 편집 동작을 우선한다.
                    if not isinstance(focus, QLineEdit) and not (isinstance(focus, QTableWidget) and focus.state() == QAbstractItemView.State.EditingState):
                        self._delete_selected_opening_markers()
                        return True

                modifiers = event.modifiers()
                if modifiers & Qt.ControlModifier:
                    digit = self._numpad_digit_from_event(event)
                    if digit in (2, 5):
                        if not self.lower_info_panel.isHidden():
                            self._toggle_lower_panel()
                        return True
                    if digit == 8:
                        if self.lower_info_panel.isHidden():
                            self._toggle_lower_panel()
                        return True
                    if digit == 6:
                        if not self.finish_panel.isHidden():
                            self._toggle_finish_panel()
                        return True
                    if digit == 4:
                        if self.finish_panel.isHidden():
                            self._toggle_finish_panel()
                        return True
        except Exception:
            pass
        return super().eventFilter(watched, event)

    def _toggle_lower_panel(self):
        if not self.lower_info_panel.isHidden():
            sizes = self.center_splitter.sizes()
            if len(sizes) > 1 and sizes[1] > 0:
                self._lower_panel_last_size = sizes[1]
            self.lower_info_panel.hide()
            self.lower_panel_toggle_btn.setText("▲")
            self.lower_panel_toggle_btn.setToolTip("아래쪽 선택 실/층 정보와 재료마감 집계를 펼칩니다. · Ctrl+Num8")
            return

        self.lower_info_panel.show()
        sizes = self.center_splitter.sizes()
        total = max(sum(sizes), self.center_splitter.height(), 1)
        lower_size = max(120, int(self._lower_panel_last_size))
        self.center_splitter.setSizes([max(1, total - lower_size), lower_size])
        self.lower_panel_toggle_btn.setText("▼")
        self.lower_panel_toggle_btn.setToolTip("선택 실/층 정보와 재료마감 집계를 아래로 접습니다. · Ctrl+Num2 / Ctrl+Num5")

    def _toggle_finish_panel(self):
        if self.finish_panel.isVisible():
            sizes = self.main_splitter.sizes()
            if len(sizes) > 2 and sizes[2] > 0:
                self._finish_panel_last_size = sizes[2]
            self.finish_panel.hide()
            self.finish_panel_toggle_btn.setText("◀")
            self.finish_panel_toggle_btn.setToolTip("오른쪽 재료마감 패널을 펼칩니다. · Ctrl+Num4")
            return

        self.finish_panel.show()
        sizes = self.main_splitter.sizes()
        total = max(sum(sizes), self.main_splitter.width(), 1)
        left_size = sizes[0] if sizes else 300
        right_size = max(280, int(self._finish_panel_last_size))
        center_size = max(1, total - left_size - right_size)
        self.main_splitter.setSizes([left_size, center_size, right_size])
        self.finish_panel_toggle_btn.setText("▶")
        self.finish_panel_toggle_btn.setToolTip("오른쪽 재료마감 패널을 접습니다. · Ctrl+Num6")

    def _stop_drawing_edit_modes(self) -> None:
        """도면편집 패널 소속의 활성 편집모드를 모두 종료한다."""
        if self.boundary_split_btn.isChecked():
            self.boundary_split_btn.setChecked(False)
        if self.boundary_merge_btn.isChecked():
            self.boundary_merge_btn.setChecked(False)
        if self.opening_edit_btn.isChecked() or self._opening_edit_mode:
            self.opening_edit_btn.setChecked(False)

    def _toggle_floor_tools_overlay(self) -> None:
        collapsed = self.floor_tools_content.isVisible()
        # V23.59: 편집 중 도구모음을 접으면 숨은 모드가 남지 않도록 즉시 종료한다.
        if collapsed:
            self._stop_drawing_edit_modes()
        self.floor_tools_content.setVisible(not collapsed)
        self.floor_tools_collapse_btn.setText("▶" if collapsed else "◀")
        self.floor_tools_collapse_btn.setToolTip(
            "도면 편집 도구를 펼칩니다." if collapsed else "도면 편집 도구를 접습니다."
        )
        self.floor_tools_overlay.adjustSize()
        self.floor_tools_overlay.raise_()

    def _on_view_mode_changed(self, *_args):
        selected_mode = self.view_mode.currentText() == "선택 실만 보기"
        self.dimension_overlay.setVisible(selected_mode)
        self.floor_tools_overlay.setVisible(not selected_mode)
        if selected_mode:
            if self.boundary_split_btn.isChecked():
                self.boundary_split_btn.setChecked(False)
            if self.boundary_merge_btn.isChecked():
                self.boundary_merge_btn.setChecked(False)
            if self._opening_edit_mode:
                self.opening_edit_btn.setChecked(False)
            self.dimension_overlay.raise_()
        else:
            self.floor_tools_overlay.raise_()
        self.refresh_preview()

    def set_original_text_visible(self, enabled: bool):
        self.original_text_check.blockSignals(True)
        self.original_text_check.setChecked(bool(enabled))
        self.original_text_check.blockSignals(False)
        self.preview.set_show_original_text(bool(enabled))

    def set_cad_background(self, segments):
        self.cad_background = list(segments or [])
        self.preview.set_cad_background(self.cad_background)

    def set_background_cache_files(self, background_files: dict):
        old_floors = set(self.preview.floor_background_files)
        new_files = dict(background_files or {})
        self.preview.set_background_cache_files(new_files)
        if background_files:
            self.cad_background_status_label.setText("(준비됨)")
        # V24.07.2: 실이 일부 층에 이미 있어도 새 배경 캐시가 완성되면 층 목록을
        # 다시 구성해야 실 0개인 나머지 층이 즉시 나타난다.
        if old_floors != set(new_files) or not self.rooms:
            self.rebuild_views()

    def set_vector_display_stage(self, stage: str):
        stage = str(stage or "none").strip().lower()
        self.vector_display_stage = stage
        if stage == "raw":
            text, color = "표시: 원본 정리", "#8A5A00"
        elif stage == "optimized":
            text, color = "표시: 최적화 도면", "#0B63CE"
        else:
            text, color = "표시: 없음", "#555555"
        self.vector_stage_label.setText(text)
        self.vector_stage_label.setStyleSheet(f"font-weight: 600; color: {color};")

    def set_vector_background_cache(self, floor_vectors: dict, preserve_layout: bool = False):
        # V24.25: 레이어/그룹 표시 ON/OFF는 화면 내용만 바뀌는 작업이다.
        # 캐시 교체나 rebuild_views() 과정에서 reset_view()가 호출되더라도
        # 사용자가 보고 있던 줌 배율과 CAD/world 중심은 그대로 복원한다.
        view_state = self.preview.capture_view_state() if preserve_layout else None
        old_floors = set(self.preview.floor_vector_cache)
        new_vectors = dict(floor_vectors or {})
        self.preview.set_vector_background_cache(new_vectors, preserve_layout=preserve_layout)
        # V24.07.1의 층 union 로직만으로는 캐시 완료 후 테이블 재구성이 안 되어
        # 실이 있는 층만 계속 남았다. 캐시 층 집합이 바뀌면 실 유무와 무관하게 갱신한다.
        if old_floors != set(new_vectors) or not self.rooms:
            self.rebuild_views()
        if view_state:
            self.preview.restore_view_state(view_state)
            # 탭/테이블 재구성 직후 Qt가 scrollbar 범위를 한 번 더 정리하는 경우가 있어
            # 이벤트 루프가 돈 뒤 동일 상태를 한 번 더 복원한다.
            QTimer.singleShot(0, lambda state=view_state: self.preview.restore_view_state(state))

    def set_background_render_mode(self, mode: str):
        self.preview.set_background_render_mode(mode)
        has_ready = bool(self.preview.floor_background_files or self.preview.floor_vector_cache or self.preview.cad_background)
        self.cad_background_combo.setEnabled(has_ready)
        self.original_text_check.setEnabled(has_ready)

    def clear_background_cache_files(self):
        self.preview.clear_background_sources()
        self.cad_background_status_label.setText("(미생성)")
        self.cad_background_combo.setEnabled(False)
        self.room_line_check.setEnabled(False)
        self.original_text_check.setEnabled(False)

    def set_background_loading(self, loading: bool):
        # V18.7: 배경 생성은 백그라운드 작업. 기존 배경/화면은 유지하고
        # 재생성 버튼만 잠가 사용자가 실 선택·수량 작업 등을 계속할 수 있게 한다.
        has_ready = bool(self.preview.floor_background_files or self.preview.floor_vector_cache or self.preview.cad_background)
        self.cad_background_combo.setEnabled(has_ready)
        self.room_line_check.setEnabled(has_ready and str(self.cad_background_combo.currentData() or "none") != "none")
        self.original_text_check.setEnabled(has_ready)
        self.regenerate_background_btn.setEnabled(not loading)
        self.vector_filter_btn.setEnabled(not loading)
        if loading:
            self.cad_background_status_label.setText("(백그라운드 생성중)")
        else:
            self.cad_background_status_label.setText("(준비됨)" if has_ready else "(미생성)")

    def set_background_ready(self, ready: bool):
        self.cad_background_combo.setEnabled(bool(ready))
        self.room_line_check.setEnabled(bool(ready) and str(self.cad_background_combo.currentData() or "none") != "none")
        self.original_text_check.setEnabled(bool(ready))
        self.regenerate_background_btn.setEnabled(True)
        self.cad_background_status_label.setText("(준비됨)" if ready else "(미생성)")
        if not ready:
            self.preview.set_background_cache_files({})
            self.preview.set_vector_background_cache({})
            self.preview.set_cad_background([])

    def set_room_rules(self, rules: dict):
        self.room_rules = dict(rules or {})

    def set_room_category_rules(self, rules: dict):
        self.room_category_rules = normalize_room_category_rules(rules)
        changed = False
        for room in self.rooms:
            changed = self._classify_room_if_needed(room) or changed
        if changed:
            self.data_changed.emit()
            self.rebuild_views()

    def set_source_path(self, path: str):
        self.current_source_path = path
        self.original_source_path = path
        self.working_dxf_path = ""
        self.save_dxf_btn.setEnabled(bool(path and self.rooms))
        self.load_finish_assignments()

    def set_base_dxf_path(self, path: str):
        self.base_dxf_path = str(path or "")
        if self._opening_read_dialog is not None:
            self._opening_read_dialog.read_btn.setEnabled(bool(self.cad_connected_mode or self._base_dxf_available()))

    def _base_dxf_available(self) -> bool:
        try:
            if self.base_dxf_path and Path(self.base_dxf_path).is_file() and Path(self.base_dxf_path).suffix.lower() == ".dxf":
                return True
            current = Path(self.current_source_path) if self.current_source_path else None
            if current is not None and current.is_file() and current.suffix.lower() == ".dxf":
                self.base_dxf_path = str(current)
                return True
            # 임시 프로젝트의 베이스 DXF도 보조 검색한다.
            candidates = [p for p in TEMP_PROJECT_DIR.glob("*.dxf") if p.is_file()]
            if len(candidates) == 1:
                self.base_dxf_path = str(candidates[0])
                return True
        except Exception:
            pass
        return False

    def _operation_source_path(self) -> str:
        return self.working_dxf_path or self.current_source_path

    def _run_cached_dxf_operation(self, title, operation, targets):
        source_path = self._operation_source_path()
        if not source_path:
            return None
        source = Path(source_path)
        if source.suffix.lower() != ".dxf" or not source.exists():
            QMessageBox.warning(
                self, title,
                "현재 작업 원본이 DXF가 아닙니다. CAD 연동 상태를 다시 확인한 뒤 실행하십시오.\n"
                "이 작업에서는 DXF 저장창을 열지 않습니다."
            )
            return None
        try:
            WORKING_DXF_PATH.parent.mkdir(parents=True, exist_ok=True)
            staging = WORKING_DXF_PATH.with_name("현재작업_작성중.dxf")
            report = operation(str(source), str(staging), targets, self.room_rules)
            staging.replace(WORKING_DXF_PATH)
            self.working_dxf_path = str(WORKING_DXF_PATH)
            self.current_source_path = self.working_dxf_path
            self.rebuild_views()
            return self.working_dxf_path, report
        except Exception:
            QMessageBox.critical(self, title + " 오류", traceback.format_exc())
            return None

    def set_data(self, rooms: list[RoomRecord], errors: list[str]):
        # 창호 읽기 결과는 프로젝트 데이터로 유지하고, 현재 도곽 오버레이만 새 실목록에 맞춰 다시 동기화한다.
        self.preview.opening_test_floor = ""
        self.preview.opening_test_markers = []
        self.opening_test_btn.setText("창호 리스트")
        self.rooms = rooms
        self.errors = errors
        self.opening_test_btn.setEnabled(bool(self.rooms))
        self.boundary_split_btn.setEnabled(bool(self.rooms))
        self.boundary_merge_btn.setEnabled(bool(self.rooms))
        if not self.rooms and self.boundary_split_btn.isChecked():
            self.boundary_split_btn.setChecked(False)
        if not self.rooms and self.boundary_merge_btn.isChecked():
            self.boundary_merge_btn.setChecked(False)
        for room in self.rooms:
            # 구버전 캐시의 긴 미입력 문구도 새 기본값 '-'로 한 번에 정규화한다.
            room.room_name = normalize_room_name(room.room_name)
            self._classify_room_if_needed(room)
        self.rebuild_views()

    def _classify_room_if_needed(self, room: RoomRecord) -> bool:
        if str(getattr(room, "room_category", "") or "").strip():
            return False
        category = classify_room_name(room.room_name, self.room_category_rules)
        if not category:
            return False
        room.room_category = category
        return True

    def _room_sort_key(self, room: RoomRecord, column: int):
        if column == 0:
            return floor_sort_key(room.floor)
        if column == 1:
            return (str(room.room_number or "").casefold(),)
        if column == 2:
            return (str(room.room_name or "").casefold(),)
        if column == 3:
            return (str(room.room_category or "").casefold(),)
        if column == 4:
            return (str(room.floor_scope or room.floor).casefold(), int(room.repeat_count or 1))
        if column == 5:
            return (float(room.area_m2),)
        if column == 6:
            return (float(room.perimeter_m),)
        if column == 7:
            height = room.ceiling_height_m
            return (float(height) if height is not None else -1.0,)
        return (0,)

    @staticmethod
    def _format_ceiling_height(value) -> str:
        if value is None or str(value).strip() == "":
            return ""
        try:
            number = float(value)
        except Exception:
            return ""
        return f"{number:.3f}".rstrip("0").rstrip(".")

    @classmethod
    def _floor_height_text(cls, rooms: list[RoomRecord]) -> str:
        values = {
            round(float(room.floor_height_m), 9)
            for room in rooms
            if getattr(room, "floor_height_m", None) is not None
        }
        if len(values) != 1:
            return ""
        return cls._format_ceiling_height(next(iter(values)))

    def _sorted_room_indices(self, indices: list[int]) -> list[int]:
        if self.room_sort_column is None:
            return list(indices)

        # No(실번호) 정렬에서는 번호가 없는 실을 오름/내림차순과 관계없이 항상 맨 아래 둔다.
        # 번호가 있는 실끼리만 사용자가 선택한 방향으로 정렬한다.
        if self.room_sort_column == 1:
            numbered = [idx for idx in indices if str(self.rooms[idx].room_number or "").strip()]
            empty = [idx for idx in indices if not str(self.rooms[idx].room_number or "").strip()]
            numbered = sorted(
                numbered,
                key=lambda idx: self._room_sort_key(self.rooms[idx], 1),
                reverse=not self.room_sort_ascending,
            )
            return numbered + empty

        return sorted(
            indices,
            key=lambda idx: self._room_sort_key(self.rooms[idx], self.room_sort_column),
            reverse=not self.room_sort_ascending,
        )

    def _populate_room_table_row(self, table_row: int, room_index: int):
        room = self.rooms[room_index]
        values = [
            room.floor,
            room.room_number,
            normalize_room_name(room.room_name),
            room.room_category,
            f"{room.floor_scope or room.floor}/{max(1, int(room.repeat_count or 1))}개",
            f"{room.area_m2:.3f}",
            f"{room.perimeter_m:.3f}",
            self._format_ceiling_height(room.ceiling_height_m),
        ]
        for col, value in enumerate(values):
            item = QTableWidgetItem(str(value))
            item.setData(Qt.UserRole, room_index)
            if col >= 3 and col != 7:
                item.setFlags(item.flags() & ~Qt.ItemIsEditable)
            self.room_table.setItem(table_row, col, item)

        category_item = self.room_table.item(table_row, 3)
        if category_item is not None:
            category_item.setFlags(category_item.flags() | Qt.ItemIsEditable)
            font = category_item.font(); font.setBold(False); font.setWeight(QFont.Normal); category_item.setFont(font)

    @staticmethod
    def _quantity_room_count(rooms: list[RoomRecord]) -> int:
        """대표층 실에 적용된 층 개소를 반영한 실제 산출 실 개수."""
        return sum(max(1, int(room.repeat_count or 1)) for room in rooms)

    def _set_room_summary_row(self, table_row: int, values: list[str], background: str, editable_columns=None):
        """전층/층별 집계행을 표시한다. 층 합계의 천정고 열은 층높이(H) 입력용으로 허용한다."""
        editable = {int(value) for value in (editable_columns or [])}
        for col, value in enumerate(values):
            item = QTableWidgetItem(str(value))
            item.setData(Qt.UserRole, None)
            if col not in editable:
                item.setFlags(item.flags() & ~Qt.ItemIsEditable)
            font = item.font()
            font.setBold(True)
            item.setFont(font)
            item.setBackground(QBrush(QColor(background)))
            self.room_table.setItem(table_row, col, item)

    def rebuild_views(self):
        table_state = self._capture_room_table_state()
        self.loading_table = True
        self.table_row_room_indices = []
        self.table_row_floors = []
        self.room_category_combos = {}
        self.room_table.clearContents()

        floor_lookup = defaultdict(list)
        for room_index, room in enumerate(self.rooms):
            floor_lookup[room.floor].append(room_index)

        number_errors = validate_room_numbers(self.rooms)
        background_floors = sorted(
            {str(v) for v in (set(self.preview.floor_background_files) | set(self.preview.floor_vector_cache)) if str(v).strip()},
            key=floor_sort_key,
        )
        if not self.rooms and background_floors:
            # V23.93: 실구분선이 없어도 평면도 도곽에 지정된 층을 목록에 표시해
            # 배경만 먼저 확인/전환할 수 있게 한다. 이 행은 실제 RoomRecord를 만들지 않는다.
            self.room_table.setRowCount(len(background_floors))
            for table_row, floor in enumerate(background_floors):
                self.table_row_room_indices.append(None)
                self.table_row_floors.append(floor)
                values = [floor, "실 0개", "", "", "", "0.000", "0.000", ""]
                self._set_room_summary_row(table_row, values, "#DCEEFF")
                self.room_table.item(table_row, 0).setToolTip("실구분선 없이 도곽/층 정보로 생성된 배경입니다.")
        elif not self.rooms:
            self.room_table.setRowCount(0)
        elif self.group_by_floor:
            # V24.07.1: 실이 일부 층에만 있어도 배경 캐시에 존재하는 나머지 층을
            # 실 0개 층으로 계속 노출한다. 실 존재 여부는 배경 층 목록/전환 조건이 아니다.
            floor_names = sorted(set(floor_lookup) | set(background_floors), key=floor_sort_key)
            if self.room_sort_column == 0 and not self.room_sort_ascending:
                floor_names.reverse()
            floor_groups = [(floor, self._sorted_room_indices(floor_lookup.get(floor, []))) for floor in floor_names]
            # 맨 위에는 적용 개소를 반영한 전층 대구분 집계를 항상 둔다.
            total_rows = 1 + sum(1 + len(indices) for _, indices in floor_groups)
            self.room_table.setRowCount(total_rows)

            grand_values = [
                "전층",
                f"실 {self._quantity_room_count(self.rooms):,}개",
                "전층 대구분",
                "",
                "산출합계",
                f"{sum(room.quantity_area_m2 for room in self.rooms):.3f}",
                f"{sum(room.quantity_perimeter_m for room in self.rooms):.3f}",
                "",
            ]
            self.table_row_room_indices.append(None)
            self.table_row_floors.append(None)
            self._set_room_summary_row(0, grand_values, "#B8D9F5")
            self.room_table.item(0, 0).setToolTip("모든 층의 실 개수·면적·둘레 산출수량 합계입니다.")

            table_row = 1
            for floor, room_indices in floor_groups:
                floor_rooms = [self.rooms[index] for index in room_indices]
                is_collapsed = floor in self.collapsed_floors
                total_values = [
                    ("▶ " if is_collapsed else "▼ ") + floor,
                    f"실 {self._quantity_room_count(floor_rooms):,}개",
                    "전체",
                    "", "산출합계",
                    f"{sum(room.quantity_area_m2 for room in floor_rooms):.3f}",
                    f"{sum(room.quantity_perimeter_m for room in floor_rooms):.3f}",
                    self._floor_height_text(floor_rooms),
                ]
                self.table_row_room_indices.append(None)
                self.table_row_floors.append(floor)
                self._set_room_summary_row(table_row, total_values, "#DCEEFF", editable_columns={7} if floor_rooms else set())
                if floor_rooms:
                    self.room_table.item(table_row, 0).setToolTip("더블클릭하면 이 층 목록을 접거나 펼칩니다.")
                    self.room_table.item(table_row, 7).setToolTip("층높이 H 입력값(m) · 이 층의 모든 실 산식에 공통 적용됩니다.")
                else:
                    self.room_table.item(table_row, 0).setToolTip("실구분선 없이 도곽/층 배경만 존재하는 층입니다.")
                table_row += 1

                for room_index in room_indices:
                    self.table_row_room_indices.append(room_index)
                    self.table_row_floors.append(self.rooms[room_index].floor)
                    self._populate_room_table_row(table_row, room_index)
                    self.room_table.setRowHidden(table_row, is_collapsed)
                    table_row += 1
        else:
            room_indices = list(range(len(self.rooms)))
            if self.room_sort_column is not None:
                room_indices = self._sorted_room_indices(room_indices)
            self.room_table.setRowCount(1 + len(room_indices))
            grand_values = [
                "전층",
                f"실 {self._quantity_room_count(self.rooms):,}개",
                "전층 대구분",
                "",
                "산출합계",
                f"{sum(room.quantity_area_m2 for room in self.rooms):.3f}",
                f"{sum(room.quantity_perimeter_m for room in self.rooms):.3f}",
                "",
            ]
            self.table_row_room_indices.append(None)
            self.table_row_floors.append(None)
            self._set_room_summary_row(0, grand_values, "#B8D9F5")
            self.room_table.item(0, 0).setToolTip("모든 층의 실 개수·면적·둘레 산출수량 합계입니다.")
            for table_row, room_index in enumerate(room_indices, start=1):
                self.table_row_room_indices.append(room_index)
                self.table_row_floors.append(self.rooms[room_index].floor)
                self._populate_room_table_row(table_row, room_index)

        self.loading_table = False
        combined_errors = list(self.errors)
        for msg in number_errors:
            if msg not in combined_errors:
                combined_errors.append(msg)
        for msg in self.cad_validation_errors:
            if msg not in combined_errors:
                combined_errors.append(msg)
        for msg in self.room_wall_validation_errors:
            if msg not in combined_errors:
                combined_errors.append(msg)
        self.current_combined_errors = combined_errors
        self.update_error_button()
        self.save_dxf_btn.setEnabled(bool(self.current_source_path and self.rooms))

        if self.room_table.rowCount() and (self.rooms or background_floors):
            restored = self._restore_room_table_state(table_state)
            if not restored:
                self.room_table.selectRow(0)
            current_row = self.room_table.currentRow()
            current_room_index = self._room_index_for_table_row(current_row)
            if current_room_index is None:
                current_floor = self._floor_for_table_row(current_row)
                if current_floor:
                    self.show_floor_info(current_floor)
                elif current_row == 0:
                    self.show_all_floors_info()
            else:
                self.show_room_info(self.rooms[current_room_index])
            self.refresh_preview()
        else:
            self.preview.show_floor([])
            self.clear_info()
            self.update_finish_summary(None)

    def set_group_by_floor(self, enabled: bool):
        enabled = bool(enabled)
        self.group_by_floor = enabled
        self.group_by_floor_check.blockSignals(True)
        self.group_by_floor_check.setChecked(enabled)
        self.group_by_floor_check.blockSignals(False)
        self.expand_all_floors_btn.setEnabled(enabled)
        self.collapse_all_floors_btn.setEnabled(enabled)
        self.rebuild_views()

    def _on_group_by_floor_toggled(self, checked: bool):
        self.group_by_floor = bool(checked)
        self.expand_all_floors_btn.setEnabled(self.group_by_floor)
        self.collapse_all_floors_btn.setEnabled(self.group_by_floor)
        self.rebuild_views()
        self.group_by_floor_toggled.emit(self.group_by_floor)

    def _on_room_header_clicked(self, column: int):
        column = int(column)
        if self.room_sort_column == column:
            self.room_sort_ascending = not self.room_sort_ascending
        else:
            self.room_sort_column = column
            self.room_sort_ascending = True
        self.room_table.horizontalHeader().setSortIndicator(
            column,
            Qt.AscendingOrder if self.room_sort_ascending else Qt.DescendingOrder,
        )
        self.rebuild_views()

    def _room_index_for_table_row(self, table_row: int) -> int | None:
        if 0 <= table_row < len(self.table_row_room_indices):
            return self.table_row_room_indices[table_row]
        return None

    def _floor_for_table_row(self, table_row: int) -> str | None:
        if 0 <= table_row < len(self.table_row_floors):
            return self.table_row_floors[table_row]
        return None

    def _table_row_for_room(self, room: RoomRecord) -> int | None:
        for table_row, room_index in enumerate(self.table_row_room_indices):
            if room_index is not None and self.rooms[room_index] is room:
                return table_row
        return None

    def _selected_room_indices(self) -> list[int]:
        result: list[int] = []
        seen: set[int] = set()
        selection = self.room_table.selectionModel()
        for index in selection.selectedRows() if selection else []:
            room_index = self._room_index_for_table_row(index.row())
            if room_index is not None and room_index not in seen:
                result.append(room_index)
                seen.add(room_index)
        return result

    def _capture_room_table_state(self) -> dict:
        """목록 재구성 전 현재 행·다중 선택·스크롤 위치를 실 객체 기준으로 보관한다."""
        selected_room_ids: list[int] = []
        selected_floor_headers: list[str] = []
        selection = self.room_table.selectionModel()
        for index in selection.selectedRows() if selection else []:
            table_row = index.row()
            room_index = self._room_index_for_table_row(table_row)
            if room_index is not None and 0 <= room_index < len(self.rooms):
                selected_room_ids.append(id(self.rooms[room_index]))
            elif self._floor_for_table_row(table_row):
                selected_floor_headers.append(self._floor_for_table_row(table_row))

        current_row = self.room_table.currentRow()
        current_room_index = self._room_index_for_table_row(current_row)
        current_room_id = None
        current_floor_header = None
        if current_room_index is not None and 0 <= current_room_index < len(self.rooms):
            current_room_id = id(self.rooms[current_room_index])
        elif current_row >= 0:
            current_floor_header = self._floor_for_table_row(current_row)
        return {
            "selected_room_ids": selected_room_ids,
            "selected_floor_headers": selected_floor_headers,
            "current_room_id": current_room_id,
            "current_floor_header": current_floor_header,
            "current_column": max(0, self.room_table.currentColumn()),
            "vertical_scroll": self.room_table.verticalScrollBar().value(),
            "horizontal_scroll": self.room_table.horizontalScrollBar().value(),
        }

    def _restore_room_table_state(self, state: dict | None) -> bool:
        if not state:
            return False
        selected_ids = set(state.get("selected_room_ids") or [])
        selected_headers = set(state.get("selected_floor_headers") or [])
        current_room_id = state.get("current_room_id")
        current_floor_header = state.get("current_floor_header")
        selected_rows: list[int] = []
        current_row = None
        for table_row, room_index in enumerate(self.table_row_room_indices):
            if room_index is None:
                floor = self._floor_for_table_row(table_row)
                if floor in selected_headers:
                    selected_rows.append(table_row)
                if current_room_id is None and floor == current_floor_header:
                    current_row = table_row
                continue
            if not (0 <= room_index < len(self.rooms)):
                continue
            room_id = id(self.rooms[room_index])
            if room_id in selected_ids:
                selected_rows.append(table_row)
            if room_id == current_room_id:
                current_row = table_row

        if current_row is None and selected_rows:
            current_row = selected_rows[-1]
        if current_row is None:
            return False

        self.room_table.blockSignals(True)
        try:
            self.room_table.clearSelection()
            selection = self.room_table.selectionModel()
            for table_row in selected_rows:
                selection.select(
                    self.room_table.model().index(table_row, 0),
                    QItemSelectionModel.Select | QItemSelectionModel.Rows,
                )
            column = min(int(state.get("current_column", 0)), self.room_table.columnCount() - 1)
            self.room_table.setCurrentCell(current_row, column, QItemSelectionModel.NoUpdate)
        finally:
            self.room_table.blockSignals(False)

        vertical = int(state.get("vertical_scroll", 0))
        horizontal = int(state.get("horizontal_scroll", 0))
        self.room_table.verticalScrollBar().setValue(vertical)
        self.room_table.horizontalScrollBar().setValue(horizontal)
        QTimer.singleShot(0, lambda: self.room_table.verticalScrollBar().setValue(vertical))
        QTimer.singleShot(0, lambda: self.room_table.horizontalScrollBar().setValue(horizontal))
        return True

    def _edit_room_category_on_click(self, row: int, column: int):
        if column != 3:
            return
        room_index = self._room_index_for_table_row(row)
        if room_index is None:
            return
        item = self.room_table.item(row, column)
        if item is not None and (item.flags() & Qt.ItemIsEditable):
            self.room_table.editItem(item)

    def _sync_room_category_combo(self, room_index: int):
        # 함수명은 기존 호출부 호환을 위해 유지한다. V23.15부터 상시 콤보는 없다.
        if not (0 <= room_index < len(self.rooms)):
            return
        category = str(self.rooms[room_index].room_category or "")
        table_row = self._table_row_for_room(self.rooms[room_index])
        item = self.room_table.item(table_row, 3) if table_row is not None else None
        if item is not None and item.text() != category:
            self.loading_table = True
            try:
                item.setText(category)
            finally:
                self.loading_table = False

    def _on_room_category_changed(self, room_index: int, category: str):
        if self.loading_table or not (0 <= room_index < len(self.rooms)):
            return
        category = str(category or "").strip()
        if category and category not in ROOM_CATEGORIES:
            return

        room = self.rooms[room_index]
        if str(room.room_category or "") == category:
            self._sync_room_category_combo(room_index)
            return
        room.room_category = category
        self._sync_room_category_combo(room_index)
        self.data_changed.emit()

    def apply_selected_room_fields(self):
        """선택한 실의 층·구분·천정고를 체크 항목별로 일괄 변경한다."""
        current_row = self.room_table.currentRow()
        source_index = self._room_index_for_table_row(current_row)
        if source_index is None or not (0 <= source_index < len(self.rooms)):
            QMessageBox.information(self, "일괄변경", "변경값 기준으로 사용할 실 행에 커서를 놓아 주세요.")
            return

        selected = [index for index in self._selected_room_indices() if 0 <= index < len(self.rooms)]
        selected = list(dict.fromkeys(selected))
        if not selected:
            QMessageBox.information(self, "일괄변경", "변경할 실 행을 하나 이상 선택해 주세요.")
            return

        source = self.rooms[source_index]
        dialog = SelectedRoomsBatchChangeDialog(
            self, source, [self.rooms[index] for index in selected]
        )
        if dialog.exec() != QDialog.Accepted:
            return
        changes = dialog.changes()
        if not changes:
            return

        floor_changed = False
        changed_rooms = 0
        for target_index in selected:
            target = self.rooms[target_index]
            room_changed = False
            if "floor" in changes and target.floor != changes["floor"]:
                target.floor = changes["floor"]
                floor_changed = True
                room_changed = True
            if "category" in changes and str(target.room_category or "") != changes["category"]:
                target.room_category = changes["category"]
                room_changed = True
            if "ceiling_height_m" in changes and target.ceiling_height_m != changes["ceiling_height_m"]:
                target.ceiling_height_m = changes["ceiling_height_m"]
                room_changed = True
            if room_changed:
                target.changed = True
                changed_rooms += 1

        # 층을 변경하면 층별 묶기/정렬 구조가 달라지므로 실 객체 기준으로 목록을 재구성한다.
        # 구분/천정고만 변경한 경우도 선택상태를 보존한 채 한 번 재구성해 표시와 집계를 일치시킨다.
        self.rebuild_views()
        self.data_changed.emit()

        applied_labels = []
        if "floor" in changes:
            applied_labels.append("층")
        if "category" in changes:
            applied_labels.append("구분")
        if "ceiling_height_m" in changes:
            applied_labels.append("천정고")
        QMessageBox.information(
            self, "선택한 목록 일괄변경",
            f"선택 {len(selected):,}개 실에 {', '.join(applied_labels)} 항목을 적용했습니다.\n"
            f"실제 값이 변경된 실: {changed_rooms:,}개",
        )

    def _on_room_table_double_clicked(self, row: int, column: int):
        if self._room_index_for_table_row(row) is None:
            floor = self._floor_for_table_row(row)
            # 층 합계행 천정고 칸은 층높이(H) 입력용이므로 더블클릭해도 접기/펼치지 않는다.
            if floor and int(column) == 7:
                return
            if floor:
                self._set_floor_collapsed(floor, floor not in self.collapsed_floors)

    def _set_floor_collapsed(self, floor: str, collapsed: bool):
        if collapsed:
            self.collapsed_floors.add(floor)
        else:
            self.collapsed_floors.discard(floor)
        header_row = None
        for row, row_floor in enumerate(self.table_row_floors):
            if row_floor != floor:
                continue
            if self._room_index_for_table_row(row) is None:
                header_row = row
                item = self.room_table.item(row, 0)
                if item:
                    item.setText(("▶ " if collapsed else "▼ ") + floor)
            else:
                self.room_table.setRowHidden(row, collapsed)
        if header_row is not None:
            self.room_table.viewport().update()

    def collapse_all_floors(self):
        for floor in {r.floor for r in self.rooms}:
            self._set_floor_collapsed(floor, True)

    def expand_all_floors(self):
        for floor in list(self.collapsed_floors):
            self._set_floor_collapsed(floor, False)

    def open_batch_edit_choice(self):
        floor = self.selected_floor()
        if not floor:
            QMessageBox.information(self, "목록 일괄편집", "편집할 층을 먼저 선택해 주세요.")
            return
        dialog = QDialog(self)
        dialog.setWindowTitle("현재 층 목록 일괄편집")
        dialog.setFixedSize(360, 145)
        layout = QVBoxLayout(dialog)
        layout.addWidget(QLabel(f"{floor} 실 목록을 어떤 방식으로 편집하시겠습니까?"))
        row = QHBoxLayout()
        excel_btn = QPushButton("Excel")
        temp_btn = QPushButton("임시 엑셀")
        cancel_btn = QPushButton("취소")
        row.addWidget(excel_btn); row.addWidget(temp_btn); row.addWidget(cancel_btn)
        layout.addStretch(1); layout.addLayout(row)
        choice = {"value": None}
        def choose(value):
            choice["value"] = value
            dialog.accept()
        excel_btn.clicked.connect(lambda: choose("excel"))
        temp_btn.clicked.connect(lambda: choose("temp"))
        cancel_btn.clicked.connect(dialog.reject)
        if dialog.exec() != QDialog.Accepted:
            return
        if choice["value"] == "excel":
            self.open_excel_batch_editor(floor)
        elif choice["value"] == "temp":
            self.open_internal_batch_editor(floor)

    def _room_entries_for_floor(self, floor: str):
        entries = [(idx, room) for idx, room in enumerate(self.rooms) if room.floor == floor]
        entries.sort(key=lambda pair: (-pair[1].tag_y, pair[1].tag_x, pair[1].room_name.casefold()))
        return entries

    def _apply_batch_room_updates(self, updates):
        changed = 0
        for room_index, room_number, room_name in updates:
            if not (0 <= room_index < len(self.rooms)):
                continue
            room = self.rooms[room_index]
            new_name = normalize_room_name(room_name)
            if room.room_number != room_number or room.room_name != new_name:
                room.room_number = room_number
                room.room_name = new_name
                self._classify_room_if_needed(room)
                room.changed = True
                changed += 1
        if changed:
            self.data_changed.emit()
            self.rebuild_views()
        return changed

    def open_internal_batch_editor(self, floor: str):
        entries = self._room_entries_for_floor(floor)
        if not entries:
            return
        dialog = RoomBatchEditDialog(self, floor, entries)
        if dialog.exec() == QDialog.Accepted:
            changed = self._apply_batch_room_updates(dialog.updates())
            QMessageBox.information(self, "목록 일괄편집", f"{changed}개 실의 No/실명을 반영했습니다.")

    def _excel_temp_path(self, floor: str) -> Path:
        TEMP_PROJECT_DIR.mkdir(parents=True, exist_ok=True)
        safe_floor = re.sub(r'[^0-9A-Za-z가-힣_-]+', '_', floor or '층')
        stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        return TEMP_PROJECT_DIR / f"실목록_{safe_floor}_{stamp}.xlsx"

    def _write_excel_batch_file(self, floor: str, entries, path: Path):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Protection
        wb = Workbook()
        ws = wb.active
        ws.title = "실목록"
        headers = ["_ROOM_ID", "층", "No", "실명", "면적(M2)", "둘레(m)"]
        ws.append(headers)
        for room_index, room in entries:
            ws.append([room_index, room.floor, room.room_number, normalize_room_name(room.room_name), room.area_m2, room.perimeter_m])
        ws.column_dimensions['A'].hidden = True
        ws.column_dimensions['B'].width = 10
        ws.column_dimensions['C'].width = 14
        ws.column_dimensions['D'].width = 28
        ws.column_dimensions['E'].width = 14
        ws.column_dimensions['F'].width = 14
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill('solid', fgColor='D9EAF7')
            cell.alignment = Alignment(horizontal='center')
        for row in range(2, ws.max_row + 1):
            ws.cell(row, 5).number_format = '0.000'
            ws.cell(row, 6).number_format = '0.000'
            for col in (2, 5, 6):
                ws.cell(row, col).fill = PatternFill('solid', fgColor='F2F2F2')
        ws.freeze_panes = 'B2'
        ws.auto_filter.ref = f"B1:F{ws.max_row}"
        wb.save(path)

    def _read_excel_batch_updates(self, path: Path):
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True)
        ws = wb['실목록'] if '실목록' in wb.sheetnames else wb.active
        updates = []
        for row in range(2, ws.max_row + 1):
            room_id = ws.cell(row, 1).value
            if room_id is None:
                continue
            try:
                room_index = int(room_id)
            except Exception:
                continue
            room_number = str(ws.cell(row, 3).value or '').strip()
            room_name = str(ws.cell(row, 4).value or '').strip()
            updates.append((room_index, room_number, room_name))
        wb.close()
        return updates

    def open_excel_batch_editor(self, floor: str):
        entries = self._room_entries_for_floor(floor)
        if not entries:
            return
        path = self._excel_temp_path(floor)
        try:
            self._write_excel_batch_file(floor, entries, path)
            import pythoncom
            import win32com.client
            pythoncom.CoInitialize()
            excel = win32com.client.DispatchEx('Excel.Application')
            excel.Visible = True
            excel.DisplayAlerts = False
            book = excel.Workbooks.Open(str(path))
        except Exception as exc:
            try:
                if path.exists():
                    path.unlink()
            except Exception:
                pass
            answer = QMessageBox.question(
                self, "Microsoft Excel 실행 실패",
                "Microsoft Excel을 실행할 수 없습니다.\n\n"
                f"{exc}\n\n임시 엑셀로 편집하시겠습니까?",
                QMessageBox.Yes | QMessageBox.No, QMessageBox.Yes
            )
            if answer == QMessageBox.Yes:
                self.open_internal_batch_editor(floor)
            return

        controller = QDialog(self)
        controller.setWindowTitle(f"{floor} Excel 편집")
        controller.setFixedSize(390, 155)
        layout = QVBoxLayout(controller)
        layout.addWidget(QLabel("Excel에서 No/실명을 수정한 뒤 아래 버튼을 누르세요.\n행을 정렬해도 숨김 ID로 원래 실과 매칭됩니다."))
        buttons = QHBoxLayout()
        import_btn = QPushButton("수정내용 가져오기")
        cancel_btn = QPushButton("취소")
        buttons.addStretch(1); buttons.addWidget(import_btn); buttons.addWidget(cancel_btn)
        layout.addStretch(1); layout.addLayout(buttons)
        state = {"imported": False}
        def import_changes():
            try:
                book.Save()
                updates = self._read_excel_batch_updates(path)
                changed = self._apply_batch_room_updates(updates)
                state["imported"] = True
                QMessageBox.information(controller, "목록 일괄편집", f"{changed}개 실의 No/실명을 반영했습니다.")
                controller.accept()
            except Exception:
                QMessageBox.critical(controller, "Excel 가져오기 오류", traceback.format_exc())
        import_btn.clicked.connect(import_changes)
        cancel_btn.clicked.connect(controller.reject)
        controller.exec()
        try:
            book.Close(SaveChanges=bool(state["imported"]))
        except Exception:
            pass
        try:
            excel.Quit()
        except Exception:
            pass
        try:
            if path.exists():
                path.unlink()
        except Exception:
            pass


    def set_cad_validation_errors(self, errors: list[str]):
        """CAD 실명블록 재검사 결과를 오류/확인사항에 반영한다."""
        self.cad_validation_errors = list(errors or [])
        self.rebuild_views()

    def update_error_button(self):
        count = len(self.current_combined_errors)
        if count:
            self.error_check_btn.setText(f"⚠ 오류/확인사항 ({count})")
            self.error_check_btn.setStyleSheet(
                "QPushButton { background:#ffe5e5; color:#a00000; border:1px solid #d96b6b; "
                "padding:4px 10px; font-weight:bold; }"
                "QPushButton:hover { background:#ffd3d3; }"
            )
        else:
            self.error_check_btn.setText("오류/확인사항")
            self.error_check_btn.setStyleSheet("")

    def show_error_dialog(self):
        dialog = QDialog(self)
        dialog.setWindowTitle("오류 / 확인사항")
        dialog.resize(760, 500)
        layout = QVBoxLayout(dialog)
        table = QTableWidget(0, 1)
        table.setHorizontalHeaderLabels(["오류 / 확인사항"])
        table.horizontalHeader().setSectionResizeMode(0, QHeaderView.Stretch)
        table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        table.setAlternatingRowColors(True)
        if self.current_combined_errors:
            table.setRowCount(len(self.current_combined_errors))
            for row, msg in enumerate(self.current_combined_errors):
                table.setItem(row, 0, QTableWidgetItem(msg))
        else:
            table.setRowCount(1)
            table.setItem(0, 0, QTableWidgetItem("오류 또는 확인사항이 없습니다."))
        layout.addWidget(table)
        close_btn = QPushButton("닫기")
        close_btn.clicked.connect(dialog.accept)
        buttons = QHBoxLayout()
        buttons.addStretch(1)
        buttons.addWidget(close_btn)
        layout.addLayout(buttons)
        dialog.exec()

    def clear_info(self):
        for label in [
            self.floor_label, self.number_label, self.name_label,
            self.perimeter_label, self.area_label,
            self.boundary_handle_label, self.tag_handle_label, self.closed_label
        ]:
            label.setText("-")

    def selected_room(self) -> RoomRecord | None:
        # V19.0: selectionModel().selectedRows()[0]는 이전 선택행을 먼저 반환할 수 있다.
        # 사용자가 방금 클릭한 현재 행(currentRow)을 기준으로 실/층을 판정한다.
        row = self.room_table.currentRow()
        if row >= 0:
            room_index = self._room_index_for_table_row(row)
            if room_index is not None and 0 <= room_index < len(self.rooms):
                return self.rooms[room_index]
        return None

    def selected_floor(self) -> str | None:
        # 층 헤더를 눌렀을 때도 반드시 현재 행의 층을 사용한다.
        row = self.room_table.currentRow()
        if row >= 0:
            floor = self._floor_for_table_row(row)
            if floor:
                return floor
        room = self.selected_room()
        if room:
            return room.floor
        return self.rooms[0].floor if self.rooms else None

    def on_preview_room_clicked(self, room: RoomRecord):
        if self._direct_number_mode:
            if self._direct_number_floor and room.floor != self._direct_number_floor:
                return
            self._assign_direct_number(room)
            return
        if room.floor in self.collapsed_floors:
            self._set_floor_collapsed(room.floor, False)
        row = self._table_row_for_room(room)
        if row is None:
            return

        self.room_table.clearSelection()
        self.room_table.selectRow(row)
        self.room_table.scrollToItem(
            self.room_table.item(row, 0),
            QAbstractItemView.PositionAtCenter
        )
        self.show_room_info(room)
        self.refresh_preview()

    def on_preview_empty_clicked(self):
        """도곽 빈 공간 클릭 시 현재 층의 '전체' 행으로 선택을 되돌린다."""
        if self._direct_number_mode or self.preview.wall_finish_mode or self.preview.room_group_mode:
            return
        if not self.rooms or self.room_table.rowCount() <= 0:
            return

        current_room = self.selected_room()
        floor = current_room.floor if current_room is not None else self.selected_floor()
        target_row = 0
        if self.group_by_floor and floor:
            for row in range(self.room_table.rowCount()):
                if self._room_index_for_table_row(row) is None and self._floor_for_table_row(row) == floor:
                    target_row = row
                    break

        # selectionChanged에 의한 중복 refresh를 피하고 한 번만 화면/하단정보를 갱신한다.
        self.room_table.blockSignals(True)
        try:
            self.room_table.clearSelection()
            self.room_table.setCurrentCell(target_row, 0)
            self.room_table.selectRow(target_row)
        finally:
            self.room_table.blockSignals(False)

        if target_row == 0:
            self.show_all_floors_info()
        else:
            self.show_floor_info(str(floor or ""))
        self.refresh_preview()

    def on_preview_room_edit_requested(self, room: RoomRecord):
        """도곽 우클릭에서 실번호/실명을 빠르게 수정한다."""
        if room is None:
            return
        # 우클릭 대상이 현재 표 선택과 다르면 먼저 같은 실을 선택 상태로 맞춘다.
        row = self._table_row_for_room(room)
        if row is not None:
            self.room_table.clearSelection()
            self.room_table.selectRow(row)
            item = self.room_table.item(row, 0)
            if item is not None:
                self.room_table.scrollToItem(item, QAbstractItemView.PositionAtCenter)

        dialog = QDialog(self)
        dialog.setWindowTitle(f"실 정보 수정 - {room.floor}")
        layout = QVBoxLayout(dialog)
        form = QFormLayout()
        number_edit = QLineEdit(room.room_number or "")
        name_edit = QLineEdit("" if room_name_is_missing(room.room_name) else (room.room_name or ""))
        form.addRow("실번호", number_edit)
        form.addRow("실명", name_edit)
        layout.addLayout(form)
        buttons = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        buttons.accepted.connect(dialog.accept)
        buttons.rejected.connect(dialog.reject)
        layout.addWidget(buttons)
        number_edit.selectAll()
        number_edit.setFocus()
        if dialog.exec() != QDialog.Accepted:
            return

        new_number = number_edit.text().strip()
        new_name = normalize_room_name(name_edit.text())
        if room.room_number == new_number and room.room_name == new_name:
            return
        room.room_number = new_number
        room.room_name = new_name
        self._classify_room_if_needed(room)
        room.changed = True

        # 표의 같은 행만 갱신한다. itemChanged 재진입은 막고 기존 CAD 연동 상태를 유지한다.
        if row is not None:
            self.loading_table = True
            try:
                number_item = self.room_table.item(row, 1)
                name_item = self.room_table.item(row, 2)
                if number_item is not None:
                    number_item.setText(room.room_number)
                if name_item is not None:
                    name_item.setText(room.room_name)
            finally:
                self.loading_table = False
        room_index = next((idx for idx, candidate in enumerate(self.rooms) if candidate is room), -1)
        self._sync_room_category_combo(room_index)

        self.show_room_info(room)
        if not self.preview.update_room_label(room):
            self.refresh_preview()

        number_errors = validate_room_numbers(self.rooms)
        combined_errors = list(self.errors)
        for msg in number_errors:
            if msg not in combined_errors:
                combined_errors.append(msg)
        for msg in self.cad_validation_errors:
            if msg not in combined_errors:
                combined_errors.append(msg)
        for msg in self.room_wall_validation_errors:
            if msg not in combined_errors:
                combined_errors.append(msg)
        self.current_combined_errors = combined_errors
        self.update_error_button()
        self.data_changed.emit()

    def on_table_select(self):
        # V19.0: 다중/과거 selection 목록의 첫 행이 아니라 방금 클릭한 currentRow 사용.
        table_row = self.room_table.currentRow()
        if table_row < 0:
            return
        room_index = self._room_index_for_table_row(table_row)
        if room_index is None:
            floor = self._floor_for_table_row(table_row)
            if floor:
                self.show_floor_info(floor)
                self.finish_panel.set_selected_room_count(
                    self._quantity_room_count([r for r in self.rooms if r.floor == floor])
                )
                self.refresh_preview()
            elif table_row == 0:
                self.show_all_floors_info()
                self.refresh_preview()
            return
        if not (0 <= room_index < len(self.rooms)):
            return
        room = self.rooms[room_index]
        self.show_room_info(room)
        self.refresh_preview()

    def show_room_info(self, room: RoomRecord):
        self.floor_label.setText(f"{room.floor}  (적용 {room.floor_scope or room.floor} / {max(1, int(room.repeat_count or 1))}개층)")
        self.number_label.setText(room.room_number or "(미입력)")
        self.name_label.setText(normalize_room_name(room.room_name))
        self.perimeter_label.setText(f"{room.perimeter_m:,.3f} m")
        self.area_label.setText(f"{room.area_m2:,.3f} M2")
        self.boundary_handle_label.setText(room.boundary_handle)
        self.tag_handle_label.setText(room.tag_handle)
        self.closed_label.setText("예" if room.is_closed else "아니오")
        self.finish_panel.set_selected_room_count(len(self.selected_target_rooms()))
        self.finish_panel.set_group_wall_lengths(self._wall_group_lengths_for_room(room))
        self.finish_panel.set_current_room(
            room, self.room_finish_assignments.get(room.boundary_handle, {})
        )
        self.update_finish_summary(room.floor)

    def show_floor_info(self, floor: str):
        floor_rooms = [room for room in self.rooms if room.floor == floor]
        self.floor_label.setText(floor)
        self.number_label.setText(f"실 {self._quantity_room_count(floor_rooms):,}개")
        self.name_label.setText("층 합계")
        self.perimeter_label.setText(f"{sum(r.quantity_perimeter_m for r in floor_rooms):,.3f} m")
        self.area_label.setText(f"{sum(r.quantity_area_m2 for r in floor_rooms):,.3f} M2")
        self.boundary_handle_label.setText("-")
        self.tag_handle_label.setText("-")
        self.closed_label.setText("-")
        self.finish_panel.set_selected_room_count(self._quantity_room_count(floor_rooms))
        self.finish_panel.set_group_catalog(self._finish_groups_for_scope(floor), f"{floor} 그룹 목록")
        self.update_finish_summary(floor)

    def show_all_floors_info(self):
        """전층 대구분 행 선택 시 하단 정보에도 동일한 산출 합계를 표시한다."""
        self.floor_label.setText("전층")
        self.number_label.setText(f"실 {self._quantity_room_count(self.rooms):,}개")
        self.name_label.setText("전층 대구분")
        self.perimeter_label.setText(f"{sum(r.quantity_perimeter_m for r in self.rooms):,.3f} m")
        self.area_label.setText(f"{sum(r.quantity_area_m2 for r in self.rooms):,.3f} M2")
        self.boundary_handle_label.setText("-")
        self.tag_handle_label.setText("-")
        self.closed_label.setText("-")
        self.finish_panel.set_selected_room_count(self._quantity_room_count(self.rooms))
        self.finish_panel.set_group_catalog(self._finish_groups_for_scope(None))
        self.update_finish_summary(None)

    def _wall_group_lengths_for_room(self, room: RoomRecord) -> dict[str, float]:
        result: dict[str, float] = {}
        handle = str(room.boundary_handle or "")
        for group_id, group in self.finish_groups.items():
            length = 0.0
            for key in group.get("wall_keys", []) or []:
                if wall_key_room_handle(key) != handle:
                    continue
                length += wall_key_length_m(room, key)
            if length:
                result[str(group_id)] = length
        return result

    @staticmethod
    def _finish_material_identity(material: dict):
        material_id = str(material.get("id", "") or "").strip()
        if material_id:
            return ("id", material_id)
        return (
            "fallback",
            str(material.get("code", "") or "").strip(),
            str(material.get("name", "") or "").strip(),
            str(material.get("spec", "") or "").strip(),
            str(material.get("unit", "") or "").strip(),
        )

    @staticmethod
    def _same_finish_material(left: dict, right: dict) -> bool:
        return RoomsTab._finish_material_identity(left) == RoomsTab._finish_material_identity(right)

    @staticmethod
    def _assignment_group_formula(assignment: dict, group_id: str) -> tuple[bool, str]:
        """실별 그룹 산식을 반환한다. 공란도 명시값이므로 존재 여부를 함께 돌려준다."""
        group_id = str(group_id or "")
        normalized = normalize_room_finish_assignment(assignment)
        for entry in normalized.get("groups", []):
            if str(entry.get("id", "") or "") == group_id:
                return True, normalize_formula_text(entry.get("formula", ""))
        return False, ""

    @staticmethod
    def _set_assignment_group_formula(assignment: dict, group_id: str, formula: str) -> tuple[dict, bool]:
        """한 실의 그룹행/그룹재료 사본에 같은 그룹 산식을 저장한다."""
        group_id = str(group_id or "")
        formula = normalize_formula_text(formula)
        normalized = normalize_room_finish_assignment(assignment)
        found = False
        for entry in normalized.get("groups", []):
            if str(entry.get("id", "") or "") == group_id:
                entry["formula"] = formula
                found = True
        for materials in normalized.get("items", {}).values():
            for material in materials:
                if str(material.get("group_id", "") or "") == group_id:
                    material["group_formula"] = formula
        return normalized, found

    def _group_default_formula_for_floor(self, group: dict, floor: str | None) -> str:
        """새로 그룹에 연결되는 실이 상속할 범위 기준 산식을 구한다."""
        scope_formulas = group.get("scope_formulas", {})
        if not isinstance(scope_formulas, dict):
            scope_formulas = {}
        floor_key = str(floor or "")
        if floor_key and floor_key in scope_formulas:
            return normalize_formula_text(scope_formulas.get(floor_key, ""))
        if "__all__" in scope_formulas:
            return normalize_formula_text(scope_formulas.get("__all__", ""))
        return normalize_formula_text(group.get("formula", ""))

    def _room_floor_for_handle(self, handle: str) -> str:
        handle = str(handle or "")
        for room in self.rooms:
            if str(room.boundary_handle or "") == handle:
                return str(room.floor or "")
        return ""

    def _finish_groups_for_scope(self, floor: str | None) -> list[dict]:
        """층/전층 그룹목록용 복사본을 만든다.

        재료 산식과 그룹 산식 모두 범위 기준값을 우선 표시한다. 별도 기준값이
        없을 때만 실제 해당 범위 실들의 저장값을 확인해 모두 같으면 표시하고,
        서로 다르면 공란으로 표시한다.
        """
        target_rooms = self.rooms if floor is None else [room for room in self.rooms if room.floor == floor]
        handles = {str(room.boundary_handle or "") for room in target_rooms if str(room.boundary_handle or "")}
        source_groups = [
            group for group in self.finish_groups.values()
            if floor is None
            or handles.intersection(str(value) for value in group.get("room_handles", []) or [])
            or any(wall_key_room_handle(key) in handles for key in group.get("wall_keys", []) or [])
        ]
        formula_map: dict[tuple[str, object], set[str]] = defaultdict(set)
        group_formula_map: dict[str, set[str]] = defaultdict(set)
        for handle in handles:
            assignment = normalize_room_finish_assignment(self.room_finish_assignments.get(handle, {}))
            for entry in assignment.get("groups", []):
                group_id = str(entry.get("id", "") or "")
                if group_id:
                    group_formula_map[group_id].add(normalize_formula_text(entry.get("formula", "")))
            for values in assignment.get("items", {}).values():
                for material in values:
                    group_id = str(material.get("group_id", "") or "")
                    if not group_id:
                        continue
                    formula_map[(group_id, self._finish_material_identity(material))].add(
                        normalize_formula_text(material.get("method", ""))
                    )

        result = []
        scope_key = "__all__" if floor is None else str(floor)
        for raw_group in source_groups:
            group = deepcopy(raw_group)
            group_id = str(group.get("id", "") or "")
            scope_formulas = raw_group.get("scope_formulas", {})
            if isinstance(scope_formulas, dict) and scope_key in scope_formulas:
                group["formula"] = normalize_formula_text(scope_formulas.get(scope_key, ""))
            else:
                unique_group_formulas = group_formula_map.get(group_id, set())
                if len(unique_group_formulas) == 1:
                    group["formula"] = next(iter(unique_group_formulas))
                elif len(unique_group_formulas) > 1:
                    group["formula"] = ""
                else:
                    group["formula"] = self._group_default_formula_for_floor(raw_group, floor)
            for material in group.get("materials", []) or []:
                # 층-전체/전층에서 직접 입력한 산식은 해당 집계행의 기준 산식이다.
                # 개별 실에 덮어쓸지와 무관하게 집계행 자체에는 먼저 저장/표시한다.
                scope_methods = material.get("scope_methods", {})
                if isinstance(scope_methods, dict) and scope_key in scope_methods:
                    material["method"] = normalize_formula_text(scope_methods.get(scope_key, ""))
                    continue
                unique = formula_map.get((group_id, self._finish_material_identity(material)), set())
                if len(unique) == 1:
                    material["method"] = next(iter(unique))
                elif len(unique) > 1:
                    material["method"] = ""
            result.append(group)
        return result

    def _finish_groups_for_floor(self, floor: str) -> list[dict]:
        return self._finish_groups_for_scope(floor)

    def refresh_preview(self):
        room = self.selected_room()
        floor = room.floor if room else self.selected_floor()
        # 표에서 실제 실 행을 다중선택한 경우 현재 표시층에 속한 선택실을 모두 강조한다.
        # 층/전층 합계행 선택은 기존처럼 전체보기 상태를 유지한다.
        selected_indices = self._selected_room_indices()
        selected_rooms = [self.rooms[index] for index in selected_indices if 0 <= index < len(self.rooms)]
        selected_on_floor = [target for target in selected_rooms if not floor or target.floor == floor]

        if self.view_mode.currentText() == "선택 실만 보기":
            scope_rooms = [r for r in self.rooms if r.floor == floor] if floor else []
            self._prepare_opening_read_preview(str(floor or ""))
            self.preview.show_selected_scope(scope_rooms, room, selected_on_floor, show_edge_lengths=self.dimension_show_check.isChecked())
            return

        floor = room.floor if room else self.selected_floor()
        if not floor and self.rooms:
            floor = self.rooms[0].floor
        floor_rooms = [r for r in self.rooms if r.floor == floor]
        selected_on_floor = [target for target in selected_rooms if target.floor == floor]
        self._prepare_opening_read_preview(str(floor or ""))
        if not floor_rooms and floor and (floor in self.preview.floor_background_files or floor in self.preview.floor_vector_cache):
            self.preview.show_background_floor(floor)
        else:
            self.preview.show_floor(floor_rooms, room, selected_on_floor)

    def on_table_changed(self, item: QTableWidgetItem):
        if self.loading_table or item.column() not in (0, 1, 2, 7):
            return
        row = item.row()
        room_index = self._room_index_for_table_row(row)
        if room_index is None:
            # 층 합계행의 천정고 열은 개별 천정고가 아니라 층높이(H) 입력칸으로 사용한다.
            floor = self._floor_for_table_row(row)
            if item.column() == 7 and floor:
                new_value = item.text().strip()
                try:
                    floor_height = None if not new_value else float(new_value.replace(",", ""))
                    if floor_height is not None and floor_height < 0:
                        raise ValueError
                except Exception:
                    QMessageBox.warning(self, "층높이 입력", "층높이는 0 이상의 숫자(m) 또는 공란으로 입력해 주세요.")
                    floor_rooms = [room for room in self.rooms if room.floor == floor]
                    self.loading_table = True
                    try:
                        item.setText(self._floor_height_text(floor_rooms))
                    finally:
                        self.loading_table = False
                    return
                for target in self.rooms:
                    if target.floor == floor:
                        target.floor_height_m = floor_height
                        target.changed = True
                canonical = self._format_ceiling_height(floor_height)
                if item.text() != canonical:
                    self.loading_table = True
                    try:
                        item.setText(canonical)
                    finally:
                        self.loading_table = False
                # 산식 값만 갱신하고 도면 장면은 다시 만들지 않는다.
                self._refresh_current_finish_panel()
                self.update_finish_summary(self.selected_floor())
                self.data_changed.emit()
            return
        if not (0 <= room_index < len(self.rooms)):
            return

        room = self.rooms[room_index]
        new_value = item.text().strip()
        if item.column() == 0:
            room.floor = new_value or "(층 미입력)"
        elif item.column() == 1:
            room.room_number = new_value
        elif item.column() == 2:
            room.room_name = normalize_room_name(new_value)
            if item.text() != room.room_name:
                self.loading_table = True
                try:
                    item.setText(room.room_name)
                finally:
                    self.loading_table = False
            self._classify_room_if_needed(room)
            self._sync_room_category_combo(room_index)
        elif item.column() == 7:
            try:
                ceiling_height = None if not new_value else float(new_value.replace(",", ""))
                if ceiling_height is not None and ceiling_height < 0:
                    raise ValueError
            except Exception:
                QMessageBox.warning(self, "천정고 입력", "천정고는 0 이상의 숫자(m) 또는 공란으로 입력해 주세요.")
                self.loading_table = True
                try:
                    item.setText(self._format_ceiling_height(room.ceiling_height_m))
                finally:
                    self.loading_table = False
                return
            room.ceiling_height_m = ceiling_height
            canonical = self._format_ceiling_height(ceiling_height)
            if item.text() != canonical:
                self.loading_table = True
                try:
                    item.setText(canonical)
                finally:
                    self.loading_table = False
        room.changed = True

        # 실정보를 수정해도 현재 CAD 연동 모드는 유지한다.
        # 실번호/실명은 해당 행과 화면 라벨만 부분 갱신한다.
        # 층 변경은 층 그룹 구조가 달라지므로 그때만 목록을 다시 구성한다.
        if item.column() == 0:
            self.data_changed.emit()
            self.rebuild_views()
            return

        # 선택 정보와 미리보기 라벨만 즉시 갱신한다. CAD 조회/배경 재생성은 하지 않는다.
        selected = self.selected_room()
        if selected is room:
            self.show_room_info(room)
        if not self.preview.update_room_label(room):
            self.refresh_preview()

        # 번호 중복 등 가벼운 프로그램 내부 검사만 갱신한다.
        number_errors = validate_room_numbers(self.rooms)
        combined_errors = list(self.errors)
        for msg in number_errors:
            if msg not in combined_errors:
                combined_errors.append(msg)
        for msg in self.cad_validation_errors:
            if msg not in combined_errors:
                combined_errors.append(msg)
        for msg in self.room_wall_validation_errors:
            if msg not in combined_errors:
                combined_errors.append(msg)
        self.current_combined_errors = combined_errors
        self.update_error_button()
        self.data_changed.emit()

    @staticmethod
    def _normalize_number_entry_settings(values: dict | None = None) -> dict:
        values = dict(values or {})
        try:
            digits = int(values.get("digits", 3))
        except Exception:
            digits = 3
        return {
            "digits": max(1, min(5, digits)),
            "floor_pattern": str(values.get("floor_pattern", "") or ""),
            "separator": str(values.get("separator", "-") if values.get("separator", "-") is not None else "-"),
            "suffix": str(values.get("suffix", "") or ""),
        }

    def _floor_number_prefix(self, floor: str) -> str:
        raw = (floor or "").strip().upper().replace(" ", "")
        if raw.endswith("F") and raw not in {"RF", "PHRF"}:
            raw = raw[:-1]
        return raw

    def _floor_display_from_pattern(self, floor: str, pattern: str) -> str:
        """번호 앞 층표시 문자를 만든다. 패턴의 0이 실제 층값 자리다.

        예) 0F -> B2F, 2F / 0층 -> B2층, 1층
        빈 문자열이면 층표시를 사용하지 않는다.
        """
        pattern = str(pattern or "").strip()
        if not pattern:
            return ""
        floor_token = self._floor_number_prefix(floor)
        if "0" in pattern:
            return pattern.replace("0", floor_token)
        return pattern

    def _format_room_number(self, room: RoomRecord, seq: int, settings: dict) -> str:
        digits = max(1, min(5, int(settings.get("digits", 3))))
        number = str(max(0, int(seq))).zfill(digits)
        prefix = self._floor_display_from_pattern(room.floor, settings.get("floor_pattern", ""))
        if prefix:
            number = f"{prefix}{settings.get('separator', '-')}{number}"
        suffix = str(settings.get("suffix", "") or "")
        return number + suffix

    def _next_room_number_sequence(self, floor: str, settings: dict) -> int:
        """현재층 기존 실번호의 마지막 숫자 중 최댓값 다음 번호를 반환한다."""
        suffix = str(settings.get("suffix", "") or "")
        maximum = 0
        for room in self.rooms:
            if room.floor != floor:
                continue
            text = str(room.room_number or "").strip()
            if not text:
                continue
            if suffix and text.endswith(suffix):
                text = text[:-len(suffix)]
            matches = re.findall(r"\d+", text)
            if not matches:
                continue
            try:
                maximum = max(maximum, int(matches[-1]))
            except ValueError:
                pass
        return max(1, maximum + 1)

    def _rooms_in_current_display_order(self, floor: str | None = None) -> list[RoomRecord]:
        result = []
        seen = set()
        for row in range(self.room_table.rowCount()):
            idx = self._room_index_for_table_row(row)
            if idx is None or not (0 <= idx < len(self.rooms)):
                continue
            room = self.rooms[idx]
            if floor is not None and room.floor != floor:
                continue
            if id(room) in seen:
                continue
            result.append(room)
            seen.add(id(room))
        return result

    def _rooms_in_drawing_order(self, floor: str) -> list[RoomRecord]:
        rooms = [r for r in self.rooms if r.floor == floor]

        def room_position(room: RoomRecord):
            # 실명블록 위치가 아니라 실구분선 자체의 배치 위치를 사용한다.
            # 사용자가 CAD에서 실명블록을 옮겨도 자동번호 순서는 바뀌지 않는다.
            pts = list(getattr(room, "points", None) or [])
            if pts:
                xs = [float(p[0]) for p in pts]
                ys = [float(p[1]) for p in pts]
                return ((min(xs) + max(xs)) / 2.0, (min(ys) + max(ys)) / 2.0)
            return (float(getattr(room, "tag_x", 0.0) or 0.0), float(getattr(room, "tag_y", 0.0) or 0.0))

        # 도곽상 좌상단 → 우하단: 높은 Y 우선, 같은 높이에서는 작은 X 우선.
        rooms.sort(key=lambda r: (-room_position(r)[1], room_position(r)[0]))
        return rooms

    def open_number_entry_dialog(self):
        if not self.rooms:
            return
        floor = self.selected_floor()
        if not floor:
            QMessageBox.information(self, "번호 기입", "번호를 기입할 층을 먼저 선택해 주세요.")
            return

        dialog = QDialog(self)
        dialog.setWindowTitle("번호 기입")
        dialog.setFixedWidth(430)
        layout = QVBoxLayout(dialog)
        form = QGridLayout()
        saved_number_settings = self._normalize_number_entry_settings(self.number_entry_settings)
        digits = QSpinBox(); digits.setRange(1, 5); digits.setValue(saved_number_settings["digits"])
        floor_pattern = QLineEdit(saved_number_settings["floor_pattern"])
        floor_pattern.setPlaceholderText("예 : 0F = B2F, 2F , 2~14F")
        separator = QLineEdit(saved_number_settings["separator"]); separator.setMaximumWidth(100)
        suffix = QLineEdit(saved_number_settings["suffix"]); suffix.setPlaceholderText("예: (임시)")
        form.addWidget(QLabel("번호 자리수"), 0, 0); form.addWidget(digits, 0, 1)
        form.addWidget(QLabel("층표시 문자"), 1, 0); form.addWidget(floor_pattern, 1, 1)
        form.addWidget(QLabel("구분기호"), 2, 0); form.addWidget(separator, 2, 1)
        form.addWidget(QLabel("뒤에 추가할 문자"), 3, 0); form.addWidget(suffix, 3, 1)
        layout.addLayout(form)
        example = QLabel()
        example.setStyleSheet("color:#666666;")
        layout.addWidget(example)
        def refresh_example():
            current = {"digits": digits.value(), "floor_pattern": floor_pattern.text(), "separator": separator.text(), "suffix": suffix.text()}
            digits_text = str(1).zfill(current["digits"])
            prefix = self._floor_display_from_pattern(floor, current["floor_pattern"])
            example.setText("예시: " + (prefix + current["separator"] if prefix else "") + digits_text + current["suffix"])
        digits.valueChanged.connect(refresh_example); floor_pattern.textChanged.connect(refresh_example); separator.textChanged.connect(refresh_example); suffix.textChanged.connect(refresh_example)
        refresh_example()

        layout.addWidget(QLabel("실행"))
        buttons = QHBoxLayout()
        direct_btn = QPushButton("직접입력\n(현재층)")
        auto_floor_btn = QPushButton("자동기입\n(현재층)")
        auto_all_btn = QPushButton("자동기입\n(전층)")
        buttons.addWidget(direct_btn); buttons.addWidget(auto_floor_btn); buttons.addWidget(auto_all_btn)
        layout.addLayout(buttons)
        cancel = QPushButton("닫기"); layout.addWidget(cancel, 0, Qt.AlignRight)

        def settings():
            return self._normalize_number_entry_settings({
                "digits": digits.value(),
                "floor_pattern": floor_pattern.text(),
                "separator": separator.text(),
                "suffix": suffix.text(),
            })

        def persist_settings(_result=None):
            current = settings()
            self.number_entry_settings = dict(current)
            self.number_entry_settings_changed.emit(dict(current))

        def run_direct():
            current = settings()
            start_dialog = QDialog(dialog)
            start_dialog.setWindowTitle("직접입력 시작 번호")
            start_layout = QVBoxLayout(start_dialog)
            start_form = QFormLayout()
            start_number_spin = QSpinBox()
            start_number_spin.setRange(1, 99999)
            start_number_spin.setValue(self._next_room_number_sequence(floor, current))
            write_name_check = QCheckBox("실명 작성")
            write_name_check.setToolTip("실을 클릭해 번호를 넣을 때 실명 입력창도 함께 표시합니다.")
            start_form.addRow("시작할 번호", start_number_spin)
            start_form.addRow("", write_name_check)
            start_layout.addLayout(start_form)
            start_buttons = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
            start_buttons.accepted.connect(start_dialog.accept)
            start_buttons.rejected.connect(start_dialog.reject)
            start_layout.addWidget(start_buttons)
            start_number_spin.selectAll()
            start_number_spin.setFocus()
            if start_dialog.exec() != QDialog.Accepted:
                return
            dialog.accept()
            self.start_direct_number_mode(
                floor, current, start_number_spin.value(), write_name_check.isChecked()
            )

        def run_auto_floor():
            current = settings()
            dialog.accept()
            self.auto_number_scope(current, floor)

        def run_auto_all():
            current = settings()
            dialog.accept()
            self.auto_number_scope(current, None)

        direct_btn.clicked.connect(run_direct)
        auto_floor_btn.clicked.connect(run_auto_floor)
        auto_all_btn.clicked.connect(run_auto_all)
        cancel.clicked.connect(dialog.reject)
        # 실행/닫기/X 어느 경로로 창을 닫아도 마지막 입력값을 settings2.ini 쪽으로 전달한다.
        dialog.finished.connect(persist_settings)
        dialog.exec()

    def start_direct_number_mode(self, floor: str, settings: dict, start_number: int = 1, write_room_name: bool = False):
        self._direct_number_mode = True
        self._direct_number_settings = dict(settings)
        self._direct_number_next = max(1, int(start_number))
        self._direct_number_floor = floor
        self._direct_number_write_name = bool(write_room_name)
        self.view_mode.setCurrentText("층 전체 보기")
        # 현재 층을 확실히 표시한다.
        for row in range(self.room_table.rowCount()):
            if self._floor_for_table_row(row) == floor:
                self.room_table.selectRow(row)
                break
        self.refresh_preview()
        self.preview.setFocus()
        action_text = "번호와 실명" if self._direct_number_write_name else "번호"
        self.preview_help.setText(f"{action_text} 직접기입 중 ({floor}) : 실을 클릭하면 {self._direct_number_next}번부터 순차 기입 | ESC: 종료")

    def cancel_direct_number_mode(self):
        if not self._direct_number_mode:
            return
        self._direct_number_mode = False
        self._direct_number_settings = None
        self._direct_number_floor = None
        self._direct_number_write_name = False
        self.preview_help.setText("마우스 휠: 확대·축소 | 일반모드 드래그: 화면 이동 | 벽 선택모드 드래그: 박스 선택 | 가운데버튼 더블클릭: 전체보기")
        # ESC로 직접입력 모드를 종료했음을 커서 근처에서 즉시 확인할 수 있게 한다.
        QToolTip.showText(QCursor.pos(), "번호 직접입력 모드 종료", self.preview)

    def _assign_direct_number(self, room: RoomRecord):
        settings = self._direct_number_settings or {"digits": 3, "floor_pattern": "", "separator": "-", "suffix": ""}
        room.room_number = self._format_room_number(room, self._direct_number_next, settings)
        room.changed = True
        self._direct_number_next += 1

        if self._direct_number_write_name:
            name_dialog = QInputDialog(self)
            name_dialog.setWindowTitle(f"실명 작성 - {room.room_number}")
            name_dialog.setLabelText(f"현재 실: {room.room_number}\n실명을 입력하세요.")
            name_dialog.setTextValue("" if room_name_is_missing(room.room_name) else room.room_name)
            name_dialog.setOkButtonText("입력")
            name_dialog.setCancelButtonText("건너뛰기")
            name_dialog.adjustSize()
            pos = self.mapToGlobal(self.rect().bottomLeft())
            name_dialog.move(pos.x() + 20, max(40, pos.y() - name_dialog.height() - 80))
            if name_dialog.exec() == QDialog.Accepted:
                room.room_name = normalize_room_name(name_dialog.textValue())
                self._classify_room_if_needed(room)

        row = self._table_row_for_room(room)
        if row is not None:
            number_item = self.room_table.item(row, 1)
            name_item = self.room_table.item(row, 2)
            self.loading_table = True
            try:
                if number_item:
                    number_item.setText(room.room_number)
                if name_item:
                    name_item.setText(room.room_name)
            finally:
                self.loading_table = False
        self._sync_room_category_combo(next((idx for idx, candidate in enumerate(self.rooms) if candidate is room), -1))
        self.preview.update_room_label(room)
        self.show_room_info(room)
        self.data_changed.emit()
        action_text = "번호/실명" if self._direct_number_write_name else "번호"
        self.preview_help.setText(f"{action_text} 직접기입 중 ({room.floor}) : {room.room_number} 입력 완료, 다음 실 클릭 | ESC: 종료")

    def auto_number_scope(self, settings: dict, floor: str | None):
        floors = [floor] if floor else sorted({r.floor for r in self.rooms}, key=floor_sort_key)
        changed = 0
        for current_floor in floors:
            seq = 1
            floor_rooms = self._rooms_in_drawing_order(current_floor)
            used_numbers = {
                str(room.room_number or "").strip()
                for room in floor_rooms
                if str(room.room_number or "").strip()
            }
            for room in floor_rooms:
                # 자동기입은 사용자가 이미 입력했거나 CAD에서 읽은 기존 번호를 절대 덮어쓰지 않는다.
                if str(room.room_number or "").strip():
                    seq += 1
                    continue
                new_number = self._format_room_number(room, seq, settings)
                while new_number in used_numbers:
                    seq += 1
                    new_number = self._format_room_number(room, seq, settings)
                room.room_number = new_number
                room.changed = True
                used_numbers.add(new_number)
                changed += 1
                seq += 1
        self.data_changed.emit()
        self.rebuild_views()
        QMessageBox.information(self, "번호 기입", f"{changed}개 실의 번호를 기입했습니다.")

    def start_direct_name_input(self):
        if self._name_input_mode or not self.rooms:
            return
        floor = self.selected_floor()
        if not floor:
            QMessageBox.information(self, "실명 직접입력", "입력할 층을 먼저 선택해 주세요.")
            return
        rooms = self._rooms_in_current_display_order(floor)
        if not rooms:
            return
        self._name_input_mode = True
        try:
            for room in rooms:
                row = self._table_row_for_room(room)
                if row is not None:
                    self.room_table.clearSelection(); self.room_table.selectRow(row)
                    self.room_table.scrollToItem(self.room_table.item(row, 0), QAbstractItemView.PositionAtCenter)
                self.show_room_info(room)
                self.refresh_preview()  # 현재 입력 대상을 선택색으로 표시
                QApplication.processEvents()

                dlg = QInputDialog(self)
                dlg.setWindowTitle(f"실명 직접입력 - {floor}")
                dlg.setLabelText(f"현재 실: {room.room_number or '(번호 미입력)'}\n실명을 입력하세요.  ESC: 입력모드 종료")
                dlg.setTextValue("" if room_name_is_missing(room.room_name) else room.room_name)
                dlg.setOkButtonText("입력")
                dlg.setCancelButtonText("종료")
                dlg.adjustSize()
                # 도곽을 가리지 않도록 프로그램 좌측 하단 쪽에 배치한다.
                pos = self.mapToGlobal(self.rect().bottomLeft())
                dlg.move(pos.x() + 20, max(40, pos.y() - dlg.height() - 80))
                if dlg.exec() != QDialog.Accepted:
                    break
                name = dlg.textValue().strip()
                room.room_name = normalize_room_name(name)
                self._classify_room_if_needed(room)
                room.changed = True
                if row is not None:
                    item = self.room_table.item(row, 2)
                    if item:
                        self.loading_table = True; item.setText(room.room_name); self.loading_table = False
                self._sync_room_category_combo(next((idx for idx, candidate in enumerate(self.rooms) if candidate is room), -1))
                self.preview.update_room_label(room)
                self.data_changed.emit()
        finally:
            self._name_input_mode = False
            self.refresh_preview()

    def set_cad_connected_mode(self, enabled: bool):
        self.cad_connected_mode = bool(enabled)
        self.save_dxf_btn.setEnabled(bool(self.current_source_path) and not self.cad_connected_mode)
        self.opening_test_btn.setEnabled(bool(self.rooms))
        if self._opening_read_dialog is not None:
            self._opening_read_dialog.read_btn.setEnabled(bool(self.cad_connected_mode or self._base_dxf_available()))
        # CAD 연결 여부와 무관하게 캐시에 창호 후보가 있으면 편집은 가능해야 한다.
        self._update_opening_edit_controls()
        self.create_block_btn.setToolTip(
            "연결된 CAD 도면 내부에 실명 블록을 생성합니다." if self.cad_connected_mode
            else "새 DXF 파일로 실명 블록을 생성합니다."
        )

    def _current_floor_scan_rect(self, floor: str):
        """창호 읽기가 CAD에 요청할 층별 최소 평면 도곽 범위."""
        floor = str(floor or "")
        info = dict(self.preview.floor_background_files.get(floor, {}) or {})
        crop = info.get("crop_rect")
        if crop and len(crop) == 4 and info.get("crop_source") == "drawing_frame":
            try:
                return tuple(float(value) for value in crop)
            except (TypeError, ValueError):
                pass
        floor_rooms = [room for room in self.rooms if str(room.floor or "") == floor]
        points = [point for room in floor_rooms for point in (room.points or [])]
        if not points:
            return None
        xs = [float(point[0]) for point in points]
        ys = [float(point[1]) for point in points]
        minx, miny, maxx, maxy = min(xs), min(ys), max(xs), max(ys)
        span = max(maxx - minx, maxy - miny, 1.0)
        # 배경 도곽 캐시가 없을 때도 외벽 창호가 빠지지 않도록 실 외곽에 최소 1m 여유를 둔다.
        padding = max(1000.0, span * 0.05)
        return (minx - padding, miny - padding, maxx + padding, maxy + padding)

    def _normalize_cached_opening_results(self) -> None:
        """구버전/자동복원 캐시도 창호 편집 가능한 read_id 형태로 보정한다."""
        normalized = {}
        for floor, values in dict(self.opening_read_results or {}).items():
            floor = str(floor or "")
            if not floor:
                continue
            items = []
            for raw in values or []:
                marker = dict(raw or {})
                if not str(marker.get("read_id", "") or ""):
                    marker["read_id"] = self._opening_read_id(floor, marker)
                if not str(marker.get("read_label", "") or ""):
                    marker["read_label"] = str(marker.get("label", "") or "")
                # 캐시만 불러온 창호도 현재 실구분선에 즉시 재부착해 편집 가능하게 한다.
                if self.rooms:
                    self._refresh_opening_marker_wall_link(floor, marker)
                items.append(marker)
            normalized[floor] = items
        self.opening_read_results = normalized

    def set_window_catalog_provider(self, provider) -> None:
        self._window_catalog_provider = provider if callable(provider) else (lambda: [])
        if self._opening_catalog_dialog is not None:
            self._opening_catalog_dialog.refresh()

    def _window_catalog(self) -> list[dict]:
        try:
            values = self._window_catalog_provider() if callable(self._window_catalog_provider) else []
        except Exception:
            values = []
        result = []
        for raw in values or []:
            row = dict(raw or {})
            code = str(row.get("code") or "").strip().upper()
            if not code:
                continue
            row["code"] = code
            result.append(row)
        return sorted(result, key=lambda row: str(row.get("code") or "").casefold())

    def _window_catalog_by_code(self) -> dict[str, dict]:
        return {str(row.get("code") or "").upper(): row for row in self._window_catalog()}

    def refresh_opening_catalog_assignments(self) -> None:
        """창호 탭의 가로값이 바뀌면 이미 지정된 창호 길이도 마스터 기준으로 갱신한다."""
        catalog = self._window_catalog_by_code()
        changed = False
        for markers in self.opening_read_results.values():
            for marker in markers or []:
                code = str(marker.get("applied_code") or "").strip().upper()
                master = catalog.get(code)
                if master is None:
                    continue
                try:
                    length_mm = float(master.get("applied_length_mm", 0.0) or 0.0)
                except Exception:
                    length_mm = 0.0
                new_value = length_mm if length_mm > 0 else ""
                if marker.get("applied_length_mm") != new_value:
                    marker["applied_length_mm"] = new_value
                    changed = True
                dw = str(master.get("dw") or "").upper()
                if dw == "D":
                    marker["kind"] = "door"
                elif dw == "W":
                    marker["kind"] = "window"
        if changed:
            self.save_finish_assignments()
            self._sync_opening_read_preview()
            self.data_changed.emit()
            if self._opening_read_dialog is not None:
                self._opening_read_dialog.refresh("창호 리스트 길이 반영")
            if self._opening_assign_dialog is not None:
                self._opening_assign_dialog.refresh("창호 리스트 길이 반영")
        if self._opening_catalog_dialog is not None:
            self._opening_catalog_dialog.refresh()

    def _handle_f5_shortcut(self) -> None:
        # 보조 창이 활성화된 상태의 F5는 그 창의 지정 흐름을 우선한다.
        active = QApplication.activeWindow()
        if isinstance(active, OpeningReadDialog):
            active.open_catalog_for_selected()
            return
        if self._opening_edit_mode:
            self.open_selected_opening_assign_dialog()
        else:
            self.finish_panel.open_material_picker()

    def _position_opening_catalog_dialog(self, dialog, anchor=None) -> None:
        anchor = anchor or (self._opening_read_dialog if self._opening_read_dialog is not None and self._opening_read_dialog.isVisible() else None)
        try:
            if anchor is not None:
                frame = anchor.frameGeometry()
                screen = anchor.screen() or QApplication.primaryScreen()
                available = screen.availableGeometry() if screen is not None else QApplication.primaryScreen().availableGeometry()
                x = frame.right() + 8
                y = frame.top()
                if x + dialog.width() > available.right():
                    x = frame.left() - dialog.width() - 8
                x = max(available.left(), min(x, available.right() - dialog.width() + 1))
                y = max(available.top(), min(y, available.bottom() - dialog.height() + 1))
                dialog.move(x, y)
                return
            pos = self.preview.mapToGlobal(self.preview.rect().topRight())
            dialog.move(pos.x() + 8, pos.y() + 40)
        except Exception:
            pass

    def open_opening_catalog_dialog(self, anchor=None) -> None:
        pending_ids = set(self._opening_edit_selected_ids or set())
        if not self._opening_edit_mode:
            if self._opening_result_count() <= 0:
                QMessageBox.information(self, "창호 지정", "편집할 창호가 없습니다.")
                return
            self.opening_edit_btn.setChecked(True)
            if not self._opening_edit_mode:
                return
            if pending_ids:
                self._opening_edit_selected_ids = set(pending_ids)
                self.preview.set_selected_opening_ids(pending_ids)
        if not self._window_catalog():
            QMessageBox.information(self, "창호 지정", "창호 탭에 창호 리스트를 먼저 작성해 주세요.")
            return
        if anchor is None and self._opening_read_dialog is not None and self._opening_read_dialog.isVisible():
            anchor = self._opening_read_dialog
        # 부모 창이 바뀌면 기존 선택창을 닫고 새 부모에 종속해 다시 만든다.
        if self._opening_catalog_dialog is not None:
            current_parent = getattr(self._opening_catalog_dialog, "parent_dialog", None)
            if current_parent is not anchor:
                self._opening_catalog_dialog.close()
                self._opening_catalog_dialog = None
        if self._opening_catalog_dialog is None:
            self._opening_catalog_dialog = OpeningCatalogDialog(self, anchor)
            self._opening_catalog_dialog.finished.connect(lambda _result: setattr(self, "_opening_catalog_dialog", None))
        self._opening_catalog_dialog.refresh()
        self._position_opening_catalog_dialog(self._opening_catalog_dialog, anchor)
        self._opening_catalog_dialog.show(); self._opening_catalog_dialog.raise_(); self._opening_catalog_dialog.activateWindow()
        self.preview.setFocus(Qt.OtherFocusReason)

    def _marker_for_read_id(self, read_id: str):
        target = str(read_id or "")
        for floor, markers in self.opening_read_results.items():
            for marker in markers or []:
                if str(marker.get("read_id") or "") == target:
                    return str(floor), marker
        return "", None

    def apply_opening_read_pending(self, pending: dict) -> int:
        """창호 편집창에서 대기 중인 적용부호/길이를 한 번에 확정한다."""
        catalog = self._window_catalog_by_code()
        changed = 0
        touched_ids = set()
        for (floor, read_id), values in dict(pending or {}).items():
            marker_floor, marker = self._marker_for_read_id(str(read_id or ""))
            if marker is None:
                continue
            before_code = str(marker.get("applied_code") or "").strip().upper()
            before_length = marker.get("applied_length_mm")
            code = str(values.get("applied_code", before_code) or "").strip().upper()
            if code and code not in catalog:
                continue
            length = values.get("applied_length_mm", before_length)
            if length not in ("", None):
                try:
                    length = max(0.0, float(str(length).replace(",", "")))
                except Exception:
                    length = before_length
            else:
                length = ""
            if before_code != code or before_length != length:
                marker["applied_code"] = code
                marker["applied_length_mm"] = length
                master = catalog.get(code) if code else None
                if master is not None:
                    dw = str(master.get("dw") or "").upper()
                    if dw == "D":
                        marker["kind"] = "door"
                    elif dw == "W":
                        marker["kind"] = "window"
                self._refresh_opening_marker_wall_link(str(marker_floor or floor or self.selected_floor() or ""), marker)
                changed += 1
                touched_ids.add(str(read_id or ""))
        if changed:
            self.save_finish_assignments()
            self._sync_opening_read_preview()
            if touched_ids:
                self._opening_edit_selected_ids = set(touched_ids)
                self.preview.set_selected_opening_ids(touched_ids)
            self.data_changed.emit()
        return changed

    def apply_opening_catalog_code(self, code: str, ids=None) -> None:
        code = str(code or "").strip().upper()
        catalog = self._window_catalog_by_code()
        master = catalog.get(code)
        if master is None:
            QMessageBox.warning(self, "창호 지정", "창호 리스트에 없는 부호는 적용할 수 없습니다.")
            return
        targets = {str(value or "") for value in (ids if ids is not None else self._opening_edit_selected_ids) if str(value or "")}
        if not targets:
            QMessageBox.information(self, "창호 지정", "도면에서 적용할 창호를 하나 이상 선택해 주세요.")
            return
        try:
            length_mm = float(master.get("applied_length_mm", 0.0) or 0.0)
        except Exception:
            length_mm = 0.0
        changed = 0
        for floor, markers in self.opening_read_results.items():
            for marker in markers or []:
                if str(marker.get("read_id") or "") not in targets:
                    continue
                marker["applied_code"] = code
                marker["applied_length_mm"] = length_mm if length_mm > 0 else ""
                dw = str(master.get("dw") or "").upper()
                if dw == "D":
                    marker["kind"] = "door"
                elif dw == "W":
                    marker["kind"] = "window"
                self._refresh_opening_marker_wall_link(str(floor or self.selected_floor() or ""), marker)
                changed += 1
        if changed <= 0:
            return
        self.save_finish_assignments()
        self._sync_opening_read_preview()
        self.preview.set_selected_opening_ids(targets)
        self._opening_edit_selected_ids = set(targets)
        self.data_changed.emit()
        if self._opening_read_dialog is not None:
            self._opening_read_dialog.refresh(f"{code} · {changed:,}개 적용")
        if self._opening_catalog_dialog is not None:
            self._opening_catalog_dialog.refresh()
        self.opening_edit_status_label.setText(
            f"창호 지정 완료 | {code} · {changed:,}개 | 길이 {'목록 '+format(length_mm, ',.0f')+' mm' if length_mm > 0 else '검출길이 사용'}"
        )
        self.opening_edit_overlay.raise_()
        self.preview.setFocus(Qt.OtherFocusReason)

    def _edit_selected_opening_codes(self, ids) -> None:
        """F2: 창호 마스터 적용부호가 아니라 도면에서 자동 생성된 읽은 창호명을 수정한다."""
        targets = {str(value or "") for value in (ids or []) if str(value or "")}
        if not targets:
            return
        rows = []
        for read_id in targets:
            floor, marker = self._marker_for_read_id(read_id)
            if marker is not None:
                rows.append((floor_sort_key(floor), -float(marker.get("y", 0.0) or 0.0), float(marker.get("x", 0.0) or 0.0), floor, read_id, marker))
        rows.sort()
        changed = 0
        for _fk, _ny, _x, floor, read_id, marker in rows:
            # F2는 현재 도면 위치를 유지한 채 이름만 수정한다.
            # 선택 창호를 화면 중앙으로 강제 이동시키지 않는다.
            current = str(marker.get("read_label", marker.get("label", "")) or "").strip().upper()
            applied = str(marker.get("applied_code", "") or "").strip().upper()
            text, ok = QInputDialog.getText(
                self,
                "F2 창호 자동이름 수정",
                (
                    f"자동지정 부호명: {current or '(없음)'}\n"
                    f"적용 창호명: {applied or '없음'}\n"
                    "새 자동지정 부호명"
                ),
                QLineEdit.Normal,
                current,
            )
            if not ok:
                break
            value = str(text or "").strip().upper()
            if not value or value == current:
                continue
            marker["read_label"] = value
            marker["manual_read_label"] = True
            changed += 1
        if changed:
            self.save_finish_assignments()
            self._sync_opening_read_preview()
            self.data_changed.emit()
            if self._opening_read_dialog is not None:
                self._opening_read_dialog.refresh(f"자동 창호명 {changed:,}개 수정")
            if self._opening_assign_dialog is not None:
                self._opening_assign_dialog.refresh(f"자동 창호명 {changed:,}개 수정")
        self.preview.set_selected_opening_ids(targets)
        self._opening_edit_selected_ids = set(targets)
        self.preview.setFocus(Qt.OtherFocusReason)

    def _opening_result_count(self) -> int:
        return sum(len(values or []) for values in dict(self.opening_read_results or {}).values())

    def _update_opening_edit_controls(self) -> None:
        has_items = self._opening_result_count() > 0
        self.opening_edit_btn.setEnabled(has_items or self._opening_edit_mode)
        self._update_opening_edit_status(self._opening_edit_selected_ids)

    def _toggle_opening_edit_mode(self, enabled: bool) -> None:
        enabled = bool(enabled)
        if enabled and self.boundary_split_btn.isChecked():
            self.boundary_split_btn.setChecked(False)
        if enabled and self.boundary_merge_btn.isChecked():
            self.boundary_merge_btn.setChecked(False)
        if enabled and self._opening_result_count() <= 0:
            self.opening_edit_btn.blockSignals(True)
            self.opening_edit_btn.setChecked(False)
            self.opening_edit_btn.blockSignals(False)
            QMessageBox.information(self, "창호 편집", "현재 프로젝트/캐시에 창호 후보가 없습니다. 창호 읽기를 먼저 실행해 주세요.")
            return
        self._opening_edit_mode = enabled
        self._opening_last_move = None
        if hasattr(self, "opening_move_distance_spin"):
            self.opening_move_distance_spin.setEnabled(False)
            self.opening_move_exact_btn.setEnabled(False)
        if not enabled:
            self._opening_edit_selected_ids.clear()
            if self._opening_catalog_dialog is not None:
                self._opening_catalog_dialog.close()
            if self._opening_assign_dialog is not None:
                self._opening_assign_dialog.close()
        self.preview.set_opening_edit_mode(enabled)
        self.opening_edit_overlay.setVisible(enabled)
        self._update_opening_edit_status(self._opening_edit_selected_ids)
        if enabled:
            # V23.59: 도면편집 패널 등 같은 셀의 다른 오버레이보다 항상 앞에 둔다.
            self.opening_edit_overlay.adjustSize()
            self.opening_edit_overlay.raise_()
            QTimer.singleShot(0, self.opening_edit_overlay.raise_)
            self.preview.setFocus(Qt.OtherFocusReason)

    def _on_preview_opening_selection_changed(self, ids) -> None:
        self._opening_edit_selected_ids = {str(value or "") for value in (ids or []) if str(value or "")}
        self._update_opening_edit_status(self._opening_edit_selected_ids)
        if self._opening_catalog_dialog is not None:
            self._opening_catalog_dialog.refresh()

    def _on_preview_opening_delete_requested(self, ids) -> None:
        self._opening_edit_selected_ids = {str(value or "") for value in (ids or []) if str(value or "")}
        self._delete_selected_opening_markers()

    def _update_opening_edit_status(self, ids=None) -> None:
        count = len(ids if ids is not None else self._opening_edit_selected_ids)
        if hasattr(self, "opening_edit_status_label"):
            self.opening_edit_status_label.setText(
                f"창호 편집 | 선택 {count:,}개 | 드래그/방향키 10mm 이동 · F2 자동이름 수정 · F5 지정목록 · Delete 삭제 · ESC 종료"
            )
        if hasattr(self, "opening_edit_delete_btn"):
            self.opening_edit_delete_btn.setEnabled(count > 0)
        if self._opening_edit_mode and hasattr(self, "opening_edit_overlay"):
            self.opening_edit_overlay.raise_()

    def _on_preview_opening_move_preview(self, payload) -> None:
        data = dict(payload or {})
        if not data:
            self._update_opening_edit_status(self._opening_edit_selected_ids)
            return
        dx = float(data.get("dx", 0.0) or 0.0)
        dy = float(data.get("dy", 0.0) or 0.0)
        distance = float(data.get("distance", math.hypot(dx, dy)) or 0.0)
        if hasattr(self, "opening_edit_status_label"):
            self.opening_edit_status_label.setText(
                f"창호 이동 | ΔX {dx:+,.0f} · ΔY {dy:+,.0f} · 거리 {distance:,.0f} mm | 10mm 스냅"
            )
        if hasattr(self, "opening_move_distance_spin"):
            self.opening_move_distance_spin.blockSignals(True)
            self.opening_move_distance_spin.setValue(round(distance / 10.0) * 10.0)
            self.opening_move_distance_spin.blockSignals(False)

    def _refresh_opening_marker_wall_link(self, floor: str, marker: dict) -> None:
        """수동/캐시 창호를 현재 실구분선 벽체에 다시 부착한다."""
        try:
            px, py = self._opening_marker_wall_center(marker)
        except (TypeError, ValueError):
            return
        length_mm = self._opening_effective_length_mm(marker)
        attachment = self._opening_wall_attachment(floor, px, py, length_mm, 10.0)
        self._apply_opening_attachment(marker, attachment)

    def _on_preview_opening_move_requested(self, ids, dx: float, dy: float) -> None:
        ids = {str(value or "") for value in (ids or []) if str(value or "")}
        if not ids:
            return
        floor = str(self.selected_floor() or self.preview.opening_test_floor or "")
        markers = list(self.opening_read_results.get(floor, []) or [])
        originals = {}
        directions = {}
        actual_moves = {}
        moved = 0
        rejected = 0
        dx = round(float(dx) / 10.0) * 10.0
        dy = round(float(dy) / 10.0) * 10.0
        for marker in markers:
            rid = str(marker.get("read_id", "") or "")
            if rid not in ids:
                continue
            try:
                ox, oy = self._opening_marker_wall_center(marker)
            except (TypeError, ValueError):
                continue
            originals[rid] = (ox, oy)
            proposed_x, proposed_y = ox + dx, oy + dy
            length_mm = self._opening_effective_length_mm(marker)
            attachment = self._opening_wall_attachment(floor, proposed_x, proposed_y, length_mm, 10.0)
            # 창호는 벽 소속 객체다. 붙일 벽을 찾지 못하면 자유공간에 확정하지 않고 원위치한다.
            if attachment is None:
                rejected += 1
                continue
            self._apply_opening_attachment(marker, attachment)
            marker["manual_moved"] = True
            nx, ny = self._opening_marker_wall_center(marker)
            adx = float(nx) - ox
            ady = float(ny) - oy
            adist = math.hypot(adx, ady)
            actual_moves[rid] = (adx, ady, adist)
            if adist > 1e-9:
                directions[rid] = (adx / adist, ady / adist)
            moved += 1
        if moved <= 0:
            self._sync_opening_read_preview()
            self.preview.set_selected_opening_ids(ids)
            self.opening_edit_status_label.setText("창호 이동 취소 | 붙일 수 있는 벽체를 찾지 못했습니다. 창호는 실내 자유좌표에 배치되지 않습니다.")
            self.opening_edit_overlay.raise_()
            self.preview.setFocus(Qt.OtherFocusReason)
            return

        first_id = next((rid for rid in ids if rid in actual_moves), "")
        first_dx, first_dy, distance = actual_moves.get(first_id, (dx, dy, math.hypot(dx, dy)))
        if distance > 1e-9:
            self._opening_last_move = {
                "floor": floor, "ids": set(ids), "originals": originals,
                "directions": directions, "distance": distance,
            }
            self.opening_move_distance_spin.blockSignals(True)
            self.opening_move_distance_spin.setValue(round(distance / 10.0) * 10.0)
            self.opening_move_distance_spin.blockSignals(False)
            self.opening_move_distance_spin.setEnabled(True)
            self.opening_move_exact_btn.setEnabled(True)
        self.save_finish_assignments()
        self._sync_opening_read_preview()
        self.preview.set_selected_opening_ids(ids)
        self._opening_edit_selected_ids = set(ids)
        self.data_changed.emit()
        if self._opening_read_dialog is not None:
            self._opening_read_dialog.refresh("창호 위치 이동 반영")
        reject_text = f" · 벽 미검출 {rejected:,}개 원위치" if rejected else ""
        self.opening_edit_status_label.setText(
            f"창호 이동 완료 | {moved:,}개 · 벽체 스냅 · ΔX {first_dx:+,.0f} · ΔY {first_dy:+,.0f} · 거리 {distance:,.0f} mm | 10mm{reject_text}"
        )
        self.opening_edit_overlay.raise_()
        self.preview.setFocus(Qt.OtherFocusReason)

    def _apply_opening_exact_distance(self) -> None:
        move = dict(self._opening_last_move or {})
        if not move:
            return
        desired = round(float(self.opening_move_distance_spin.value()) / 10.0) * 10.0
        floor = str(move.get("floor", "") or "")
        ids = set(move.get("ids", set()) or set())
        originals = dict(move.get("originals", {}) or {})
        directions = dict(move.get("directions", {}) or {})
        changed = 0
        rejected = 0
        for marker in self.opening_read_results.get(floor, []) or []:
            rid = str(marker.get("read_id", "") or "")
            if rid not in ids or rid not in originals:
                continue
            ox, oy = originals[rid]
            ux, uy = directions.get(rid, (0.0, 0.0))
            if abs(ux) <= 1e-12 and abs(uy) <= 1e-12:
                continue
            proposed_x = float(ox) + float(ux) * desired
            proposed_y = float(oy) + float(uy) * desired
            length_mm = self._opening_effective_length_mm(marker)
            attachment = self._opening_wall_attachment(floor, proposed_x, proposed_y, length_mm, 10.0)
            if attachment is None:
                rejected += 1
                continue
            self._apply_opening_attachment(marker, attachment)
            marker["manual_moved"] = True
            changed += 1
        if changed <= 0:
            return
        move["distance"] = desired
        self._opening_last_move = move
        self.save_finish_assignments()
        self._sync_opening_read_preview()
        self.preview.set_selected_opening_ids(ids)
        self.data_changed.emit()
        if self._opening_read_dialog is not None:
            self._opening_read_dialog.refresh("정확 이동거리 반영")
        reject_text = f" · 벽 미검출 {rejected:,}개 원위치" if rejected else ""
        self.opening_edit_status_label.setText(
            f"정확 이동 적용 | {changed:,}개 · 벽체 스냅 · 거리 {desired:,.0f} mm · 10mm 단위{reject_text}"
        )
        self.opening_edit_overlay.raise_()
        self.preview.setFocus(Qt.OtherFocusReason)

    def _delete_selected_opening_markers(self) -> None:
        ids = set(self._opening_edit_selected_ids or self.preview.selected_opening_ids or set())
        if not ids:
            return
        floor = str(self.selected_floor() or self.preview.opening_test_floor or "")
        before = list(self.opening_read_results.get(floor, []) or [])
        after = [marker for marker in before if str(marker.get("read_id", "") or "") not in ids]
        deleted = len(before) - len(after)
        if deleted <= 0:
            return
        if after:
            self.opening_read_results[floor] = after
        else:
            self.opening_read_results.pop(floor, None)
        self._opening_edit_selected_ids.clear()
        self.save_finish_assignments()
        self._sync_opening_read_preview()
        self.preview.set_selected_opening_ids(set())
        self._update_opening_edit_controls()
        if self._opening_read_dialog is not None:
            self._opening_read_dialog.refresh(f"도곽 편집에서 {deleted:,}개 삭제")
        # 삭제 후 편집모드는 의도적으로 유지한다.
        if self._opening_edit_mode:
            self.opening_edit_overlay.show()
            self.preview.setFocus(Qt.OtherFocusReason)

    def _on_preview_escape_pressed(self) -> None:
        if self.boundary_split_btn.isChecked():
            self.boundary_split_btn.setChecked(False)
            return
        if self.boundary_merge_btn.isChecked():
            self.boundary_merge_btn.setChecked(False)
            return
        if self._opening_edit_mode:
            self.opening_edit_btn.setChecked(False)
            return
        if self._room_group_dialog is not None or self._wall_group_dialog is not None:
            self._cancel_finish_group_overlay()
            return
        self.cancel_direct_number_mode()

    def _toggle_boundary_split_mode(self, enabled: bool) -> None:
        enabled = bool(enabled)
        if enabled and self.view_mode.currentText() != "층 전체 보기":
            self.view_mode.setCurrentText("층 전체 보기")
        # 다른 도곽 편집모드와 동시에 켜지지 않게 한다.
        if enabled and self._opening_edit_mode:
            self.opening_edit_btn.setChecked(False)
        if enabled and self.boundary_merge_btn.isChecked():
            self.boundary_merge_btn.setChecked(False)
        self.preview.set_boundary_split_mode(enabled)
        if enabled:
            self.boundary_split_btn.setText("구획벽 나누기 [ON]")
            self.boundary_split_status_label.setText("마우스 위치 10mm 스냅")
            self.boundary_split_status_label.show()
            self.preview_help.setText("구획벽 나누기: 선분 위 마우스 이동 → 10mm 단위 거리 실시간 표시 → 클릭 즉시 절점 추가 | 가운데버튼 이동 | ESC 종료")
            self.floor_tools_overlay.raise_()
            self.preview.setFocus(Qt.OtherFocusReason)
        else:
            self.boundary_split_btn.setText("구획벽 나누기")
            self.boundary_split_status_label.hide()
            self.preview.clear_boundary_split_highlight()
            self.preview_help.setText("마우스 휠: 확대·축소 | 일반모드 드래그: 화면 이동 | 벽 선택모드 드래그: 박스 선택 | 가운데버튼 더블클릭: 전체보기")

    def _toggle_boundary_merge_mode(self, enabled: bool) -> None:
        enabled = bool(enabled)
        if enabled and self.view_mode.currentText() != "층 전체 보기":
            self.view_mode.setCurrentText("층 전체 보기")
        if enabled and self._opening_edit_mode:
            self.opening_edit_btn.setChecked(False)
        if enabled and self.boundary_split_btn.isChecked():
            self.boundary_split_btn.setChecked(False)
        self.preview.set_boundary_merge_mode(enabled)
        if enabled:
            self.boundary_merge_btn.setText("구획벽 합치기 [ON]")
            self.boundary_split_status_label.setText("일직선 공유 절점 선택")
            self.boundary_split_status_label.show()
            self.preview_help.setText("구획벽 합치기: 일직선으로 이어진 두 선의 공유 절점에 마우스 이동 → 클릭하여 한 선으로 합치기 | 가운데버튼 이동 | ESC 종료")
            self.floor_tools_overlay.raise_()
            self.preview.setFocus(Qt.OtherFocusReason)
        else:
            self.boundary_merge_btn.setText("구획벽 합치기")
            self.boundary_split_status_label.hide()
            self.preview.clear_boundary_split_highlight()
            self.preview_help.setText("마우스 휠: 확대·축소 | 일반모드 드래그: 화면 이동 | 벽 선택모드 드래그: 박스 선택 | 가운데버튼 더블클릭: 전체보기")

    @staticmethod
    def _remap_wall_key_after_split(key: str, room_handle: str, split_index: int) -> list[str]:
        """외곽 벽키 handle:index를 절점 추가 뒤 새 인덱스로 옮긴다.

        나눈 기존 구간은 두 구간 모두 같은 그룹/재료를 승계하고, 그 뒤 구간은 +1 이동한다.
        공제벽 키(handle:dN:index)와 다른 실 키는 그대로 둔다.
        """
        text = str(key or "")
        parts = text.split(":")
        if len(parts) != 2 or parts[0] != str(room_handle or ""):
            return [text] if text else []
        try:
            index = int(parts[1])
        except (TypeError, ValueError):
            return [text]
        if index < split_index:
            return [text]
        if index == split_index:
            return [f"{room_handle}:{split_index}", f"{room_handle}:{split_index + 1}"]
        return [f"{room_handle}:{index + 1}"]

    def _remap_wall_assignments_after_split(self, room: RoomRecord, split_index: int) -> None:
        handle = str(room.boundary_handle or "")
        if not handle:
            return

        old_assignments = dict(self.wall_finish_assignments or {})
        new_assignments = {}
        for key, material in old_assignments.items():
            for new_key in self._remap_wall_key_after_split(str(key), handle, split_index):
                new_assignments[new_key] = deepcopy(material)
        self.wall_finish_assignments = new_assignments

        for group in self.finish_groups.values():
            remapped = []
            for key in list(group.get("wall_keys", []) or []):
                remapped.extend(self._remap_wall_key_after_split(str(key), handle, split_index))
            group["wall_keys"] = list(dict.fromkeys(value for value in remapped if value))

        def remap_set(values):
            result = set()
            for key in set(values or set()):
                result.update(self._remap_wall_key_after_split(str(key), handle, split_index))
            return result

        self.selected_wall_keys = remap_set(self.selected_wall_keys)
        self._group_mode_blocked_wall_keys = remap_set(self._group_mode_blocked_wall_keys)

    def _on_boundary_split_hover_changed(self, payload) -> None:
        if not hasattr(self, "boundary_split_status_label") or not self.boundary_split_btn.isChecked():
            return
        data = dict(payload or {})
        if not data:
            self.boundary_split_status_label.setText("마우스 위치 10mm 스냅")
            return
        self.boundary_split_status_label.setText(
            f"① {float(data.get('distance1', 0.0)):,.0f} / ② {float(data.get('distance2', 0.0)):,.0f} mm"
        )

    def _on_boundary_segment_split_requested(self, room: RoomRecord, segment_index: int, distance: float) -> None:
        if not self.boundary_split_btn.isChecked() or room not in self.rooms:
            return
        points = list(getattr(room, "points", []) or [])
        if len(points) < 2 or not (0 <= int(segment_index) < len(points)):
            return
        segment_index = int(segment_index)
        p1 = points[segment_index]
        p2 = points[(segment_index + 1) % len(points)]
        length_mm = math.hypot(float(p2[0]) - float(p1[0]), float(p2[1]) - float(p1[1]))
        distance = round(float(distance or 0.0) / 10.0) * 10.0
        if length_mm < 20.0 or not (0.0 < distance < length_mm):
            QToolTip.showText(QCursor.pos(), "끝점에는 절점을 추가할 수 없습니다.", self.preview)
            return
        ratio = distance / length_mm
        new_point = (
            float(p1[0]) + (float(p2[0]) - float(p1[0])) * ratio,
            float(p1[1]) + (float(p2[1]) - float(p1[1])) * ratio,
        )

        # 나눈 기존 벽의 마감/그룹은 두 새 구간에 모두 승계한다.
        self._remap_wall_assignments_after_split(room, segment_index)
        room.points.insert(segment_index + 1, new_point)
        room.changed = True
        room.geometry_changed = True

        self.preview.set_wall_finish_assignments(self.wall_finish_assignments)
        self.preview.set_finish_groups(self.finish_groups)
        self.preview.set_selected_wall_keys(self.selected_wall_keys)
        self.preview.refresh_current_geometry()
        # 추가 직후에는 나뉜 양쪽 구간/공유 절점을 동시에 임시 표시한다.
        self.preview.show_boundary_split_confirmation(room, segment_index)
        self.save_finish_assignments()
        self.data_changed.emit()
        d2 = max(0.0, length_mm - distance)
        self.boundary_split_status_label.setText(f"분할 완료 ① {distance:,.0f} / ② {d2:,.0f} mm")
        self.preview_help.setText(
            f"구획벽 나누기: ① {distance:,.0f} / ② {d2:,.0f} mm 위치 분할 완료 | 다음 선분에 마우스를 올려 계속 작업 | ESC 종료"
        )
        QToolTip.showText(QCursor.pos(), f"절점 추가 · ① {distance:,.0f} / ② {d2:,.0f} mm", self.preview)
        self.preview.setFocus(Qt.OtherFocusReason)

    @staticmethod
    def _wall_key_after_merge(key: str, room_handle: str, vertex_index: int, old_segment_count: int) -> str:
        text = str(key or "")
        parts = text.split(":")
        if len(parts) != 2 or parts[0] != str(room_handle or ""):
            return text
        try:
            index = int(parts[1])
        except (TypeError, ValueError):
            return text
        vertex_index = int(vertex_index)
        old_segment_count = int(old_segment_count)
        if vertex_index == 0:
            if index in (0, old_segment_count - 1):
                new_index = old_segment_count - 2
            else:
                new_index = index - 1
        else:
            if index in (vertex_index - 1, vertex_index):
                new_index = vertex_index - 1
            elif index > vertex_index:
                new_index = index - 1
            else:
                new_index = index
        return f"{room_handle}:{new_index}"

    def _wall_group_signature(self, wall_key: str) -> tuple:
        signature = []
        key = str(wall_key or "")
        for group_id, group in dict(self.finish_groups or {}).items():
            if key not in set(group.get("wall_keys", []) or []):
                continue
            signature.append((
                normalize_finish_group_type(group.get("group_type", "")),
                normalize_finish_group_division(group.get("group_division", ""), group.get("group_type", "")),
                str(group_id),
            ))
        return tuple(sorted(signature))

    def _wall_group_ids_for_key(self, wall_key: str) -> tuple[str, ...]:
        key = str(wall_key or "")
        values = []
        for group_id, group in dict(self.finish_groups or {}).items():
            if key in set(group.get("wall_keys", []) or []):
                values.append(str(group_id))
        return tuple(sorted(values))

    def _wall_finish_snapshot(self, wall_key: str) -> dict:
        key = str(wall_key or "")
        material = self.wall_finish_assignments.get(key, None)
        return {
            "material": deepcopy(material) if isinstance(material, dict) else None,
            "group_ids": self._wall_group_ids_for_key(key),
        }

    def _wall_finish_description(self, wall_key: str) -> str:
        key = str(wall_key or "")
        group_lines = []
        for group_id in self._wall_group_ids_for_key(key):
            group = self.finish_groups.get(group_id, {})
            group_type = normalize_finish_group_type(group.get("group_type", ""))
            division = normalize_finish_group_division(group.get("group_division", ""), group_type)
            name = str(group.get("name", "") or "그룹")
            group_lines.append(f"{group_type}/{division} · {name}")

        material = self.wall_finish_assignments.get(key, None)
        material_text = "미지정"
        if isinstance(material, dict) and material:
            name = str(material.get("name", "") or "미지정").strip() or "미지정"
            spec = str(material.get("spec", "") or "").strip()
            material_text = f"{name} / {spec}" if spec else name

        if group_lines:
            return "그룹: " + "\n      ".join(group_lines) + f"\n재료: {material_text}"
        return f"그룹: 없음\n재료: {material_text}"

    def _choose_wall_finish_for_merge(self, key1: str, key2: str, length1: float, length2: float):
        direct1 = self.wall_finish_assignments.get(str(key1), None)
        direct2 = self.wall_finish_assignments.get(str(key2), None)
        groups1 = self._wall_group_signature(key1)
        groups2 = self._wall_group_signature(key2)
        if direct1 == direct2 and groups1 == groups2:
            return self._wall_finish_snapshot(key1)

        box = QMessageBox(self)
        box.setWindowTitle("구획벽 합치기 - 마감 선택")
        box.setIcon(QMessageBox.Question)
        box.setText("두 구획벽의 마감 지정이 서로 다릅니다.\n합친 벽에 사용할 마감을 선택해 주세요.")
        box.setInformativeText(
            f"① 구간 ({float(length1):,.0f} mm)\n{self._wall_finish_description(key1)}\n\n"
            f"② 구간 ({float(length2):,.0f} mm)\n{self._wall_finish_description(key2)}"
        )
        first_btn = box.addButton("① 마감으로 합치기", QMessageBox.AcceptRole)
        second_btn = box.addButton("② 마감으로 합치기", QMessageBox.ActionRole)
        box.addButton("취소", QMessageBox.RejectRole)
        box.exec()
        clicked = box.clickedButton()
        if clicked is first_btn:
            return self._wall_finish_snapshot(key1)
        if clicked is second_btn:
            return self._wall_finish_snapshot(key2)
        return None

    def _apply_wall_finish_after_merge(self, room: RoomRecord, merged_key: str, snapshot: dict, affected_group_ids) -> None:
        merged_key = str(merged_key or "")
        snapshot = dict(snapshot or {})
        selected_group_ids = {str(value) for value in snapshot.get("group_ids", ()) if str(value)}

        material = snapshot.get("material", None)
        if isinstance(material, dict) and material:
            self.wall_finish_assignments[merged_key] = deepcopy(material)
        else:
            self.wall_finish_assignments.pop(merged_key, None)

        # 기존 양쪽 선이 서로 다른 그룹이었다면 remap 과정에서 한 벽키에 두 그룹이
        # 겹칠 수 있으므로, 합치기에서 선택한 쪽의 그룹 구성만 남긴다.
        for group_id, group in self.finish_groups.items():
            keys = [str(value) for value in list(group.get("wall_keys", []) or []) if str(value)]
            if merged_key in keys:
                group["wall_keys"] = [value for value in keys if value != merged_key]
            if str(group_id) in selected_group_ids:
                group["wall_keys"] = list(dict.fromkeys(list(group.get("wall_keys", []) or []) + [merged_key]))

        # 벽 그룹의 실 연결/재료목록도 실제 남은 wall_keys와 동기화한다.
        for group_id in {str(value) for value in affected_group_ids if str(value)} | selected_group_ids:
            group = self.finish_groups.get(group_id)
            if not group:
                continue
            handles = sorted({
                wall_key_room_handle(key) for key in list(group.get("wall_keys", []) or [])
                if wall_key_room_handle(key)
            })
            self._set_rooms_for_finish_group(group_id, handles)

    def _can_merge_wall_vertex(self, room: RoomRecord, vertex_index: int) -> tuple[bool, str]:
        points = list(getattr(room, "points", []) or [])
        if len(points) < 4 or not (0 <= int(vertex_index) < len(points)):
            return False, "합칠 수 있는 절점이 아닙니다."
        if not RoomPreview._merge_vertex_is_collinear(points, int(vertex_index)):
            return False, "두 구획벽이 일직선이 아닙니다."
        return True, ""

    def _remap_wall_assignments_after_merge(self, room: RoomRecord, vertex_index: int, old_segment_count: int) -> None:
        handle = str(room.boundary_handle or "")
        if not handle:
            return

        new_assignments = {}
        for key, material in dict(self.wall_finish_assignments or {}).items():
            new_key = self._wall_key_after_merge(str(key), handle, vertex_index, old_segment_count)
            if new_key not in new_assignments:
                new_assignments[new_key] = deepcopy(material)
        self.wall_finish_assignments = new_assignments

        for group in self.finish_groups.values():
            remapped = [
                self._wall_key_after_merge(str(key), handle, vertex_index, old_segment_count)
                for key in list(group.get("wall_keys", []) or [])
            ]
            group["wall_keys"] = list(dict.fromkeys(value for value in remapped if value))

        def remap_set(values):
            return {
                self._wall_key_after_merge(str(key), handle, vertex_index, old_segment_count)
                for key in set(values or set()) if str(key or "")
            }

        self.selected_wall_keys = remap_set(self.selected_wall_keys)
        self._group_mode_blocked_wall_keys = remap_set(self._group_mode_blocked_wall_keys)

    def _on_boundary_merge_hover_changed(self, payload) -> None:
        if not hasattr(self, "boundary_split_status_label") or not self.boundary_merge_btn.isChecked():
            return
        data = dict(payload or {})
        if not data:
            self.boundary_split_status_label.setText("일직선 공유 절점 선택")
            return
        self.boundary_split_status_label.setText(
            f"{float(data.get('length1', 0.0)):,.0f} + {float(data.get('length2', 0.0)):,.0f} = {float(data.get('length', 0.0)):,.0f} mm"
        )

    def _on_boundary_vertex_merge_requested(self, room: RoomRecord, vertex_index: int) -> None:
        if not self.boundary_merge_btn.isChecked() or room not in self.rooms:
            return
        points = list(getattr(room, "points", []) or [])
        vertex_index = int(vertex_index)
        ok, reason = self._can_merge_wall_vertex(room, vertex_index)
        if not ok:
            QToolTip.showText(QCursor.pos(), reason, self.preview)
            return

        old_count = len(points)
        removed_point = points[vertex_index]
        prev_point = points[(vertex_index - 1) % old_count]
        next_point = points[(vertex_index + 1) % old_count]
        length1 = math.hypot(float(removed_point[0]) - float(prev_point[0]), float(removed_point[1]) - float(prev_point[1]))
        length2 = math.hypot(float(next_point[0]) - float(removed_point[0]), float(next_point[1]) - float(removed_point[1]))
        merged_length = math.hypot(float(next_point[0]) - float(prev_point[0]), float(next_point[1]) - float(prev_point[1]))

        handle = str(room.boundary_handle or "")
        prev_index = (vertex_index - 1) % old_count
        next_index = vertex_index
        key1 = f"{handle}:{prev_index}" if handle else ""
        key2 = f"{handle}:{next_index}" if handle else ""
        affected_group_ids = set(self._wall_group_ids_for_key(key1)) | set(self._wall_group_ids_for_key(key2))
        finish_snapshot = self._choose_wall_finish_for_merge(key1, key2, length1, length2) if handle else {"material": None, "group_ids": ()}
        if finish_snapshot is None:
            self.preview.setFocus(Qt.OtherFocusReason)
            return

        self._remap_wall_assignments_after_merge(room, vertex_index, old_count)
        room.points.pop(vertex_index)
        room.changed = True
        room.geometry_changed = True
        merged_index = old_count - 2 if vertex_index == 0 else vertex_index - 1
        merged_key = f"{handle}:{merged_index}" if handle else ""
        if merged_key:
            self._apply_wall_finish_after_merge(room, merged_key, finish_snapshot, affected_group_ids)

        self.preview.set_wall_finish_assignments(self.wall_finish_assignments)
        self.preview.set_room_finish_assignments(self.room_finish_assignments)
        self.preview.set_finish_groups(self.finish_groups)
        self.preview.set_selected_wall_keys(self.selected_wall_keys)
        self.preview.refresh_current_geometry()
        self.preview.show_boundary_merge_confirmation(room, merged_index, removed_point)
        self.save_finish_assignments()
        self._refresh_current_finish_panel()
        self.update_finish_summary(self.selected_floor())
        self.data_changed.emit()
        self.boundary_split_status_label.setText(f"합치기 완료 {merged_length:,.0f} mm")
        self.preview_help.setText(f"구획벽 합치기: 일직선 2구간 합치기 완료 · {merged_length:,.0f} mm | 다음 절점을 선택해 계속 작업 | ESC 종료")
        QToolTip.showText(QCursor.pos(), f"구획벽 합치기 완료 · {merged_length:,.0f} mm", self.preview)
        self.preview.setFocus(Qt.OtherFocusReason)

    def open_selected_opening_assign_dialog(self) -> None:
        """도면 창호 편집에서 F5/창호지정: 현재 선택 창호만 보이는 지정목록을 연다."""
        ids = {str(v or "") for v in self._opening_edit_selected_ids if str(v or "")}
        if not ids:
            QMessageBox.information(self, "창호 지정", "도면에서 지정할 창호를 먼저 선택해 주세요.")
            return
        if self._opening_assign_dialog is not None:
            self._opening_assign_dialog.close()
            self._opening_assign_dialog = None
        dialog = OpeningReadDialog(self, filter_ids=ids, assignment_mode=True)
        self._opening_assign_dialog = dialog
        dialog.finished.connect(lambda _result, d=dialog: self._clear_opening_assign_dialog(d))
        dialog.show()
        dialog.raise_()
        dialog.activateWindow()

    def _clear_opening_assign_dialog(self, dialog) -> None:
        if self._opening_assign_dialog is dialog:
            self._opening_assign_dialog = None
        if self._opening_catalog_dialog is not None and getattr(self._opening_catalog_dialog, "parent_dialog", None) is dialog:
            self._opening_catalog_dialog.close()

    def open_opening_read_dialog(self) -> None:
        if not self.rooms:
            QMessageBox.information(self, "창호 읽기", "먼저 평면 실목록을 불러와 주세요.")
            return
        if self._opening_read_dialog is None:
            self._opening_read_dialog = OpeningReadDialog(self)
            self._opening_read_dialog.finished.connect(lambda _result: setattr(self, "_opening_read_dialog", None))
        self._opening_read_dialog.read_btn.setEnabled(bool(self.cad_connected_mode or self._base_dxf_available()))
        self._opening_read_dialog.refresh()
        self._opening_read_dialog.show()
        self._opening_read_dialog.raise_()
        self._opening_read_dialog.activateWindow()

    def request_opening_read_from_dialog(self) -> None:
        current_floor = str(self.selected_floor() or "")
        cad_available = bool(self.cad_connected_mode)
        dxf_available = bool(self._base_dxf_available())
        if not cad_available and not dxf_available:
            QMessageBox.information(self, "창호 읽기", "읽을 수 있는 연동 CAD 또는 베이스 DXF가 없습니다.")
            return

        dialog = QDialog(self._opening_read_dialog or self)
        dialog.setWindowTitle("창호 읽기")
        dialog.setModal(True)
        dialog.resize(390, 245)
        root = QVBoxLayout(dialog)

        source_box = QGroupBox("읽기 소스")
        source_layout = QVBoxLayout(source_box)
        cad_radio = QRadioButton("연동된 CAD에서 읽기")
        dxf_radio = QRadioButton("베이스 DXF에서 읽기")
        cad_radio.setEnabled(cad_available)
        dxf_radio.setEnabled(dxf_available)
        if cad_available:
            cad_radio.setChecked(True)
        else:
            dxf_radio.setChecked(True)
        if not cad_available:
            cad_radio.setToolTip("현재 연동된 CAD가 없어 선택할 수 없습니다.")
        if not dxf_available:
            dxf_radio.setToolTip("프로젝트의 베이스 DXF가 없어 선택할 수 없습니다.")
        source_layout.addWidget(cad_radio); source_layout.addWidget(dxf_radio)
        if dxf_available:
            source_path_label = QLabel(f"베이스 DXF: {Path(self.base_dxf_path).name}")
            source_path_label.setStyleSheet("color:#666;")
            source_layout.addWidget(source_path_label)
        root.addWidget(source_box)

        scope_box = QGroupBox("읽기 범위")
        scope_layout = QVBoxLayout(scope_box)
        current_label = f"현재층 ({current_floor})" if current_floor else "현재층 (선택된 층 없음)"
        current_radio = QRadioButton(current_label)
        all_radio = QRadioButton("전층")
        current_radio.setEnabled(bool(current_floor))
        if current_floor:
            current_radio.setChecked(True)
        else:
            all_radio.setChecked(True)
        scope_layout.addWidget(current_radio); scope_layout.addWidget(all_radio)
        root.addWidget(scope_box)

        buttons = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        if buttons.button(QDialogButtonBox.Ok) is not None:
            buttons.button(QDialogButtonBox.Ok).setText("읽기")
            buttons.button(QDialogButtonBox.Ok).setDefault(True)
        buttons.accepted.connect(dialog.accept); buttons.rejected.connect(dialog.reject)
        root.addWidget(buttons)
        if dialog.exec() != QDialog.Accepted:
            return

        source = "cad" if cad_radio.isChecked() else "dxf"
        if current_radio.isChecked():
            floors = [current_floor]
            scope = "current"
        else:
            floors = sorted({str(room.floor or "") for room in self.rooms if str(room.floor or "")}, key=floor_sort_key)
            scope = "all"

        requests = []
        missing = []
        for floor in floors:
            rect = self._current_floor_scan_rect(floor)
            if rect:
                requests.append({"floor": floor, "rect": rect})
            else:
                missing.append(floor)
        if not requests:
            QMessageBox.information(self, "창호 읽기", "선택 범위의 평면 도곽을 확인하지 못했습니다.")
            return
        if self._opening_read_dialog is not None:
            source_text = "연동 CAD" if source == "cad" else "베이스 DXF"
            self._opening_read_dialog.status_label.setText(f"{source_text} 읽는 중... {len(requests)}개 층")
        self.opening_read_requested.emit({
            "scope": scope,
            "source": source,
            "base_dxf_path": self.base_dxf_path if source == "dxf" else "",
            "requests": requests,
            "window_acis": [4, 5],
            "missing_floors": missing,
        })

    def request_opening_test(self) -> None:
        """구버전 호출 호환: 새 창호 읽기 관리창을 연다."""
        self.open_opening_read_dialog()

    @staticmethod
    def _project_opening_point_to_segment(px: float, py: float, p1, p2):
        x1, y1 = float(p1[0]), float(p1[1])
        x2, y2 = float(p2[0]), float(p2[1])
        dx, dy = x2 - x1, y2 - y1
        length_sq = dx * dx + dy * dy
        if length_sq <= 1e-9:
            return x1, y1, math.hypot(px - x1, py - y1)
        t = ((px - x1) * dx + (py - y1) * dy) / length_sq
        t = max(0.0, min(1.0, t))
        sx, sy = x1 + t * dx, y1 + t * dy
        return sx, sy, math.hypot(px - sx, py - sy)

    @staticmethod
    def _opening_effective_length_mm(marker: dict) -> float:
        """도면 표시/벽 부착에 사용할 창호 길이. 적용길이가 없으면 검출길이를 사용한다."""
        for key in ("applied_length_mm", "detected_length_mm"):
            try:
                value = float(marker.get(key, 0.0) or 0.0)
            except (TypeError, ValueError):
                value = 0.0
            if value > 0.0:
                return value
        return 0.0

    @staticmethod
    def _project_opening_point_to_segment_full(px: float, py: float, p1, p2):
        x1, y1 = float(p1[0]), float(p1[1])
        x2, y2 = float(p2[0]), float(p2[1])
        dx, dy = x2 - x1, y2 - y1
        length_sq = dx * dx + dy * dy
        if length_sq <= 1e-9:
            return x1, y1, 0.0, 0.0, math.hypot(px - x1, py - y1)
        length = math.sqrt(length_sq)
        t_raw = ((px - x1) * dx + (py - y1) * dy) / length_sq
        t = max(0.0, min(1.0, t_raw))
        sx, sy = x1 + t * dx, y1 + t * dy
        return sx, sy, t_raw, length, math.hypot(px - sx, py - sy)

    @staticmethod
    def _opening_marker_wall_center(marker: dict) -> tuple[float, float]:
        """표시용 부호 좌표와 별개인 실제 벽 부착 중심을 반환한다."""
        if bool(marker.get("wall_snapped")):
            try:
                return float(marker.get("wall_center_x")), float(marker.get("wall_center_y"))
            except (TypeError, ValueError):
                pass
        return float(marker.get("x", 0.0)), float(marker.get("y", 0.0))

    def _opening_point_inside_any_room(self, floor: str, x: float, y: float) -> bool:
        for room in self.rooms:
            if str(room.floor or "") != str(floor or ""):
                continue
            points = list(room.points or [])
            if len(points) >= 3 and point_in_polygon(float(x), float(y), points):
                return True
        return False

    def _opening_display_coordinate(
        self, floor: str, center_x: float, center_y: float, ux: float, uy: float,
        primary_x: float, primary_y: float, pair_gap: float = 0.0,
    ) -> tuple[float, float, float]:
        """창호 부호/좌표점을 실 내부가 아닌 벽 또는 외측에 둔다.

        양면 벽이 검출되면 벽체 중간축을 그대로 사용한다. 한쪽 실구분선만 있는
        외곽벽은 벽 법선 양쪽을 검사해 실이 없는 쪽으로 기본 180mm 띄운다.
        양쪽 모두 실인 단일 공용선은 실내로 밀지 않고 경계선 위를 사용한다.
        """
        cx, cy = float(center_x), float(center_y)
        # 양쪽 벽 사이의 중심은 벽체 내부이므로 실내 좌표가 아니면 그대로 둔다.
        if float(pair_gap or 0.0) > 0.0 and not self._opening_point_inside_any_room(floor, cx, cy):
            return cx, cy, 0.0

        norm_len = math.hypot(float(ux), float(uy))
        if norm_len <= 1e-9:
            return cx, cy, 0.0
        ux, uy = float(ux) / norm_len, float(uy) / norm_len
        nx, ny = -uy, ux
        bx, by = float(primary_x), float(primary_y)
        # 작은 벽두께/도면 오차를 넘기도록 180mm부터 검사한다.
        for offset in (180.0, 250.0, 350.0, 500.0):
            candidates = [
                (bx + nx * offset, by + ny * offset, offset),
                (bx - nx * offset, by - ny * offset, -offset),
            ]
            for tx, ty, signed in candidates:
                if not self._opening_point_inside_any_room(floor, tx, ty):
                    return tx, ty, signed

        # 단일 경계 양쪽이 모두 실이면 임의로 어느 실 안에 넣지 않고 벽선 자체를 사용한다.
        return bx, by, 0.0

    def _opening_wall_attachment(
        self, floor: str, px: float, py: float, opening_length_mm: float = 0.0, snap_step: float = 10.0
    ) -> dict | None:
        """창호 중심을 가장 가까운 벽에 붙이고, 가까운 평행벽이 있으면 두 벽의 중간축에 배치한다.

        `opening_length_mm`는 적용 창길이 우선/검출길이 대체값으로 전달한다.
        창호 길이가 벽 구간보다 짧으면 중심이 벽 끝을 넘어가지 않도록 반길이만큼 여유를 둔다.
        """
        segments = self._opening_wall_segments(floor)
        if not segments:
            return None

        nearest = None
        for segment in segments:
            sx, sy, t_raw, seg_len, distance = self._project_opening_point_to_segment_full(
                px, py, segment["p1"], segment["p2"]
            )
            if seg_len <= 1e-9:
                continue
            if nearest is None or distance < nearest[0]:
                nearest = (distance, sx, sy, t_raw, seg_len, segment)
        if nearest is None:
            return None

        raw_distance, _sx, _sy, _t_raw, seg_len, primary = nearest
        x1, y1 = primary["p1"]
        x2, y2 = primary["p2"]
        ux, uy = (float(x2) - float(x1)) / seg_len, (float(y2) - float(y1)) / seg_len
        along = (float(px) - float(x1)) * ux + (float(py) - float(y1)) * uy
        if snap_step > 0:
            along = round(along / float(snap_step)) * float(snap_step)
        along = max(0.0, min(seg_len, along))

        try:
            use_length = max(0.0, float(opening_length_mm or 0.0))
        except (TypeError, ValueError):
            use_length = 0.0
        if 0.0 < use_length <= seg_len:
            half = use_length / 2.0
            along = max(half, min(seg_len - half, along))

        q1x = float(x1) + ux * along
        q1y = float(y1) + uy * along

        # 실구분선이 벽 양면으로 존재하는 경우 20~800mm 범위의 가장 가까운 평행선을
        # 반대쪽 벽면으로 잡는다. 복도 건너편 등 먼 평행선은 벽쌍으로 오인하지 않는다.
        secondary = None
        best_pair_gap = float("inf")
        max_pair_gap = 800.0
        sin_limit = math.sin(math.radians(7.0))
        for candidate in segments:
            if candidate is primary:
                continue
            c1x, c1y = candidate["p1"]
            c2x, c2y = candidate["p2"]
            cdx, cdy = float(c2x) - float(c1x), float(c2y) - float(c1y)
            clen = math.hypot(cdx, cdy)
            if clen <= 1e-9:
                continue
            vux, vuy = cdx / clen, cdy / clen
            cross = abs(ux * vuy - uy * vux)
            if cross > sin_limit:
                continue
            q2x, q2y, t2_raw, _cseg_len, _dist2 = self._project_opening_point_to_segment_full(
                q1x, q1y, candidate["p1"], candidate["p2"]
            )
            if t2_raw < -0.08 or t2_raw > 1.08:
                continue
            gap_dx, gap_dy = q2x - q1x, q2y - q1y
            gap = math.hypot(gap_dx, gap_dy)
            if gap < 20.0 or gap > max_pair_gap:
                continue
            # 두 평행선이 길이방향으로 크게 어긋난 경우는 동일 벽 양면으로 보지 않는다.
            along_offset = abs(gap_dx * ux + gap_dy * uy)
            if along_offset > max(80.0, gap * 0.35):
                continue
            if gap < best_pair_gap:
                best_pair_gap = gap
                secondary = (candidate, q2x, q2y)

        result = {
            "x": q1x, "y": q1y,
            "wall_center_x": q1x, "wall_center_y": q1y,
            "primary_x": q1x, "primary_y": q1y,
            "primary_p1": (float(primary["p1"][0]), float(primary["p1"][1])),
            "primary_p2": (float(primary["p2"][0]), float(primary["p2"][1])),
            "primary_refs": list(primary.get("refs", []) or []),
            "primary_geom_key": repr(primary.get("geom_key")),
            "ux": ux, "uy": uy,
            "wall_distance": float(raw_distance),
            "pair_gap": 0.0,
            "secondary_refs": [],
        }
        if secondary is not None:
            second_segment, q2x, q2y = secondary
            center_x, center_y = (q1x + q2x) / 2.0, (q1y + q2y) / 2.0
            result.update({
                "x": center_x,
                "y": center_y,
                "wall_center_x": center_x,
                "wall_center_y": center_y,
                "secondary_x": q2x, "secondary_y": q2y,
                "secondary_p1": (float(second_segment["p1"][0]), float(second_segment["p1"][1])),
                "secondary_p2": (float(second_segment["p2"][0]), float(second_segment["p2"][1])),
                "secondary_refs": list(second_segment.get("refs", []) or []),
                "secondary_geom_key": repr(second_segment.get("geom_key")),
                "pair_gap": float(best_pair_gap),
            })

        display_x, display_y, offset_mm = self._opening_display_coordinate(
            floor, float(result["wall_center_x"]), float(result["wall_center_y"]), ux, uy,
            q1x, q1y, float(result.get("pair_gap", 0.0) or 0.0),
        )
        result["x"] = float(display_x)
        result["y"] = float(display_y)
        result["coordinate_offset_mm"] = float(offset_mm)
        return result

    def _apply_opening_attachment(self, marker: dict, attachment: dict | None) -> None:
        if not attachment:
            marker["wall_snapped"] = False
            marker["wall_refs"] = []
            marker["wall_geom_key"] = ""
            marker.pop("wall2_geom_key", None)
            marker.pop("wall_attach_primary", None)
            marker.pop("wall_attach_secondary", None)
            marker.pop("wall_center_x", None)
            marker.pop("wall_center_y", None)
            marker.pop("coordinate_offset_mm", None)
            return
        marker["x"] = float(attachment.get("x", marker.get("x", 0.0)))
        marker["y"] = float(attachment.get("y", marker.get("y", 0.0)))
        marker["wall_center_x"] = float(attachment.get("wall_center_x", attachment.get("x", marker.get("x", 0.0))))
        marker["wall_center_y"] = float(attachment.get("wall_center_y", attachment.get("y", marker.get("y", 0.0))))
        marker["coordinate_offset_mm"] = float(attachment.get("coordinate_offset_mm", 0.0) or 0.0)
        marker["wall_snapped"] = True
        marker["wall_distance"] = float(attachment.get("wall_distance", 0.0) or 0.0)
        primary_refs = list(attachment.get("primary_refs", []) or [])
        secondary_refs = list(attachment.get("secondary_refs", []) or [])
        marker["wall_refs"] = primary_refs + [ref for ref in secondary_refs if ref not in primary_refs]
        marker["wall_geom_key"] = str(attachment.get("primary_geom_key", "") or "")
        marker["wall2_geom_key"] = str(attachment.get("secondary_geom_key", "") or "")
        marker["wall_direction"] = [float(attachment.get("ux", 1.0)), float(attachment.get("uy", 0.0))]
        marker["wall_pair_gap"] = float(attachment.get("pair_gap", 0.0) or 0.0)
        marker["wall_attach_primary"] = [
            float(attachment.get("primary_x", marker["x"])), float(attachment.get("primary_y", marker["y"]))
        ]
        if "secondary_x" in attachment and "secondary_y" in attachment:
            marker["wall_attach_secondary"] = [
                float(attachment.get("secondary_x")), float(attachment.get("secondary_y"))
            ]
        else:
            marker.pop("wall_attach_secondary", None)

    def _opening_wall_segments(self, floor: str):
        """현재층 실 외곽 벽구간을 형상별로 묶는다.

        문/창 TEST 중심을 공제 가능한 벽 위치에 붙이기 위한 임시 인덱스다.
        CAD를 다시 읽지 않고 이미 캐시에 있는 실 폴리곤만 사용한다.
        """
        segments = {}
        for room in self.rooms:
            if str(room.floor or "") != str(floor or ""):
                continue
            points = list(room.points or [])
            if len(points) < 2:
                continue
            rid = room_program_id(room)
            for index, p1 in enumerate(points):
                p2 = points[(index + 1) % len(points)]
                try:
                    a = (round(float(p1[0]), 3), round(float(p1[1]), 3))
                    b = (round(float(p2[0]), 3), round(float(p2[1]), 3))
                except (TypeError, ValueError, IndexError):
                    continue
                geom_key = tuple(sorted((a, b)))
                entry = segments.setdefault(geom_key, {
                    "p1": (float(p1[0]), float(p1[1])),
                    "p2": (float(p2[0]), float(p2[1])),
                    "refs": [],
                    "geom_key": geom_key,
                })
                entry["refs"].append({
                    "room_id": rid,
                    "room_handle": str(room.boundary_handle or ""),
                    "room_number": str(room.room_number or ""),
                    "room_name": str(room.room_name or ""),
                    "segment_index": int(index),
                    "wall_id": wall_program_id(rid, p1, p2),
                })
        return list(segments.values())

    def _snap_opening_test_markers_to_walls(self, floor: str, markers):
        """TEST 후보를 가장 가까운 현재층 실 벽구간으로 스냅하고 과검출을 거른다."""
        segments = self._opening_wall_segments(floor)
        if not segments:
            return [dict(value or {}) for value in (markers or [])], 0, 0

        snapped = []
        rejected = 0
        for raw in markers or []:
            marker = dict(raw or {})
            try:
                px, py = float(marker.get("x")), float(marker.get("y"))
            except (TypeError, ValueError):
                continue
            nearest = None
            for segment in segments:
                sx, sy, distance = self._project_opening_point_to_segment(px, py, segment["p1"], segment["p2"] )
                if nearest is None or distance < nearest[0]:
                    nearest = (distance, sx, sy, segment)
            if nearest is None:
                continue

            distance, sx, sy, segment = nearest
            # 블록 중심은 문짝 회전호 때문에 벽에서 조금 떨어질 수 있다.
            # 창은 보통 벽 중심에 더 가까워야 하므로 문보다 엄격하게 제한한다.
            kind = str(marker.get("kind", "") or "")
            bounds = marker.get("bounds")
            longest = 0.0
            if isinstance(bounds, (list, tuple)) and len(bounds) == 4:
                try:
                    longest = max(abs(float(bounds[2]) - float(bounds[0])), abs(float(bounds[3]) - float(bounds[1])))
                except (TypeError, ValueError):
                    longest = 0.0
            base_limit = 1800.0 if kind == "door" else 1200.0
            if longest > 0:
                base_limit = max(base_limit, min(2400.0, longest * (0.85 if kind == "door" else 0.55)))
            if distance > base_limit:
                rejected += 1
                continue

            marker["source_x"] = px
            marker["source_y"] = py
            marker["x"] = sx
            marker["y"] = sy
            marker["wall_snapped"] = True
            marker["wall_distance"] = float(distance)
            marker["wall_refs"] = list(segment.get("refs", []) or [])
            marker["wall_geom_key"] = repr(segment.get("geom_key"))
            # 검출 창길이는 가능하면 해당 벽 방향으로 후보 외곽을 투영한 길이를 사용한다.
            # 블록/선군의 축정렬 BoundingBox 최대값보다 회전된 창호에서 덜 과대평가된다.
            if isinstance(bounds, (list, tuple)) and len(bounds) == 4:
                try:
                    x1, y1 = segment["p1"]
                    x2, y2 = segment["p2"]
                    dx, dy = float(x2) - float(x1), float(y2) - float(y1)
                    seg_len = math.hypot(dx, dy)
                    if seg_len > 1e-9:
                        ux, uy = dx / seg_len, dy / seg_len
                        corners = [
                            (float(bounds[0]), float(bounds[1])),
                            (float(bounds[0]), float(bounds[3])),
                            (float(bounds[2]), float(bounds[1])),
                            (float(bounds[2]), float(bounds[3])),
                        ]
                        projections = [cx * ux + cy * uy for cx, cy in corners]
                        projected = max(projections) - min(projections)
                        if 50.0 <= projected <= 30000.0 and str(marker.get("source", "")) != "arc90":
                            marker["detected_length_mm"] = float(projected)
                except Exception:
                    pass
            attachment = self._opening_wall_attachment(
                floor, px, py, self._opening_effective_length_mm(marker), 10.0
            )
            if attachment is not None:
                # 검출점이 벽 사이에 있으면 두 평행 벽면의 중간축으로 두고, 양쪽 벽 참조를 함께 보존한다.
                self._apply_opening_attachment(marker, attachment)
                marker["source_x"] = px
                marker["source_y"] = py
                marker["wall_distance"] = float(distance)
            snapped.append(marker)

        # 블록/보조선 등으로 같은 개구부가 중복 검출되어도 벽에 스냅한 뒤 가까우면 하나만 유지한다.
        snapped.sort(key=lambda value: -int(value.get("confidence", 0) or 0))
        deduped = []
        duplicate_count = 0
        for marker in snapped:
            duplicate = False
            for existing in deduped:
                if str(existing.get("kind")) != str(marker.get("kind")):
                    continue
                same_wall = bool(existing.get("wall_geom_key")) and existing.get("wall_geom_key") == marker.get("wall_geom_key")
                mx = float(marker.get("wall_center_x", marker.get("x", 0.0)) or 0.0)
                my = float(marker.get("wall_center_y", marker.get("y", 0.0)) or 0.0)
                ex = float(existing.get("wall_center_x", existing.get("x", 0.0)) or 0.0)
                ey = float(existing.get("wall_center_y", existing.get("y", 0.0)) or 0.0)
                distance = math.hypot(mx - ex, my - ey)
                if same_wall and distance <= 450.0:
                    duplicate = True
                    duplicate_count += 1
                    break
            if not duplicate:
                deduped.append(marker)

        # TEST 자동부호는 필터링 후 번호가 끊기지 않도록 다시 붙인다. 실제 블록부호(AW01 등)는 유지한다.
        deduped.sort(key=lambda value: (
            -float(value.get("wall_center_y", value.get("y", 0.0)) or 0.0),
            float(value.get("wall_center_x", value.get("x", 0.0)) or 0.0),
            str(value.get("kind", "")),
        ))
        door_no = 0
        window_no = 0
        for marker in deduped:
            if not bool(marker.get("auto_label")):
                continue
            if marker.get("kind") == "door":
                door_no += 1
                marker["label"] = f"D{door_no:02d}"
            else:
                window_no += 1
                marker["label"] = f"W{window_no:02d}"
        return deduped, rejected, duplicate_count

    @staticmethod
    def _opening_read_id(floor: str, marker: dict) -> str:
        try:
            x = round(float(marker.get("wall_center_x", marker.get("x", 0.0)) or 0.0) / 10.0) * 10.0
            y = round(float(marker.get("wall_center_y", marker.get("y", 0.0)) or 0.0) / 10.0) * 10.0
        except (TypeError, ValueError):
            x = y = 0.0
        seed = "|".join([
            str(floor or ""), str(marker.get("kind", "") or ""),
            str(marker.get("wall_geom_key", "") or ""), f"{x:.0f}", f"{y:.0f}",
            str(marker.get("source", "") or ""), str(marker.get("block_name", "") or ""),
        ])
        return hashlib.sha1(seed.encode("utf-8")).hexdigest()[:20]

    def _opening_display_markers(self, floor: str) -> list[dict]:
        result = []
        floor = str(floor or "")
        for raw in self.opening_read_results.get(floor, []) or []:
            # 적용 창길이가 있으면 그것을 사용하고, 없거나 0이면 검출 창길이를 사용한다.
            marker = dict(raw or {})
            marker["label"] = str(marker.get("applied_code", "") or marker.get("read_label", "") or marker.get("label", ""))
            marker["display_length_mm"] = self._opening_effective_length_mm(marker)
            result.append(marker)
        return result

    def _prepare_opening_read_preview(self, floor: str) -> None:
        floor = str(floor or "")
        markers = self._opening_display_markers(floor)
        self.preview.opening_test_floor = floor if markers else ""
        self.preview.opening_test_markers = markers

    def _sync_opening_read_preview(self) -> None:
        floor = str(self.selected_floor() or "")
        self.preview.set_opening_test_markers(floor, self._opening_display_markers(floor))

    def set_opening_read_results(self, results: dict, scope: str = "current", errors=None, missing_floors=None) -> None:
        total_scanned = total_candidates = total_rejected = total_duplicates = 0
        old_by_floor = {}
        for floor, markers in dict(self.opening_read_results or {}).items():
            old_by_floor[str(floor)] = {str(item.get("read_id", "") or ""): dict(item) for item in (markers or [])}

        for floor, result in dict(results or {}).items():
            floor = str(floor or "")
            payload = dict(result or {})
            raw_markers = list(payload.get("candidates", []) or [])
            markers, wall_rejected, duplicate_count = self._snap_opening_test_markers_to_walls(floor, raw_markers)
            previous = old_by_floor.get(floor, {})
            normalized = []
            for raw in markers:
                marker = dict(raw or {})
                read_id = self._opening_read_id(floor, marker)
                marker["read_id"] = read_id
                marker["read_label"] = str(marker.get("label", "") or "")
                old = previous.get(read_id, {})
                if not old and previous:
                    # 벽 양면 중간축 부착 방식 변경으로 read_id 좌표가 달라져도
                    # 같은 CAD 후보의 기존 적용 부호/길이/수동이동 정보는 최대한 승계한다.
                    best_old = None
                    best_distance = float("inf")
                    try:
                        nx = float(marker.get("source_x", marker.get("x", 0.0)) or 0.0)
                        ny = float(marker.get("source_y", marker.get("y", 0.0)) or 0.0)
                    except (TypeError, ValueError):
                        nx = ny = 0.0
                    for candidate_old in previous.values():
                        if str(candidate_old.get("kind", "") or "") != str(marker.get("kind", "") or ""):
                            continue
                        if str(candidate_old.get("source", "") or "") != str(marker.get("source", "") or ""):
                            continue
                        if str(candidate_old.get("block_name", "") or "") != str(marker.get("block_name", "") or ""):
                            continue
                        try:
                            ox = float(candidate_old.get("source_x", candidate_old.get("x", 0.0)) or 0.0)
                            oy = float(candidate_old.get("source_y", candidate_old.get("y", 0.0)) or 0.0)
                        except (TypeError, ValueError):
                            continue
                        dist = math.hypot(nx - ox, ny - oy)
                        if dist < best_distance:
                            best_distance = dist; best_old = candidate_old
                    if best_old is not None and best_distance <= 600.0:
                        old = dict(best_old)
                if bool(old.get("manual_read_label")) and str(old.get("read_label", "") or "").strip():
                    marker["read_label"] = str(old.get("read_label") or "").strip().upper()
                    marker["manual_read_label"] = True
                marker["applied_code"] = str(old.get("applied_code", "") or "")
                marker["applied_length_mm"] = old.get("applied_length_mm", "")
                # 같은 CAD 후보를 다시 읽어도 사용자가 창호 편집에서 수동 이동한 위치는 유지한다.
                if bool(old.get("manual_moved")):
                    try:
                        marker["x"] = float(old.get("x", marker.get("x", 0.0)))
                        marker["y"] = float(old.get("y", marker.get("y", 0.0)))
                        marker["wall_center_x"] = float(old.get("wall_center_x", old.get("x", marker.get("x", 0.0))))
                        marker["wall_center_y"] = float(old.get("wall_center_y", old.get("y", marker.get("y", 0.0))))
                        marker["wall_snapped"] = bool(old.get("wall_snapped", True))
                        marker["manual_moved"] = True
                        self._refresh_opening_marker_wall_link(floor, marker)
                    except (TypeError, ValueError):
                        pass
                normalized.append(marker)
            self.opening_read_results[floor] = normalized
            total_scanned += int(payload.get("scanned", 0) or 0)
            total_candidates += len(normalized)
            total_rejected += int(wall_rejected)
            total_duplicates += int(duplicate_count)

        self.save_finish_assignments()
        self._sync_opening_read_preview()
        self._update_opening_edit_controls()
        error_list = [str(value) for value in (errors or []) if str(value)]
        missing = [str(value) for value in (missing_floors or []) if str(value)]
        parts = [f"검출 {total_candidates:,}개", f"CAD 객체 {total_scanned:,}개"]
        if total_duplicates:
            parts.append(f"중복정리 {total_duplicates:,}개")
        if total_rejected:
            parts.append(f"벽거리 제외 {total_rejected:,}개")
        if error_list:
            parts.append(f"오류 {len(error_list):,}개 층")
        if missing:
            parts.append(f"도곽누락 {len(missing):,}개 층")
        status = " / ".join(parts)
        if self._opening_read_dialog is not None:
            self._opening_read_dialog.refresh(status)
        if self._opening_assign_dialog is not None:
            self._opening_assign_dialog.refresh(status)

    def set_opening_test_results(self, floor: str, result) -> None:
        """구버전 호출 호환."""
        self.set_opening_read_results({str(floor or ""): result}, "current")

    def _find_opening_marker(self, floor: str, read_id: str):
        for marker in self.opening_read_results.get(str(floor or ""), []) or []:
            if str(marker.get("read_id", "") or "") == str(read_id or ""):
                return marker
        return None

    def apply_opening_read_value(self, dialog_rows, field: str, value, dialog: OpeningReadDialog | None = None) -> None:
        target_dialog = dialog or self._opening_read_dialog
        changed = 0
        if target_dialog is None:
            return
        master = None
        if field == "applied_code":
            code = str(value or "").strip().upper()
            if code:
                master = self._window_catalog_by_code().get(code)
                if master is None:
                    QMessageBox.warning(target_dialog, "창호 지정", "창호 리스트에 없는 부호는 입력할 수 없습니다.")
                    target_dialog.refresh()
                    return
            value = code
        for row in dialog_rows or []:
            floor, read_id = target_dialog._row_key(int(row))
            marker = self._find_opening_marker(floor, read_id)
            if marker is None:
                continue
            marker[field] = value
            if field == "applied_code":
                if master is not None:
                    try:
                        length_mm = float(master.get("applied_length_mm", 0.0) or 0.0)
                    except Exception:
                        length_mm = 0.0
                    marker["applied_length_mm"] = length_mm if length_mm > 0 else ""
                    dw = str(master.get("dw") or "").upper()
                    if dw == "D":
                        marker["kind"] = "door"
                    elif dw == "W":
                        marker["kind"] = "window"
                else:
                    marker["applied_length_mm"] = ""
                self._refresh_opening_marker_wall_link(floor, marker)
            elif field == "applied_length_mm":
                self._refresh_opening_marker_wall_link(floor, marker)
            changed += 1
        if not changed:
            return
        self.save_finish_assignments()
        self._sync_opening_read_preview()
        self.data_changed.emit()
        target_dialog.refresh(f"선택 {changed:,}개 항목 변경")

    def delete_opening_read_rows(self, dialog_rows, dialog: OpeningReadDialog | None = None) -> None:
        target_dialog = dialog or self._opening_read_dialog
        if target_dialog is None:
            return
        keys = {target_dialog._row_key(int(row)) for row in (dialog_rows or [])}
        keys.discard(("", ""))
        if not keys:
            return
        deleted = 0
        for floor in list(self.opening_read_results):
            before = list(self.opening_read_results.get(floor, []) or [])
            after = [marker for marker in before if (str(floor), str(marker.get("read_id", "") or "")) not in keys]
            deleted += len(before) - len(after)
            if after:
                self.opening_read_results[floor] = after
            else:
                self.opening_read_results.pop(floor, None)
        if deleted:
            self.save_finish_assignments()
            self._sync_opening_read_preview()
            self._update_opening_edit_controls()
            target_dialog.refresh(f"선택 {deleted:,}개 항목 삭제")

    def selected_rooms(self):
        selected = []
        seen = set()
        rows = sorted({idx.row() for idx in self.room_table.selectionModel().selectedRows()})
        for table_row in rows:
            room_index = self._room_index_for_table_row(table_row)
            if room_index is not None:
                room = self.rooms[room_index]
                if id(room) not in seen:
                    selected.append(room)
                    seen.add(id(room))
                continue
            floor = self._floor_for_table_row(table_row)
            for room in self.rooms:
                if room.floor == floor and id(room) not in seen:
                    selected.append(room)
                    seen.add(id(room))
        return selected

    def request_create_blocks(self):
        if not self.current_source_path or not self.rooms:
            return

        current_floor = self.selected_floor()
        current_floor_rooms = [room for room in self.rooms if room.floor == current_floor]
        selected_indices = self._selected_room_indices()
        selected_rooms = [self.rooms[index] for index in selected_indices]

        dialog = QDialog(self)
        dialog.setWindowTitle("실명블록 생성")
        dialog.setModal(True)
        dialog.setFixedWidth(360)
        layout = QVBoxLayout(dialog)
        scope_box = QGroupBox("생성 범위")
        scope_layout = QVBoxLayout(scope_box)
        scope_group = QButtonGroup(dialog)
        current_floor_radio = QRadioButton(f"현재층 ({current_floor or '-'}, {len(current_floor_rooms)}개)")
        all_floors_radio = QRadioButton(f"전체층 ({len(self.rooms)}개)")
        selected_radio = QRadioButton(f"선택실 ({len(selected_rooms)}개)")
        scope_group.addButton(current_floor_radio)
        scope_group.addButton(all_floors_radio)
        scope_group.addButton(selected_radio)
        scope_layout.addWidget(current_floor_radio)
        scope_layout.addWidget(all_floors_radio)
        scope_layout.addWidget(selected_radio)
        current_floor_radio.setEnabled(bool(current_floor_rooms))
        selected_radio.setEnabled(bool(selected_rooms))
        if current_floor_rooms:
            current_floor_radio.setChecked(True)
        elif selected_rooms:
            selected_radio.setChecked(True)
        else:
            all_floors_radio.setChecked(True)
        layout.addWidget(scope_box)
        note = QLabel("※ 선택 범위 중 실명블록이 없는 실만 생성합니다.")
        note.setStyleSheet("color:#666666;")
        layout.addWidget(note)
        buttons = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        buttons.button(QDialogButtonBox.Ok).setText("생성")
        buttons.button(QDialogButtonBox.Cancel).setText("취소")
        buttons.accepted.connect(dialog.accept)
        buttons.rejected.connect(dialog.reject)
        layout.addWidget(buttons)
        if dialog.exec() != QDialog.Accepted:
            return

        if selected_radio.isChecked():
            scope_rooms = selected_rooms
        elif all_floors_radio.isChecked():
            scope_rooms = list(self.rooms)
        else:
            scope_rooms = current_floor_rooms
        targets = [room for room in scope_rooms if not room.tag_handle]
        if not targets:
            QMessageBox.information(self, "실명블록 생성", "선택한 범위에는 실명블록이 없는 실이 없습니다.")
            return

        table_state = self._capture_room_table_state()
        try:
            if self.cad_connected_mode:
                self.cad_create_blocks_requested.emit(targets)
                return
            result = self._run_cached_dxf_operation("실명블록 생성", create_room_blocks, targets)
            if result:
                _out, report = result
                QMessageBox.information(
                    self,
                    "실명블록 생성 완료",
                    f"생성: {report['created']}/{report['requested']}개\n"
                    f"기존 블록 있음: {report['already_exists']}개\n"
                    f"경계 누락: {report['boundary_missing']}개\n\n"
                    "내부 작업 캐시에 반영했습니다.",
                )
        finally:
            # CAD 생성 과정에서 목록이 재구성되어도 원래 행·다중 선택·스크롤을 되돌린다.
            self._restore_room_table_state(table_state)

    def _execute_refresh_blocks(self, targets, scope: str, current_floor: str, apply_room_tags: bool, apply_finishes: bool):
        """CAD 적용창과 통합 CAD 연동창에서 공통으로 사용하는 실제 적용 처리."""
        targets = list(targets or [])
        if not targets:
            QMessageBox.information(self, 'CAD 연동', '선택한 범위에 적용할 실이 없습니다.')
            return
        if not apply_room_tags and not apply_finishes:
            QMessageBox.information(self, 'CAD 연동', '적용할 항목을 하나 이상 선택하십시오.')
            return

        payload = {
            'targets': targets,
            'scope': str(scope or ''),
            'floor': current_floor if scope == 'floor' else '',
            'apply_room_tags': bool(apply_room_tags),
            'apply_finishes': bool(apply_finishes),
        }
        table_state = self._capture_room_table_state()
        if self.cad_connected_mode:
            try:
                self.cad_refresh_blocks_requested.emit(payload)
            finally:
                # CAD 적용 중 목록/배경이 갱신돼도 실행 전 행과 스크롤 위치로 돌아간다.
                self._restore_room_table_state(table_state)
            return

        # DXF 캐시 모드에서는 기존 실명 속성 갱신만 지원한다.
        if apply_finishes and not apply_room_tags:
            QMessageBox.information(self, 'CAD 연동', '마감 적용은 CAD 연결 상태에서 사용할 수 있습니다.')
            self._restore_room_table_state(table_state)
            return
        room_targets = [room for room in targets if room.tag_handle]
        if not room_targets:
            QMessageBox.information(self, 'CAD 연동', 'CAD에 적용할 실명블록이 없습니다.')
            self._restore_room_table_state(table_state)
            return
        result = self._run_cached_dxf_operation('수정사항 CAD에 적용', refresh_room_blocks, room_targets)
        if result:
            _out, report = result
            QMessageBox.information(
                self, 'CAD 적용 완료',
                f"새로고침: {report['refreshed']}/{report['requested']}개\n"
                f"블록 누락: {report['missing_insert']}개\n\n내부 작업 캐시에 반영했습니다."
            )
        self._restore_room_table_state(table_state)

    def open_cad_link_dialog(self):
        """수정사항 적용과 실/벽 연동을 한 창의 탭으로 통합한다."""
        if not self.rooms:
            QMessageBox.information(self, "CAD 연동", "먼저 평면 실목록을 읽어 주세요.")
            return

        dialog = QDialog(self)
        dialog.setWindowTitle("CAD 연동")
        dialog.resize(680, 540)
        root = QVBoxLayout(dialog)
        tabs = QTabWidget()
        root.addWidget(tabs, 1)

        # ------------------------------------------------------------------
        # 1) 수정사항 적용
        apply_page = QWidget()
        apply_layout = QVBoxLayout(apply_page)
        current_floor = str(self.selected_floor() or "")
        current_floor_targets = [room for room in self.rooms if str(room.floor or "") == current_floor]
        selected_indices = self._selected_room_indices()
        selected_targets = [self.rooms[index] for index in selected_indices]

        scope_box = QGroupBox("적용 범위")
        scope_layout = QVBoxLayout(scope_box)
        scope_group = QButtonGroup(dialog)
        current_floor_radio = QRadioButton(f"현재층 ({current_floor or '-'}, {len(current_floor_targets)}개)")
        all_floors_radio = QRadioButton(f"전체층 ({len(self.rooms)}개)")
        current_floor_radio.setEnabled(bool(current_floor_targets))
        scope_group.addButton(current_floor_radio)
        scope_group.addButton(all_floors_radio)
        scope_layout.addWidget(current_floor_radio)
        scope_layout.addWidget(all_floors_radio)
        selected_radio = None
        if selected_targets:
            selected_radio = QRadioButton(f"선택목록 ({len(selected_targets)}개)")
            scope_group.addButton(selected_radio)
            scope_layout.addWidget(selected_radio)
        if current_floor_targets:
            current_floor_radio.setChecked(True)
        elif selected_radio is not None:
            selected_radio.setChecked(True)
        else:
            all_floors_radio.setChecked(True)
        apply_layout.addWidget(scope_box)

        room_tag_check = QCheckBox("실명 적용")
        room_tag_check.setChecked(True)
        finish_check = QCheckBox("마감 적용")
        finish_check.setChecked(False)
        finish_check.setToolTip("현재는 벽체 구간선 CAD 적용을 지원합니다.")
        apply_layout.addWidget(room_tag_check)
        apply_layout.addWidget(finish_check)
        apply_note = QLabel("※ 실명/마감 수정사항을 선택 범위의 연결 CAD에 반영합니다.")
        apply_note.setStyleSheet("color:#666666;")
        apply_layout.addWidget(apply_note)
        apply_layout.addStretch(1)
        apply_button = QPushButton("선택 내용 CAD에 적용")
        apply_layout.addWidget(apply_button)
        tabs.addTab(apply_page, "수정사항 적용")

        def apply_changes():
            if selected_radio is not None and selected_radio.isChecked():
                targets = selected_targets
                scope = "selected"
            elif all_floors_radio.isChecked():
                targets = list(self.rooms)
                scope = "all"
            else:
                targets = current_floor_targets
                scope = "floor"
            self._execute_refresh_blocks(
                targets, scope, current_floor,
                room_tag_check.isChecked(), finish_check.isChecked(),
            )

        apply_button.clicked.connect(apply_changes)

        # ------------------------------------------------------------------
        # 2) 실/벽 연동
        link_page = QWidget()
        link_layout = QVBoxLayout(link_page)
        link_layers = room_wall_link_layers(self.room_rules)
        info = QLabel(
            "프로그램에서 지정한 그룹을 CAD/DXF에 덧그리거나, CAD/DXF의 색상별 구획선을 그룹으로 가져옵니다.\n"
            "벽 연동선은 실구획선의 절점·길이와 1:1로 맞아야 하며 완전중복/부분중복도 검사합니다."
        )
        info.setWordWrap(True)
        link_layout.addWidget(info)

        layer_box = QGroupBox("CAD 연동 레이어")
        layer_layout = QGridLayout(layer_box)
        for row, section in enumerate(("벽", "바닥", "천정", "기타")):
            layer_layout.addWidget(QLabel(section), row, 0)
            label = QLabel(link_layers[section])
            label.setTextInteractionFlags(Qt.TextSelectableByMouse)
            layer_layout.addWidget(label, row, 1)
        link_layout.addWidget(layer_box)

        link_scope_box = QGroupBox("내보내기 범위")
        link_scope_layout = QHBoxLayout(link_scope_box)
        link_scope_layout.addWidget(QLabel("층:"))
        scope_combo = QComboBox()
        scope_combo.addItem("전체층", "")
        floors = []
        for room in self.rooms:
            floor = str(room.floor or "")
            if floor and floor not in floors:
                floors.append(floor)
        for floor in sorted(floors, key=floor_sort_key):
            scope_combo.addItem(floor, floor)
        if current_floor:
            index = scope_combo.findData(current_floor)
            if index >= 0:
                scope_combo.setCurrentIndex(index)
        link_scope_layout.addWidget(scope_combo)
        section_checks = {}
        for section in ("벽", "바닥", "천정", "기타"):
            check = QCheckBox(section)
            check.setChecked(section in {"벽", "바닥", "천정"})
            section_checks[section] = check
            link_scope_layout.addWidget(check)
        link_scope_layout.addStretch(1)
        link_layout.addWidget(link_scope_box)

        actions = QGridLayout()
        export_cad = QPushButton("프로그램 → 연결 CAD")
        export_dxf = QPushButton("프로그램 → DXF 저장")
        import_cad = QPushButton("연결 CAD → 프로그램")
        import_dxf = QPushButton("DXF → 프로그램")
        validate_cad = QPushButton("연결 CAD 구간검증")
        validate_dxf = QPushButton("DXF 구간검증")
        export_cad.setEnabled(bool(self.cad_connected_mode))
        import_cad.setEnabled(bool(self.cad_connected_mode))
        validate_cad.setEnabled(bool(self.cad_connected_mode))
        actions.addWidget(export_cad, 0, 0)
        actions.addWidget(export_dxf, 0, 1)
        actions.addWidget(import_cad, 1, 0)
        actions.addWidget(import_dxf, 1, 1)
        actions.addWidget(validate_cad, 2, 0)
        actions.addWidget(validate_dxf, 2, 1)
        link_layout.addLayout(actions)
        link_note = QLabel(
            "※ 프로그램에서 내보낸 객체는 그룹 ID도 함께 기록하므로 다시 읽으면 원래 그룹을 우선 인식합니다.\n"
            "※ CAD에서 새로 그린 객체는 색상별로 묶어 기존 그룹 또는 새 그룹에 연결합니다."
        )
        link_note.setWordWrap(True)
        link_layout.addWidget(link_note)
        link_layout.addStretch(1)
        tabs.addTab(link_page, "실/벽 연동")

        def selected_sections():
            return [name for name, check in section_checks.items() if check.isChecked()]

        def build_items():
            sections = selected_sections()
            if not sections:
                QMessageBox.information(dialog, "CAD 연동", "내보낼 부위를 하나 이상 선택해 주세요.")
                return []
            items = self._room_wall_link_export_items(str(scope_combo.currentData() or ""), sections)
            if not items:
                QMessageBox.information(dialog, "CAD 연동", "선택 범위에 CAD로 보낼 그룹 지정 객체가 없습니다.")
            return items

        def export_to_cad():
            items = build_items()
            if items:
                self.room_wall_link_cad_requested.emit({"action": "export", "items": items, "layers": link_layers})

        def export_to_dxf():
            items = build_items()
            if not items or not self.current_source_path:
                return
            source = Path(self.current_source_path)
            suggested = str(source.with_name(source.stem + "_실벽연동.dxf"))
            out, _ = QFileDialog.getSaveFileName(dialog, "실/벽 연동 DXF 저장", suggested, "DXF 파일 (*.dxf)")
            if not out:
                return
            try:
                report = write_room_wall_link_dxf(self.current_source_path, out, items, link_layers)
                QMessageBox.information(
                    dialog, "실/벽 연동 DXF 저장 완료",
                    f"생성: {report['created']:,}/{report['requested']:,}개\n"
                    f"ID 갱신: {report.get('updated', 0):,}개\n"
                    f"기존 동일객체 제외: {report.get('skipped_existing', 0):,}개\n"
                    f"실패: {report['failed']:,}개\n\n{out}"
                )
            except Exception:
                QMessageBox.critical(dialog, "실/벽 연동 DXF 저장 오류", traceback.format_exc())

        def import_from_dxf():
            start = str(Path(self.current_source_path).parent) if self.current_source_path else ""
            path, _ = QFileDialog.getOpenFileName(dialog, "실/벽 연동 DXF 읽기", start, "DXF 파일 (*.dxf)")
            if not path:
                return
            try:
                result = read_room_wall_link_dxf(path, link_layers)
                if not result.get("items"):
                    QMessageBox.information(dialog, "CAD 연동", "설정된 실/벽 연동 레이어에서 읽을 객체가 없습니다.")
                    return
                self._import_room_wall_link_items(result["items"], f"DXF: {Path(path).name}")
            except Exception:
                QMessageBox.critical(dialog, "실/벽 연동 DXF 읽기 오류", traceback.format_exc())

        def validate_dxf_file():
            start = str(Path(self.current_source_path).parent) if self.current_source_path else ""
            path, _ = QFileDialog.getOpenFileName(dialog, "실/벽 연동 DXF 구간검증", start, "DXF 파일 (*.dxf)")
            if not path:
                return
            try:
                result = read_room_wall_link_dxf(path, link_layers)
                report = validate_room_wall_link_items(result.get("items", []), self.rooms, 2.0)
                self._show_room_wall_link_validation(report, f"DXF: {Path(path).name}")
            except Exception:
                QMessageBox.critical(dialog, "실/벽 연동 구간검증 오류", traceback.format_exc())

        export_cad.clicked.connect(export_to_cad)
        export_dxf.clicked.connect(export_to_dxf)
        import_cad.clicked.connect(lambda: self.room_wall_link_cad_requested.emit({"action": "read", "layers": link_layers}))
        import_dxf.clicked.connect(import_from_dxf)
        validate_cad.clicked.connect(lambda: self.room_wall_link_cad_requested.emit({"action": "validate", "layers": link_layers}))
        validate_dxf.clicked.connect(validate_dxf_file)

        close_buttons = QDialogButtonBox(QDialogButtonBox.Close)
        close_buttons.button(QDialogButtonBox.Close).setText("닫기")
        close_buttons.rejected.connect(dialog.reject)
        close_buttons.button(QDialogButtonBox.Close).clicked.connect(dialog.accept)
        root.addWidget(close_buttons)
        dialog.exec()

    def request_refresh_blocks(self):
        if not self.current_source_path:
            return

        current_floor = self.selected_floor()
        current_floor_targets = [room for room in self.rooms if room.floor == current_floor]
        selected_indices = self._selected_room_indices()
        selected_targets = [self.rooms[index] for index in selected_indices]
        if not self.rooms:
            QMessageBox.information(self, '수정사항 CAD에 적용', 'CAD에 적용할 실이 없습니다.')
            return

        dialog = QDialog(self)
        dialog.setWindowTitle('수정사항 CAD에 적용')
        dialog.setModal(True)
        dialog.setFixedWidth(360)
        layout = QVBoxLayout(dialog)

        scope_box = QGroupBox('적용 범위')
        scope_layout = QVBoxLayout(scope_box)
        scope_group = QButtonGroup(dialog)
        current_floor_radio = QRadioButton(f'현재층 ({current_floor or "-"}, {len(current_floor_targets)}개)')
        all_floors_radio = QRadioButton(f'전체층 ({len(self.rooms)}개)')
        current_floor_radio.setChecked(True)
        scope_group.addButton(current_floor_radio)
        scope_group.addButton(all_floors_radio)
        scope_layout.addWidget(current_floor_radio)
        scope_layout.addWidget(all_floors_radio)
        selected_radio = None
        if selected_targets:
            selected_radio = QRadioButton(f'선택목록 ({len(selected_targets)}개)')
            scope_group.addButton(selected_radio)
            scope_layout.addWidget(selected_radio)
        layout.addWidget(scope_box)

        room_tag_check = QCheckBox('실명 적용')
        room_tag_check.setChecked(True)
        finish_check = QCheckBox('마감 적용')
        finish_check.setChecked(False)
        finish_check.setToolTip('현재는 벽체 구간선만 CAD에 적용합니다. 바닥/천장은 추후 이 항목에 확장합니다.')
        layout.addWidget(room_tag_check)
        layout.addWidget(finish_check)

        note = QLabel('※ 마감 적용: 현재 벽체 구간선만 적용')
        note.setStyleSheet('color:#666666;')
        layout.addWidget(note)

        buttons = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        buttons.accepted.connect(dialog.accept)
        buttons.rejected.connect(dialog.reject)
        layout.addWidget(buttons)
        if dialog.exec() != QDialog.Accepted:
            return

        if selected_radio is not None and selected_radio.isChecked():
            targets = selected_targets
            scope = 'selected'
        elif all_floors_radio.isChecked():
            targets = list(self.rooms)
            scope = 'all'
        else:
            targets = current_floor_targets
            scope = 'floor'

        if not targets:
            QMessageBox.information(self, '수정사항 CAD에 적용', '선택한 범위에 적용할 실이 없습니다.')
            return

        apply_room_tags = room_tag_check.isChecked()
        apply_finishes = finish_check.isChecked()
        if not apply_room_tags and not apply_finishes:
            QMessageBox.information(self, '수정사항 CAD에 적용', '적용할 항목을 하나 이상 선택하십시오.')
            return

        payload = {
            'targets': targets,
            'scope': scope,
            'floor': current_floor if scope == 'floor' else '',
            'apply_room_tags': apply_room_tags,
            'apply_finishes': apply_finishes,
        }
        table_state = self._capture_room_table_state()
        if self.cad_connected_mode:
            try:
                self.cad_refresh_blocks_requested.emit(payload)
            finally:
                # CAD 적용 중 목록/배경이 갱신돼도 실행 전 행과 스크롤 위치로 돌아간다.
                self._restore_room_table_state(table_state)
            return

        # DXF 캐시 모드에서는 기존 실명 속성 갱신만 지원한다.
        # 마감은 연결 CAD에서 직접 벽체 구간선을 생성하는 방식으로 처리한다.
        if apply_finishes and not apply_room_tags:
            QMessageBox.information(self, '수정사항 CAD에 적용', '마감 적용은 CAD 연동 상태에서 사용할 수 있습니다.')
            self._restore_room_table_state(table_state)
            return
        room_targets = [r for r in targets if r.tag_handle]
        if not room_targets:
            QMessageBox.information(self, '수정사항 CAD에 적용', 'CAD에 적용할 실명블록이 없습니다.')
            self._restore_room_table_state(table_state)
            return
        result = self._run_cached_dxf_operation('수정사항 CAD에 적용', refresh_room_blocks, room_targets)
        if result:
            out, report = result
            QMessageBox.information(self, '수정사항 CAD 적용 완료', f"새로고침: {report['refreshed']}/{report['requested']}개\n블록 누락: {report['missing_insert']}개\n\n내부 작업 캐시에 반영했습니다.")
        self._restore_room_table_state(table_state)

    def _wall_key_points(self, key: str):
        """벽 지정키를 현재 실구획선의 실제 두 절점으로 변환한다."""
        parts = str(key or "").split(":")
        if not parts:
            return None
        room = next((value for value in self.rooms if str(value.boundary_handle or "") == parts[0]), None)
        if room is None:
            return None
        try:
            if len(parts) == 2:
                points = list(room.points or [])
                index = int(parts[1])
            elif len(parts) == 3 and parts[1].lower().startswith("d"):
                deduction_index = int(parts[1][1:])
                points = list((room.deduction_points or [])[deduction_index] or [])
                index = int(parts[2])
            else:
                return None
            if len(points) < 2:
                return None
            return (
                (float(points[index][0]), float(points[index][1])),
                (float(points[(index + 1) % len(points)][0]), float(points[(index + 1) % len(points)][1])),
            )
        except (ValueError, TypeError, IndexError):
            return None

    def _room_wall_link_export_items(self, floor: str = "", sections=None) -> list[dict]:
        """현재 그룹지정 자료를 CAD/DXF 연동 도형으로 만든다.

        V23.45부터 CAD HANDLE은 호환/진단용만 기록하고 room_id/wall_id/object_id를 주 식별자로 사용한다.
        공용벽 양쪽 면이 같은 그룹에 지정된 경우 CAD에는 중복선을 만들지 않고 한 객체 XData에 여러 ID를 함께 기록한다.
        """
        sections = set(sections or ROOM_WALL_LINK_LAYERS.keys())
        room_by_handle = {str(room.boundary_handle or ""): room for room in self.rooms if room.boundary_handle}
        items = []
        for group_id, raw_group in self.finish_groups.items():
            group = dict(raw_group or {})
            section = normalize_room_finish_section(group.get("section", ""))
            if section not in sections or section not in ROOM_WALL_LINK_LAYERS:
                continue
            group_type = normalize_finish_group_type(group.get("group_type", ""))
            group_division = normalize_finish_group_division(group.get("group_division", ""), group_type)
            group_aci = finish_group_aci(group_id, group.get("color", ""), group.get("aci"))
            common = {
                "section": section,
                "color": aci_to_hex(group_aci),
                "aci": group_aci,
                "group_id": str(group_id),
                "group_type": group_type,
                "group_division": group_division,
                "group_name": str(group.get("name", "") or "그룹"),
            }
            if section == "벽":
                merged = {}
                for wall_key in group.get("wall_keys", []) or []:
                    points = self._wall_key_points(str(wall_key))
                    if not points:
                        continue
                    handle = wall_key_room_handle(str(wall_key))
                    room = room_by_handle.get(handle)
                    if room is None or (floor and str(room.floor or "") != floor):
                        continue
                    rid = room_program_id(room)
                    parts = str(wall_key).split(":")
                    role = "DEDUCTION" if len(parts) >= 3 and str(parts[1]).lower().startswith("d") else "OUTER"
                    wid = wall_program_id(rid, points[0], points[1], role)
                    a = (round(float(points[0][0]), 3), round(float(points[0][1]), 3))
                    b = (round(float(points[1][0]), 3), round(float(points[1][1]), 3))
                    geom = tuple(sorted((a, b)))
                    entry = merged.setdefault(geom, {
                        **common, "points": [points[0], points[1]], "closed": False,
                        "room_ids": set(), "wall_ids": set(), "wall_keys": set(), "room_handles": set(),
                    })
                    entry["room_ids"].add(rid)
                    entry["wall_ids"].add(wid)
                    entry["wall_keys"].add(str(wall_key))
                    if handle:
                        entry["room_handles"].add(handle)
                for entry in merged.values():
                    room_ids = sorted(entry.pop("room_ids"))
                    wall_ids = sorted(entry.pop("wall_ids"))
                    wall_keys = sorted(entry.pop("wall_keys"))
                    room_handles = sorted(entry.pop("room_handles"))
                    rid_text = "|".join(room_ids)
                    wid_text = "|".join(wall_ids)
                    entry.update({
                        "object_id": link_object_id(section, str(group_id), rid_text, wid_text),
                        "room_id": rid_text,
                        "wall_id": wid_text,
                        # 아래는 구버전 프로젝트 호환/진단값. 복수면 | 구분.
                        "wall_key": "|".join(wall_keys),
                        "room_handle": "|".join(room_handles),
                    })
                    items.append(entry)
            else:
                for handle in group.get("room_handles", []) or []:
                    room = room_by_handle.get(str(handle))
                    if room is None or (floor and str(room.floor or "") != floor):
                        continue
                    points = [(float(x), float(y)) for x, y in (room.points or [])]
                    if len(points) < 3:
                        continue
                    rid = room_program_id(room)
                    items.append({
                        **common,
                        "points": points,
                        "closed": True,
                        "object_id": link_object_id(section, str(group_id), rid, ""),
                        "room_id": rid,
                        "wall_id": "",
                        "wall_key": "",
                        "room_handle": str(handle),
                    })
        return items


    def open_room_wall_link_dialog(self):
        """프로그램 그룹과 CAD/DXF 구획선을 왕복하고 1:1 구간을 검사한다."""
        if not self.rooms:
            QMessageBox.information(self, "실/벽 연동", "먼저 내부 실목록을 읽어 주세요.")
            return
        link_layers = room_wall_link_layers(self.room_rules)
        dialog = QDialog(self)
        dialog.setWindowTitle("실/벽 연동")
        dialog.resize(610, 430)
        layout = QVBoxLayout(dialog)

        info = QLabel(
            "프로그램에서 지정한 그룹을 CAD/DXF에 덧그리거나, CAD/DXF의 색상별 구획선을 그룹으로 가져옵니다.\n"
            "벽 연동선은 실구획선의 절점·길이와 1:1로 맞아야 하며 완전중복/부분중복도 검사합니다."
        )
        info.setWordWrap(True)
        layout.addWidget(info)

        layer_box = QGroupBox("CAD 연동 레이어")
        layer_layout = QGridLayout(layer_box)
        for row, section in enumerate(("벽", "바닥", "천정", "기타")):
            layer_layout.addWidget(QLabel(section), row, 0)
            label = QLabel(link_layers[section])
            label.setTextInteractionFlags(Qt.TextSelectableByMouse)
            layer_layout.addWidget(label, row, 1)
        layout.addWidget(layer_box)

        scope_box = QGroupBox("내보내기 범위")
        scope_layout = QHBoxLayout(scope_box)
        scope_layout.addWidget(QLabel("층:"))
        scope_combo = QComboBox()
        scope_combo.addItem("전체층", "")
        current_floor = str(self.selected_floor() or "")
        floors = []
        for room in self.rooms:
            floor = str(room.floor or "")
            if floor and floor not in floors:
                floors.append(floor)
        for floor in sorted(floors, key=floor_sort_key):
            scope_combo.addItem(floor, floor)
        if current_floor:
            index = scope_combo.findData(current_floor)
            if index >= 0:
                scope_combo.setCurrentIndex(index)
        scope_layout.addWidget(scope_combo)
        section_checks = {}
        for section in ("벽", "바닥", "천정", "기타"):
            check = QCheckBox(section)
            check.setChecked(section in {"벽", "바닥", "천정"})
            section_checks[section] = check
            scope_layout.addWidget(check)
        scope_layout.addStretch(1)
        layout.addWidget(scope_box)

        actions = QGridLayout()
        export_cad = QPushButton("프로그램 → 연결 CAD")
        export_dxf = QPushButton("프로그램 → DXF 저장")
        import_cad = QPushButton("연결 CAD → 프로그램")
        import_dxf = QPushButton("DXF → 프로그램")
        validate_cad = QPushButton("연결 CAD 구간검증")
        validate_dxf = QPushButton("DXF 구간검증")
        export_cad.setEnabled(bool(self.cad_connected_mode))
        import_cad.setEnabled(bool(self.cad_connected_mode))
        validate_cad.setEnabled(bool(self.cad_connected_mode))
        actions.addWidget(export_cad, 0, 0)
        actions.addWidget(export_dxf, 0, 1)
        actions.addWidget(import_cad, 1, 0)
        actions.addWidget(import_dxf, 1, 1)
        actions.addWidget(validate_cad, 2, 0)
        actions.addWidget(validate_dxf, 2, 1)
        layout.addLayout(actions)

        note = QLabel(
            "※ 프로그램에서 내보낸 객체는 그룹 ID도 함께 기록하므로 다시 읽으면 원래 그룹을 우선 인식합니다.\n"
            "※ CAD에서 새로 그린 객체는 색상별로 묶어 기존 그룹 또는 새 그룹에 연결합니다."
        )
        note.setWordWrap(True)
        layout.addWidget(note)
        close_buttons = QDialogButtonBox(QDialogButtonBox.Close)
        close_buttons.rejected.connect(dialog.reject)
        close_buttons.button(QDialogButtonBox.Close).clicked.connect(dialog.accept)
        layout.addWidget(close_buttons)

        def selected_sections():
            return [name for name, check in section_checks.items() if check.isChecked()]

        def build_items():
            sections = selected_sections()
            if not sections:
                QMessageBox.information(dialog, "실/벽 연동", "내보낼 부위를 하나 이상 선택해 주세요.")
                return []
            items = self._room_wall_link_export_items(str(scope_combo.currentData() or ""), sections)
            if not items:
                QMessageBox.information(dialog, "실/벽 연동", "선택 범위에 CAD로 보낼 그룹 지정 객체가 없습니다.")
            return items

        def export_to_cad():
            items = build_items()
            if items:
                self.room_wall_link_cad_requested.emit({"action": "export", "items": items, "layers": link_layers})

        def export_to_dxf():
            items = build_items()
            if not items or not self.current_source_path:
                return
            source = Path(self.current_source_path)
            suggested = str(source.with_name(source.stem + "_실벽연동.dxf"))
            out, _ = QFileDialog.getSaveFileName(dialog, "실/벽 연동 DXF 저장", suggested, "DXF 파일 (*.dxf)")
            if not out:
                return
            try:
                report = write_room_wall_link_dxf(self.current_source_path, out, items, link_layers)
                QMessageBox.information(
                    dialog, "실/벽 연동 DXF 저장 완료",
                    f"생성: {report['created']:,}/{report['requested']:,}개\nID 갱신: {report.get('updated', 0):,}개\n기존 동일객체 제외: {report.get('skipped_existing', 0):,}개\n실패: {report['failed']:,}개\n\n{out}"
                )
            except Exception:
                QMessageBox.critical(dialog, "실/벽 연동 DXF 저장 오류", traceback.format_exc())

        def import_from_dxf():
            start = str(Path(self.current_source_path).parent) if self.current_source_path else ""
            path, _ = QFileDialog.getOpenFileName(dialog, "실/벽 연동 DXF 읽기", start, "DXF 파일 (*.dxf)")
            if not path:
                return
            try:
                result = read_room_wall_link_dxf(path, link_layers)
                if not result.get("items"):
                    QMessageBox.information(dialog, "실/벽 연동", "설정된 실/벽 연동 레이어에서 읽을 객체가 없습니다.")
                    return
                self._import_room_wall_link_items(result["items"], f"DXF: {Path(path).name}")
            except Exception:
                QMessageBox.critical(dialog, "실/벽 연동 DXF 읽기 오류", traceback.format_exc())

        def validate_dxf_file():
            start = str(Path(self.current_source_path).parent) if self.current_source_path else ""
            path, _ = QFileDialog.getOpenFileName(dialog, "실/벽 연동 DXF 구간검증", start, "DXF 파일 (*.dxf)")
            if not path:
                return
            try:
                result = read_room_wall_link_dxf(path, link_layers)
                report = validate_room_wall_link_items(result.get("items", []), self.rooms, 2.0)
                self._show_room_wall_link_validation(report, f"DXF: {Path(path).name}")
            except Exception:
                QMessageBox.critical(dialog, "실/벽 연동 구간검증 오류", traceback.format_exc())

        export_cad.clicked.connect(export_to_cad)
        export_dxf.clicked.connect(export_to_dxf)
        import_cad.clicked.connect(lambda: self.room_wall_link_cad_requested.emit({"action": "read", "layers": link_layers}))
        import_dxf.clicked.connect(import_from_dxf)
        validate_cad.clicked.connect(lambda: self.room_wall_link_cad_requested.emit({"action": "validate", "layers": link_layers}))
        validate_dxf.clicked.connect(validate_dxf_file)
        dialog.exec()

    def _show_room_wall_link_validation(self, report: dict, source_label: str = "") -> None:
        issues = list(report.get("issues", []) or [])
        self.room_wall_validation_errors = [f"[실/벽 연동] {value}" for value in issues]
        self.rebuild_views()
        summary = report.get("summary", {}) or {}
        text = (
            f"대상: {source_label or '실/벽 연동 객체'}\n"
            f"허용오차: ±{float(report.get('tolerance', 2.0)):.1f}mm\n\n"
            f"1:1 미대응: {report.get('unmatched', 0):,}건\n"
            f"완전중복: {report.get('exact_duplicate_pairs', 0):,}쌍\n"
            f"부분중복: {report.get('partial_overlap_pairs', 0):,}쌍\n"
            f"프로그램 ID 직접매칭: {summary.get('id', 0):,}건\n"
            f"구버전 핸들/벽키 호환매칭: {summary.get('legacy_metadata', 0):,}건\n"
            f"공용벽 다중매칭: {summary.get('shared_wall', 0):,}건\n"
            f"전체 확인사항: {report.get('issue_count', 0):,}건"
        )
        if issues:
            text += "\n\n상세\n" + "\n".join(issues[:25])
            if len(issues) > 25:
                text += f"\n외 {len(issues)-25:,}건"
        message_type = QMessageBox.warning if issues else QMessageBox.information
        message_type(self, "실/벽 연동 구간검증", text)

    def _import_room_wall_link_items(self, raw_items, source_label: str = "CAD") -> None:
        """색상별 CAD/DXF 객체를 사용자가 지정한 탭/구분/그룹에 연결한다."""
        matched = match_room_wall_link_items(raw_items, self.rooms, 2.0)
        items = list(matched.get("items", []) or [])
        if not items:
            QMessageBox.information(self, "실/벽 연동", "가져올 객체가 없습니다.")
            return

        buckets = defaultdict(list)
        for item in items:
            metadata_group_id = str(item.get("group_id", "") or "")
            aci = color_to_aci(item.get("color", "#FFFFFF"), item.get("aci"))
            item["aci"] = aci
            item["color"] = aci_to_hex(aci)
            bucket_key = (
                str(item.get("section", "") or ""),
                int(aci),
                metadata_group_id if metadata_group_id else "",
            )
            buckets[bucket_key].append(item)

        dialog = QDialog(self)
        dialog.setWindowTitle("CAD 색상별 그룹 연결")
        dialog.resize(1040, min(680, 240 + len(buckets) * 42))
        layout = QVBoxLayout(dialog)
        info = QLabel(
            f"읽은 대상: {source_label}\n"
            "색상별 객체를 기존 그룹에 연결하거나 새 그룹을 만드세요. 새 그룹일 때만 탭/구분/그룹명을 수정합니다."
        )
        info.setWordWrap(True)
        layout.addWidget(info)
        table = QTableWidget(len(buckets), 8)
        table.setHorizontalHeaderLabels(["색상", "부위", "개체", "그룹", "탭", "구분", "새 그룹명", "매칭"])
        table.verticalHeader().setVisible(False)
        table.setSelectionBehavior(QAbstractItemView.SelectRows)
        table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        table.horizontalHeader().setStretchLastSection(True)
        layout.addWidget(table, 1)

        rows = []
        for row, (bucket_key, bucket_items) in enumerate(sorted(buckets.items(), key=lambda value: (value[0][0], value[0][1], value[0][2]))):
            section, aci, metadata_group_id = bucket_key
            color = aci_to_hex(aci)
            color_item = QTableWidgetItem(f"ACI {aci}")
            qcolor = QColor(color)
            if qcolor.isValid():
                color_item.setBackground(QBrush(qcolor))
                # 밝은 색 배경에서는 검정, 어두운 배경에서는 흰색 글자를 사용한다.
                brightness = qcolor.red() * 0.299 + qcolor.green() * 0.587 + qcolor.blue() * 0.114
                color_item.setForeground(QBrush(QColor("#000000" if brightness > 150 else "#FFFFFF")))
            table.setItem(row, 0, color_item)
            table.setItem(row, 1, QTableWidgetItem(section))
            table.setItem(row, 2, QTableWidgetItem(f"{len(bucket_items):,}"))

            group_combo = QComboBox()
            group_combo.addItem("새 그룹 생성", "")
            for group_id, group in self.finish_groups.items():
                if normalize_room_finish_section(group.get("section", "")) != section:
                    continue
                gtype = normalize_finish_group_type(group.get("group_type", ""))
                gdivision = normalize_finish_group_division(group.get("group_division", ""), gtype)
                label = f"{self.finish_panel.group_type_label(gtype)} > {gdivision} > {group.get('name', '그룹')}"
                group_combo.addItem(label, str(group_id))
            table.setCellWidget(row, 3, group_combo)

            type_combo = QComboBox()
            for group_type in FINISH_GROUP_TYPES:
                type_combo.addItem(self.finish_panel.group_type_label(group_type), group_type)
            table.setCellWidget(row, 4, type_combo)
            division_combo = QComboBox()
            table.setCellWidget(row, 5, division_combo)
            name_edit = QLineEdit()
            table.setCellWidget(row, 6, name_edit)

            preferred_type = normalize_finish_group_type(bucket_items[0].get("group_type", ""))
            if preferred_type not in FINISH_GROUP_TYPES:
                preferred_type = self.finish_panel.current_group_type()
            type_index = type_combo.findData(preferred_type)
            type_combo.setCurrentIndex(max(0, type_index))

            preferred_division = str(bucket_items[0].get("group_division", "") or "").strip()
            preferred_name = str(bucket_items[0].get("group_name", "") or "").strip()
            name_edit.setText(preferred_name or f"CAD ACI {aci}")

            matched_count = sum(1 for value in bucket_items if value.get("matched_wall_keys") or value.get("matched_room_handles"))
            unmatched_count = len(bucket_items) - matched_count
            shared_count = sum(1 for value in bucket_items if len(value.get("matched_wall_keys", []) or []) > 1)
            match_text = f"정상 {matched_count}"
            if shared_count:
                match_text += f" / 공용벽 {shared_count}"
            if unmatched_count:
                match_text += f" / 미대응 {unmatched_count}"
            table.setItem(row, 7, QTableWidgetItem(match_text))

            def refresh_divisions(_index=0, tcombo=type_combo, dcombo=division_combo, preferred=preferred_division):
                current = str(dcombo.currentText() or "")
                dcombo.blockSignals(True)
                dcombo.clear()
                values = self.finish_panel.group_divisions(str(tcombo.currentData() or ""))
                for value in values:
                    dcombo.addItem(value, value)
                target = preferred if preferred and dcombo.findText(preferred) >= 0 else current
                index = dcombo.findText(target) if target else -1
                if index >= 0:
                    dcombo.setCurrentIndex(index)
                dcombo.blockSignals(False)

            type_combo.currentIndexChanged.connect(refresh_divisions)
            refresh_divisions()

            def sync_existing(_index=0, gcombo=group_combo, tcombo=type_combo, dcombo=division_combo, nedit=name_edit):
                group_id = str(gcombo.currentData() or "")
                group = self.finish_groups.get(group_id)
                existing = bool(group)
                if group:
                    gtype = normalize_finish_group_type(group.get("group_type", ""))
                    gdivision = normalize_finish_group_division(group.get("group_division", ""), gtype)
                    ti = tcombo.findData(gtype)
                    if ti >= 0:
                        tcombo.setCurrentIndex(ti)
                    if dcombo.findText(gdivision) < 0:
                        dcombo.addItem(gdivision, gdivision)
                    dcombo.setCurrentText(gdivision)
                    nedit.setText(str(group.get("name", "") or "그룹"))
                tcombo.setEnabled(not existing)
                dcombo.setEnabled(not existing)
                nedit.setEnabled(not existing)

            group_combo.currentIndexChanged.connect(sync_existing)
            if metadata_group_id and metadata_group_id in self.finish_groups:
                gi = group_combo.findData(metadata_group_id)
                if gi >= 0:
                    group_combo.setCurrentIndex(gi)
            else:
                # 메타데이터가 없더라도 같은 이름/탭/구분 그룹이 있으면 우선 제안한다.
                for i in range(1, group_combo.count()):
                    candidate = self.finish_groups.get(str(group_combo.itemData(i) or ""), {})
                    if preferred_name and str(candidate.get("name", "") or "").casefold() == preferred_name.casefold():
                        group_combo.setCurrentIndex(i)
                        break
            sync_existing()
            rows.append({
                "items": bucket_items,
                "section": section,
                "color": color,
                "aci": int(aci),
                "group_combo": group_combo,
                "type_combo": type_combo,
                "division_combo": division_combo,
                "name_edit": name_edit,
            })

        table.resizeColumnsToContents()
        table.setColumnWidth(3, max(250, table.columnWidth(3)))
        table.setColumnWidth(6, max(160, table.columnWidth(6)))
        buttons = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        buttons.button(QDialogButtonBox.Ok).setText("가져오기")
        buttons.accepted.connect(dialog.accept)
        buttons.rejected.connect(dialog.reject)
        layout.addWidget(buttons)
        if dialog.exec() != QDialog.Accepted:
            return

        created_groups = 0
        assigned_entities = 0
        skipped_unmatched = 0
        skipped_conflict = 0
        shared_wall_entities = 0
        for row in rows:
            section = row["section"]
            group_id = str(row["group_combo"].currentData() or "")
            group = self.finish_groups.get(group_id)
            if group is None:
                group_type = normalize_finish_group_type(row["type_combo"].currentData())
                group_division = normalize_finish_group_division(row["division_combo"].currentText(), group_type)
                name = str(row["name_edit"].text() or "").strip()
                if not name:
                    QMessageBox.warning(self, "실/벽 연동", f"ACI {row['aci']} {section}: 새 그룹명이 비어 있어 건너뜁니다.")
                    continue
                self.finish_panel.ensure_group_division(group_type, group_division)
                group_id = uuid.uuid4().hex
                group = {
                    "id": group_id, "name": name, "section": section,
                    "room_handles": [], "wall_keys": [], "materials": [],
                    "color": aci_to_hex(row["aci"]), "aci": int(row["aci"]), "formula": "",
                    "group_type": group_type, "group_division": group_division,
                }
                self.finish_groups[group_id] = group
                created_groups += 1
            else:
                group_type = normalize_finish_group_type(group.get("group_type", ""))
                group_division = normalize_finish_group_division(group.get("group_division", ""), group_type)
                # CAD에서 읽은 색상으로 기존 그룹을 연결하면 그 ACI 번호를 그룹 기준색으로 동기화한다.
                group["aci"] = int(row["aci"])
                group["color"] = aci_to_hex(row["aci"])

            if section == "벽":
                target_keys = set()
                for item in row["items"]:
                    keys = {str(value) for value in (item.get("matched_wall_keys", []) or []) if value}
                    if not keys:
                        skipped_unmatched += 1
                        continue
                    if len(keys) > 1:
                        shared_wall_entities += 1
                    target_keys.update(keys)
                    assigned_entities += 1
                occupied = self._occupied_wall_group_keys(group_type, group_division, exclude_group_id=group_id)
                conflicts = target_keys & occupied
                skipped_conflict += len(conflicts)
                target_keys.difference_update(conflicts)
                group["wall_keys"] = sorted(set(group.get("wall_keys", []) or []) | target_keys)
                handles = sorted({wall_key_room_handle(key) for key in target_keys if wall_key_room_handle(key)})
                if handles:
                    self._attach_rooms_to_finish_group(group_id, handles)
                first_material = deepcopy((group.get("materials", []) or [None])[0])
                if first_material:
                    for key in target_keys:
                        self.wall_finish_assignments[str(key)] = deepcopy(first_material)
            else:
                target_handles = set()
                for item in row["items"]:
                    handles = {str(value) for value in (item.get("matched_room_handles", []) or []) if value}
                    if not handles:
                        skipped_unmatched += 1
                        continue
                    target_handles.update(handles)
                    assigned_entities += 1
                occupied = self._occupied_room_group_handles(section, group_type, group_division, exclude_group_id=group_id)
                conflicts = target_handles & occupied
                skipped_conflict += len(conflicts)
                target_handles.difference_update(conflicts)
                if target_handles:
                    self._attach_rooms_to_finish_group(group_id, target_handles)

        self._ensure_finish_group_colors()
        self.save_finish_assignments()
        self.preview.set_wall_finish_assignments(self.wall_finish_assignments)
        self.preview.set_room_finish_assignments(self.room_finish_assignments)
        self.preview.set_finish_groups(self.finish_groups)
        self._rebuild_finish_group_filter_combo()
        self._refresh_current_finish_panel()
        self.update_finish_summary(self.selected_floor())
        self.data_changed.emit()
        QMessageBox.information(
            self, "실/벽 연동 가져오기 완료",
            f"대상: {source_label}\n"
            f"새 그룹: {created_groups:,}개\n"
            f"연결 처리 객체: {assigned_entities:,}개\n"
            f"1:1 미대응 건너뜀: {skipped_unmatched:,}개\n"
            f"동일 구분 중복 충돌 제외: {skipped_conflict:,}개\n"
            f"공용벽(한 선이 양쪽 실구획과 일치): {shared_wall_entities:,}개\n\n"
            "공용벽은 CAD 선만으로 어느 한쪽 면인지 구분할 수 없으므로, 메타데이터가 없는 수동 CAD 선은 일치하는 양쪽 벽키에 적용됩니다."
        )

    def handle_room_wall_link_cad_result(self, action: str, payload) -> None:
        """MainWindow가 연결 CAD 작업을 끝낸 뒤 결과를 다시 내부탭에 전달한다."""
        action = str(action or "")
        if action == "export":
            report = dict(payload or {})
            message_type = QMessageBox.warning if report.get("failed") else QMessageBox.information
            message_type(
                self, "실/벽 연동 CAD 반영",
                f"생성: {report.get('created', 0):,}/{report.get('requested', 0):,}개\n"
                f"ID 갱신: {report.get('updated', 0):,}개\n"
                f"기존 동일객체 제외: {report.get('skipped_existing', 0):,}개\n"
                f"실패: {report.get('failed', 0):,}개"
                + (("\n\n" + "\n".join(report.get("details", [])[:20])) if report.get("details") else "")
            )
        elif action == "read":
            data = dict(payload or {})
            items = list(data.get("items", []) or [])
            if not items:
                QMessageBox.information(self, "실/벽 연동", "연결 CAD의 실/벽 연동 레이어에서 읽을 객체가 없습니다.")
                return
            self._import_room_wall_link_items(items, "연결 CAD")
        elif action == "validate":
            data = dict(payload or {})
            report = validate_room_wall_link_items(data.get("items", []) or [], self.rooms, 2.0)
            self._show_room_wall_link_validation(report, "연결 CAD")

    def request_create_wall_lines(self):
        """구버전 호출 호환: 기존 버튼 기능은 새 실/벽 연동창으로 연결한다."""
        self.open_room_wall_link_dialog()

    def _finish_assignment_path(self) -> Path | None:
        if not self.current_source_path:
            return None
        source = Path(self.current_source_path).resolve()
        digest = hashlib.sha1(str(source).encode("utf-8")).hexdigest()[:12]
        folder = WALL_FINISH_ASSIGNMENTS_DIR
        return folder / f"{source.stem}_{digest}.json"

    def _ensure_finish_group_colors(self):
        """그룹 표시색을 ACI 기준으로 정규화하고 실별 참조에도 동일 번호를 복사한다."""
        for group_id, group in self.finish_groups.items():
            group["id"] = str(group_id)
            source_color = finish_group_color(group_id, group.get("color", ""))
            group_aci = finish_group_aci(group_id, source_color, group.get("aci", group.get("color_aci")))
            group["aci"] = group_aci
            group["color"] = aci_to_hex(group_aci)
            raw_group_type = str(group.get("group_type", "") or "").strip()
            raw_group_division = group.get("group_division", "")
            if not str(raw_group_division or "").strip() and raw_group_type == "바탕":
                raw_group_division = "바탕"
            group["group_type"] = normalize_finish_group_type(raw_group_type)
            group["group_division"] = normalize_finish_group_division(
                raw_group_division, group["group_type"]
            )
            group["formula"] = normalize_formula_text(group.get("formula", ""))
            if isinstance(group.get("scope_formulas"), dict):
                group["scope_formulas"] = {
                    str(key): normalize_formula_text(value) for key, value in group.get("scope_formulas", {}).items()
                }
            self.finish_panel.ensure_group_division(group["group_type"], group["group_division"])
            for material in group.get("materials", []) or []:
                if isinstance(material, dict):
                    material["method"] = normalize_formula_text(material.get("method", ""))
                    material["group_formula"] = normalize_formula_text(material.get("group_formula", group.get("formula", "")))
                    if isinstance(material.get("scope_methods"), dict):
                        material["scope_methods"] = {
                            str(key): normalize_formula_text(value) for key, value in material.get("scope_methods", {}).items()
                        }
                    material["group_color"] = group["color"]
                    material["group_aci"] = group["aci"]
                    material["group_type"] = group["group_type"]
                    material["group_division"] = group["group_division"]
        for handle, raw in list(self.room_finish_assignments.items()):
            assignment = normalize_room_finish_assignment(raw)
            for entry in assignment.get("groups", []):
                group = self.finish_groups.get(str(entry.get("id", "") or ""))
                if group:
                    entry["color"] = group["color"]
                    entry["aci"] = group["aci"]
                    entry["group_type"] = group["group_type"]
                    entry["group_division"] = group["group_division"]
            for materials in assignment.get("items", {}).values():
                for material in materials:
                    group = self.finish_groups.get(str(material.get("group_id", "") or ""))
                    if group:
                        material["group_color"] = group["color"]
                        material["group_aci"] = group["aci"]
                        material["group_type"] = group["group_type"]
                        material["group_division"] = group["group_division"]
            self.room_finish_assignments[str(handle)] = assignment
        for key, material in list(self.wall_finish_assignments.items()):
            if not isinstance(material, dict):
                continue
            material["method"] = normalize_formula_text(material.get("method", ""))
            material["group_formula"] = normalize_formula_text(material.get("group_formula", ""))
            group = self.finish_groups.get(str(material.get("group_id", "") or ""))
            if group:
                material["group_color"] = group["color"]
                material["group_aci"] = group["aci"]
                material["group_type"] = group["group_type"]
                material["group_division"] = group["group_division"]

    def load_finish_assignments(self):
        path = self._finish_assignment_path()
        self.wall_finish_assignments = {}
        self.room_finish_assignments = {}
        self.finish_groups = {}
        self.summary_markups = {}
        self.opening_read_results = {}
        self.selected_wall_keys.clear()
        if path and path.exists():
            try:
                data = json.loads(path.read_text(encoding="utf-8"))
                self.wall_finish_assignments = dict(data.get("wall_assignments", data.get("assignments", {})))
                self.room_finish_assignments = {
                    str(handle): normalize_room_finish_assignment(assignment)
                    for handle, assignment in dict(data.get("room_assignments", {})).items()
                }
                self.finish_groups = {
                    str(group_id): dict(group)
                    for group_id, group in dict(data.get("finish_groups", {})).items()
                }
                self.summary_markups = {
                    str(key): float(value or 0.0)
                    for key, value in dict(data.get("summary_markups", {})).items()
                    if str(key)
                }
                self.opening_read_results = {
                    str(floor): [dict(item or {}) for item in (items or [])]
                    for floor, items in dict(data.get("opening_read_results", {})).items()
                    if str(floor)
                }
            except Exception:
                self.wall_finish_assignments = {}
                self.room_finish_assignments = {}
                self.finish_groups = {}
                self.summary_markups = {}
                self.opening_read_results = {}
        self._ensure_finish_group_colors()
        self._rebuild_finish_group_filter_combo()
        self.preview.set_wall_finish_assignments(self.wall_finish_assignments)
        self.preview.set_room_finish_assignments(self.room_finish_assignments)
        self.preview.set_finish_groups(self.finish_groups)
        self.preview.set_selected_wall_keys(self.selected_wall_keys)
        self.finish_panel.set_selected_wall_count(0)
        self.finish_panel.set_selected_room_count(0)
        selected_room = self.selected_room()
        if selected_room is not None:
            self.finish_panel.set_group_wall_lengths(self._wall_group_lengths_for_room(selected_room))
            self.finish_panel.set_current_room(
                selected_room,
                self.room_finish_assignments.get(selected_room.boundary_handle, {}),
            )
        else:
            self._refresh_current_finish_panel()
        self._normalize_cached_opening_results()
        self._update_opening_edit_controls()
        self._sync_opening_read_preview()
        self.update_finish_summary(self.selected_floor())

    def restore_finish_assignments(self, wall_assignments=None, room_assignments=None, finish_groups=None, summary_markups=None, opening_read_results=None):
        """프로젝트 스냅샷의 마감 지정을 현재 화면에 적용한다."""
        self.wall_finish_assignments = dict(wall_assignments or {})
        self.room_finish_assignments = {
            str(handle): normalize_room_finish_assignment(assignment)
            for handle, assignment in dict(room_assignments or {}).items()
        }
        self.finish_groups = {
            str(group_id): dict(group)
            for group_id, group in dict(finish_groups or {}).items()
        }
        if summary_markups is not None:
            self.summary_markups = {
                str(key): float(value or 0.0)
                for key, value in dict(summary_markups or {}).items()
                if str(key)
            }
        if opening_read_results is not None:
            self.opening_read_results = {
                str(floor): [dict(item or {}) for item in (items or [])]
                for floor, items in dict(opening_read_results or {}).items()
                if str(floor)
            }
        self._ensure_finish_group_colors()
        self._rebuild_finish_group_filter_combo()
        self.selected_wall_keys.clear()
        self.preview.set_wall_finish_assignments(self.wall_finish_assignments)
        self.preview.set_room_finish_assignments(self.room_finish_assignments)
        self.preview.set_finish_groups(self.finish_groups)
        self.preview.set_selected_wall_keys(self.selected_wall_keys)
        selected_room = self.selected_room()
        if selected_room is not None:
            self.finish_panel.set_group_wall_lengths(self._wall_group_lengths_for_room(selected_room))
            self.finish_panel.set_current_room(
                selected_room,
                self.room_finish_assignments.get(selected_room.boundary_handle, {}),
            )
        else:
            self._refresh_current_finish_panel()
        self.update_finish_summary(self.selected_floor())
        self._normalize_cached_opening_results()
        self._update_opening_edit_controls()
        self._sync_opening_read_preview()


    def _on_room_finish_panel_changed(self, boundary_handle: str, assignment):
        handle = str(boundary_handle or "").strip()
        if not handle:
            return
        previous = normalize_room_finish_assignment(self.room_finish_assignments.get(handle, {}))
        normalized = normalize_room_finish_assignment(assignment)
        self.room_finish_assignments[handle] = normalized
        previous_ids = {str(group.get("id", "") or "") for group in previous.get("groups", [])}
        current_ids = {str(group.get("id", "") or "") for group in normalized.get("groups", [])}
        for group_id in previous_ids - current_ids:
            group = self.finish_groups.get(group_id)
            if not group:
                continue
            group["room_handles"] = [value for value in group.get("room_handles", []) if value != handle]
            own_keys = [key for key in group.get("wall_keys", []) if wall_key_room_handle(key) == handle]
            group["wall_keys"] = [key for key in group.get("wall_keys", []) if key not in own_keys]
            for key in own_keys:
                self.wall_finish_assignments.pop(str(key), None)
        for group in normalized.get("groups", []):
            group_id = str(group.get("id", "") or "")
            if group_id in self.finish_groups and handle in list(self.finish_groups[group_id].get("room_handles", []) or []):
                self._propagate_finish_group(group_id, normalized)
        self.save_finish_assignments()
        self.preview.set_wall_finish_assignments(self.wall_finish_assignments)
        self.preview.set_room_finish_assignments(self.room_finish_assignments)
        self.preview.set_finish_groups(self.finish_groups)
        self.update_finish_summary(self.selected_floor())
        self.data_changed.emit()

    def save_finish_assignments(self):
        path = self._finish_assignment_path()
        if path is None:
            return
        path.parent.mkdir(parents=True, exist_ok=True)
        data = {
            "source_path": str(Path(self.current_source_path).resolve()),
            "wall_assignments": self.wall_finish_assignments,
            "room_assignments": self.room_finish_assignments,
            "finish_groups": self.finish_groups,
            "summary_markups": self.summary_markups,
            "opening_read_results": self.opening_read_results,
        }
        path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")

    def selected_target_rooms(self) -> list[RoomRecord]:
        selected = self.selected_rooms()
        if selected:
            return selected
        room = self.selected_room()
        if room:
            return [room]
        floor = self.selected_floor()
        return [r for r in self.rooms if r.floor == floor] if floor else []

    def _refresh_current_finish_panel(self):
        room = self.selected_room()
        if room is None:
            if self.room_table.currentRow() == 0:
                self.finish_panel.set_group_catalog(self._finish_groups_for_scope(None))
            else:
                floor = self.selected_floor()
                if floor:
                    self.finish_panel.set_group_catalog(self._finish_groups_for_scope(floor), f"{floor} 그룹 목록")
                else:
                    self.finish_panel.clear_current_room()
            return
        self.finish_panel.set_group_wall_lengths(self._wall_group_lengths_for_room(room))
        self.finish_panel.set_current_room(
            room, self.room_finish_assignments.get(room.boundary_handle, {})
        )

    def _choose_group_rooms(self, single_room_only: bool = False) -> list[RoomRecord]:
        current = self.selected_room()
        if single_room_only:
            return [current] if current is not None else []
        if not self.rooms:
            return []

        selected_handles = {room.boundary_handle for room in self.selected_rooms()}
        if current is not None:
            selected_handles.add(current.boundary_handle)
        elif self.room_table.currentRow() == 0:
            selected_handles = {room.boundary_handle for room in self.rooms}
        else:
            floor = self._floor_for_table_row(self.room_table.currentRow())
            if floor:
                selected_handles = {room.boundary_handle for room in self.rooms if room.floor == floor}

        dialog = QDialog(self)
        dialog.setWindowTitle("마감그룹 적용 실 선택")
        dialog.resize(390, 420)
        layout = QVBoxLayout(dialog)
        layout.addWidget(QLabel("같은 재료 구성을 공유할 실을 체크하세요."))
        room_list = QListWidget()
        # 바닥/천정 마감그룹 지정창도 내부 실목록과 동일한 건축 층 순서를 사용한다.
        # 원본 self.rooms 배열 순서(DXF 읽기 순서)에 의존하면 1F부터 나오고 지하층이 뒤로 밀릴 수 있다.
        # 지정창에서는 항상 깊은 지하층 -> B1F -> 1F -> 2F -> ... -> RF/PH 순으로 고정한다.
        dialog_rooms = sorted(
            self.rooms,
            key=lambda room: (
                floor_sort_key(room.floor),
                str(room.room_number or "").casefold(),
                normalize_room_name(room.room_name).casefold(),
            ),
        )
        for room in dialog_rooms:
            label = f"{room.floor} / {room.room_number or '-'} / {normalize_room_name(room.room_name)}"
            item = QListWidgetItem(label)
            item.setData(Qt.UserRole, room.boundary_handle)
            item.setFlags(item.flags() | Qt.ItemIsUserCheckable)
            item.setCheckState(Qt.Checked if room.boundary_handle in selected_handles else Qt.Unchecked)
            room_list.addItem(item)
        layout.addWidget(room_list, 1)
        buttons = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        buttons.button(QDialogButtonBox.Ok).setText("확인")
        buttons.button(QDialogButtonBox.Cancel).setText("취소")
        buttons.accepted.connect(dialog.accept)
        buttons.rejected.connect(dialog.reject)
        layout.addWidget(buttons)
        if dialog.exec() != QDialog.Accepted:
            return []
        handles = {
            str(room_list.item(row).data(Qt.UserRole) or "")
            for row in range(room_list.count())
            if room_list.item(row).checkState() == Qt.Checked
        }
        return [room for room in self.rooms if room.boundary_handle in handles]

    def _select_or_create_group(
        self, section: str, parent=None, preferred_type: str | None = None,
        preferred_division: str | None = None,
    ):
        """현재 재료마감 탭에서 기존 그룹을 고르거나 새 그룹을 만든다.

        group_type은 마감/단열/방수/외부/기타의 고정 내부탭키이고,
        group_division은 각 탭의 사용자 정의 구분(예: 바탕/마감)이다.
        """
        section = normalize_room_finish_section(section)
        preferred_type = (
            normalize_finish_group_type(preferred_type)
            if str(preferred_type or "").strip()
            else None
        )
        if preferred_type is None:
            preferred_type = self.finish_panel.current_group_type()
        divisions = self.finish_panel.group_divisions(preferred_type)
        preferred_division = str(preferred_division or "").strip()
        if preferred_division:
            preferred_division = self.finish_panel.ensure_group_division(preferred_type, preferred_division)

        candidates = [
            group for group in self.finish_groups.values()
            if normalize_room_finish_section(group.get("section", "")) == section
            and normalize_finish_group_type(group.get("group_type", "")) == preferred_type
            and (
                not preferred_division
                or normalize_finish_group_division(group.get("group_division", ""), preferred_type) == preferred_division
            )
        ]
        dialog = QDialog(parent or self)
        dialog.setWindowTitle(f"{section} 마감그룹 선택")
        dialog.setFixedWidth(380)
        layout = QVBoxLayout(dialog)
        layout.addWidget(QLabel("기존 그룹을 선택하거나 새 그룹명을 입력하세요."))

        tab_row = QHBoxLayout()
        tab_row.addWidget(QLabel("탭"))
        tab_value = QLineEdit(self.finish_panel.group_type_label(preferred_type))
        tab_value.setReadOnly(True)
        tab_row.addWidget(tab_value, 1)
        layout.addLayout(tab_row)

        division_row = QHBoxLayout()
        division_row.addWidget(QLabel("구분"))
        division_combo = QComboBox()
        for division in divisions:
            division_combo.addItem(division, division)
        wanted_division = preferred_division or (divisions[0] if divisions else "")
        division_index = division_combo.findData(wanted_division)
        division_combo.setCurrentIndex(max(0, division_index))
        if preferred_division:
            division_combo.setEnabled(False)
        division_row.addWidget(division_combo, 1)
        layout.addLayout(division_row)

        combo = QComboBox()
        combo.setEditable(True)
        combo.addItem("새 그룹 생성", "")
        for group in sorted(candidates, key=lambda value: str(value.get("name", ""))):
            division = normalize_finish_group_division(group.get("group_division", ""), preferred_type)
            combo.addItem(
                f"[{division}] {str(group.get('name', '') or '그룹')}",
                str(group.get("id", "") or ""),
            )
        combo.setCurrentIndex(0)
        combo.lineEdit().clear()
        combo.lineEdit().setPlaceholderText("새 그룹명 또는 기존 그룹 선택")
        layout.addWidget(combo)

        def sync_group_division(_index=0):
            current_id = str(combo.currentData() or "")
            current_group = self.finish_groups.get(current_id)
            if current_group:
                division = normalize_finish_group_division(current_group.get("group_division", ""), preferred_type)
                index = division_combo.findData(division)
                if index < 0:
                    division_combo.addItem(division, division)
                    index = division_combo.findData(division)
                division_combo.setCurrentIndex(max(0, index))
                division_combo.setEnabled(False)
            else:
                division_combo.setEnabled(not bool(preferred_division))
                if preferred_division:
                    index = division_combo.findData(preferred_division)
                    if index >= 0:
                        division_combo.setCurrentIndex(index)

        combo.currentIndexChanged.connect(sync_group_division)
        sync_group_division()
        buttons = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        buttons.accepted.connect(dialog.accept)
        buttons.rejected.connect(dialog.reject)
        layout.addWidget(buttons)
        if dialog.exec() != QDialog.Accepted:
            return None
        group_id = str(combo.currentData() or "")
        name = str(combo.currentText() or "").strip()
        if group_id and group_id in self.finish_groups:
            return group_id, str(self.finish_groups[group_id].get("name", "") or "그룹")
        group_division = normalize_finish_group_division(division_combo.currentData(), preferred_type)
        self.finish_panel.ensure_group_division(preferred_type, group_division)
        for group in self.finish_groups.values():
            if (
                normalize_room_finish_section(group.get("section", "")) == section
                and normalize_finish_group_type(group.get("group_type", "")) == preferred_type
                and normalize_finish_group_division(group.get("group_division", ""), preferred_type) == group_division
                and str(group.get("name", "") or "").casefold() == name.casefold()
                and name
            ):
                return str(group.get("id", "") or ""), str(group.get("name", "") or name)
        if not name or name == "새 그룹 생성":
            QMessageBox.information(dialog, "마감그룹", "새 그룹명을 입력해 주세요.")
            return None
        group_id = uuid.uuid4().hex
        group_aci = finish_group_aci(group_id, finish_group_color(group_id))
        self.finish_groups[group_id] = {
            "id": group_id, "name": name, "section": section,
            "room_handles": [], "wall_keys": [], "materials": [],
            "color": aci_to_hex(group_aci), "aci": group_aci, "formula": "",
            "group_type": preferred_type, "group_division": group_division,
        }
        return group_id, name

    def _choose_group_division_for_mode(self, group_type: str, parent=None) -> str:
        group_type = normalize_finish_group_type(group_type)
        current = self._current_finish_group_display_division(group_type)
        if current:
            return current
        values = self.finish_panel.group_divisions(group_type)
        if len(values) == 1:
            return values[0]
        label = self.finish_panel.group_type_label(group_type)
        value, ok = QInputDialog.getItem(
            parent or self, f"{label} 지정 구분", "지정할 구분:", values, 0, False
        )
        return str(value or "").strip() if ok else ""

    def _group_materials_preserve_room_formulas(self, assignment: dict, group_id: str, base_materials) -> list[dict]:
        """그룹 구성은 공통으로 맞추되 각 실에서 직접 입력한 재료 산식(method)은 보존한다."""
        normalized = normalize_room_finish_assignment(assignment)
        existing = []
        for values in normalized.get("items", {}).values():
            existing.extend(
                material for material in values
                if str(material.get("group_id", "") or "") == str(group_id or "")
            )
        result = []
        for raw in base_materials or []:
            value = deepcopy(raw)
            for old in existing:
                if self._same_finish_material(old, value):
                    value["method"] = normalize_formula_text(old.get("method", ""))
                    break
            result.append(value)
        return result

    def _attach_rooms_to_finish_group(self, group_id: str, handles) -> None:
        group = self.finish_groups.get(str(group_id or ""))
        if not group:
            return
        section = normalize_room_finish_section(group.get("section", ""))
        name = str(group.get("name", "") or "그룹")
        group_aci = finish_group_aci(group_id, group.get("color", ""), group.get("aci"))
        color = aci_to_hex(group_aci)
        group["aci"] = group_aci
        group["color"] = color
        formula = str(group.get("formula", "") or "")
        group_type = normalize_finish_group_type(group.get("group_type", ""))
        group_division = normalize_finish_group_division(group.get("group_division", ""), group_type)
        group["group_type"] = group_type
        group["group_division"] = group_division
        self.finish_panel.ensure_group_division(group_type, group_division)
        group["color"] = color
        merged_handles = set(group.get("room_handles", []) or []) | {str(value) for value in handles if value}
        group["room_handles"] = sorted(merged_handles)
        materials = []
        for material in group.get("materials", []) or []:
            value = deepcopy(material)
            value["group_id"] = group_id
            value["group_name"] = name
            value["group_color"] = color
            value["group_aci"] = group_aci
            value["group_formula"] = formula
            value["group_type"] = group_type
            value["group_division"] = group_division
            materials.append(value)
        for handle in sorted(merged_handles):
            assignment = normalize_room_finish_assignment(self.room_finish_assignments.get(handle, {}))
            had_room_formula, room_formula = self._assignment_group_formula(assignment, group_id)
            if not had_room_formula:
                room_formula = self._group_default_formula_for_floor(group, self._room_floor_for_handle(handle))
            if section not in assignment["sections"]:
                assignment["sections"].append(section)
                assignment["items"].setdefault(section, [])
            assignment["groups"] = [entry for entry in assignment.get("groups", []) if entry.get("id") != group_id]
            assignment["groups"].append({
                "id": group_id, "name": name, "section": section,
                "color": color, "aci": group_aci, "formula": room_formula, "group_type": group_type,
                "group_division": group_division,
            })
            room_materials = self._group_materials_preserve_room_formulas(assignment, group_id, materials)
            for material in room_materials:
                material["group_formula"] = room_formula
            for part in assignment["items"]:
                assignment["items"][part] = [
                    material for material in assignment["items"][part]
                    if str(material.get("group_id", "") or "") != group_id
                ]
            assignment["items"].setdefault(section, []).extend(room_materials)
            self.room_finish_assignments[handle] = assignment

    def on_finish_group_selected(self, group_id: str):
        """마감그룹 행 선택을 도면 강조 및 이후 '지정' 편집 대상과 연동한다."""
        group_id = str(group_id or "")
        self._active_finish_group_id = group_id if group_id in self.finish_groups else ""
        group = self.finish_groups.get(self._active_finish_group_id)
        if group:
            # 그룹을 집은 순간 도면 표시도 해당 탭의 해당 구분으로 맞춘다.
            group_type = normalize_finish_group_type(group.get("group_type", ""))
            group_division = normalize_finish_group_division(group.get("group_division", ""), group_type)
            self._set_finish_group_display_filter(group_type, group_division)
        self.preview.set_focused_finish_group(self._active_finish_group_id)

    def _active_group_for_section(self, section: str):
        """현재 재료마감 목록에서 선택된 그룹이 요청 부위와 같으면 반환한다."""
        group = self.finish_groups.get(str(self._active_finish_group_id or ""))
        if not group:
            return None
        if normalize_room_finish_section(group.get("section", "")) != normalize_room_finish_section(section):
            return None
        return group

    def _occupied_room_group_handles(
        self, section: str, group_type: str, group_division: str, exclude_group_id: str = ""
    ) -> set[str]:
        section = normalize_room_finish_section(section)
        group_type = normalize_finish_group_type(group_type)
        group_division = normalize_finish_group_division(group_division, group_type)
        occupied: set[str] = set()
        for other_id, other in self.finish_groups.items():
            if str(other_id) == str(exclude_group_id or ""):
                continue
            if normalize_room_finish_section(other.get("section", "")) != section:
                continue
            if normalize_finish_group_type(other.get("group_type", "")) != group_type:
                continue
            if normalize_finish_group_division(other.get("group_division", ""), group_type) != group_division:
                continue
            occupied.update(str(value) for value in (other.get("room_handles", []) or []) if value)
        return occupied

    def _occupied_wall_group_keys(
        self, group_type: str, group_division: str, exclude_group_id: str = ""
    ) -> set[str]:
        group_type = normalize_finish_group_type(group_type)
        group_division = normalize_finish_group_division(group_division, group_type)
        occupied: set[str] = set()
        for other_id, other in self.finish_groups.items():
            if str(other_id) == str(exclude_group_id or ""):
                continue
            if normalize_room_finish_section(other.get("section", "")) != "벽":
                continue
            if normalize_finish_group_type(other.get("group_type", "")) != group_type:
                continue
            if normalize_finish_group_division(other.get("group_division", ""), group_type) != group_division:
                continue
            occupied.update(str(value) for value in (other.get("wall_keys", []) or []) if value)
        return occupied

    def _room_group_conflicts(self, group_id: str, handles) -> set[str]:
        """바닥/천정 그룹의 동일 타입 중복 실을 반환한다."""
        group_id = str(group_id or "")
        group = self.finish_groups.get(group_id)
        if not group:
            return set()
        section = normalize_room_finish_section(group.get("section", ""))
        if section not in {"바닥", "천정"}:
            return set()
        group_type = normalize_finish_group_type(group.get("group_type", ""))
        group_division = normalize_finish_group_division(group.get("group_division", ""), group_type)
        targets = {str(value) for value in (handles or []) if value}
        if not targets:
            return set()
        occupied = self._occupied_room_group_handles(
            section, group_type, group_division, exclude_group_id=group_id
        )
        return targets & occupied

    def _wall_group_conflicts(self, group_id: str, wall_keys) -> set[str]:
        """벽 그룹의 동일 타입 중복 벽체키를 반환한다."""
        group_id = str(group_id or "")
        group = self.finish_groups.get(group_id)
        if not group:
            return set()
        group_type = normalize_finish_group_type(group.get("group_type", ""))
        group_division = normalize_finish_group_division(group.get("group_division", ""), group_type)
        targets = {str(value) for value in (wall_keys or []) if value}
        occupied = self._occupied_wall_group_keys(group_type, group_division, exclude_group_id=group_id)
        return targets & occupied

    def _show_group_conflict(self, parent, group: dict, count: int) -> None:
        section = normalize_room_finish_section(group.get("section", ""))
        group_type = normalize_finish_group_type(group.get("group_type", ""))
        group_division = normalize_finish_group_division(group.get("group_division", ""), group_type)
        group_label = f"{self.finish_panel.group_type_label(group_type)} > {group_division}"
        target_name = "벽체" if section == "벽" else "실"
        QMessageBox.warning(
            parent or self,
            "재료마감 그룹 중복 지정",
            f"선택한 {target_name} 중 {count:,}개가 이미 다른 '{group_label}' 그룹에 지정되어 있습니다.\n\n"
            f"같은 개체에는 {group_label} 그룹을 1개만 지정할 수 있습니다.\n"
            "기존 지정에서 해제한 뒤 다시 지정해 주세요.",
        )

    def _set_rooms_for_finish_group(self, group_id: str, handles) -> None:
        """그룹의 실 연결을 합집합이 아니라 지정 결과 그대로 교체한다.

        기존 지정 편집에서 체크/도면 선택을 빼면 해당 실 연결도 해제되고,
        추가하면 새로 연결된다. 그룹 자체와 재료 구성은 유지한다.
        """
        group_id = str(group_id or "")
        group = self.finish_groups.get(group_id)
        if not group:
            return
        target_handles = {str(value) for value in (handles or []) if value}
        old_handles = {str(value) for value in (group.get("room_handles", []) or []) if value}
        section = normalize_room_finish_section(group.get("section", ""))
        name = str(group.get("name", "") or "그룹")
        group_aci = finish_group_aci(group_id, group.get("color", ""), group.get("aci"))
        color = aci_to_hex(group_aci)
        group["aci"] = group_aci
        group["color"] = color
        formula = str(group.get("formula", "") or "")
        group_type = normalize_finish_group_type(group.get("group_type", ""))
        group_division = normalize_finish_group_division(group.get("group_division", ""), group_type)
        group["group_type"] = group_type
        group["group_division"] = group_division
        group["color"] = color
        materials = []
        for material in group.get("materials", []) or []:
            value = deepcopy(material)
            value["group_id"] = group_id
            value["group_name"] = name
            value["group_color"] = color
            value["group_aci"] = group_aci
            value["group_formula"] = formula
            value["group_type"] = group_type
            value["group_division"] = group_division
            materials.append(value)

        for handle in sorted(old_handles | target_handles):
            assignment = normalize_room_finish_assignment(self.room_finish_assignments.get(handle, {}))
            had_room_formula, room_formula = self._assignment_group_formula(assignment, group_id)
            if not had_room_formula:
                room_formula = self._group_default_formula_for_floor(group, self._room_floor_for_handle(handle))
            assignment["groups"] = [
                entry for entry in assignment.get("groups", [])
                if str(entry.get("id", "") or "") != group_id
            ]
            room_materials = self._group_materials_preserve_room_formulas(assignment, group_id, materials)
            for material in room_materials:
                material["group_formula"] = room_formula
            for part in assignment["items"]:
                assignment["items"][part] = [
                    material for material in assignment["items"][part]
                    if str(material.get("group_id", "") or "") != group_id
                ]
            if handle in target_handles:
                if section not in assignment["sections"]:
                    assignment["sections"].append(section)
                assignment["items"].setdefault(section, []).extend(room_materials)
                assignment["groups"].append({
                    "id": group_id, "name": name, "section": section,
                    "color": color, "aci": group_aci, "formula": room_formula, "group_type": group_type,
                    "group_division": group_division,
                })
            self.room_finish_assignments[handle] = assignment
        group["room_handles"] = sorted(target_handles)

    def add_materials_to_finish_group(self, group_id: str, materials) -> None:
        """층/전층 그룹목록에서 F5로 선택한 재료를 그룹 공통자료에 등록한다."""
        group_id = str(group_id or "")
        group = self.finish_groups.get(group_id)
        if not group:
            return
        group_name = str(group.get("name", "") or "그룹")
        group_aci = finish_group_aci(group_id, group.get("color", ""), group.get("aci"))
        color = aci_to_hex(group_aci)
        group["aci"] = group_aci
        group["color"] = color
        formula = str(group.get("formula", "") or "")
        group_type = normalize_finish_group_type(group.get("group_type", ""))
        group_division = normalize_finish_group_division(group.get("group_division", ""), group_type)
        group["group_type"] = group_type
        group["group_division"] = group_division
        group["color"] = color

        current = [deepcopy(value) for value in group.get("materials", []) or [] if isinstance(value, dict)]
        existing_ids = {str(value.get("id", "") or "") for value in current if str(value.get("id", "") or "")}
        added = 0
        for raw in materials or []:
            if not isinstance(raw, dict):
                continue
            value = deepcopy(raw)
            material_id = str(value.get("id", "") or "")
            if material_id and material_id in existing_ids:
                continue
            value["group_id"] = group_id
            value["group_name"] = group_name
            value["group_color"] = color
            value["group_aci"] = group_aci
            value["group_formula"] = formula
            value["group_type"] = group_type
            value["group_division"] = group_division
            current.append(value)
            if material_id:
                existing_ids.add(material_id)
            added += 1
        if not added:
            return
        group["materials"] = current

        # 그룹에 연결된 모든 실/벽에 동일 재료구성을 즉시 동기화한다.
        self._attach_rooms_to_finish_group(group_id, list(group.get("room_handles", []) or []))
        first_material = deepcopy(current[0]) if current else None
        for key in list(group.get("wall_keys", []) or []):
            if first_material:
                self.wall_finish_assignments[str(key)] = deepcopy(first_material)
            else:
                self.wall_finish_assignments.pop(str(key), None)

        self.save_finish_assignments()
        self.preview.set_wall_finish_assignments(self.wall_finish_assignments)
        self.preview.set_room_finish_assignments(self.room_finish_assignments)
        self.preview.set_finish_groups(self.finish_groups)
        self._refresh_current_finish_panel()
        self.update_finish_summary(self.selected_floor())
        self.data_changed.emit()

    def create_finish_group(self, section: str, multi_room: bool = False):
        section = normalize_room_finish_section(section)
        if multi_room:
            if section == "벽":
                self.start_wall_group_mode()
            else:
                self.start_room_group_mode(section)
            return
        room = self.selected_room()
        # 그룹 정의는 전층 공통이다. 전층/층 합계 어디에서 만들더라도
        # 실 연결 없이 그룹 자체를 먼저 등록할 수 있다.
        if room is None:
            selected = self._select_or_create_group(section, preferred_type=self.finish_panel.current_group_type())
            if selected:
                self.save_finish_assignments()
                self.preview.set_finish_groups(self.finish_groups)
                self._refresh_current_finish_panel()
                self.data_changed.emit()
            return
        selected = self._select_or_create_group(section, preferred_type=self.finish_panel.current_group_type())
        if not selected:
            return
        group_id, _name = selected
        conflicts = self._room_group_conflicts(group_id, [room.boundary_handle])
        if conflicts:
            self._show_group_conflict(self, self.finish_groups[group_id], len(conflicts))
            return
        self._attach_rooms_to_finish_group(group_id, [room.boundary_handle])
        self.save_finish_assignments()
        self.preview.set_room_finish_assignments(self.room_finish_assignments)
        self.preview.set_finish_groups(self.finish_groups)
        self._refresh_current_finish_panel()
        self.data_changed.emit()

    def _propagate_finish_group(self, group_id: str, source_assignment: dict):
        group = self.finish_groups.get(group_id)
        if not group:
            return
        section = normalize_room_finish_section(group.get("section", ""))
        name = str(group.get("name", "") or "그룹")
        group_aci = finish_group_aci(group_id, group.get("color", ""), group.get("aci"))
        color = aci_to_hex(group_aci)
        group["aci"] = group_aci
        group["color"] = color
        formula = str(group.get("formula", "") or "")
        group_type = normalize_finish_group_type(group.get("group_type", ""))
        group_division = normalize_finish_group_division(group.get("group_division", ""), group_type)
        group["group_type"] = group_type
        group["group_division"] = group_division
        group["color"] = color
        source = normalize_room_finish_assignment(source_assignment)
        previous_group_materials = [deepcopy(value) for value in group.get("materials", []) or []]
        materials = []
        for values in source["items"].values():
            for material in values:
                if str(material.get("group_id", "") or "") == group_id:
                    clean = deepcopy(material)
                    clean["group_id"] = group_id
                    clean["group_name"] = name
                    clean["group_color"] = color
                    clean["group_aci"] = group_aci
                    clean["group_formula"] = formula
                    clean["group_type"] = group_type
                    clean["group_division"] = group_division
                    # 실별 산식은 그룹 기준자료로 역전파하지 않는다.
                    for base in previous_group_materials:
                        if self._same_finish_material(base, clean):
                            clean["method"] = str(base.get("method", "") or "")
                            scope_methods = base.get("scope_methods", {})
                            if isinstance(scope_methods, dict) and scope_methods:
                                clean["scope_methods"] = deepcopy(scope_methods)
                            break
                    materials.append(clean)
        group["materials"] = deepcopy(materials)

        for handle in list(group.get("room_handles", []) or []):
            assignment = normalize_room_finish_assignment(self.room_finish_assignments.get(handle, {}))
            had_room_formula, room_formula = self._assignment_group_formula(assignment, group_id)
            if not had_room_formula:
                room_formula = self._group_default_formula_for_floor(group, self._room_floor_for_handle(handle))
            room_materials = self._group_materials_preserve_room_formulas(assignment, group_id, materials)
            for material in room_materials:
                material["group_formula"] = room_formula
            if section not in assignment["sections"]:
                assignment["sections"].append(section)
            for part in assignment["items"]:
                assignment["items"][part] = [
                    material for material in assignment["items"][part]
                    if str(material.get("group_id", "") or "") != group_id
                ]
            assignment["items"].setdefault(section, []).extend(room_materials)
            assignment["groups"] = [entry for entry in assignment.get("groups", []) if entry.get("id") != group_id]
            assignment["groups"].append({
                "id": group_id, "name": name, "section": section,
                "color": color, "aci": group_aci, "formula": room_formula, "group_type": group_type,
                "group_division": group_division,
            })
            self.room_finish_assignments[handle] = assignment

        first_material = deepcopy(materials[0]) if materials else None
        for key in list(group.get("wall_keys", []) or []):
            if first_material:
                self.wall_finish_assignments[str(key)] = deepcopy(first_material)
            else:
                self.wall_finish_assignments.pop(str(key), None)

    def reorder_finish_groups(self, section: str, group_type: str, ordered_ids) -> None:
        """같은 부위 안에서 드래그한 그룹의 표시 순서를 프로젝트 전체에 반영한다."""
        section = normalize_room_finish_section(section)
        group_type = normalize_finish_group_type(group_type)
        requested = [str(value or "") for value in list(ordered_ids or []) if str(value or "")]
        requested = [
            group_id for group_id in requested
            if group_id in self.finish_groups
            and normalize_room_finish_section(self.finish_groups[group_id].get("section", "")) == section
            and normalize_finish_group_type(self.finish_groups[group_id].get("group_type", "")) == group_type
        ]
        if len(requested) < 2:
            return
        requested_set = set(requested)

        # finish_groups의 dict 순서를 바꿔 전층/층 그룹목록에서도 동일 순서로 보이게 한다.
        old_items = list(self.finish_groups.items())
        target_positions = [
            index for index, (group_id, _group) in enumerate(old_items)
            if group_id in requested_set
        ]
        if len(target_positions) >= 2:
            by_id = {group_id: group for group_id, group in old_items if group_id in requested_set}
            new_items = list(old_items)
            for index, group_id in zip(target_positions, requested):
                new_items[index] = (group_id, by_id[group_id])
            self.finish_groups = dict(new_items)

        # 실별 저장본의 그룹 배열도 같은 상대순서로 맞춘다. 다른 그룹 위치는 건드리지 않는다.
        for handle, raw_assignment in list(self.room_finish_assignments.items()):
            assignment = normalize_room_finish_assignment(raw_assignment)
            groups = list(assignment.get("groups", []) or [])
            positions = [
                index for index, entry in enumerate(groups)
                if str(entry.get("id", "") or "") in requested_set
            ]
            if len(positions) < 2:
                continue
            entries = {str(entry.get("id", "") or ""): entry for entry in groups if str(entry.get("id", "") or "") in requested_set}
            ordered_present = [group_id for group_id in requested if group_id in entries]
            for index, group_id in zip(positions, ordered_present):
                groups[index] = entries[group_id]
            assignment["groups"] = groups
            self.room_finish_assignments[str(handle)] = assignment

        self.save_finish_assignments()
        self.preview.set_finish_groups(self.finish_groups)
        self._refresh_current_finish_panel()
        self.data_changed.emit()

    def move_finish_group_section(self, group_id: str, source_section: str, target_section: str) -> None:
        """그룹을 다른 부위로 옮긴다. 부위가 바뀌면 기존 도곽 객체지정은 모두 해제한다."""
        group_id = str(group_id or "")
        group = self.finish_groups.get(group_id)
        if not group:
            self._refresh_current_finish_panel()
            return
        source_section = normalize_room_finish_section(source_section or group.get("section", ""))
        target_section = normalize_room_finish_section(target_section)
        if not target_section or source_section == target_section:
            self._refresh_current_finish_panel()
            return

        # 확인창은 실제 드롭 위치를 가진 FinishAssignmentList에서 먼저 처리한다.
        # 여기서는 확인된 이동만 받아 데이터/객체 연결을 일괄 갱신한다.

        old_wall_keys = {str(value) for value in group.get("wall_keys", []) or [] if str(value)}

        # 객체 연결을 끊기 위해 모든 실 저장본에서 해당 그룹/재료를 제거한다.
        for handle, raw_assignment in list(self.room_finish_assignments.items()):
            assignment = normalize_room_finish_assignment(raw_assignment)
            changed = False
            old_count = len(assignment.get("groups", []) or [])
            assignment["groups"] = [
                entry for entry in assignment.get("groups", []) or []
                if str(entry.get("id", "") or "") != group_id
            ]
            changed = changed or len(assignment["groups"]) != old_count
            for section in list(assignment.get("items", {})):
                values = list(assignment["items"].get(section, []) or [])
                kept = [value for value in values if str(value.get("group_id", "") or "") != group_id]
                if len(kept) != len(values):
                    assignment["items"][section] = kept
                    changed = True
            if changed:
                self.room_finish_assignments[str(handle)] = assignment

        for key in old_wall_keys:
            self.wall_finish_assignments.pop(key, None)
        self.selected_wall_keys.difference_update(old_wall_keys)

        group["section"] = target_section
        group["room_handles"] = []
        group["wall_keys"] = []

        # 대상 부위의 마지막 그룹 뒤에 보이도록 전역 dict 순서도 옮긴다.
        moved = self.finish_groups.pop(group_id)
        items = list(self.finish_groups.items())
        insert_at = len(items)
        moved_type = normalize_finish_group_type(moved.get("group_type", ""))
        for index, (_other_id, other) in enumerate(items):
            if (
                normalize_room_finish_section(other.get("section", "")) == target_section
                and normalize_finish_group_type(other.get("group_type", "")) == moved_type
            ):
                insert_at = index + 1
        items.insert(insert_at, (group_id, moved))
        self.finish_groups = dict(items)

        self.save_finish_assignments()
        self.preview.set_wall_finish_assignments(self.wall_finish_assignments)
        self.preview.set_room_finish_assignments(self.room_finish_assignments)
        self.preview.set_finish_groups(self.finish_groups)
        self.preview.set_selected_wall_keys(self.selected_wall_keys)
        self._refresh_current_finish_panel()
        self.update_finish_summary(self.selected_floor())
        self.data_changed.emit()

    def delete_finish_group(self, group_id: str):
        group = self.finish_groups.get(str(group_id or ""))
        if not group:
            return
        room = self.selected_room()
        table_row = self.room_table.currentRow()
        global_delete = table_row == 0
        floor = "" if global_delete or room is not None else str(self._floor_for_table_row(table_row) or "")
        if global_delete:
            target_handles = set(group.get("room_handles", []) or [])
            question = f"'{group.get('name', '그룹')}' 그룹과 모든 층·실 연결 자료를 완전히 삭제할까요?"
        elif room is not None:
            target_handles = {str(room.boundary_handle or "")}
            question = f"'{group.get('name', '그룹')}' 그룹을 현재 실에서만 해제할까요?\n전체 그룹과 다른 실 연결은 유지됩니다."
        elif floor:
            target_handles = {str(value.boundary_handle or "") for value in self.rooms if value.floor == floor}
            question = f"'{group.get('name', '그룹')}' 그룹을 {floor}에서만 해제할까요?\n전체 그룹과 다른 층 연결은 유지됩니다."
        else:
            return
        if QMessageBox.question(
            self, "마감그룹 삭제",
            question,
            QMessageBox.Yes | QMessageBox.No, QMessageBox.No,
        ) != QMessageBox.Yes:
            return
        group_id = str(group_id)
        linked_handles = set(group.get("room_handles", []) or [])
        affected_handles = linked_handles & target_handles
        for handle in affected_handles:
            assignment = normalize_room_finish_assignment(self.room_finish_assignments.get(handle, {}))
            assignment["groups"] = [entry for entry in assignment.get("groups", []) if entry.get("id") != group_id]
            for section in assignment["items"]:
                assignment["items"][section] = [
                    material for material in assignment["items"][section]
                    if str(material.get("group_id", "") or "") != group_id
                ]
            self.room_finish_assignments[handle] = assignment
        removed_wall_keys = [
            str(key) for key in list(group.get("wall_keys", []) or [])
            if global_delete or wall_key_room_handle(key) in target_handles
        ]
        for key in removed_wall_keys:
            self.wall_finish_assignments.pop(key, None)
        if global_delete:
            self.finish_groups.pop(group_id, None)
        else:
            group["room_handles"] = sorted(linked_handles - target_handles)
            group["wall_keys"] = [key for key in group.get("wall_keys", []) or [] if str(key) not in removed_wall_keys]
        self.save_finish_assignments()
        self.preview.set_wall_finish_assignments(self.wall_finish_assignments)
        self.preview.set_room_finish_assignments(self.room_finish_assignments)
        self.preview.set_finish_groups(self.finish_groups)
        self._refresh_current_finish_panel()
        self.data_changed.emit()

    def change_finish_material_formula(self, material_ref: dict, formula: str) -> None:
        """재료행 산식을 실/층/전층 범위에 맞게 저장한다.

        실 행에서는 해당 실만 수정한다. 층 합계/전층 행에서는 먼저 현재
        집계행의 기준 산식을 저장한 뒤, 연결된 개별 실에 덮어쓸지 별도로 묻는다.
        따라서 '아니오'를 눌러도 방금 입력한 층-전체/전층 산식 자체는 유지된다.
        """
        material_ref = dict(material_ref or {})
        formula = normalize_formula_text(formula)
        group_id = str(material_ref.get("group_id", "") or "")
        current_room = self.selected_room()
        current_row = self.room_table.currentRow()
        scope_key = ""

        if current_room is not None:
            target_rooms = [current_room]
            apply_mode = "overwrite"
            scope_text = f"{current_room.floor} / {current_room.room_number or '-'} / {normalize_room_name(current_room.room_name)}"
        else:
            if current_row == 0:
                target_rooms = list(self.rooms)
                scope_text = "전층"
                scope_key = "__all__"
            else:
                floor = str(self._floor_for_table_row(current_row) or "")
                if not floor:
                    return
                target_rooms = [room for room in self.rooms if room.floor == floor]
                scope_text = f"{floor} 전체"
                scope_key = floor

            # 사용자가 편집한 '전체' 행의 기준 산식을 먼저 저장한다.
            # 이후 선택하는 덮어쓰기/빈칸만/아니오는 개별 실로의 전파 범위만 결정한다.
            group = self.finish_groups.get(group_id)
            if group is not None:
                for material in group.get("materials", []) or []:
                    if not self._same_finish_material(material, material_ref):
                        continue
                    scope_methods = material.get("scope_methods", {})
                    if not isinstance(scope_methods, dict):
                        scope_methods = {}
                    scope_methods = dict(scope_methods)
                    scope_methods[scope_key] = formula
                    material["scope_methods"] = scope_methods
                    # 전층 기준 산식은 기존 데이터/신규 연결 호환을 위해 method에도 같이 둔다.
                    if scope_key == "__all__":
                        material["method"] = formula

            box = QMessageBox(self)
            box.setWindowTitle("산식 공통 적용")
            box.setIcon(QMessageBox.Question)
            box.setText(f"{scope_text}의 같은 재료에 이 산식을 공통 적용할까요?")
            box.setInformativeText(
                f"산식: {formula or '(공란)'}\n\n"
                "산식 덮어씌우기: 기존 산식도 모두 변경\n"
                "빈칸만 적용: 산식이 없는 항목에만 입력\n"
                "아니오: 전체 기준 산식만 변경하고 개별 실 산식은 유지"
            )
            overwrite_btn = box.addButton("산식 덮어씌우기", QMessageBox.AcceptRole)
            blank_btn = box.addButton("빈칸만 적용", QMessageBox.ActionRole)
            no_btn = box.addButton("아니오", QMessageBox.RejectRole)
            box.setDefaultButton(blank_btn)
            box.exec()
            clicked = box.clickedButton()
            if clicked is overwrite_btn:
                apply_mode = "overwrite"
            elif clicked is blank_btn:
                apply_mode = "blank"
            else:
                apply_mode = "none"

        target_handles = {str(room.boundary_handle or "") for room in target_rooms if str(room.boundary_handle or "")}
        changed = 0
        if apply_mode != "none":
            for handle in target_handles:
                assignment = normalize_room_finish_assignment(self.room_finish_assignments.get(handle, {}))
                assignment_changed = False
                for values in assignment.get("items", {}).values():
                    for material in values:
                        if group_id and str(material.get("group_id", "") or "") != group_id:
                            continue
                        if not group_id and str(material.get("group_id", "") or ""):
                            continue
                        if not self._same_finish_material(material, material_ref):
                            continue
                        if apply_mode == "blank" and normalize_formula_text(material.get("method", "")):
                            continue
                        if normalize_formula_text(material.get("method", "")) == formula:
                            continue
                        material["method"] = formula
                        assignment_changed = True
                        changed += 1
                if assignment_changed:
                    self.room_finish_assignments[handle] = assignment

            # 벽 그룹에 보관된 표시/연동 재료 사본도 실제 전파를 선택한 범위에서는 같이 맞춘다.
            for key, material in list(self.wall_finish_assignments.items()):
                if wall_key_room_handle(key) not in target_handles or not isinstance(material, dict):
                    continue
                if group_id and str(material.get("group_id", "") or "") != group_id:
                    continue
                if not self._same_finish_material(material, material_ref):
                    continue
                if apply_mode == "blank" and normalize_formula_text(material.get("method", "")):
                    continue
                material["method"] = formula

        self.save_finish_assignments()
        self._refresh_current_finish_panel()
        self.update_finish_summary(self.selected_floor())
        self.data_changed.emit()

    def change_finish_group_formula(self, group_id: str, formula: str) -> None:
        """그룹 산식을 개별실/층-전체/전층 범위에 맞게 저장한다.

        개별실에서는 현재 실에만 저장한다. 층-전체/전층에서는 먼저 해당
        집계행의 기준 산식을 저장하고, 실제 실별 그룹 산식을 덮어쓸지 묻는다.
        """
        group_id = str(group_id or "")
        group = self.finish_groups.get(group_id)
        if not group:
            return
        formula = normalize_formula_text(formula)
        current_room = self.selected_room()
        current_row = self.room_table.currentRow()

        if current_room is not None:
            handle = str(current_room.boundary_handle or "")
            if not handle:
                return
            assignment = normalize_room_finish_assignment(self.room_finish_assignments.get(handle, {}))
            assignment, found = self._set_assignment_group_formula(assignment, group_id, formula)
            if not found:
                return
            self.room_finish_assignments[handle] = assignment
            target_handles = {handle}
        else:
            if current_row == 0:
                floor = ""
                target_rooms = list(self.rooms)
                scope_text = "전체층"
                scope_key = "__all__"
            else:
                floor = str(self._floor_for_table_row(current_row) or "")
                if not floor:
                    return
                target_rooms = [room for room in self.rooms if room.floor == floor]
                scope_text = f"{floor} 전체"
                scope_key = floor

            # 집계행 자체의 기준 산식은 질문 결과와 관계없이 먼저 저장한다.
            scope_formulas = group.get("scope_formulas", {})
            if not isinstance(scope_formulas, dict):
                scope_formulas = {}
            scope_formulas = dict(scope_formulas)
            scope_formulas[scope_key] = formula
            group["scope_formulas"] = scope_formulas
            if scope_key == "__all__":
                # 전층 기준은 앞으로 새로 연결되는 실의 기본값으로도 사용한다.
                group["formula"] = formula
                for material in group.get("materials", []) or []:
                    material["group_formula"] = formula

            box = QMessageBox(self)
            box.setWindowTitle("그룹 산식 적용")
            box.setIcon(QMessageBox.Question)
            box.setText(f"{scope_text}의 '{group.get('name', '그룹')}' 그룹 산식을 적용할 범위를 선택해 주세요.")
            if current_row == 0:
                box.setInformativeText(
                    f"산식: {formula or '(공란)'}\n\n"
                    "전체층 적용: 전체층의 기존 그룹 산식까지 모두 변경\n"
                    "빈칸 적용: 전체층 중 그룹 산식이 비어 있는 실에만 입력\n"
                    "아니요: 전체층 기준 산식만 변경하고 각 실 산식은 유지"
                )
                current_floor_btn = None
                all_floor_btn = box.addButton("전체층 적용", QMessageBox.AcceptRole)
            else:
                box.setInformativeText(
                    f"산식: {formula or '(공란)'}\n\n"
                    "현재층 적용: 현재층의 기존 그룹 산식까지 모두 변경\n"
                    "전체층 적용: 모든 층의 기존 그룹 산식까지 모두 변경\n"
                    "빈칸 적용: 현재층에서 그룹 산식이 비어 있는 실에만 입력\n"
                    f"아니요: {scope_text} 기준 산식만 변경하고 각 실 산식은 유지"
                )
                current_floor_btn = box.addButton("현재층 적용", QMessageBox.AcceptRole)
                all_floor_btn = box.addButton("전체층 적용", QMessageBox.ActionRole)
            blank_btn = box.addButton("빈칸 적용", QMessageBox.ActionRole)
            no_btn = box.addButton("아니요", QMessageBox.RejectRole)
            box.setDefaultButton(blank_btn)
            box.exec()
            clicked = box.clickedButton()
            apply_mode = "none"
            apply_all_floors = clicked is all_floor_btn
            if (current_floor_btn is not None and clicked is current_floor_btn) or clicked is all_floor_btn:
                apply_mode = "overwrite"
            elif clicked is blank_btn:
                apply_mode = "blank"

            if apply_all_floors:
                target_rooms = list(self.rooms)
                scope_formulas["__all__"] = formula
                group["scope_formulas"] = scope_formulas
                group["formula"] = formula
                for material in group.get("materials", []) or []:
                    material["group_formula"] = formula

            target_handles = set()
            if apply_mode != "none":
                for room in target_rooms:
                    handle = str(room.boundary_handle or "")
                    if not handle:
                        continue
                    assignment = normalize_room_finish_assignment(self.room_finish_assignments.get(handle, {}))
                    found, current_formula = self._assignment_group_formula(assignment, group_id)
                    if not found:
                        continue
                    if apply_mode == "blank" and normalize_formula_text(current_formula):
                        continue
                    assignment, _ = self._set_assignment_group_formula(assignment, group_id, formula)
                    self.room_finish_assignments[handle] = assignment
                    target_handles.add(handle)

        # 벽 그룹의 별도 사본에도 같은 실별 그룹 산식을 맞춰 둔다.
        if target_handles:
            for key in list(group.get("wall_keys", []) or []):
                handle = wall_key_room_handle(key)
                if handle not in target_handles:
                    continue
                material = self.wall_finish_assignments.get(str(key))
                if isinstance(material, dict):
                    material["group_formula"] = formula

        self.save_finish_assignments()
        self.preview.set_room_finish_assignments(self.room_finish_assignments)
        self.preview.set_wall_finish_assignments(self.wall_finish_assignments)
        self.preview.set_finish_groups(self.finish_groups)
        self._refresh_current_finish_panel()
        self.update_finish_summary(self.selected_floor())
        self.data_changed.emit()

    def rename_finish_group(self, group_id: str) -> None:
        group_id = str(group_id or "")
        group = self.finish_groups.get(group_id)
        if not group:
            return
        current_name = str(group.get("name", "") or "그룹")
        name, ok = QInputDialog.getText(
            self, "그룹 이름 변경", "새 그룹 이름:", QLineEdit.Normal, current_name
        )
        name = str(name or "").strip()
        if not ok or not name or name == current_name:
            return

        group["name"] = name
        for material in group.get("materials", []) or []:
            if isinstance(material, dict):
                material["group_name"] = name

        for handle, raw in list(self.room_finish_assignments.items()):
            assignment = normalize_room_finish_assignment(raw)
            changed = False
            for entry in assignment.get("groups", []) or []:
                if str(entry.get("id", "") or "") == group_id:
                    entry["name"] = name
                    changed = True
            for materials in assignment.get("items", {}).values():
                for material in materials:
                    if str(material.get("group_id", "") or "") == group_id:
                        material["group_name"] = name
                        changed = True
            if changed:
                self.room_finish_assignments[handle] = assignment

        for key, material in list(self.wall_finish_assignments.items()):
            if isinstance(material, dict) and str(material.get("group_id", "") or "") == group_id:
                material["group_name"] = name

        self.save_finish_assignments()
        self.preview.set_wall_finish_assignments(self.wall_finish_assignments)
        self.preview.set_room_finish_assignments(self.room_finish_assignments)
        self.preview.set_finish_groups(self.finish_groups)
        self._refresh_current_finish_panel()
        self.update_finish_summary(self.selected_floor())
        self.data_changed.emit()

    def change_finish_group_division(self, group_id: str) -> None:
        """기존 그룹의 지정 객체/재료를 유지한 채 현재 대분류 탭 안의 구분만 변경한다."""
        group_id = str(group_id or "")
        group = self.finish_groups.get(group_id)
        if not group:
            return

        group_type = normalize_finish_group_type(group.get("group_type", ""))
        current_division = normalize_finish_group_division(group.get("group_division", ""), group_type)
        divisions = self.finish_panel.group_divisions(group_type)
        if current_division not in divisions:
            divisions = [current_division] + [value for value in divisions if value != current_division]
        if len(divisions) <= 1:
            QMessageBox.information(
                self, "그룹 구분 변경",
                f"'{self.finish_panel.group_type_label(group_type)}' 탭에 변경할 다른 구분이 없습니다.\n"
                "탭별 설정에서 구분을 먼저 추가해 주세요.",
            )
            return

        current_index = max(0, divisions.index(current_division)) if current_division in divisions else 0
        new_division, ok = QInputDialog.getItem(
            self,
            "그룹 구분 변경",
            f"[{self.finish_panel.group_type_label(group_type)}] 변경할 구분:",
            divisions,
            current_index,
            False,
        )
        new_division = str(new_division or "").strip()
        if not ok or not new_division or new_division == current_division:
            return
        new_division = self.finish_panel.ensure_group_division(group_type, new_division)

        # 바닥/천정/벽 등 이미 지정된 객체가 새 구분의 다른 그룹과 겹치면 변경하지 않는다.
        section = normalize_room_finish_section(group.get("section", ""))
        if section == "벽":
            targets = {str(value) for value in (group.get("wall_keys", []) or []) if value}
            occupied = self._occupied_wall_group_keys(group_type, new_division, exclude_group_id=group_id)
            conflicts = targets & occupied
        else:
            targets = {str(value) for value in (group.get("room_handles", []) or []) if value}
            occupied = self._occupied_room_group_handles(
                section, group_type, new_division, exclude_group_id=group_id
            )
            conflicts = targets & occupied
        if conflicts:
            target_name = "벽체" if section == "벽" else "실"
            QMessageBox.warning(
                self,
                "그룹 구분 변경",
                f"이 그룹의 {target_name} 중 {len(conflicts):,}개가 이미 다른 "
                f"'{self.finish_panel.group_type_label(group_type)} > {new_division}' 그룹에 지정되어 있습니다.\n\n"
                "중복되는 기존 지정을 먼저 해제한 뒤 구분을 변경해 주세요.",
            )
            return

        group["group_division"] = new_division
        for material in group.get("materials", []) or []:
            if isinstance(material, dict):
                material["group_type"] = group_type
                material["group_division"] = new_division

        # 실별 그룹 사본과 재료 사본도 같은 구분으로 동기화한다.
        for handle, raw in list(self.room_finish_assignments.items()):
            assignment = normalize_room_finish_assignment(raw)
            changed = False
            for entry in assignment.get("groups", []) or []:
                if str(entry.get("id", "") or "") == group_id:
                    entry["group_type"] = group_type
                    entry["group_division"] = new_division
                    changed = True
            for materials in assignment.get("items", {}).values():
                for material in materials:
                    if str(material.get("group_id", "") or "") == group_id:
                        material["group_type"] = group_type
                        material["group_division"] = new_division
                        changed = True
            if changed:
                self.room_finish_assignments[handle] = assignment

        # 벽 지정은 별도 저장 사본도 사용하므로 같이 갱신한다.
        for key, material in list(self.wall_finish_assignments.items()):
            if isinstance(material, dict) and str(material.get("group_id", "") or "") == group_id:
                material["group_type"] = group_type
                material["group_division"] = new_division

        self.save_finish_assignments()
        self.preview.set_wall_finish_assignments(self.wall_finish_assignments)
        self.preview.set_room_finish_assignments(self.room_finish_assignments)
        self.preview.set_finish_groups(self.finish_groups)
        self._set_finish_group_display_filter(group_type, new_division)
        self.preview.set_focused_finish_group(group_id)
        self._refresh_current_finish_panel()
        self.update_finish_summary(self.selected_floor())
        self.data_changed.emit()

    def change_finish_group_color(self, group_id: str):
        if self.room_table.currentRow() != 0:
            QMessageBox.information(self, "그룹 색상 변경", "그룹 색상은 전층 그룹 목록에서 변경해 주세요.")
            return
        group = self.finish_groups.get(str(group_id or ""))
        if not group:
            return
        current_aci = finish_group_aci(group_id, group.get("color", ""), group.get("aci"))
        selected_aci = get_aci_color(current_aci, self, f"{group.get('name', '그룹')} 색상 선택")
        if selected_aci is None:
            return
        value = aci_to_hex(selected_aci)
        group["aci"] = int(selected_aci)
        group["color"] = value
        for material in group.get("materials", []) or []:
            material["group_color"] = value
            material["group_aci"] = int(selected_aci)
        for handle, raw in list(self.room_finish_assignments.items()):
            assignment = normalize_room_finish_assignment(raw)
            for entry in assignment.get("groups", []):
                if str(entry.get("id", "") or "") == str(group_id):
                    entry["color"] = value
                    entry["aci"] = int(selected_aci)
            for materials in assignment.get("items", {}).values():
                for material in materials:
                    if str(material.get("group_id", "") or "") == str(group_id):
                        material["group_color"] = value
                        material["group_aci"] = int(selected_aci)
            self.room_finish_assignments[handle] = assignment
        for key, material in list(self.wall_finish_assignments.items()):
            if isinstance(material, dict) and str(material.get("group_id", "") or "") == str(group_id):
                material["group_color"] = value
                material["group_aci"] = int(selected_aci)
        self.save_finish_assignments()
        self.preview.set_finish_groups(self.finish_groups)
        self.preview.set_room_finish_assignments(self.room_finish_assignments)
        self.preview.set_wall_finish_assignments(self.wall_finish_assignments)
        self._refresh_current_finish_panel()
        self.data_changed.emit()

    def break_finish_group(self, group_id: str):
        room = self.selected_room()
        group = self.finish_groups.get(str(group_id or ""))
        if room is None or not group:
            return
        handle = str(room.boundary_handle or "")
        if handle not in list(group.get("room_handles", []) or []):
            return
        new_id = uuid.uuid4().hex
        new_group = deepcopy(group)
        new_group.update({"id": new_id, "name": f"{group.get('name', '그룹')} (개별)", "room_handles": [handle]})
        for material in new_group.get("materials", []) or []:
            material["group_id"] = new_id
            material["group_name"] = new_group["name"]
        own_wall_keys = [key for key in group.get("wall_keys", []) or [] if wall_key_room_handle(key) == handle]
        new_group["wall_keys"] = own_wall_keys
        group["room_handles"] = [value for value in group.get("room_handles", []) if value != handle]
        group["wall_keys"] = [value for value in group.get("wall_keys", []) if value not in own_wall_keys]
        self.finish_groups[new_id] = new_group

        assignment = normalize_room_finish_assignment(self.room_finish_assignments.get(handle, {}))
        for entry in assignment.get("groups", []):
            if entry.get("id") == group_id:
                entry.update({"id": new_id, "name": new_group["name"]})
        for materials in assignment["items"].values():
            for material in materials:
                if str(material.get("group_id", "") or "") == group_id:
                    material["group_id"] = new_id
                    material["group_name"] = new_group["name"]
        self.room_finish_assignments[handle] = assignment
        self.save_finish_assignments()
        self.preview.set_room_finish_assignments(self.room_finish_assignments)
        self.preview.set_finish_groups(self.finish_groups)
        self._refresh_current_finish_panel()
        self.data_changed.emit()

    def _sync_room_group_checklist(self):
        room_list = self._room_group_list
        if room_list is None:
            return
        room_list.blockSignals(True)
        try:
            for row in range(room_list.count()):
                item = room_list.item(row)
                handle = str(item.data(Qt.UserRole) or "")
                item.setCheckState(Qt.Checked if handle in self._selected_group_room_handles else Qt.Unchecked)
        finally:
            room_list.blockSignals(False)

    def _show_finish_group_edit_overlay(
        self, section: str, group_label: str, kind: str, count: int,
        confirm_callback, cancel_callback, detail_dialog=None,
    ) -> None:
        """바닥/천정/벽 지정용 도곽 상단 비차단 상태창을 표시한다."""
        if self._opening_edit_mode:
            self.opening_edit_btn.setChecked(False)
        self._finish_group_overlay_section = str(section or "그룹")
        self._finish_group_overlay_label = str(group_label or "")
        self._finish_group_overlay_kind = str(kind or "room")
        self._finish_group_overlay_confirm_callback = confirm_callback
        self._finish_group_overlay_cancel_callback = cancel_callback
        self._finish_group_overlay_detail_dialog = detail_dialog
        self.finish_group_edit_list_btn.setVisible(detail_dialog is not None and self._finish_group_overlay_kind == "room")
        self.finish_group_edit_overlay.show()
        self.finish_group_edit_overlay.raise_()
        self._update_finish_group_edit_status(count)
        self.preview.setFocus(Qt.OtherFocusReason)

    def _update_finish_group_edit_status(self, count: int | None = None) -> None:
        if count is None:
            count = len(self.selected_wall_keys) if self._finish_group_overlay_kind == "wall" else len(self._selected_group_room_handles)
        section = self._finish_group_overlay_section or "그룹"
        label = self._finish_group_overlay_label
        prefix = f"{section} 지정" + (f" | {label}" if label else "")
        target = "벽" if self._finish_group_overlay_kind == "wall" else "실"
        self.finish_group_edit_status_label.setText(
            f"{prefix} | 선택 {int(count or 0):,}개 | {target} 클릭/드래그 선택 · Ctrl+드래그 해제 · ESC 취소"
        )

    def _hide_finish_group_edit_overlay(self) -> None:
        self.finish_group_edit_overlay.hide()
        self._finish_group_overlay_section = ""
        self._finish_group_overlay_label = ""
        self._finish_group_overlay_kind = ""
        self._finish_group_overlay_confirm_callback = None
        self._finish_group_overlay_cancel_callback = None
        self._finish_group_overlay_detail_dialog = None

    def _toggle_finish_group_detail_dialog(self) -> None:
        dialog = self._finish_group_overlay_detail_dialog
        if dialog is None:
            return
        if dialog.isVisible():
            dialog.hide()
            self.preview.setFocus(Qt.OtherFocusReason)
            return
        dialog.show()
        dialog.raise_()
        dialog.activateWindow()
        anchor = self.preview.mapToGlobal(self.preview.rect().topRight())
        dialog.move(anchor.x() + 8, max(20, anchor.y()))

    def _confirm_finish_group_overlay(self) -> None:
        callback = self._finish_group_overlay_confirm_callback
        if callable(callback):
            callback()

    def _cancel_finish_group_overlay(self) -> None:
        callback = self._finish_group_overlay_cancel_callback
        if callable(callback):
            callback()

    def _finish_room_group_mode(self):
        self._hide_finish_group_edit_overlay()
        self.preview.set_room_group_mode(False)
        self.preview.set_selected_group_room_handles(set())
        self._selected_group_room_handles.clear()
        self._group_mode_blocked_room_handles.clear()
        self._room_group_list = None
        self._room_group_dialog = None

    def start_room_group_mode(self, section: str):
        if self.boundary_split_btn.isChecked():
            self.boundary_split_btn.setChecked(False)
        if self.boundary_merge_btn.isChecked():
            self.boundary_merge_btn.setChecked(False)
        """바닥/천정 그룹을 도면 클릭·드래그와 옆 체크목록으로 지정한다."""
        if not self.rooms:
            return
        section = normalize_room_finish_section(section)
        if self._wall_group_dialog is not None:
            self._wall_group_dialog.close()
        if self._room_group_dialog is not None:
            self._room_group_dialog.close()
        active_group = self._active_group_for_section(section)
        mode_group_type = (
            normalize_finish_group_type(active_group.get("group_type", ""))
            if active_group is not None
            else self.finish_panel.current_group_type()
        )
        mode_group_division = (
            normalize_finish_group_division(active_group.get("group_division", ""), mode_group_type)
            if active_group is not None
            else self._choose_group_division_for_mode(mode_group_type, self)
        )
        if not mode_group_division:
            return
        # 지정모드 진입과 동시에 같은 탭/같은 구분의 기지정 객체만 그룹색으로 보이게 한다.
        self._set_finish_group_display_filter(mode_group_type, mode_group_division)
        mode_group_label = f"{self.finish_panel.group_type_label(mode_group_type)} > {mode_group_division}"
        self.group_overlay_check.setChecked(True)
        active_id = str(active_group.get("id", "") or "") if active_group is not None else ""
        self._group_mode_blocked_room_handles = self._occupied_room_group_handles(
            section, mode_group_type, mode_group_division, exclude_group_id=active_id
        )
        if active_group is not None:
            # 선택된 그룹을 편집할 때는 현재 실 선택과 무관하게 기존 지정 상태 그대로 시작한다.
            selected = {str(value) for value in (active_group.get("room_handles", []) or []) if value}
        else:
            selected = {str(room.boundary_handle or "") for room in self.selected_rooms() if room.boundary_handle}
            current = self.selected_room()
            if current is not None and current.boundary_handle:
                selected.add(str(current.boundary_handle))
            selected.difference_update(self._group_mode_blocked_room_handles)
        self._selected_group_room_handles = selected

        self.view_mode.setCurrentText("층 전체 보기")
        self.refresh_preview()
        self.preview.set_room_group_mode(True)
        self.preview.set_selected_group_room_handles(selected)

        dialog = QDialog(self)
        dialog.setWindowTitle(f"{section} 마감그룹 지정")
        dialog.setModal(False)
        dialog.resize(350, 520)
        layout = QVBoxLayout(dialog)
        guide = QLabel(
            f"현재 지정 구분: {mode_group_label}\n"
            "같은 구분으로 이미 지정된 실은 화면에 그룹색으로 표시되며 재선택되지 않습니다.\n"
            "도면에서 실을 클릭하거나 드래그해 선택하세요.  Ctrl+드래그: 선택 해제"
        )
        guide.setWordWrap(True)
        layout.addWidget(guide)
        room_list = QListWidget()
        for room in self.rooms:
            label = f"{room.floor} / {room.room_number or '-'} / {normalize_room_name(room.room_name)}"
            item = QListWidgetItem(label)
            item.setData(Qt.UserRole, str(room.boundary_handle or ""))
            item.setFlags(item.flags() | Qt.ItemIsUserCheckable)
            handle = str(room.boundary_handle or "")
            if handle in self._group_mode_blocked_room_handles:
                item.setFlags(item.flags() & ~Qt.ItemIsEnabled)
                item.setToolTip(f"이미 다른 {mode_group_label} 그룹에 지정된 실")
            item.setCheckState(Qt.Checked if handle in selected else Qt.Unchecked)
            room_list.addItem(item)
        layout.addWidget(room_list, 1)
        count_label = QLabel()
        layout.addWidget(count_label)
        buttons = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        buttons.button(QDialogButtonBox.Ok).setText("확인")
        buttons.button(QDialogButtonBox.Cancel).setText("취소")
        layout.addWidget(buttons)

        def update_count():
            count = len(self._selected_group_room_handles)
            count_label.setText(f"선택 실: {count:,}개")
            self._update_finish_group_edit_status(count)

        def checklist_changed(item):
            handle = str(item.data(Qt.UserRole) or "")
            if handle in self._group_mode_blocked_room_handles:
                item.setCheckState(Qt.Unchecked)
                return
            if item.checkState() == Qt.Checked:
                self._selected_group_room_handles.add(handle)
            else:
                self._selected_group_room_handles.discard(handle)
            self.preview.set_selected_group_room_handles(self._selected_group_room_handles)
            update_count()

        def finish():
            # 기존 그룹 편집은 0개 선택도 허용한다. 전부 빼면 그룹은 남고 지정만 비워진다.
            if active_group is not None:
                group_id = str(active_group.get("id", "") or "")
                conflicts = self._room_group_conflicts(group_id, self._selected_group_room_handles)
                if conflicts:
                    self._show_group_conflict(dialog, self.finish_groups[group_id], len(conflicts))
                    return
                self._set_rooms_for_finish_group(group_id, self._selected_group_room_handles)
            else:
                if not self._selected_group_room_handles:
                    QMessageBox.information(dialog, "마감그룹 지정", "실을 하나 이상 선택해 주세요.")
                    return
                selected_group = self._select_or_create_group(
                    section, dialog, preferred_type=mode_group_type, preferred_division=mode_group_division
                )
                if not selected_group:
                    return
                group_id, _name = selected_group
                conflicts = self._room_group_conflicts(group_id, self._selected_group_room_handles)
                if conflicts:
                    self._show_group_conflict(dialog, self.finish_groups[group_id], len(conflicts))
                    return
                self._attach_rooms_to_finish_group(group_id, self._selected_group_room_handles)
            self.save_finish_assignments()
            self.preview.set_room_finish_assignments(self.room_finish_assignments)
            self.preview.set_finish_groups(self.finish_groups)
            dialog.accept()
            self._refresh_current_finish_panel()
            self.data_changed.emit()

        room_list.itemChanged.connect(checklist_changed)
        buttons.accepted.connect(finish)
        buttons.rejected.connect(dialog.reject)
        dialog.finished.connect(lambda _result: self._finish_room_group_mode())
        self._room_group_dialog = dialog
        self._room_group_list = room_list
        self._show_finish_group_edit_overlay(
            section, mode_group_label, "room", len(self._selected_group_room_handles),
            finish, dialog.reject, detail_dialog=dialog,
        )
        update_count()

    def on_group_room_clicked(self, handle: str):
        handle = str(handle or "")
        if not handle:
            return
        if handle in self._selected_group_room_handles:
            self._selected_group_room_handles.remove(handle)
        elif handle not in self._group_mode_blocked_room_handles:
            self._selected_group_room_handles.add(handle)
        self.preview.set_selected_group_room_handles(self._selected_group_room_handles)
        self._sync_room_group_checklist()
        self._update_finish_group_edit_status(len(self._selected_group_room_handles))

    def on_group_rooms_box_selected(self, handles, additive: bool, remove: bool):
        handles = {str(value) for value in (handles or []) if value}
        if not remove:
            handles.difference_update(self._group_mode_blocked_room_handles)
        if remove:
            self._selected_group_room_handles.difference_update(handles)
        elif additive:
            self._selected_group_room_handles.update(handles)
        else:
            self._selected_group_room_handles = handles
        self.preview.set_selected_group_room_handles(self._selected_group_room_handles)
        self._sync_room_group_checklist()
        self._update_finish_group_edit_status(len(self._selected_group_room_handles))

    def start_wall_group_mode(self):
        if self.boundary_split_btn.isChecked():
            self.boundary_split_btn.setChecked(False)
        if self.boundary_merge_btn.isChecked():
            self.boundary_merge_btn.setChecked(False)
        if not self.rooms:
            return
        if self._room_group_dialog is not None:
            self._room_group_dialog.close()
        active_group = self._active_group_for_section("벽")
        mode_group_type = (
            normalize_finish_group_type(active_group.get("group_type", ""))
            if active_group is not None
            else self.finish_panel.current_group_type()
        )
        mode_group_division = (
            normalize_finish_group_division(active_group.get("group_division", ""), mode_group_type)
            if active_group is not None
            else self._choose_group_division_for_mode(mode_group_type, self)
        )
        if not mode_group_division:
            return
        # 벽지정에서도 같은 탭/같은 구분의 기지정 벽만 그룹색+선택불가로 보여준다.
        self._set_finish_group_display_filter(mode_group_type, mode_group_division)
        mode_group_label = f"{self.finish_panel.group_type_label(mode_group_type)} > {mode_group_division}"
        self.group_overlay_check.setChecked(True)
        active_id = str(active_group.get("id", "") or "") if active_group is not None else ""
        self._group_mode_blocked_wall_keys = self._occupied_wall_group_keys(
            mode_group_type, mode_group_division, exclude_group_id=active_id
        )
        if active_group is not None:
            # 선택된 벽 그룹의 기존 지정선을 그대로 불러와 추가/해제 편집한다.
            self.selected_wall_keys = {str(value) for value in (active_group.get("wall_keys", []) or []) if value}
        else:
            self.selected_wall_keys.clear()
        self.preview.set_wall_finish_mode(True)
        self.view_mode.setCurrentText("층 전체 보기")
        self.refresh_preview()
        self.preview.set_selected_wall_keys(self.selected_wall_keys)

        if self._wall_group_dialog is not None:
            self._wall_group_dialog.close()
        dialog = QDialog(self)
        dialog.setWindowTitle("벽 마감그룹 지정")
        dialog.setModal(False)
        dialog.setFixedWidth(360)
        layout = QVBoxLayout(dialog)
        label = QLabel(
            f"현재 지정 구분: {mode_group_label}\n"
            "같은 구분으로 이미 지정된 벽은 그룹색으로 표시되며 재선택되지 않습니다.\n"
            "도면에서 벽체를 클릭하거나 드래그로 선택한 뒤 확인을 누르세요."
        )
        label.setWordWrap(True)
        layout.addWidget(label)
        count_label = QLabel("선택 벽체: 0개")
        layout.addWidget(count_label)
        buttons = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        buttons.button(QDialogButtonBox.Ok).setText("확인")
        buttons.button(QDialogButtonBox.Cancel).setText("취소")
        layout.addWidget(buttons)

        def refresh_count():
            count = len(self.selected_wall_keys)
            count_label.setText(f"선택 벽체: {count:,}개")
            self._update_finish_group_edit_status(count)
            if dialog.isVisible():
                QTimer.singleShot(150, refresh_count)

        def finish():
            # 기존 그룹 편집은 지정선을 전부 제거하는 것도 허용한다.
            if active_group is not None:
                group_id = str(active_group.get("id", "") or "")
                group = self.finish_groups[group_id]
                old_wall_keys = {str(value) for value in (group.get("wall_keys", []) or []) if value}
                new_wall_keys = {str(value) for value in self.selected_wall_keys if value}
                conflicts = self._wall_group_conflicts(group_id, new_wall_keys)
                if conflicts:
                    self._show_group_conflict(dialog, group, len(conflicts))
                    return
                for key in old_wall_keys - new_wall_keys:
                    self.wall_finish_assignments.pop(key, None)
                group["wall_keys"] = sorted(new_wall_keys)
                handles = sorted({wall_key_room_handle(key) for key in new_wall_keys if wall_key_room_handle(key)})
                self._set_rooms_for_finish_group(group_id, handles)
            else:
                if not self.selected_wall_keys:
                    QMessageBox.information(dialog, "벽 마감그룹 지정", "벽체를 하나 이상 선택해 주세요.")
                    return
                selected_group = self._select_or_create_group(
                    "벽", dialog, preferred_type=mode_group_type, preferred_division=mode_group_division
                )
                if not selected_group:
                    return
                group_id, _name = selected_group
                group = self.finish_groups[group_id]
                conflicts = self._wall_group_conflicts(group_id, self.selected_wall_keys)
                if conflicts:
                    self._show_group_conflict(dialog, group, len(conflicts))
                    return
                group["wall_keys"] = sorted(set(group.get("wall_keys", []) or []) | self.selected_wall_keys)
                handles = sorted({wall_key_room_handle(key) for key in self.selected_wall_keys if wall_key_room_handle(key)})
                self._attach_rooms_to_finish_group(group_id, handles)

            first_material = deepcopy((group.get("materials", []) or [None])[0])
            if first_material:
                for key in self.selected_wall_keys:
                    self.wall_finish_assignments[str(key)] = deepcopy(first_material)
            self.preview.set_wall_finish_mode(False)
            self.selected_wall_keys.clear()
            self.preview.set_selected_wall_keys(set())
            dialog.accept()
            self.save_finish_assignments()
            self.preview.set_room_finish_assignments(self.room_finish_assignments)
            self.preview.set_finish_groups(self.finish_groups)
            self._refresh_current_finish_panel()
            self.data_changed.emit()

        def cancel():
            self.preview.set_wall_finish_mode(False)
            self.selected_wall_keys.clear()
            self.preview.set_selected_wall_keys(set())
            dialog.reject()

        buttons.accepted.connect(finish)
        buttons.rejected.connect(cancel)
        def cleanup(_result):
            self._hide_finish_group_edit_overlay()
            self.preview.set_wall_finish_mode(False)
            self.selected_wall_keys.clear()
            self._group_mode_blocked_wall_keys.clear()
            self.preview.set_selected_wall_keys(set())
            self._wall_group_dialog = None
        dialog.finished.connect(cleanup)
        self._wall_group_dialog = dialog
        self._show_finish_group_edit_overlay(
            "벽", mode_group_label, "wall", len(self.selected_wall_keys),
            finish, cancel, detail_dialog=None,
        )
        refresh_count()

    def on_wall_segment_clicked(self, key: str, additive: bool):
        key = str(key or "")
        if key in self._group_mode_blocked_wall_keys:
            return
        if additive:
            if key in self.selected_wall_keys:
                self.selected_wall_keys.remove(key)
            else:
                self.selected_wall_keys.add(key)
        else:
            self.selected_wall_keys = {key}
        self.preview.set_selected_wall_keys(self.selected_wall_keys)
        self.finish_panel.set_selected_wall_count(len(self.selected_wall_keys))
        self._update_finish_group_edit_status(len(self.selected_wall_keys))

    def on_wall_segments_box_selected(self, keys, additive: bool, remove: bool):
        keys = {str(value) for value in (keys or []) if value}
        if not remove:
            keys.difference_update(self._group_mode_blocked_wall_keys)
        if remove:
            self.selected_wall_keys.difference_update(keys)
        elif additive:
            self.selected_wall_keys.update(keys)
        else:
            self.selected_wall_keys = keys
        self.preview.set_selected_wall_keys(self.selected_wall_keys)
        self.finish_panel.set_selected_wall_count(len(self.selected_wall_keys))
        self._update_finish_group_edit_status(len(self.selected_wall_keys))

    def assign_wall_finish(self, material: dict):
        category = material.get("target_category", material.get("category", "벽"))
        if category == "벽":
            if not self.selected_wall_keys:
                QMessageBox.information(self, "벽 마감 지정", "벽 구간 선택 모드에서 벽 구간을 먼저 선택하십시오.")
                return
            for key in self.selected_wall_keys:
                self.wall_finish_assignments[key] = dict(material)
        elif category and category != "복합재료":
            rooms = self.selected_target_rooms()
            if not rooms:
                QMessageBox.information(self, "재료마감 지정", "적용할 실 또는 층을 먼저 선택하십시오.")
                return
            for room in rooms:
                room_data = normalize_room_finish_assignment(
                    self.room_finish_assignments.get(room.boundary_handle, {})
                )
                section = normalize_room_finish_section(category)
                if section not in room_data["sections"]:
                    room_data["sections"].append(section)
                    room_data["items"][section] = []
                room_data["items"].setdefault(section, []).append(dict(material))
                self.room_finish_assignments[room.boundary_handle] = room_data
        else:
            return
        self.save_finish_assignments()
        self.preview.set_wall_finish_assignments(self.wall_finish_assignments)
        self.preview.set_room_finish_assignments(self.room_finish_assignments)
        self.update_finish_summary(self.selected_floor())

    def clear_finish_assignment(self, category: str):
        if category == "벽":
            if not self.selected_wall_keys:
                QMessageBox.information(self, "벽 마감 해제", "해제할 벽 구간을 먼저 선택하십시오.")
                return
            for key in list(self.selected_wall_keys):
                self.wall_finish_assignments.pop(key, None)
        elif category and category != "복합재료":
            rooms = self.selected_target_rooms()
            if not rooms:
                QMessageBox.information(self, "재료마감 해제", "해제할 실 또는 층을 먼저 선택하십시오.")
                return
            for room in rooms:
                room_data = normalize_room_finish_assignment(
                    self.room_finish_assignments.get(room.boundary_handle, {})
                )
                section = normalize_room_finish_section(category)
                room_data["items"].pop(section, None)
                room_data["sections"] = [name for name in room_data["sections"] if name != section]
                if not any(room_data["items"].values()):
                    self.room_finish_assignments.pop(room.boundary_handle, None)
                else:
                    self.room_finish_assignments[room.boundary_handle] = room_data
        self.save_finish_assignments()
        self.preview.set_wall_finish_assignments(self.wall_finish_assignments)
        self.preview.set_room_finish_assignments(self.room_finish_assignments)
        self.update_finish_summary(self.selected_floor())

    def handle_material_deleted(self, deleted_id: str, replacement):
        changed = False
        for group in self.finish_groups.values():
            updated_group_materials = []
            for material in list(group.get("materials", []) or []):
                if material.get("id") != deleted_id:
                    updated_group_materials.append(material)
                    continue
                if replacement is not None:
                    new_material = dict(replacement)
                    if material.get("group_id"):
                        new_material["group_id"] = material.get("group_id")
                        new_material["group_name"] = material.get("group_name", "그룹")
                    updated_group_materials.append(new_material)
                changed = True
            group["materials"] = updated_group_materials
        for key, material in list(self.wall_finish_assignments.items()):
            if material.get("id") == deleted_id:
                if replacement is None:
                    self.wall_finish_assignments.pop(key, None)
                else:
                    self.wall_finish_assignments[key] = dict(replacement)
                changed = True
        for handle, assignment in list(self.room_finish_assignments.items()):
            normalized = normalize_room_finish_assignment(assignment)
            for category, materials in list(normalized["items"].items()):
                updated = []
                for material in materials:
                    if material.get("id") != deleted_id:
                        updated.append(material)
                        continue
                    if replacement is not None:
                        new_material = dict(replacement)
                        if material.get("group_id"):
                            new_material["group_id"] = material.get("group_id")
                            new_material["group_name"] = material.get("group_name", "그룹")
                        updated.append(new_material)
                    changed = True
                normalized["items"][category] = updated
            if not any(normalized["items"].values()):
                self.room_finish_assignments.pop(handle, None)
            else:
                self.room_finish_assignments[handle] = normalized
        if changed:
            self.save_finish_assignments()
            self.preview.set_wall_finish_assignments(self.wall_finish_assignments)
            self.preview.set_room_finish_assignments(self.room_finish_assignments)
            current_room = self.selected_room()
            if current_room is not None:
                self.finish_panel.set_current_room(
                    current_room,
                    self.room_finish_assignments.get(current_room.boundary_handle, {}),
                )
            self.update_finish_summary(self.selected_floor())

    @staticmethod
    def _summary_material_key(material: dict) -> str:
        return str(
            material.get("id") or material.get("code") or material.get("name") or "미지정"
        )

    def set_summary_markup(self, material_key: str, percent: float):
        key = str(material_key or "").strip()
        if not key:
            return
        try:
            value = max(-100.0, float(percent or 0.0))
        except (TypeError, ValueError):
            value = 0.0
        if abs(value) < 1e-12:
            self.summary_markups.pop(key, None)
        else:
            self.summary_markups[key] = value
        self.save_finish_assignments()
        self.data_changed.emit()

    def finish_aggregate_rows(self, scope_rooms=None) -> list[dict]:
        """재료마감 지정값을 재료 단위로 합산한다.

        calculate_room_finish_value의 기본값(apply_repeat_count=True)을 사용하므로
        층범위/반복개소까지 반영된 최종 수량이 집계된다.
        """
        rooms = list(self.rooms if scope_rooms is None else scope_rooms)
        if not rooms:
            return []

        quantities = defaultdict(float)
        labels: dict[tuple[str, str], dict] = {}
        grouped_wall_keys = {
            str(key)
            for group in self.finish_groups.values()
            for key in list(group.get("wall_keys", []) or [])
        }

        def add_quantity(material: dict, quantity: float, default_unit: str):
            material_key = self._summary_material_key(material)
            unit = normalize_material_unit(material.get("unit", "") or default_unit)
            key = (material_key, unit)
            quantities[key] += float(quantity or 0.0)
            labels[key] = {
                "key": material_key,
                "id": str(material.get("id", "") or ""),
                "code": str(material.get("code", "") or ""),
                "name": str(material.get("name", "") or "미지정"),
                "spec": str(material.get("spec", "") or ""),
                "unit": unit,
                "markup": str(material.get("markup", "") or ""),
                "trade_code": str(material.get("trade_code", "") or ""),
                "trade": str(material.get("trade", material.get("trade_name", "")) or ""),
            }

        for room in rooms:
            pts = room.points
            # 공유 벽그룹은 실별 그룹 재료 산식에서 계산하므로 기존 단일 벽지정만 별도 합산한다.
            for idx in range(len(pts)):
                wall_key = f"{room.boundary_handle}:{idx}"
                if wall_key in grouped_wall_keys:
                    continue
                material = self.wall_finish_assignments.get(wall_key)
                if not material:
                    continue
                x1, y1 = pts[idx]; x2, y2 = pts[(idx + 1) % len(pts)]
                segment_length = ((x2 - x1) ** 2 + (y2 - y1) ** 2) ** 0.5 / 1000.0
                quantity = calculate_room_finish_value(room, material, wall_length=segment_length)
                add_quantity(material, 0.0 if quantity is None else quantity, "M2")

            assigned = normalize_room_finish_assignment(
                self.room_finish_assignments.get(room.boundary_handle, {})
            )
            room_group_formulas = {
                str(entry.get("id", "") or ""): str(entry.get("formula", "") or "")
                for entry in assigned.get("groups", [])
                if str(entry.get("id", "") or "")
            }
            group_wall_lengths = self._wall_group_lengths_for_room(room)
            for section, materials in assigned["items"].items():
                for material in materials:
                    group_id = str(material.get("group_id", "") or "")
                    wall_length = group_wall_lengths.get(group_id) if section == "벽" and group_id else None
                    group_formula = ""
                    if group_id:
                        if group_id in room_group_formulas:
                            group_formula = room_group_formulas[group_id]
                        elif group_id in self.finish_groups:
                            group_formula = self._group_default_formula_for_floor(
                                self.finish_groups[group_id], room.floor
                            )
                    effective_group_formula = "" if normalize_formula_text(material.get("method", "")) else group_formula
                    quantity = calculate_room_finish_value(
                        room, material, wall_length=wall_length, formula_override=effective_group_formula
                    )
                    default_unit = "M2" if section in {"바닥", "벽", "천정"} else ""
                    add_quantity(material, 0.0 if quantity is None else quantity, default_unit)

        rows = []
        for key, quantity in quantities.items():
            record = dict(labels[key])
            record["quantity"] = float(quantity)
            rows.append(record)
        return sorted(rows, key=lambda row: (
            str(row.get("trade_code", "")), str(row.get("trade", "")),
            str(row.get("code", "")), str(row.get("name", ""))
        ))

    def update_finish_summary(self, floor: str | None):
        table = getattr(self, "finish_summary_table", None)
        if table is None:
            return
        table.setRowCount(0)
        selected_room = self.selected_room()
        if selected_room is not None:
            scope_rooms = [selected_room]
        else:
            scope_rooms = [room for room in self.rooms if room.floor == floor] if floor else list(self.rooms)
        if not scope_rooms:
            return
        if hasattr(self, "finish_summary_group"):
            if selected_room is not None:
                room_number = str(selected_room.room_number or "-")
                room_name = str(normalize_room_name(selected_room.room_name) or "실명없음")
                self.finish_summary_group.setTitle(
                    f"{selected_room.floor} / {room_number} / {room_name} 재료마감 집계"
                )
            else:
                self.finish_summary_group.setTitle(
                    f"{floor} 재료마감 집계" if floor else "전층 재료마감 집계"
                )

        rows = self.finish_aggregate_rows(scope_rooms)
        table.setRowCount(len(rows))
        for row_index, record in enumerate(rows):
            values = [
                str(record.get("name", "") or ""),
                str(record.get("spec", "") or ""),
                str(record.get("unit", "") or ""),
                f"{float(record.get('quantity', 0.0) or 0.0):,.3f}",
            ]
            for col, text in enumerate(values):
                item = QTableWidgetItem(text)
                if col == 3:
                    item.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
                table.setItem(row_index, col, item)

    def request_save_dxf(self):
        if not self.current_source_path:
            return
        base_path = self.original_source_path or self.current_source_path
        suggested = str(
            Path(base_path).with_name(
                Path(base_path).stem + "_실정보수정.dxf"
            )
        )
        out, _ = QFileDialog.getSaveFileName(
            self, "수정내용을 DXF에 저장", suggested, "DXF 파일 (*.dxf)"
        )
        if not out:
            return
        try:
            source_path = self._operation_source_path()
            report = save_rooms_to_dxf(source_path, out, self.rooms, self.room_rules)
            for room in self.rooms:
                room.changed = False
            self.rebuild_views()
            QMessageBox.information(
                self, "DXF 저장 완료",
                f"DXF 저장 완료\n\n{out}\n\n"
                f"실 블록 처리: {report['updated']:,}/{report['rooms']:,}개\n"
                f"구획벽 형상 반영: {report.get('boundary_updated', 0):,}개\n"
                f"구획벽 반영 누락/실패: {report.get('boundary_missing', 0) + report.get('boundary_failed', 0):,}개\n"
                f"블록 핸들 누락: {report['missing_insert']:,}개\n"
                f"No 속성 누락: {report['missing_number_attr']:,}개\n"
                f"FLOOR/ROOM_NAME 속성 누락: {report['missing_basic_attr']:,}개"
            )
        except Exception:
            QMessageBox.critical(self, "DXF 저장 오류", traceback.format_exc())
