from __future__ import annotations

import json
import re
import uuid
from pathlib import Path

from PySide6.QtCore import Qt, Signal, QEvent, QItemSelectionModel, QTimer
from PySide6.QtGui import QColor, QKeySequence, QShortcut
from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QTabWidget, QTableWidget,
    QTableWidgetItem, QHeaderView, QAbstractItemView, QPushButton, QColorDialog,
    QMessageBox, QInputDialog, QLineEdit, QFileDialog, QComboBox, QDialog,
    QDialogButtonBox, QTreeWidget, QTreeWidgetItem, QStyledItemDelegate
)

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.common.finish_calculations import (
    canonical_finish_category,
    formula_for_category,
    normalize_finish_categories,
    normalize_formula_codes,
    normalize_formula_text,
)
from app.common.material_trades import (
    material_trade_display,
    normalize_material_trades,
)
from app.common.material_units import normalize_material_unit, normalize_material_units

DEFAULT_COLORS = {
    "바닥": "#5B8FF9",
    "벽": "#F6BD16",
    "천정": "#5AD8A6",
    "천장": "#5AD8A6",
    "기타": "#808080",
    "복합재료": "#9270CA",
}
LEGACY_BASE_CATEGORIES = ("바닥", "벽", "천장", "천정", "기타")



def _new_id() -> str:
    return uuid.uuid4().hex[:12].upper()


MATERIAL_CODE_LENGTH = 12
COMPOSITE_CODE_LENGTH = 8
ROLE_INTERNAL_ID = Qt.UserRole
ROLE_COMPONENT_IDS = Qt.UserRole + 1


def normalize_user_code(value: str, minimum_length: int, first_group_count: int) -> str:
    """사용자 코드를 4자리 구간으로 표시하고 마지막 구간은 필요하면 계속 늘린다."""
    compact = re.sub(r"[^0-9A-Za-z]", "", str(value or "")).upper()
    if not compact:
        return ""
    compact = compact.zfill(minimum_length)
    fixed_length = first_group_count * 4
    groups = [compact[index:index + 4] for index in range(0, fixed_length, 4)]
    groups.append(compact[fixed_length:])
    return "-".join(group for group in groups if group)


def increment_user_code(value: str, minimum_length: int, first_group_count: int) -> str:
    compact = re.sub(r"[^0-9A-Za-z]", "", str(value or "")).upper()
    if not compact:
        compact = "0" * (minimum_length - 1) + "1"
    elif compact.isdigit():
        compact = str(int(compact) + 1).zfill(max(minimum_length, len(compact)))
    else:
        match = re.search(r"(\d+)$", compact)
        if match:
            digits = match.group(1)
            compact = compact[:match.start()] + str(int(digits) + 1).zfill(len(digits))
        else:
            compact += "1"
    return normalize_user_code(compact, minimum_length, first_group_count)


class RequiredCodeDelegate(QStyledItemDelegate):
    def __init__(self, panel: "FinishMaterialsPanel", key: str, parent=None):
        super().__init__(parent)
        self.panel = panel
        self.key = key

    def setModelData(self, editor, model, index):
        raw = editor.text() if hasattr(editor, "text") else ""
        code = self.panel.normalize_code(self.key, raw)
        old = str(index.data(Qt.EditRole) or "")
        if not code:
            QMessageBox.warning(self.panel, "코드 입력", "코드는 공란으로 둘 수 없습니다.")
            model.setData(index, old, Qt.EditRole)
            return
        if not self.panel.is_code_available(self.key, code, index):
            QMessageBox.warning(self.panel, "코드 중복", f"이미 사용 중인 코드입니다.\n{code}")
            model.setData(index, old, Qt.EditRole)
            return
        model.setData(index, code, Qt.EditRole)


class MaterialUnitDelegate(QStyledItemDelegate):
    """단위 마스터를 사용하는 편집 콤보. 미등록 입력은 리스트 추가 여부를 확인한다."""

    def __init__(self, panel: "FinishMaterialsPanel", parent=None):
        super().__init__(parent)
        self.panel = panel

    def createEditor(self, parent, option, index):
        combo = QComboBox(parent)
        combo.setEditable(True)
        combo.setInsertPolicy(QComboBox.NoInsert)
        combo.addItems(self.panel.material_units)
        combo.setMaxVisibleItems(18)
        return combo

    def setEditorData(self, editor, index):
        value = normalize_material_unit(index.data(Qt.EditRole) or "")
        if isinstance(editor, QComboBox):
            idx = editor.findText(value, Qt.MatchFixedString)
            if idx >= 0:
                editor.setCurrentIndex(idx)
            else:
                editor.setEditText(value)
            if editor.lineEdit() is not None:
                editor.lineEdit().selectAll()

    def setModelData(self, editor, model, index):
        old = normalize_material_unit(index.data(Qt.EditRole) or "")
        raw = editor.currentText() if isinstance(editor, QComboBox) else ""
        unit = normalize_material_unit(raw)
        if not unit:
            model.setData(index, "", Qt.EditRole)
            return
        known = {value.casefold() for value in self.panel.material_units}
        if unit.casefold() not in known:
            answer = QMessageBox.question(
                self.panel, "단위 추가",
                f"단위 리스트에 없는 '{unit}'입니다.\n단위 리스트에 추가할까요?",
                QMessageBox.Yes | QMessageBox.No, QMessageBox.Yes,
            )
            if answer != QMessageBox.Yes:
                model.setData(index, old, Qt.EditRole)
                return
            self.panel.add_material_unit(unit)
        model.setData(index, unit, Qt.EditRole)


class MaterialCategoryDelegate(QStyledItemDelegate):
    """재료마감 구분은 평소 일반 셀로 표시하고 편집할 때만 드롭다운을 연다."""

    def __init__(self, panel: "FinishMaterialsPanel", table: QTableWidget, parent=None):
        super().__init__(parent or table)
        self.panel = panel
        self.table = table

    def createEditor(self, parent, option, index):
        combo = QComboBox(parent)
        combo.addItem("")
        combo.addItems(self.panel.category_names())
        combo.setMaxVisibleItems(18)
        font = combo.font(); font.setBold(False); combo.setFont(font)
        return combo

    def setEditorData(self, editor, index):
        value = str(index.data(Qt.EditRole) or "").strip()
        if isinstance(editor, QComboBox):
            pos = editor.findText(value, Qt.MatchFixedString)
            if pos < 0 and value:
                editor.addItem(value); pos = editor.count() - 1
            editor.setCurrentIndex(max(0, pos))
            editor.showPopup()

    def setModelData(self, editor, model, index):
        category = editor.currentText().strip() if isinstance(editor, QComboBox) else ""
        self.panel._set_material_category(self.table, index.row(), category)


class MaterialTradeDelegate(QStyledItemDelegate):
    """공종도 상시 셀 위젯을 쓰지 않고 실제 편집 순간에만 드롭다운을 만든다."""

    def __init__(self, panel: "FinishMaterialsPanel", table: QTableWidget, parent=None):
        super().__init__(parent or table)
        self.panel = panel
        self.table = table

    def createEditor(self, parent, option, index):
        combo = QComboBox(parent)
        combo.addItem("", {"code": "", "name": "", "parent_code": ""})
        for trade in self.panel.material_trades:
            combo.addItem(material_trade_display(trade["code"], trade["name"]), dict(trade))
        combo.setMaxVisibleItems(20)
        font = combo.font(); font.setBold(False); combo.setFont(font)
        return combo

    def setEditorData(self, editor, index):
        if not isinstance(editor, QComboBox):
            return
        raw = index.data(Qt.UserRole)
        trade = self.panel._trade_record(
            raw.get("code", "") if isinstance(raw, dict) else "",
            raw.get("name", "") if isinstance(raw, dict) else "",
        )
        selected = 0
        for pos in range(editor.count()):
            data = editor.itemData(pos)
            if isinstance(data, dict) and str(data.get("code", "")).casefold() == trade["code"].casefold() and trade["code"]:
                selected = pos; break
        editor.setCurrentIndex(selected)
        editor.showPopup()

    def setModelData(self, editor, model, index):
        raw = editor.currentData() if isinstance(editor, QComboBox) else {}
        raw = raw if isinstance(raw, dict) else {}
        self.panel._set_material_trade(self.table, index.row(), raw.get("code", ""), raw.get("name", ""))


class CompositeMaterialsTree(QTreeWidget):
    f5_pressed = Signal()
    duplicate_pressed = Signal()
    append_pressed = Signal()
    delete_pressed = Signal()

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._down_append_guard = False
        self.itemChanged.connect(lambda *_args: setattr(self, "_down_append_guard", False))

    def _last_top_level_blank(self) -> bool:
        if self.topLevelItemCount() <= 0:
            return True
        item = self.topLevelItem(self.topLevelItemCount() - 1)
        # 복합재료 신규행은 코드/구분/단위/기본산식이 자동 생성되므로
        # 사용자가 직접 작성하는 '항목명'이 비어 있으면 빈 신규행으로 본다.
        return not str(item.text(3) or "").strip()

    def keyPressEvent(self, event):
        if event.key() == Qt.Key_F5:
            self.f5_pressed.emit(); return
        if event.key() == Qt.Key_F3:
            self.duplicate_pressed.emit(); return
        if event.key() == Qt.Key_F4:
            self._down_append_guard = False
            self._delete_requested_by_f4 = True
            self.delete_pressed.emit(); return
        if event.key() == Qt.Key_Down:
            item = self.currentItem()
            if item is not None:
                top = item if item.parent() is None else item.parent()
                if self.indexOfTopLevelItem(top) == self.topLevelItemCount() - 1:
                    if self._last_top_level_blank():
                        return
                    self.append_pressed.emit()
                    self._down_append_guard = self._last_top_level_blank()
                    return
        super().keyPressEvent(event)


class ComponentMaterialDialog(QDialog):
    """복합재료에 넣을 일반 재료를 여러 개 고르는 작은 선택창."""

    def __init__(self, materials: list[dict], existing_ids=(), parent=None):
        super().__init__(parent)
        self.setWindowTitle("복합재료 구성항목 선택")
        self.resize(820, 560)
        self.selected_materials: list[dict] = []
        self.materials = list(materials or [])
        existing = {str(value) for value in existing_ids or []}
        layout = QVBoxLayout(self)
        guide = QLabel("재료마감 항목을 여러 개 선택한 뒤 등록하십시오. 이미 포함된 항목은 목록에서 제외됩니다.")
        guide.setWordWrap(True)
        layout.addWidget(guide)
        self.search_edit = QLineEdit()
        self.search_edit.setPlaceholderText("재료코드·구분·품명·규격 검색")
        layout.addWidget(self.search_edit)
        self.table = QTableWidget(0, 7)
        self.table.setHorizontalHeaderLabels(["재료코드", "구분", "공종", "품명", "규격", "단위", "산출식"])
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.ExtendedSelection)
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.verticalHeader().setVisible(False)
        self.table.horizontalHeader().setStretchLastSection(True)
        for material in self.materials:
            if str(material.get("id", "")) in existing:
                continue
            row = self.table.rowCount(); self.table.insertRow(row)
            values = [material.get("code", ""), material.get("category", ""), material.get("trade", ""),
                      material.get("name", ""), material.get("spec", ""), normalize_material_unit(material.get("unit", "")), normalize_formula_text(material.get("method", ""))]
            for column, value in enumerate(values):
                item = QTableWidgetItem(str(value or ""))
                if column == 0:
                    item.setData(ROLE_INTERNAL_ID, material.get("id", ""))
                self.table.setItem(row, column, item)
        self.table.setColumnWidth(0, 140); self.table.setColumnWidth(1, 75); self.table.setColumnWidth(2, 120)
        self.table.setColumnWidth(3, 180); self.table.setColumnWidth(4, 130); self.table.setColumnWidth(5, 55); self.table.setColumnWidth(6, 280)
        layout.addWidget(self.table, 1)
        self.search_edit.textChanged.connect(self._filter)
        buttons = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        buttons.button(QDialogButtonBox.Ok).setText("등록")
        buttons.button(QDialogButtonBox.Cancel).setText("취소")
        buttons.accepted.connect(self._accept)
        buttons.rejected.connect(self.reject)
        layout.addWidget(buttons)

    def _filter(self, text: str):
        keyword = str(text or "").strip().casefold()
        for row in range(self.table.rowCount()):
            line = " ".join(self.table.item(row, col).text() for col in range(self.table.columnCount()) if self.table.item(row, col))
            self.table.setRowHidden(row, bool(keyword) and keyword not in line.casefold())

    def _accept(self):
        rows = sorted({index.row() for index in self.table.selectionModel().selectedRows()})
        if not rows:
            QMessageBox.information(self, "구성항목 등록", "등록할 재료를 하나 이상 선택해 주세요.")
            return
        selected_ids = {str(self.table.item(row, 0).data(ROLE_INTERNAL_ID) or "") for row in rows}
        self.selected_materials = [material for material in self.materials if str(material.get("id", "")) in selected_ids]
        self.accept()


class MaterialTableWidget(QTableWidget):
    """재료마감 관리표에서 hover/창 활성화만으로 행 선택이 바뀌지 않게 한다."""
    duplicate_pressed = Signal()
    append_pressed = Signal()
    delete_pressed = Signal()

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._down_append_guard = False
        self.itemChanged.connect(lambda *_args: setattr(self, "_down_append_guard", False))

    def _last_row_blank(self) -> bool:
        if self.rowCount() <= 0:
            return True
        row = self.rowCount() - 1
        # 재료코드/단위 등 자동 생성값은 제외하고 실제 작성 칸만 검사한다.
        for col in (2, 3, 4, 5, 8):
            if col >= self.columnCount():
                continue
            item = self.item(row, col)
            if item is not None and str(item.text() or "").strip():
                return False
        return True

    def keyPressEvent(self, event):
        row, col = self.currentRow(), self.currentColumn()
        if event.key() == Qt.Key_F3 and row >= 0:
            self.duplicate_pressed.emit(); return
        if event.key() == Qt.Key_F4 and row >= 0:
            self._down_append_guard = False
            self._delete_requested_by_f4 = True
            self.delete_pressed.emit(); return
        if event.key() == Qt.Key_Down and row >= 0 and row == self.rowCount() - 1:
            if self._last_row_blank():
                return
            self.append_pressed.emit()
            self._down_append_guard = self._last_row_blank()
            return
        if event.key() in (Qt.Key_Return, Qt.Key_Enter):
            super().keyPressEvent(event)
            if row >= 0 and col >= 0:
                QTimer.singleShot(0, lambda r=row,c=col: self.setCurrentCell(min(r,self.rowCount()-1), min(c,self.columnCount()-1)) if self.rowCount() and self.columnCount() else None)
            return
        super().keyPressEvent(event)

    def selectionCommand(self, index, event=None):
        if event is not None:
            try:
                if event.type() in (
                    QEvent.MouseMove,
                    QEvent.HoverMove,
                    QEvent.HoverEnter,
                    QEvent.Enter,
                    QEvent.FocusIn,
                    QEvent.WindowActivate,
                    QEvent.ActivationChange,
                ):
                    return QItemSelectionModel.NoUpdate
            except Exception:
                pass
        return super().selectionCommand(index, event)


class FinishCategoryComboBox(QComboBox):
    """재료마감 표 안에서 실제 클릭 때만 행을 선택하는 구분/공종 콤보."""

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setFrame(False)
        # 상시 셀 위젯이 창 활성화/hover만으로 부모 표의 current row를 바꾸는 경로를 차단한다.
        self.setFocusPolicy(Qt.NoFocus)
        self.setAttribute(Qt.WA_NoMousePropagation, True)
        self._row_selected = False
        self._apply_row_style(False)
        popup = self.view()
        popup.setMouseTracking(True)
        popup.viewport().setMouseTracking(True)
        popup.setStyleSheet(
            "QAbstractItemView { border:1px solid #A0A0A0; background:#FFFFFF; color:#111111; "
            "selection-background-color:#17365D; selection-color:#FFFFFF; outline:0; }"
            "QAbstractItemView::item { background:#FFFFFF; color:#111111; padding:2px 4px; min-height:20px; }"
            "QAbstractItemView::item:hover { background:#FFF2CC; color:#111111; }"
            "QAbstractItemView::item:selected { background:#17365D; color:#FFFFFF; }"
        )
        font = self.font()
        font.setBold(False)
        self.setFont(font)
        popup.setFont(font)

    def _apply_row_style(self, selected: bool):
        self._row_selected = bool(selected)
        background = "#17365D" if self._row_selected else "transparent"
        color = "#FFFFFF" if self._row_selected else "#111111"
        self.setStyleSheet(
            f"QComboBox {{ border:none; outline:0; background:{background}; color:{color}; padding:0 16px 0 3px; font-weight:normal; }}"
            f"QComboBox:focus, QComboBox:on {{ border:none; outline:0; background:{background}; color:{color}; }}"
            "QComboBox::drop-down { border:none; background:transparent; width:16px; }"
        )

    def set_row_selected(self, selected: bool):
        selected = bool(selected)
        if selected != self._row_selected:
            self._apply_row_style(selected)

    def mousePressEvent(self, event):
        # 마우스 이동/창 활성화는 선택을 바꾸지 않는다. 실제 클릭 시에만 해당 행을 선택한다.
        table = self.parentWidget()
        while table is not None and not isinstance(table, QTableWidget):
            table = table.parentWidget()
        if table is not None:
            try:
                point = table.viewport().mapFromGlobal(event.globalPosition().toPoint())
                index = table.indexAt(point)
                if index.isValid():
                    table.setCurrentCell(index.row(), index.column(), QItemSelectionModel.NoUpdate)
                    table.clearSelection()
                    table.selectRow(index.row())
            except Exception:
                pass
        super().mousePressEvent(event)


class MaterialTradeComboBox(FinishCategoryComboBox):
    """재료마감 표 안에서 테두리와 굵은 글씨 없이 표시하는 공종 선택 콤보."""


class FinishMaterialsPanel(QWidget):
    material_assign_requested = Signal(dict)
    material_clear_requested = Signal(str)
    wall_mode_changed = Signal(bool)
    materials_changed = Signal()
    material_delete_requested = Signal(str, object)
    material_units_changed = Signal(object)

    BASE_HEADERS = ["★", "재료코드", "구분", "공종", "품명", "규격", "단위", "할증", "산출식", "색상"]
    COMPOSITE_HEADERS = ["★", "복합재료 코드", "구분", "항목명", "단위", "산출식", "색상"]

    def __init__(self, data_path: Path, assignment_enabled: bool = True, formula_codes=None, finish_categories=None, material_trades=None, material_units=None, picker_mode: bool = False):
        super().__init__()
        self.data_path = Path(data_path)
        self.loading = False
        self.assignment_enabled = bool(assignment_enabled)
        # 재료 선택창에서는 구분/공종 셀마다 QComboBox를 만들지 않는다.
        # 수백 개 행에서 최초 F5 창 생성이 느려지는 주원인을 제거하고,
        # 셀 배경 선택 강조도 모든 열에 동일하게 보이게 한다.
        self.picker_mode = bool(picker_mode)
        self.formula_codes = normalize_formula_codes(formula_codes)
        self.finish_categories = normalize_finish_categories(finish_categories)
        self.material_trades = normalize_material_trades(material_trades)
        self.material_units = normalize_material_units(material_units)
        self.tables: dict[str, QWidget] = {}

        root = QVBoxLayout(self)
        title = QLabel("재료마감")
        font = title.font(); font.setBold(True); title.setFont(font)
        root.addWidget(title)

        self.search_edit = QLineEdit()
        self.search_edit.setPlaceholderText("재료코드·품명·규격 검색")
        self.sort_btn = QPushButton("가나다순 정렬")
        search_row = QHBoxLayout()
        search_row.addWidget(self.search_edit, 1)
        search_row.addWidget(self.sort_btn)
        self.search_widget = QWidget()
        self.search_widget.setLayout(search_row)
        self.search_widget.setVisible(not self.assignment_enabled)
        root.addWidget(self.search_widget)

        self.wall_mode_btn = QPushButton("벽 구간 선택 모드")
        self.wall_mode_btn.setCheckable(True)
        self.wall_mode_btn.setVisible(self.assignment_enabled)
        root.addWidget(self.wall_mode_btn)

        self.tabs = QTabWidget()
        if self.assignment_enabled:
            self._build_assignment_tabs()
        else:
            self._build_admin_tabs()
        root.addWidget(self.tabs, 1)

        self.selection_label = QLabel("선택 대상: 0개")
        self.selection_label.setVisible(self.assignment_enabled)
        root.addWidget(self.selection_label)
        assign_row = QHBoxLayout()
        self.assign_btn = QPushButton("선택 대상에 지정")
        self.clear_btn = QPushButton("선택 대상 마감 해제")
        assign_row.addWidget(self.assign_btn)
        assign_row.addWidget(self.clear_btn)
        assign_row_widget = QWidget()
        assign_row_widget.setLayout(assign_row)
        assign_row_widget.setVisible(self.assignment_enabled)
        root.addWidget(assign_row_widget)

        self.wall_mode_btn.toggled.connect(self.wall_mode_changed.emit)
        self.assign_btn.clicked.connect(self.emit_assign)
        self.clear_btn.clicked.connect(lambda: self.material_clear_requested.emit(self.current_target_category()))
        self.tabs.currentChanged.connect(self._on_category_changed)
        self.search_edit.textChanged.connect(self._apply_search)
        self.sort_btn.clicked.connect(self._sort_current_table)
        if not self.assignment_enabled and not self.picker_mode:
            self.add_material_shortcut = QShortcut(QKeySequence(Qt.Key_F9), self)
            self.add_material_shortcut.setContext(Qt.WidgetWithChildrenShortcut)
            self.add_material_shortcut.activated.connect(self._add_current_row_f9)
        self.load()

    def _new_table(self, headers: list[str], widths: list[int]) -> QTableWidget:
        table = MaterialTableWidget(0, len(headers))
        table.setHorizontalHeaderLabels(headers)
        table.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
        table.horizontalHeader().setStretchLastSection(True)
        table.setSelectionBehavior(QAbstractItemView.SelectRows)
        table.setSelectionMode(QAbstractItemView.SingleSelection)
        table.setAlternatingRowColors(False)
        table.verticalHeader().setVisible(False)
        for col, width in enumerate(widths):
            table.setColumnWidth(col, width)
        table.setColumnHidden(len(headers) - 1, True)
        table.itemChanged.connect(self.save)
        table.cellClicked.connect(lambda row, col, t=table: self._toggle_star(t, row) if col == 0 else None)
        table.cellDoubleClicked.connect(lambda row, col, t=table: self._pick_color_from_table(t, row, col))
        table.itemSelectionChanged.connect(lambda t=table: self._sync_embedded_combo_selection(t))
        table.duplicate_pressed.connect(lambda k=headers: self._duplicate_current_row(self.current_category()))
        table.append_pressed.connect(lambda: self.add_row(self.current_category()))
        table.delete_pressed.connect(lambda: self.delete_selected(self.current_category()))
        if not self.assignment_enabled and "단위" in headers:
            unit_col = headers.index("단위")
            table.setItemDelegateForColumn(unit_col, MaterialUnitDelegate(self, table))
            if not self.picker_mode:
                table.cellClicked.connect(
                    lambda row, col, t=table, u=unit_col: t.editItem(t.item(row, col))
                    if col == u and t.item(row, col) is not None else None
                )
        # V23.15: 구분/공종은 상시 QComboBox를 셀 위에 얹지 않는다.
        # 일반 셀로 표시하고 실제 클릭 편집 때만 delegate 콤보를 생성해
        # hover/드래그/창 활성화로 값이 바뀌는 문제와 선택색 누락을 원천 차단한다.
        if not self.assignment_enabled and not self.picker_mode:
            if "구분" in headers:
                category_col = headers.index("구분")
                table.setItemDelegateForColumn(category_col, MaterialCategoryDelegate(self, table))
                table.cellClicked.connect(
                    lambda row, col, t=table, c=category_col: t.editItem(t.item(row, col))
                    if col == c and t.item(row, col) is not None else None
                )
            if "공종" in headers:
                trade_col = headers.index("공종")
                table.setItemDelegateForColumn(trade_col, MaterialTradeDelegate(self, table))
                table.cellClicked.connect(
                    lambda row, col, t=table, c=trade_col: t.editItem(t.item(row, col))
                    if col == c and t.item(row, col) is not None else None
                )
        return table

    def _sync_embedded_combo_selection(self, table: QTableWidget):
        """선택행 배경색을 구분/공종 셀 위젯에도 동일하게 표시한다."""
        if not isinstance(table, QTableWidget):
            return
        selection_model = table.selectionModel()
        for row in range(table.rowCount()):
            selected = bool(selection_model and selection_model.isRowSelected(row, table.rootIndex()))
            for column in (2, 3):
                combo = table.cellWidget(row, column)
                if isinstance(combo, FinishCategoryComboBox):
                    combo.set_row_selected(selected)

    def _new_composite_tree(self) -> CompositeMaterialsTree:
        tree = CompositeMaterialsTree()
        tree.setColumnCount(len(self.COMPOSITE_HEADERS))
        tree.setHeaderLabels(self.COMPOSITE_HEADERS)
        tree.header().setSectionResizeMode(QHeaderView.Interactive)
        tree.header().setStretchLastSection(True)
        tree.setSelectionMode(QAbstractItemView.SingleSelection)
        tree.setSelectionBehavior(QAbstractItemView.SelectRows)
        tree.setAlternatingRowColors(False)
        tree.setRootIsDecorated(True)
        tree.setIndentation(18)
        widths = [55, 145, 80, 230, 60, 300, 70]
        for column, width in enumerate(widths):
            tree.setColumnWidth(column, width)
        tree.duplicate_pressed.connect(lambda: self._duplicate_current_row("복합재료"))
        tree.append_pressed.connect(lambda: self.add_row("복합재료"))
        tree.delete_pressed.connect(lambda: self.delete_selected("복합재료"))
        tree.setColumnHidden(len(self.COMPOSITE_HEADERS) - 1, True)
        tree.setItemDelegateForColumn(1, RequiredCodeDelegate(self, "복합재료", tree))
        if not self.assignment_enabled:
            tree.setItemDelegateForColumn(4, MaterialUnitDelegate(self, tree))
            if not self.picker_mode:
                tree.itemClicked.connect(
                    lambda item, col, t=tree: t.editItem(item, col)
                    if col == 4 and item is not None and item.parent() is None else None
                )
        if self.assignment_enabled:
            tree.setEditTriggers(QAbstractItemView.NoEditTriggers)
        else:
            tree.setEditTriggers(QAbstractItemView.DoubleClicked | QAbstractItemView.EditKeyPressed | QAbstractItemView.AnyKeyPressed)
        tree.itemChanged.connect(self.save)
        tree.itemDoubleClicked.connect(self._on_composite_item_clicked)
        tree.f5_pressed.connect(self.open_composite_component_picker)
        return tree

    @staticmethod
    def _cell_text(table: QTableWidget, row: int, col: int) -> str:
        widget = table.cellWidget(row, col)
        if isinstance(widget, QComboBox):
            return widget.currentText().strip()
        item = table.item(row, col)
        return item.text().strip() if item else ""

    def normalize_code(self, key: str, value: str) -> str:
        if key == "복합재료":
            return normalize_user_code(value, COMPOSITE_CODE_LENGTH, 1)
        return normalize_user_code(value, MATERIAL_CODE_LENGTH, 2)

    def _codes(self, key: str) -> list[str]:
        widget = self.tables.get(key)
        if key == "복합재료" and isinstance(widget, QTreeWidget):
            return [widget.topLevelItem(row).text(1).strip() for row in range(widget.topLevelItemCount())]
        if isinstance(widget, QTableWidget):
            return [self._cell_text(widget, row, 1) for row in range(widget.rowCount())]
        return []

    def is_code_available(self, key: str, code: str, index=None) -> bool:
        marker = str(code or "").casefold()
        widget = self.tables.get(key)
        for position, current in enumerate(self._codes(key)):
            if current.casefold() != marker:
                continue
            if index is not None:
                if key == "복합재료" and isinstance(widget, QTreeWidget):
                    item = widget.topLevelItem(position)
                    if widget.indexFromItem(item, 1) == index:
                        continue
                elif isinstance(widget, QTableWidget) and position == index.row():
                    continue
            return False
        return True

    def next_code(self, key: str) -> str:
        codes = self._codes(key)
        if key == "복합재료":
            return increment_user_code(codes[-1] if codes else "", COMPOSITE_CODE_LENGTH, 1)
        return increment_user_code(codes[-1] if codes else "", MATERIAL_CODE_LENGTH, 2)

    def request_new_code(self, key: str) -> str:
        title = "복합재료 코드" if key == "복합재료" else "재료코드"
        value = self.next_code(key)
        while True:
            text, ok = QInputDialog.getText(self, f"{title} 입력", f"{title}:", QLineEdit.Normal, value)
            if not ok:
                return ""
            code = self.normalize_code(key, text)
            if not code:
                QMessageBox.warning(self, title, "코드는 공란으로 둘 수 없습니다.")
                continue
            if not self.is_code_available(key, code):
                QMessageBox.warning(self, title, f"이미 사용 중인 코드입니다.\n{code}")
                value = code
                continue
            return code

    def _install_category_combo(self, table: QTableWidget, row: int, category: str):
        """구분은 일반 QTableWidgetItem으로 유지한다. 편집 콤보는 delegate가 필요할 때만 만든다."""
        raw_category = str(category or "").strip()
        display = self._display_category(raw_category) if raw_category else ""
        item = table.item(row, 2)
        if item is None:
            item = QTableWidgetItem(display)
            table.setItem(row, 2, item)
        else:
            item.setText(display)
        font = item.font(); font.setBold(False); item.setFont(font)
        if self.assignment_enabled or self.picker_mode:
            item.setFlags(item.flags() & ~Qt.ItemIsEditable)
        else:
            item.setFlags(item.flags() | Qt.ItemIsEditable)

    def _trade_record(self, code: str = "", name: str = "") -> dict[str, str]:
        code = str(code or "").strip()
        name = str(name or "").strip()
        for trade in self.material_trades:
            if code and trade["code"].casefold() == code.casefold():
                return dict(trade)
        for trade in self.material_trades:
            if name and trade["name"].casefold() == name.casefold():
                return dict(trade)
        return {"code": code, "name": name, "parent_code": ""}

    def _install_trade_combo(self, table: QTableWidget, row: int, code: str = "", name: str = ""):
        """공종도 일반 셀로 표시하고 실제 편집 순간에만 delegate 콤보를 사용한다."""
        trade = self._trade_record(code, name)
        display = material_trade_display(trade["code"], trade["name"])
        item = table.item(row, 3)
        if item is None:
            item = QTableWidgetItem(display)
            table.setItem(row, 3, item)
        else:
            item.setText(display)
        item.setData(Qt.UserRole, trade)
        font = item.font(); font.setBold(False); item.setFont(font)
        if self.assignment_enabled or self.picker_mode:
            item.setFlags(item.flags() & ~Qt.ItemIsEditable)
        else:
            item.setFlags(item.flags() | Qt.ItemIsEditable)

    def _set_material_trade(self, table: QTableWidget, row: int, code: str = "", name: str = ""):
        if self.loading or self.assignment_enabled or not (0 <= row < table.rowCount()):
            return
        trade = self._trade_record(code, name)
        item = table.item(row, 3)
        self.loading = True
        try:
            if item is None:
                item = QTableWidgetItem()
                table.setItem(row, 3, item)
            item.setText(material_trade_display(trade["code"], trade["name"]))
            item.setData(Qt.UserRole, trade)
            font = item.font(); font.setBold(False); item.setFont(font)
            item.setFlags(item.flags() | Qt.ItemIsEditable)
        finally:
            self.loading = False
        self.save()

    def _on_material_trade_changed(self, table: QTableWidget, combo: QComboBox):
        # 구버전 상시 콤보 경로 호환용. V23.15 관리표에서는 delegate 경로를 사용한다.
        if self.loading or self.assignment_enabled:
            return
        row = table.indexAt(combo.pos()).row()
        if row < 0:
            return
        raw = combo.currentData() if isinstance(combo.currentData(), dict) else {}
        self._set_material_trade(table, row, raw.get("code", ""), raw.get("name", ""))

    def _trade_from_row(self, table: QTableWidget, row: int) -> dict[str, str]:
        widget = table.cellWidget(row, 3)
        if isinstance(widget, QComboBox) and isinstance(widget.currentData(), dict):
            raw = widget.currentData()
            return self._trade_record(raw.get("code", ""), raw.get("name", ""))
        item = table.item(row, 3)
        if item is not None and isinstance(item.data(Qt.UserRole), dict):
            raw = item.data(Qt.UserRole)
            return self._trade_record(raw.get("code", ""), raw.get("name", ""))
        text = item.text().strip() if item else ""
        code, _separator, name = text.partition(" ")
        return self._trade_record(code, name)

    def _set_material_category(self, table: QTableWidget, row: int, category: str):
        if self.loading or self.assignment_enabled or not (0 <= row < table.rowCount()):
            return
        category = str(category or "").strip()
        item = table.item(row, 2)
        self.loading = True
        try:
            if item is None:
                item = QTableWidgetItem(category)
                table.setItem(row, 2, item)
            else:
                item.setText(category)
            font = item.font(); font.setBold(False); item.setFont(font)
            item.setFlags(item.flags() | Qt.ItemIsEditable)
        finally:
            self.loading = False

        code, formula = formula_for_category(category, self.formula_codes, self.finish_categories)
        method_item = table.item(row, 8)
        current_formula = method_item.text().strip() if method_item else ""
        apply_formula = False
        if formula and not current_formula:
            apply_formula = True
        elif formula and current_formula != formula:
            formula_source = f"산식 코드 {code}" if code else "직접 산식"
            answer = QMessageBox.question(
                self,
                "산출식 적용",
                f"'{category}' 구분에는 {formula_source}이(가) 연결되어 있습니다.\n\n"
                f"{formula}\n\n현재 산출식을 위 산식으로 변경할까요?",
                QMessageBox.Yes | QMessageBox.No,
                QMessageBox.Yes,
            )
            apply_formula = answer == QMessageBox.Yes
        if apply_formula:
            self.loading = True
            try:
                if method_item is None:
                    method_item = QTableWidgetItem(formula)
                    table.setItem(row, 8, method_item)
                else:
                    method_item.setText(formula)
            finally:
                self.loading = False
        self.save()

    def _on_material_category_changed(self, table: QTableWidget, combo: QComboBox, category: str):
        # 구버전 상시 콤보 경로 호환용. V23.15 관리표에서는 delegate 경로를 사용한다.
        if self.loading or self.assignment_enabled:
            return
        row = table.indexAt(combo.pos()).row()
        if row >= 0:
            self._set_material_category(table, row, category)

    def _install_composite_category_combo(self, tree: QTreeWidget, item: QTreeWidgetItem, category: str):
        display = self._display_category(category)
        item.setText(2, display)
        if self.assignment_enabled or self.picker_mode:
            return
        combo = FinishCategoryComboBox()
        combo.addItems(self.category_names())
        index = combo.findText(display, Qt.MatchFixedString)
        if index < 0 and display:
            combo.addItem(display); index = combo.count() - 1
        combo.setCurrentIndex(max(0, index))
        combo.currentTextChanged.connect(
            lambda text, current_item=item, current_combo=combo:
            self._on_composite_category_changed(current_item, current_combo, text)
        )
        tree.setItemWidget(item, 2, combo)

    def _on_composite_category_changed(self, item: QTreeWidgetItem, combo: QComboBox, category: str):
        if self.loading or self.assignment_enabled or item.parent() is not None:
            return
        self.loading = True
        try:
            item.setText(2, category)
        finally:
            self.loading = False
        code, formula = formula_for_category(category, self.formula_codes, self.finish_categories)
        current = item.text(5).strip()
        apply_formula = bool(formula and not current)
        if formula and current and current != formula:
            source = f"산식 코드 {code}" if code else "직접 산식"
            answer = QMessageBox.question(
                self, "산출식 적용",
                f"'{category}' 구분에는 {source}이(가) 연결되어 있습니다.\n\n{formula}\n\n현재 산출식을 변경할까요?",
                QMessageBox.Yes | QMessageBox.No, QMessageBox.Yes,
            )
            apply_formula = answer == QMessageBox.Yes
        if apply_formula:
            self.loading = True
            try:
                item.setText(5, formula)
            finally:
                self.loading = False
        self.save()

    def _composite_category(self, item: QTreeWidgetItem) -> str:
        tree = self.tables.get("복합재료")
        if isinstance(tree, QTreeWidget):
            combo = tree.itemWidget(item, 2)
            if isinstance(combo, QComboBox):
                return combo.currentText().strip()
        return item.text(2).strip()

    def _composite_item_material(self, item: QTreeWidgetItem | None) -> dict | None:
        if item is None:
            return None
        if item.parent() is not None:
            item = item.parent()
        code = item.text(1).strip()
        name = item.text(3).strip()
        if not code or not name:
            return None
        category = self._configured_category(self._composite_category(item)) or self._display_category(self._composite_category(item))
        internal_id = str(item.data(1, ROLE_INTERNAL_ID) or "") or _new_id()
        component_ids = list(item.data(1, ROLE_COMPONENT_IDS) or [])
        return {
            "favorite": item.text(0).strip() == "★",
            "id": internal_id,
            "code": code,
            "category": category,
            "target_category": canonical_finish_category(category),
            "name": name,
            "component_ids": component_ids,
            "spec": "",
            "unit": normalize_material_unit(item.text(4)),
            "method": normalize_formula_text(item.text(5)),
            "color": item.text(6).strip() or DEFAULT_COLORS["복합재료"],
            "is_composite": True,
        }

    def _page_with_buttons(self, table: QTableWidget, key: str, excel_label: str | None = None) -> QWidget:
        page = QWidget()
        layout = QVBoxLayout(page)
        layout.addWidget(table, 1)
        buttons = QHBoxLayout()
        add_btn = QPushButton("추가")
        del_btn = QPushButton("삭제")
        color_btn = QPushButton("색상")
        add_btn.clicked.connect(lambda: self._add_row_at_current(key))
        del_btn.clicked.connect(lambda: self.delete_selected(key))
        color_btn.clicked.connect(lambda: self.pick_selected_color(key))
        buttons.addWidget(add_btn)
        buttons.addWidget(del_btn)
        buttons.addWidget(color_btn)
        if key == "복합재료":
            component_btn = QPushButton("구성재료 등록 (F5)")
            component_btn.clicked.connect(self.open_composite_component_picker)
            buttons.addWidget(component_btn)
        if excel_label:
            buttons.addStretch(1)
            export_btn = QPushButton(f"{excel_label} Excel 내보내기")
            import_btn = QPushButton(f"{excel_label} Excel 불러오기")
            export_btn.clicked.connect(lambda checked=False, k=key: self.export_excel(k))
            import_btn.clicked.connect(lambda checked=False, k=key: self.import_excel(k))
            buttons.addWidget(export_btn)
            buttons.addWidget(import_btn)
        layout.addLayout(buttons)
        return page

    def _build_admin_tabs(self):
        composite = self._new_composite_tree()
        materials = self._new_table(self.BASE_HEADERS, [36, 145, 70, 175, 160, 120, 55, 60, 300, 70])
        materials.setItemDelegateForColumn(1, RequiredCodeDelegate(self, "재료마감", materials))
        self.tables["복합재료"] = composite
        self.tables["재료마감"] = materials
        self.tabs.addTab(self._page_with_buttons(composite, "복합재료", "복합재료"), "복합재료")
        self.tabs.addTab(self._page_with_buttons(materials, "재료마감", "재료마감"), "재료마감")

    def _build_assignment_tabs(self):
        for category in [*self.category_names(), "복합재료"]:
            if category == "복합재료":
                table = self._new_composite_tree()
            else:
                headers = self.BASE_HEADERS
                widths = [36, 130, 65, 150, 130, 90, 50, 55, 90, 60]
                table = self._new_table(headers, widths)
                table.setEditTriggers(QAbstractItemView.NoEditTriggers)
            self.tables[category] = table
            page = QWidget(); layout = QVBoxLayout(page); layout.addWidget(table)
            self.tabs.addTab(page, category)

    def category_names(self) -> list[str]:
        return [row["category"] for row in self.finish_categories]

    def _category_color(self, category: str) -> str:
        return DEFAULT_COLORS.get(category, DEFAULT_COLORS.get(canonical_finish_category(category), "#808080"))

    def _display_category(self, category: str) -> str:
        configured = self._configured_category(category)
        if configured:
            return configured
        for name in self.category_names():
            if name == "기타":
                return name
        return self.category_names()[0] if self.category_names() else "기타"

    def _configured_category(self, category: str) -> str:
        canonical = canonical_finish_category(category)
        for name in self.category_names():
            if canonical_finish_category(name) == canonical:
                return name
        return ""

    def add_material_unit(self, value: str) -> bool:
        unit = normalize_material_unit(value)
        if not unit:
            return False
        if unit.casefold() in {item.casefold() for item in self.material_units}:
            return False
        self.material_units.append(unit)
        self.material_units_changed.emit(list(self.material_units))
        return True

    def _ensure_units_registered(self, values, title: str = "단위 추가") -> bool:
        unknown = []
        known = {item.casefold() for item in self.material_units}
        for raw in values:
            unit = normalize_material_unit(raw)
            if unit and unit.casefold() not in known and unit.casefold() not in {v.casefold() for v in unknown}:
                unknown.append(unit)
        if not unknown:
            return True
        preview = ", ".join(unknown[:12])
        if len(unknown) > 12:
            preview += f" 외 {len(unknown) - 12}개"
        answer = QMessageBox.question(
            self, title,
            f"단위 리스트에 없는 단위가 있습니다.\n{preview}\n\n단위 리스트에 추가하고 계속할까요?",
            QMessageBox.Yes | QMessageBox.No, QMessageBox.Yes,
        )
        if answer != QMessageBox.Yes:
            return False
        changed = False
        for unit in unknown:
            if unit.casefold() not in {item.casefold() for item in self.material_units}:
                self.material_units.append(unit)
                changed = True
        if changed:
            self.material_units_changed.emit(list(self.material_units))
        return True

    def set_calculation_settings(self, formula_codes=None, finish_categories=None, material_trades=None, material_units=None):
        self.formula_codes = normalize_formula_codes(formula_codes)
        self.finish_categories = normalize_finish_categories(finish_categories)
        self.material_trades = normalize_material_trades(material_trades)
        self.material_units = normalize_material_units(material_units)
        self.loading = True
        self.tabs.blockSignals(True)
        try:
            self.tabs.clear()
            self.tables.clear()
            if self.assignment_enabled:
                self._build_assignment_tabs()
            else:
                self._build_admin_tabs()
        finally:
            self.tabs.blockSignals(False)
            self.loading = False
        self.load()

    def _excel_headers(self, key: str) -> list[str]:
        return list(self.COMPOSITE_HEADERS if key == "복합재료" else self.BASE_HEADERS)

    def _excel_sheet_name(self, key: str) -> str:
        return "복합재료" if key == "복합재료" else "재료마감"

    def export_excel(self, key: str):
        default_name = f"{self._excel_sheet_name(key)}_목록.xlsx"
        path, _ = QFileDialog.getSaveFileName(self, "Excel 내보내기", default_name, "Excel 통합문서 (*.xlsx)")
        if not path:
            return
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = self._excel_sheet_name(key)
            headers = self._excel_headers(key)
            ws.append(headers)
            table = self.tables[key]
            if key == "복합재료" and isinstance(table, QTreeWidget):
                for row in range(table.topLevelItemCount()):
                    item = table.topLevelItem(row)
                    values = [item.text(0), item.text(1), self._composite_category(item), item.text(3), item.text(4), item.text(5), item.text(6)]
                    if values[3]:
                        ws.append(values)
                component_sheet = wb.create_sheet("복합재료 구성")
                component_sheet.append(["복합재료 코드", "재료코드"])
                by_id = {str(material.get("id", "")): material for material in self._base_materials()}
                for row in range(table.topLevelItemCount()):
                    item = table.topLevelItem(row)
                    for component_id in list(item.data(1, ROLE_COMPONENT_IDS) or []):
                        material = by_id.get(str(component_id), {})
                        component_sheet.append([item.text(1), material.get("code", component_id)])
            elif isinstance(table, QTableWidget):
                for row in range(table.rowCount()):
                    values = [self._cell_text(table, row, col) for col in range(table.columnCount())]
                    if any(values[2:]):
                        ws.append(values)

            header_fill = PatternFill("solid", fgColor="D9EAF7")
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
            widths = [6, 20, 14, 28, 30, 20, 10, 10, 24, 14] if key == "재료마감" else [6, 18, 14, 32, 10, 24, 14]
            for idx, width in enumerate(widths, 1):
                ws.column_dimensions[get_column_letter(idx)].width = width
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions

            guide = wb.create_sheet("작성안내")
            guide_rows = [
                ["항목", "작성 기준"],
                ["★", "즐겨찾기는 ★ 또는 빈칸으로 작성"],
                ["재료코드", "필수값. 기본 0000-0000-0001 형식이며 마지막 구간은 길어질 수 있음"],
                ["복합재료 코드", "필수값. 기본 0000-0001 형식이며 마지막 구간은 길어질 수 있음"],
                ["구분", "설정 > 내부 > 재료마감 > 구분 리스트에 등록된 항목 사용"],
                ["공종", "설정 > 재료 > 공종 리스트의 코드와 공종을 선택"],
                ["할증", "백분율 숫자 입력. 예: 3 = 3% / 기본 공란"],
                ["산출식", "구분에 연결된 산식 코드 적용 또는 직접 입력"],
                ["복합재료 구성", "복합재료 Excel의 '복합재료 구성' 시트에서 복합재료 코드와 재료코드로 연결"],
                ["색상", "#RRGGBB 형식. 빈칸이면 구분별 기본색 적용"],
                ["불러오기", "전체 교체 또는 재료코드 기준 병합을 선택"],
            ]
            for row in guide_rows:
                guide.append(row)
            for cell in guide[1]:
                cell.font = Font(bold=True)
                cell.fill = header_fill
            guide.column_dimensions["A"].width = 18
            guide.column_dimensions["B"].width = 62
            wb.save(path)
            QMessageBox.information(self, "Excel 내보내기", f"내보내기를 완료했습니다.\n{path}")
        except Exception as exc:
            QMessageBox.warning(self, "Excel 내보내기", f"Excel 파일을 저장하지 못했습니다.\n{exc}")

    def import_excel(self, key: str):
        path, _ = QFileDialog.getOpenFileName(self, "Excel 불러오기", "", "Excel 통합문서 (*.xlsx)")
        if not path:
            return
        try:
            wb = load_workbook(path, data_only=True, read_only=True)
            sheet_name = self._excel_sheet_name(key)
            ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
            expected = self._excel_headers(key)
            header_row = None
            header_map = {}
            for row_idx in range(1, min(ws.max_row, 10) + 1):
                current = {str(ws.cell(row_idx, c).value or "").strip(): c for c in range(1, ws.max_column + 1)}
                if "산출식" not in current and "산출방법" in current:
                    current["산출식"] = current["산출방법"]
                required_headers = [h for h in expected if h != "할증"]
                if all(h in current for h in required_headers):
                    header_row = row_idx
                    header_map = current
                    break
            if header_row is None:
                raise ValueError("필수 열을 찾지 못했습니다: " + ", ".join(expected))
            existing_records = (
                [self._composite_item_material(self.tables[key].topLevelItem(row)) for row in range(self.tables[key].topLevelItemCount())]
                if key == "복합재료" and isinstance(self.tables.get(key), QTreeWidget)
                else self._base_materials()
            )
            existing_records = [record for record in existing_records if record]
            existing_by_code = {str(record.get("code", "")).casefold(): record for record in existing_records}
            records = []
            seen_codes = set()
            for row_idx in range(header_row + 1, ws.max_row + 1):
                values = [
                    str(ws.cell(row_idx, header_map[h]).value or "").strip() if h in header_map else ""
                    for h in expected
                ]
                name_col = 3 if key == "복합재료" else 4
                if not values[name_col]:
                    continue
                code = self.normalize_code(key, values[1])
                if not code:
                    raise ValueError(f"{row_idx}행 코드는 공란으로 둘 수 없습니다.")
                if code.casefold() in seen_codes:
                    raise ValueError(f"{row_idx}행 코드가 중복되었습니다: {code}")
                seen_codes.add(code.casefold())
                old = existing_by_code.get(code.casefold(), {})
                raw_category = values[2].strip()
                category = self._configured_category(raw_category) if raw_category else ""
                if raw_category and not category:
                    allowed = ", ".join(self.category_names()) or "(등록된 구분 없음)"
                    raise ValueError(f"{row_idx}행 구분은 설정된 목록에서 선택해야 합니다: {allowed}")
                favorite = values[0] in ("★", "1", "Y", "y", "TRUE", "True", "true")
                if key == "복합재료":
                    records.append({
                        "favorite": favorite, "id": old.get("id") or _new_id(), "code": code,
                        "category": category, "target_category": canonical_finish_category(category),
                        "name": values[3], "unit": normalize_material_unit(values[4]) or "M2", "method": normalize_formula_text(values[5]),
                        "component_ids": list(old.get("component_ids", [])),
                        "color": values[6] or old.get("color") or DEFAULT_COLORS["복합재료"],
                        "is_composite": True,
                    })
                else:
                    trade_code, _separator, trade_name = values[3].partition(" ")
                    trade = self._trade_record(trade_code, trade_name)
                    records.append({
                        "favorite": favorite, "id": old.get("id") or _new_id(), "code": code,
                        "category": category, "target_category": canonical_finish_category(category),
                        "trade_code": trade["code"], "trade": trade["name"], "name": values[4],
                        "spec": values[5], "unit": normalize_material_unit(values[6]) or "M2",
                        "markup": values[7], "method": normalize_formula_text(values[8]),
                        "component_ids": [], "color": values[9] or old.get("color") or self._category_color(category),
                        "is_composite": False,
                    })

            imported_units = [record.get("unit", "") for record in records]
            if not self._ensure_units_registered(imported_units, "Excel 단위 추가"):
                return

            if key == "복합재료" and "복합재료 구성" in wb.sheetnames:
                component_ws = wb["복합재료 구성"]
                base_by_code = {str(material.get("code", "")).casefold(): material for material in self._base_materials()}
                composite_by_code = {record["code"].casefold(): record for record in records}
                imported_components: dict[str, list[str]] = {}
                for row_idx in range(2, component_ws.max_row + 1):
                    composite_code = self.normalize_code("복합재료", component_ws.cell(row_idx, 1).value or "")
                    material_code = self.normalize_code("재료마감", component_ws.cell(row_idx, 2).value or "")
                    if not composite_code and not material_code:
                        continue
                    parent = composite_by_code.get(composite_code.casefold())
                    component = base_by_code.get(material_code.casefold())
                    if parent is None:
                        raise ValueError(f"복합재료 구성 {row_idx}행의 복합재료 코드를 찾을 수 없습니다: {composite_code}")
                    if component is None:
                        raise ValueError(f"복합재료 구성 {row_idx}행의 재료코드를 찾을 수 없습니다: {material_code}")
                    values = imported_components.setdefault(composite_code.casefold(), [])
                    component_id = str(component.get("id", ""))
                    if component_id and component_id not in values:
                        values.append(component_id)
                for code, component_ids in imported_components.items():
                    composite_by_code[code]["component_ids"] = component_ids

            box = QMessageBox(self)
            box.setWindowTitle("Excel 불러오기")
            box.setText(f"{len(records):,}개 항목을 불러옵니다.\n기존 목록을 어떻게 처리하시겠습니까?")
            replace_btn = box.addButton("전체 교체", QMessageBox.DestructiveRole)
            merge_btn = box.addButton("코드 기준 병합", QMessageBox.AcceptRole)
            cancel_btn = box.addButton("취소", QMessageBox.RejectRole)
            box.exec()
            clicked = box.clickedButton()
            if clicked == cancel_btn:
                return
            table = self.tables[key]
            if clicked == merge_btn:
                incoming = {record["code"].casefold(): record for record in records}
                order = [record["code"].casefold() for record in existing_records]
                for record in records:
                    if record["code"].casefold() not in order:
                        order.append(record["code"].casefold())
                final_records = [incoming.get(code) or existing_by_code[code] for code in order]
            else:
                final_records = records
            self.loading = True
            try:
                self._load_rows(table, final_records, composite=(key == "복합재료"))
            finally:
                self.loading = False
            self.save()
            self._apply_search()
            QMessageBox.information(self, "Excel 불러오기", f"{len(records):,}개 항목을 반영했습니다.")
        except Exception as exc:
            self.loading = False
            QMessageBox.warning(self, "Excel 불러오기", f"Excel 파일을 읽지 못했습니다.\n{exc}")

    def _on_category_changed(self):
        self._apply_search()
        is_wall = self.current_target_category() == "벽"
        self.wall_mode_btn.setEnabled(is_wall)
        if not is_wall and self.wall_mode_btn.isChecked():
            self.wall_mode_btn.setChecked(False)

    def _apply_search(self):
        if self.assignment_enabled:
            return
        keyword = self.search_edit.text().strip().casefold()
        table = self.tables.get(self.current_category())
        if table is None:
            return
        if isinstance(table, QTreeWidget):
            for row in range(table.topLevelItemCount()):
                item = table.topLevelItem(row)
                texts = [item.text(column) for column in range(table.columnCount())]
                for child_index in range(item.childCount()):
                    child = item.child(child_index)
                    texts.extend(child.text(column) for column in range(table.columnCount()))
                item.setHidden(bool(keyword) and keyword not in " ".join(texts).casefold())
            return
        for row in range(table.rowCount()):
            texts = [table.item(row, c).text() for c in range(table.columnCount()) if table.item(row, c)]
            table.setRowHidden(row, bool(keyword) and keyword not in " ".join(texts).casefold())

    def _sort_current_table(self):
        if self.assignment_enabled:
            return
        table = self.tables.get(self.current_category())
        if table is None:
            return
        if isinstance(table, QTreeWidget):
            table.sortItems(3, Qt.AscendingOrder)
            self.save(); self._apply_search(); return
        name_col = 4
        table.setSortingEnabled(True)
        table.sortItems(name_col, Qt.AscendingOrder)
        table.setSortingEnabled(False)
        self.save(); self._apply_search()

    def reload(self):
        self.load()

    def set_selected_wall_count(self, count: int):
        if self.current_target_category() == "벽":
            self.selection_label.setText(f"선택 벽 구간: {count:,}개")

    def set_selected_room_count(self, count: int):
        if self.current_target_category() != "벽":
            self.selection_label.setText(f"선택 실: {count:,}개")

    def current_category(self) -> str:
        return self.tabs.tabText(self.tabs.currentIndex())

    def current_target_category(self) -> str:
        material = self.current_material()
        if material:
            return material.get("target_category", material.get("category", "벽"))
        category = self.current_category()
        if category in ("복합재료", "재료마감"):
            return "벽"
        return canonical_finish_category(category)

    def all_materials(self) -> list[dict]:
        result = []
        for key, table in self.tables.items():
            if key == "복합재료" and isinstance(table, QTreeWidget):
                for row in range(table.topLevelItemCount()):
                    material = self._composite_item_material(table.topLevelItem(row))
                    if material and material.get("name"):
                        result.append(material)
                continue
            for row in range(table.rowCount()):
                material = self._row_material(key, row)
                if material and material.get("name"):
                    result.append(material)
        return result

    def _row_material(self, key: str, row: int) -> dict | None:
        table = self.tables[key]
        if key == "복합재료" and isinstance(table, QTreeWidget):
            if row < 0 or row >= table.topLevelItemCount():
                return None
            return self._composite_item_material(table.topLevelItem(row))
        if not isinstance(table, QTableWidget):
            return None
        values = [self._cell_text(table, row, col) for col in range(table.columnCount())]
        if key == "재료마감":
            if not values[1] or not values[4]:
                return None
            code_item = table.item(row, 1)
            raw_category = values[2].strip()
            category = (self._configured_category(raw_category) or self._display_category(raw_category)) if raw_category else ""
            target_category = canonical_finish_category(category)
            trade = self._trade_from_row(table, row)
            return {
                "favorite": values[0] == "★", "id": (str(code_item.data(ROLE_INTERNAL_ID) or "") if code_item else "") or _new_id(),
                "code": values[1],
                "category": category, "target_category": target_category,
                "trade_code": trade["code"], "trade": trade["name"], "name": values[4],
                "spec": values[5], "unit": normalize_material_unit(values[6]),
                "markup": values[7], "method": normalize_formula_text(values[8]),
                "component_ids": [], "color": values[9] or self._category_color(category),
                "is_composite": False,
            }
        # 내부 지정 패널
        if key == "복합재료":
            return None
        if not values[1] or not values[4]:
            return None
        code_item = table.item(row, 1)
        category = key
        target_category = canonical_finish_category(category)
        trade = self._trade_from_row(table, row)
        return {
            "favorite": values[0] == "★", "id": (str(code_item.data(ROLE_INTERNAL_ID) or "") if code_item else "") or _new_id(),
            "code": values[1],
            "category": category, "target_category": target_category,
            "trade_code": trade["code"], "trade": trade["name"], "name": values[4],
            "spec": values[5], "unit": normalize_material_unit(values[6]),
            "markup": values[7], "method": normalize_formula_text(values[8]),
            "component_ids": [], "color": values[9] or self._category_color(category),
            "is_composite": False,
        }

    def current_material(self) -> dict | None:
        key = self.current_category()
        table = self.tables.get(key)
        if table is None:
            return None
        if key == "복합재료" and isinstance(table, QTreeWidget):
            return self._composite_item_material(table.currentItem())
        row = table.currentRow()
        return self._row_material(key, row) if row >= 0 else None

    def emit_assign(self):
        material = self.current_material()
        if material is None:
            QMessageBox.information(self, "재료마감 지정", "먼저 재료마감 목록에서 항목을 선택하십시오.")
            return
        self.material_assign_requested.emit(material)

    def _on_composite_item_clicked(self, item: QTreeWidgetItem, column: int):
        if column != 0 or item.parent() is not None:
            return
        self.loading = True
        try:
            item.setText(0, "" if item.text(0).strip() == "★" else "★")
            item.setTextAlignment(0, Qt.AlignCenter)
        finally:
            self.loading = False
        self.save()

    def _base_materials(self) -> list[dict]:
        table = self.tables.get("재료마감")
        if isinstance(table, QTableWidget):
            return [material for row in range(table.rowCount()) if (material := self._row_material("재료마감", row))]
        result = []
        for category in self.category_names():
            category_table = self.tables.get(category)
            if not isinstance(category_table, QTableWidget):
                continue
            result.extend(
                material for row in range(category_table.rowCount())
                if (material := self._row_material(category, row))
            )
        return result

    def _rebuild_composite_children(self, parent: QTreeWidgetItem):
        expanded = parent.isExpanded()
        parent.takeChildren()
        by_id = {str(material.get("id", "")): material for material in self._base_materials()}
        for component_id in list(parent.data(1, ROLE_COMPONENT_IDS) or []):
            material = by_id.get(str(component_id))
            if material is None:
                child = QTreeWidgetItem(["", str(component_id), "", "(재료마감 목록에서 찾을 수 없음)", "", "", ""])
                child.setForeground(3, QColor("#C62828"))
            else:
                child = QTreeWidgetItem([
                    "", str(material.get("code", "")), str(material.get("category", "")),
                    str(material.get("name", "")), normalize_material_unit(material.get("unit", "")),
                    normalize_formula_text(material.get("method", "")), str(material.get("color", "")),
                ])
            child.setData(1, ROLE_INTERNAL_ID, str(component_id))
            flags = Qt.ItemIsEnabled | Qt.ItemIsSelectable
            child.setFlags(flags)
            parent.addChild(child)
        parent.setExpanded(expanded or parent.childCount() > 0)

    def open_composite_component_picker(self):
        if self.assignment_enabled:
            return
        tree = self.tables.get("복합재료")
        if not isinstance(tree, QTreeWidget):
            return
        parent = tree.currentItem()
        if parent is not None and parent.parent() is not None:
            parent = parent.parent()
        if parent is None:
            QMessageBox.information(self, "복합재료 구성", "먼저 복합재료 항목을 선택해 주세요.")
            return
        materials = self._base_materials()
        if not materials:
            QMessageBox.information(self, "복합재료 구성", "먼저 재료마감 탭에 재료를 등록해 주세요.")
            return
        existing = list(parent.data(1, ROLE_COMPONENT_IDS) or [])
        dialog = ComponentMaterialDialog(materials, existing, self)
        if dialog.exec() != QDialog.Accepted:
            return
        for material in dialog.selected_materials:
            internal_id = str(material.get("id", ""))
            if internal_id and internal_id not in existing:
                existing.append(internal_id)
        parent.setData(1, ROLE_COMPONENT_IDS, existing)
        self._rebuild_composite_children(parent)
        tree.setCurrentItem(parent)
        self.save()

    def _replace_composite_component_reference(self, old_id: str, replacement_id: str = ""):
        tree = self.tables.get("복합재료")
        if not isinstance(tree, QTreeWidget):
            return
        for row in range(tree.topLevelItemCount()):
            parent = tree.topLevelItem(row)
            values = []
            for component_id in list(parent.data(1, ROLE_COMPONENT_IDS) or []):
                value = str(component_id)
                if value == str(old_id):
                    value = str(replacement_id or "")
                if value and value not in values:
                    values.append(value)
            parent.setData(1, ROLE_COMPONENT_IDS, values)
            self._rebuild_composite_children(parent)

    def _toggle_star(self, table: QTableWidget, row: int):
        item = table.item(row, 0)
        if item is None:
            item = QTableWidgetItem(""); table.setItem(row, 0, item)
        item.setText("" if item.text() == "★" else "★")
        item.setTextAlignment(Qt.AlignCenter)
        self.save()

    def _current_insert_row(self, key: str) -> int | None:
        table = self.tables.get(key)
        if key == "복합재료" and isinstance(table, QTreeWidget):
            current = table.currentItem()
            if current is not None:
                if current.parent() is not None:
                    current = current.parent()
                index = table.indexOfTopLevelItem(current)
                return index if index >= 0 else None
        elif isinstance(table, QTableWidget):
            current_row = table.currentRow()
            return current_row if current_row >= 0 else None
        return None

    def _add_row_at_current(self, key: str):
        self.add_row(key, insert_row=self._current_insert_row(key))

    def _add_current_row_f9(self):
        """F9: 현재 행 자리에 새 항목을 삽입하고 기존 현재 행부터 아래로 민다."""
        self._add_row_at_current(self.current_category())

    def add_row(self, key: str, insert_row: int | None = None):
        code = self.request_new_code(key)
        if not code:
            return
        table = self.tables[key]
        if key == "복합재료":
            if not isinstance(table, QTreeWidget):
                return
            category = self.category_names()[0] if self.category_names() else "기타"
            _formula_code, default_formula = formula_for_category(category, self.formula_codes, self.finish_categories)
            item = QTreeWidgetItem(["", code, category, "", "M2", default_formula, DEFAULT_COLORS["복합재료"]])
            item.setData(1, ROLE_INTERNAL_ID, _new_id())
            item.setData(1, ROLE_COMPONENT_IDS, [])
            item.setFlags(item.flags() | Qt.ItemIsEditable)
            item.setTextAlignment(0, Qt.AlignCenter)
            if insert_row is None or insert_row < 0 or insert_row >= table.topLevelItemCount():
                table.addTopLevelItem(item)
            else:
                table.insertTopLevelItem(insert_row, item)
            self._install_composite_category_combo(table, item, category)
            table.setCurrentItem(item, 3)
            table.editItem(item, 3)
            self.save()
            return
        else:
            if not isinstance(table, QTableWidget):
                return
            row = table.rowCount() if insert_row is None or insert_row < 0 else min(insert_row, table.rowCount())
            table.insertRow(row)
            if key == "재료마감":
                # 신규 재료는 사용자가 부위를 결정하기 전까지 구분/산출식을 비워 둔다.
                category = ""
                default_formula = ""
            else:
                category = key
                _code, default_formula = formula_for_category(category, self.formula_codes, self.finish_categories)
            defaults = ["", code, category, "", "", "", "M2", "", default_formula, self._category_color(category)]
            name_col = 4
        for col, value in enumerate(defaults):
            item = QTableWidgetItem(value)
            if col == 1:
                item.setData(ROLE_INTERNAL_ID, _new_id())
            if self.assignment_enabled:
                item.setFlags(item.flags() & ~Qt.ItemIsEditable)
            table.setItem(row, col, item)
        if key == "재료마감":
            self._install_category_combo(table, row, category)
            self._install_trade_combo(table, row)
        table.setCurrentCell(row, name_col)
        if not self.assignment_enabled:
            table.editItem(table.item(row, name_col))
        self.save()

    def _duplicate_current_row(self, key: str):
        table = self.tables.get(key)
        if key == "복합재료" and isinstance(table, QTreeWidget):
            current = table.currentItem()
            if current is None:
                return
            if current.parent() is not None:
                current = current.parent()
            src_index = table.indexOfTopLevelItem(current)
            code = self.request_new_code(key)
            if not code:
                return
            values = [current.text(col) for col in range(len(self.COMPOSITE_HEADERS))]
            values[1] = code
            item = QTreeWidgetItem(values)
            item.setData(1, ROLE_INTERNAL_ID, _new_id())
            item.setData(1, ROLE_COMPONENT_IDS, list(current.data(1, ROLE_COMPONENT_IDS) or []))
            item.setFlags(current.flags())
            table.insertTopLevelItem(src_index + 1, item)
            self._install_composite_category_combo(table, item, values[2])
            self._rebuild_composite_children(item)
            table.setCurrentItem(item, 3)
            self.save()
            return
        if not isinstance(table, QTableWidget):
            return
        row = table.currentRow()
        if row < 0:
            return
        code = self.request_new_code(key)
        if not code:
            return
        values = [table.item(row, col).text() if table.item(row, col) else "" for col in range(table.columnCount())]
        insert = row + 1
        table.insertRow(insert)
        for col, value in enumerate(values):
            if col == 1:
                value = code
            item = QTableWidgetItem(value)
            if col == 1:
                item.setData(ROLE_INTERNAL_ID, _new_id())
            if self.assignment_enabled:
                item.setFlags(item.flags() & ~Qt.ItemIsEditable)
            table.setItem(insert, col, item)
        if key == "재료마감":
            self._install_category_combo(table, insert, values[2] if len(values) > 2 else "")
            self._install_trade_combo(table, insert)
        table.setCurrentCell(insert, min(4, table.columnCount() - 1))
        self.save()

    def delete_selected(self, key: str):
        table = self.tables[key]
        by_f4 = bool(getattr(table, "_delete_requested_by_f4", False))
        if by_f4:
            table._delete_requested_by_f4 = False
        if key == "복합재료" and isinstance(table, QTreeWidget):
            item = table.currentItem()
            if item is None:
                return
            if item.parent() is not None:
                if by_f4:
                    name = str(item.text(4) or item.text(3) or "").strip()
                    text = f"'{name}' 항목을 삭제하시겠습니까?" if name else "현재 항목을 삭제하시겠습니까?"
                    if QMessageBox.question(self, "목록 삭제", text, QMessageBox.Yes | QMessageBox.No, QMessageBox.Yes) != QMessageBox.Yes:
                        return
                parent = item.parent()
                component_id = str(item.data(1, ROLE_INTERNAL_ID) or "")
                values = [str(value) for value in list(parent.data(1, ROLE_COMPONENT_IDS) or [])]
                parent.setData(1, ROLE_COMPONENT_IDS, [value for value in values if value != component_id])
                self._rebuild_composite_children(parent)
                self.save()
                return
            material = self._composite_item_material(item)
            if material is None:
                table.takeTopLevelItem(table.indexOfTopLevelItem(item)); self.save(); return
            answer = QMessageBox.question(self, "복합재료 삭제", f"'{material['name']}' 복합재료를 삭제할까요?", QMessageBox.Yes | QMessageBox.No, QMessageBox.No)
            if answer != QMessageBox.Yes:
                return
            self.material_delete_requested.emit(material["id"], None)
            table.takeTopLevelItem(table.indexOfTopLevelItem(item)); self.save(); return
        if not isinstance(table, QTableWidget):
            return
        row = table.currentRow()
        if row < 0: return
        material = self._row_material(key, row)
        if material is None:
            if by_f4 and QMessageBox.question(
                self, "목록 삭제", "현재 행을 삭제하시겠습니까?",
                QMessageBox.Yes | QMessageBox.No, QMessageBox.Yes
            ) != QMessageBox.Yes:
                return
            table.removeRow(row); self.save(); return
        box = QMessageBox(self)
        box.setWindowTitle("재료마감 삭제")
        box.setText(f"'{material['name']}'을(를) 삭제합니다.\n기존에 지정된 대상은 어떻게 처리하시겠습니까?")
        replace_btn = box.addButton("다른 재료로 변경", QMessageBox.AcceptRole)
        clear_btn = box.addButton("기존 지정 삭제", QMessageBox.DestructiveRole)
        cancel_btn = box.addButton("취소", QMessageBox.RejectRole)
        box.exec(); clicked = box.clickedButton()
        if clicked == cancel_btn: return
        replacement = None
        if clicked == replace_btn:
            candidates = [
                m for m in self.all_materials()
                if m["id"] != material["id"]
                and not m.get("is_composite")
                and m.get("target_category") == material.get("target_category")
            ]
            if not candidates:
                QMessageBox.information(self, "대체 재료", "같은 구분의 대체 재료가 없습니다."); return
            labels = [f"{m['name']} / {m.get('spec','')} [{m.get('code', '')}]" for m in candidates]
            choice, ok = QInputDialog.getItem(self, "대체 재료 선택", "변경할 재료:", labels, 0, False)
            if not ok: return
            replacement = candidates[labels.index(choice)]
        self.material_delete_requested.emit(material["id"], replacement)
        self._replace_composite_component_reference(
            material["id"], replacement.get("id", "") if isinstance(replacement, dict) else ""
        )
        table.removeRow(row); self.save()

    def pick_selected_color(self, key: str):
        if key == "복합재료":
            tree = self.tables.get(key)
            if not isinstance(tree, QTreeWidget):
                return
            item = tree.currentItem()
            if item is not None and item.parent() is not None:
                item = item.parent()
            if item is None:
                return
            color = QColorDialog.getColor(QColor(item.text(6) or DEFAULT_COLORS["복합재료"]), self, "복합재료 표시 색상")
            if color.isValid():
                item.setText(6, color.name()); self.save()
            return
        row = self.tables[key].currentRow()
        if row >= 0:
            table = self.tables[key]
            self._pick_color_from_table(table, row, table.columnCount() - 1)

    def _pick_color_from_table(self, table: QTableWidget, row: int, col: int):
        if col == 0:
            return
        color_col = table.columnCount() - 1
        if col != color_col:
            return
        item = table.item(row, color_col)
        current = QColor(item.text() if item else "#808080")
        color = QColorDialog.getColor(current, self, "마감재 표시 색상")
        if color.isValid():
            if item is None:
                item = QTableWidgetItem(); table.setItem(row, color_col, item)
            item.setText(color.name()); self.save()

    def _normalize_data(self, raw: dict) -> tuple[list[dict], list[dict]]:
        if "materials" in raw or "composites" in raw:
            return list(raw.get("materials", [])), list(raw.get("composites", []))
        materials, composites = [], []
        for category in LEGACY_BASE_CATEGORIES:
            for record in raw.get(category, []):
                copied = dict(record); copied["category"] = category; copied["target_category"] = category
                materials.append(copied)
        for record in raw.get("복합재료", []):
            composites.append(dict(record))
        return materials, composites

    def load(self):
        self.loading = True; migrated = False
        try:
            raw = json.loads(self.data_path.read_text(encoding="utf-8")) if self.data_path.exists() else {}
            materials, composites = self._normalize_data(raw)
            if not ("materials" in raw or "composites" in raw): migrated = bool(raw)
            migrated = migrated or any(not str(record.get("code", "") or "").strip() for record in [*materials, *composites])
            migrated = migrated or any(
                normalize_material_unit(record.get("unit", "")) != str(record.get("unit", "") or "").strip()
                for record in [*materials, *composites]
            )
            for table in self.tables.values():
                if isinstance(table, QTreeWidget): table.clear()
                else: table.setRowCount(0)
            if self.assignment_enabled:
                grouped = {c: [] for c in self.category_names()}
                for record in materials:
                    raw_category = str(record.get("category") or record.get("target_category") or "").strip()
                    if not raw_category:
                        # 관리표에서 아직 구분을 정하지 않은 신규 재료는 지정 선택창에 노출하지 않는다.
                        continue
                    category = self._configured_category(raw_category)
                    if not category:
                        category = self._display_category(raw_category)
                    if category in grouped:
                        grouped[category].append(record)
                for category in self.category_names():
                    self._load_rows(self.tables[category], grouped.get(category, []), composite=False, forced_category=category)
                self._load_rows(self.tables["복합재료"], composites, composite=True)
            else:
                self._load_rows(self.tables["재료마감"], materials, composite=False)
                self._load_rows(self.tables["복합재료"], composites, composite=True)
        except Exception as exc:
            QMessageBox.warning(self, "재료마감 불러오기", f"재료마감 목록을 읽지 못했습니다.\n{exc}")
        finally:
            self.loading = False
        if migrated: self.save()
        self._on_category_changed(); self._apply_search()

    def _load_rows(self, table: QTableWidget, rows: list[dict], composite: bool, forced_category: str | None = None):
        used_codes: set[str] = set()
        previous_code = ""
        if composite and isinstance(table, QTreeWidget):
            table.clear()
            for record in rows:
                raw_code = self.normalize_code("복합재료", record.get("code", ""))
                code = raw_code if raw_code and raw_code.casefold() not in used_codes else increment_user_code(previous_code, COMPOSITE_CODE_LENGTH, 1)
                while code.casefold() in used_codes:
                    code = increment_user_code(code, COMPOSITE_CODE_LENGTH, 1)
                used_codes.add(code.casefold()); previous_code = code
                legacy_category = record.get("category") or record.get("target_category")
                category = self._configured_category(legacy_category)
                if not category:
                    category = self._display_category("기타" if legacy_category == "복합재료" else legacy_category)
                component_ids = record.get("component_ids", [])
                if isinstance(component_ids, str):
                    component_ids = [value.strip() for value in component_ids.split(",") if value.strip()]
                item = QTreeWidgetItem([
                    "★" if record.get("favorite") else "", code, category,
                    str(record.get("name", "") or ""), normalize_material_unit(record.get("unit", "")),
                    normalize_formula_text(record.get("method", "")),
                    str(record.get("color", DEFAULT_COLORS["복합재료"]) or DEFAULT_COLORS["복합재료"]),
                ])
                item.setData(1, ROLE_INTERNAL_ID, record.get("id") or _new_id())
                item.setData(1, ROLE_COMPONENT_IDS, list(component_ids))
                if not self.assignment_enabled:
                    item.setFlags(item.flags() | Qt.ItemIsEditable)
                else:
                    item.setFlags(item.flags() & ~Qt.ItemIsEditable)
                item.setTextAlignment(0, Qt.AlignCenter)
                table.addTopLevelItem(item)
                self._install_composite_category_combo(table, item, category)
                self._rebuild_composite_children(item)
            return
        table.setRowCount(len(rows))
        for row, record in enumerate(rows):
            material_id = record.get("id") or _new_id()
            raw_code = self.normalize_code("재료마감", record.get("code", ""))
            code = raw_code if raw_code and raw_code.casefold() not in used_codes else increment_user_code(previous_code, MATERIAL_CODE_LENGTH, 2)
            while code.casefold() in used_codes:
                code = increment_user_code(code, MATERIAL_CODE_LENGTH, 2)
            used_codes.add(code.casefold()); previous_code = code
            favorite = "★" if record.get("favorite") else ""
            raw_category = str(record.get("category") or record.get("target_category") or "").strip()
            category = forced_category or (self._display_category(raw_category) if raw_category else "")
            trade_code = str(record.get("trade_code", "") or "").strip()
            trade_name = str(record.get("trade", record.get("trade_name", "")) or "").strip()
            values = [
                favorite, code, category, material_trade_display(trade_code, trade_name),
                record.get("name", ""), record.get("spec", ""), normalize_material_unit(record.get("unit", "")),
                record.get("markup", ""), normalize_formula_text(record.get("method", "")), record.get("color", self._category_color(category)),
            ]
            for col, value in enumerate(values):
                item = QTableWidgetItem(str(value))
                if col == 0: item.setTextAlignment(Qt.AlignCenter)
                if col == 1: item.setData(ROLE_INTERNAL_ID, material_id)
                if self.assignment_enabled: item.setFlags(item.flags() & ~Qt.ItemIsEditable)
                table.setItem(row, col, item)
            if table is self.tables.get("재료마감"):
                self._install_category_combo(table, row, category)
                self._install_trade_combo(table, row, trade_code, trade_name)
            else:
                self._install_trade_combo(table, row, trade_code, trade_name)

    @staticmethod
    def _material_summary_key(material: dict) -> str:
        return str(material.get("id") or material.get("code") or material.get("name") or "미지정")

    def apply_summary_markups(self, markups) -> int:
        """집계 화면에서 직접 입력한 할증값을 기본 재료의 할증 열로 복사한다."""
        if self.assignment_enabled:
            return 0
        table = self.tables.get("재료마감")
        if not isinstance(table, QTableWidget):
            return 0
        source = {}
        for key, value in dict(markups or {}).items():
            try:
                source[str(key)] = float(value)
            except (TypeError, ValueError):
                continue
        if not source:
            return 0
        changed = 0
        self.loading = True
        try:
            for row in range(table.rowCount()):
                material = self._row_material("재료마감", row)
                if not material:
                    continue
                key = self._material_summary_key(material)
                if key not in source:
                    continue
                item = table.item(row, 7)
                if item is None:
                    item = QTableWidgetItem()
                    table.setItem(row, 7, item)
                text = f"{source[key]:g}"
                if item.text().strip() != text:
                    item.setText(text)
                    changed += 1
        finally:
            self.loading = False
        if changed:
            self.save()
        return changed

    def save(self):
        if self.loading or self.assignment_enabled:
            return

        # 단위는 화면/저장값 모두 M2/M3 표기로 즉시 정규화한다.
        self.loading = True
        try:
            material_table = self.tables.get("재료마감")
            if isinstance(material_table, QTableWidget):
                for row in range(material_table.rowCount()):
                    unit_item = material_table.item(row, 6)
                    if unit_item is not None:
                        normalized = normalize_material_unit(unit_item.text())
                        if unit_item.text().strip() != normalized:
                            unit_item.setText(normalized)
            composite_tree = self.tables.get("복합재료")
            if isinstance(composite_tree, QTreeWidget):
                for row in range(composite_tree.topLevelItemCount()):
                    item = composite_tree.topLevelItem(row)
                    normalized = normalize_material_unit(item.text(4))
                    if item.text(4).strip() != normalized:
                        item.setText(4, normalized)
        finally:
            self.loading = False

        materials, composites = [], []
        for row in range(self.tables["재료마감"].rowCount()):
            material = self._row_material("재료마감", row)
            if material: materials.append(material)
        composite_tree = self.tables["복합재료"]
        if isinstance(composite_tree, QTreeWidget):
            for row in range(composite_tree.topLevelItemCount()):
                material = self._composite_item_material(composite_tree.topLevelItem(row))
                if material: composites.append(material)
        data = {"schema_version": 4, "materials": materials, "composites": composites}
        self.data_path.parent.mkdir(parents=True, exist_ok=True)
        self.data_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        self.materials_changed.emit()
